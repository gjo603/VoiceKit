"""VoiceKit write-side port.

Reproduces, in pure Python, exactly what VoiceKit's AutoHotkey generators
(`macros\\NewAutomation.ahk`, `macros\\WorkflowStudio.ahk`) write to disk, so
an MCP server can create the same automations non-interactively.

This module has NO MCP dependency on purpose: the conformance test imports it
directly and feeds its output through the real AutoHotkey engine to prove the
bytes match. Keep every format function in sync with its AHK source (named in
the docstrings); the conformance test is what enforces that.
"""

from __future__ import annotations

import csv
import functools
import hashlib
import io
import itertools
import json
import os
import re
import subprocess
import tempfile
import threading
import time
from datetime import datetime
from pathlib import Path

import privacy

# ---------------------------------------------------------------------------
# Locations (self-locating; overridable by env for tests / odd installs)
# ---------------------------------------------------------------------------
REPO_ROOT = Path(os.environ.get("VOICEKIT_ROOT", Path(__file__).resolve().parent.parent))
AHK_EXE = os.environ.get(
    "VOICEKIT_AHK",
    os.path.join(os.environ.get("ProgramFiles", r"C:\Program Files"),
                 "AutoHotkey", "v2", "AutoHotkey64.exe"),
)
# (Path accessors that derive from REPO_ROOT live with the health helpers
# below and are functions, not constants, so tests can retarget REPO_ROOT.)
VOICE_MACROS = Path(os.environ.get(
    "VOICEKIT_STARTMENU",
    os.path.join(os.environ.get("APPDATA", ""), "Microsoft", "Windows",
                 "Start Menu", "Programs", "Voice Macros"),
))

# Bridge-key pool from lib\_Common.ahk BridgeKeyPool — E,N,R,X,H,I,C excluded
# (H, I and C are Workflow Studio's mark-hover, ask-for-input and collect
# keys while recording).
# (Companion hotkeys — hotkeys\<Base>.hotkey.ahk, assigned in the Voice Kit
# home window — draw from the same pool via bridge-map.txt registration.)
# Punctuation trails the letters/digits so auto-allocation still prefers a
# letter and the symbols stay free for someone who asks for one. Each symbol
# was measured usable end to end (generated module loads, Hotkey() registers
# it, a synthesized SendLevel-1 press fires it). Deliberately absent: '"'
# (breaks generated string literals), backtick (AHK's escape char), '|'
# (bridge-map delimiter), Space/Enter (register but never fire).
BRIDGE_POOL = "ABDFGJKLMOPQSTUVWYZ0123456789[];',./-=\\"
# The letters VoiceKit keeps for itself (E, N, R: the master; X, H, I, C:
# Workflow Studio while recording / looping). Never in the pool, and named in
# every message and listing that explains why a key isn't offered.
RESERVED_BRIDGE_KEYS = "ENRXHIC"

# Step types from lib\Workflow.ahk WfRunStep.
STEP_TYPES = ("run", "focus", "waitwin", "wait", "text", "keys",
              "click", "dblclick", "rclick", "hover", "move", "close",
              # drag|<window>||x1,y1,x2,y2 — press at the first point, travel,
              # release at the second (window-relative). paramB stays empty:
              # a drag has no element name; it is inherently positional.
              "drag",
              # ask|<label>|<suggested answer>| — the run collects every ask
              # input up front (one dialog per unique label) and types the
              # answer at this step's position. The loop runner can batch
              # them (typed-in rows or a CSV whose columns are the labels).
              "ask",
              # collect|<label>|<element name>| — grabs a value at this
              # position and saves it under the label (a column in the
              # workflow's <Base>.inputs.csv sheet). Empty element = copy the
              # current selection; a name = read that box's accessible value
              # in the active window.
              "collect",
              # set|<name>|<value>| — names a value the rest of the run can
              # reuse by writing {{name}}. The value may itself contain
              # {{...}}, so values compose. Every ask label and collect label
              # is a name too; {{clipboard}} / {{date}} / {{time}} /
              # {{datetime}} / {{selected_file}} / {{selected_files}} (the
              # File Explorer selection when the run starts) are built-in
              # fallbacks. An unknown name is left literal, so pre-existing
              # workflows are unaffected.
              "set",
              # capture|<name>|<command line>|<timeout seconds> — runs a cmd.exe
              # command line hidden (cwd = the VoiceKit root, UTF-8 output) and
              # keeps what it printed on stdout as {{name}}, run-local like
              # set. stderr is kept apart; a nonzero exit stops the run with
              # only "Command exited with code N" in the record (a traceback
              # can name a client's file). Timeout default 30 s; the process
              # tree is killed on timeout or Stop. {{selected_file}} /
              # {{selected_files}} arrive QUOTED in the command (and in a run
              # target), written with or without quotes around them.
              "capture",
              # fill|<window>|<label>[#N]|<value> — find an input box by its
              # label (its accessible name; "Amount#2" = the 2nd box with that
              # label, "##" a literal "#"), click it, refuse to type unless the
              # keyboard focus landed on it, type the value and READ IT BACK.
              # paramC is the value and — alone among paramC slots — takes
              # {{Name}}. "" clears the box; a line break is refused. A failure
              # reason never contains the value (it may be client data).
              "fill",
              # waitfor|<window>|<element or text>|<condType>[,<seconds>] —
              # block until something is actually true rather than betting on a
              # hand-tuned number of milliseconds, and fail naming what never
              # happened. Shares its condition vocabulary with `if`; seconds
              # default to 10 and ride in paramC beside the condType.
              "waitfor",
              # Optional branching (engine executes these; recorder never emits them).
              # if|<window>|<element>|<condType> where condType is one of
              # winexists / winnotexists / elementexists / elementnotexists /
              # textvisible / textnotvisible; paired with "else" and "endif".
              "if", "else", "endif")

# What a `wait` step's paramA may hold: whole milliseconds, or a range
# ("600-1400") the engine turns into Random(min, max). The range exists so a
# workflow looping against a website doesn't pause identically every pass —
# metronome timing is the cheapest bot signal there is. Mirrors WfWaitMs.
WAIT_MS_RE = re.compile(r"^\d+(\s*-\s*\d+)?$")

# A capture step's timeout (paramC): blank = the engine's 30 s, else a plain
# positive number of seconds (no sign, no exponent, no nan/inf — the forms
# AHK's IsNumber and Python's float() disagree on). Mirrors WfCaptureSecs.
CAPTURE_SECS_RE = re.compile(r"^\d+(\.\d+)?$")

_FILL_LABEL_RE = re.compile(r"^(.*?)(#+)([0-9]+)$", re.S)


def fill_label(raw: str) -> dict:
    """A fill step's label -> {name, occ, explicit, err}. Mirrors lib\\Workflow.ahk
    WfFillLabel exactly (tests\\engine-fill-selftest.ahk and test_conformance
    share one table): "Amount#2" is the 2nd input labelled Amount; an ODD run
    of # before trailing digits ends in the separator, the rest are "##"
    pairs meaning a literal "#"; any other "#" is literal."""
    s = (raw or "").strip(" \t")          # AHK Trim: spaces and tabs
    occ, explicit = 1, False
    m = _FILL_LABEL_RE.match(s)
    if m and len(m.group(2)) % 2 == 1:
        occ = 999999 if len(m.group(3)) > 6 else int(m.group(3))
        explicit = True
        s = m.group(1) + m.group(2)[1:]
    name = s.replace("##", "#").strip(" \t")
    err = ""
    if not name:
        err = "has no label — give the input's label as it appears on screen"
    elif explicit and occ < 1:
        err = "'#0' isn't an input number — count from #1 (the first input with that label)"
    return {"name": name, "occ": occ, "explicit": explicit, "err": err}


# Conditions shared by `if` and `waitfor` (lib\Workflow.ahk WfEvalCond).
# clipboardchanged is waitfor-only: "changed" needs a before-value, which a
# point-in-time `if` has no way to supply.
WAIT_CONDS = ("winexists", "winnotexists", "elementexists", "elementnotexists",
              "textvisible", "textnotvisible", "clipboardchanged")
IF_CONDS = tuple(c for c in WAIT_CONDS if c != "clipboardchanged")
# Conditions that need something to look FOR in paramB.
CONDS_NEEDING_ELEMENT = ("elementexists", "elementnotexists",
                         "textvisible", "textnotvisible")

# Where a `move` step can put a window (lib\Workflow.ahk WfRunStep "move").
MOVE_POSITIONS = ("left", "right", "top", "bottom", "max")
# A pointer step's window-relative position (paramC): "x,y" / "x1,y1,x2,y2".
_XY_RE = re.compile(r"^-?\d+\s*,\s*-?\d+$")
_DRAG_RE = re.compile(r"^-?\d+\s*,\s*-?\d+\s*,\s*-?\d+\s*,\s*-?\d+$")

# Marker line that identifies a generated workflow stub — the overwrite and
# delete guards on both sides look for it (lib\_Common.ahk IsWorkflowStub).
# The whole phrase, not just "Workflow Studio": several of VoiceKit's own
# tools merely MENTION the Studio, and the looser check let a workflow named
# after one of them overwrite (or delete) that tool's script.
STUDIO_MARKER = "Generated by Workflow Studio"

# Companion-hotkey module suffix from lib\_Common.ahk HotkeyCompanionRel:
# hotkeys\<Base>.hotkey.ahk is a keyboard trigger for macros\<Base>.ahk,
# assigned in the home window. The suffix doubles as the marker (scaffolded
# module names come from CleanPhrase and can never contain a dot).
COMPANION_SUFFIX = ".hotkey.ahk"


def _companion_rel(base: str) -> str:
    """Repo-relative bridge-map FILE field of <base>'s companion module."""
    return f"hotkeys\\{base}{COMPANION_SUFFIX}"


def _companion_parent(relfile: str) -> str:
    """Parent automation base for a bridge-map FILE field, or "" when the
    file isn't a companion module (mirrors _Common.ahk HotkeyCompanionParent)."""
    m = re.match(rf"(?i)^hotkeys\\(.+){re.escape(COMPANION_SUFFIX)}$", relfile)
    return m.group(1) if m else ""


class VoiceKitError(Exception):
    """A user-facing problem (bad name, collision, validation failure)."""


# ---------------------------------------------------------------------------
# Naming / encoding — byte-for-byte ports of the AHK helpers
# ---------------------------------------------------------------------------
def _title(s: str) -> str:
    """AHK StrTitle: capitalize the first letter of each SPACE-separated word,
    lowercase the rest. Word boundary is space only (NOT digits) — verified:
    'open2tabs' -> 'Open2tabs'. (Python str.title() differs on digits.)"""
    return " ".join(w[:1].upper() + w[1:].lower() for w in s.split(" "))


def clean_phrase(raw: str) -> str:
    """lib\\_Common.ahk CleanPhrase: strip all but [A-Za-z0-9 ], collapse
    whitespace, Title Case."""
    p = re.sub(r"[^A-Za-z0-9 ]", "", raw.strip())
    p = re.sub(r"\s+", " ", p)
    return _title(p.strip(" "))     # trimmed after the strip too: "! foo" -> "Foo"


def to_base(phrase: str) -> str:
    """Filename base = phrase with spaces removed (NewAutomation / SaveWorkflow)."""
    return phrase.replace(" ", "")


def is_reserved(base: str) -> bool:
    """lib\\_Common.ahk IsReservedName — Windows device names (case-insensitive)."""
    return bool(re.match(r"(?i)^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])$", base))


def space_out(base: str) -> str:
    """SpaceOut: re-insert a space between a lowercase/digit and a following
    uppercase. 'OpenPublicUserFolder' -> 'Open Public User Folder'."""
    return re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", base).strip()


def ahk_str_lit(s: str) -> str:
    """lib\\_Common.ahk AhkStrLit: render a value as an AutoHotkey v2
    double-quoted string literal (backtick doubled, quote escaped, every ';'
    escaped, CR dropped, LF -> `n). The ';' matters: one preceded by
    whitespace starts a COMMENT even inside a string, so an 'opens' target
    like 'https://x/?q=a ;b' used to generate a file that couldn't load.
    test_ahk_str_lit_matches_ahk runs the AHK original on the same inputs."""
    s = s.replace("`", "``")
    s = s.replace('"', '`"')
    s = s.replace(";", "`;")
    s = s.replace("\r", "")
    s = s.replace("\n", "`n")
    return '"' + s + '"'


def snip_encode(text: str) -> str:
    """lib\\_Common.ahk SnipEncode: a snippet's text as its one-line on-disk
    form — backticks doubled first (the escape char), every ';' escaped (a
    bare one would start a comment mid-line), newlines normalized and folded
    into a literal `n. test_generators_match_ahk runs the AHK original."""
    s = text.replace("`", "``")
    s = s.replace(";", "`;")
    s = s.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "`n")
    return s


def wf_encode(s: str) -> str:
    """lib\\Workflow.ahk WfEncode — order-sensitive: % first, then | CR LF."""
    s = s.replace("%", "%25")
    s = s.replace("|", "%7C")
    s = s.replace("\r", "%0D")
    s = s.replace("\n", "%0A")
    return s


def wf_decode(s: str) -> str:
    """lib\\Workflow.ahk WfDecode — reverse order: LF CR | then % last."""
    s = s.replace("%0A", "\n")
    s = s.replace("%0D", "\r")
    s = s.replace("%7C", "|")
    s = s.replace("%25", "%")
    return s


_COMBO_PREFIX = "Ctrl+Alt+Shift+"


def _used_bridge_keys() -> set[str]:
    """Pool keys bridge-map.txt already holds — THE rule for "is this key
    taken", shared by the allocator, get_bridge_map's free list and a
    requested key= check, and mirrored exactly by lib\\_Common.ahk
    BridgeFreeKeys (test_used_bridge_keys_match_ahk pins the two).

    A key is used when EITHER
      * the raw text contains 'Ctrl+Alt+Shift+<k>|' anywhere, caselessly
        (so a combo mentioned mid-line in a comment still counts), OR
      * a line — commented-out lines included, leading ';'/blanks stripped —
        has a first '|' field that, trimmed, is 'Ctrl+Alt+Shift+<k>'
        (so a hand-padded 'Ctrl+Alt+Shift+B | ...' counts).
    Commented lines count on purpose: a parked registration is a pairing the
    user may switch back on, and handing its key to something else would
    make the two collide the moment they do. Two rules used to disagree on
    exactly these cases and hand out one key twice."""
    mapfile = _bridge_map_file()
    text = _read_text_any(mapfile) if mapfile.exists() else ""
    used: set[str] = set()
    for k in BRIDGE_POOL:
        # re.A: ASCII-only caseless matching, like AHK's InStr — Unicode case
        # folding could otherwise turn exotic letters into a false match.
        if re.search(re.escape(_COMBO_PREFIX + k + "|"), text, re.I | re.A):
            used.add(k)
    for line in text.split("\n"):       # like AHK: Loop Parse text, "`n", "`r"
        first = line.strip("\r").lstrip("; \t").split("|", 1)[0].strip(" \t")
        if first[:len(_COMBO_PREFIX)].lower() == _COMBO_PREFIX.lower():
            k = first[len(_COMBO_PREFIX):].strip(" \t")
            if len(k) == 1 and k.upper() in BRIDGE_POOL:
                used.add(k.upper())
    return used


def allocate_bridge_key() -> str | None:
    """First pool key _used_bridge_keys() doesn't hold (the same free list
    lib\\_Common.ahk BridgeFreeKeys gives New Automation's key dropdown).
    None if the pool is exhausted."""
    used = _used_bridge_keys()
    return next((k for k in BRIDGE_POOL if k not in used), None)


# ---------------------------------------------------------------------------
# File I/O — reproduce AHK's exact bytes (BOM on create, LF endings)
# ---------------------------------------------------------------------------
def _write_new(path: Path, content: str) -> None:
    """Create a file the way AHK FileOpen(.., 'w', 'UTF-8') does: UTF-8 BOM,
    LF line endings (newline='' stops Python translating \\n to \\r\\n)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        f.write(content)


def _append(path: Path, text: str) -> None:
    """Append like AHK FileAppend to an existing file: no BOM, LF preserved."""
    with open(path, "a", encoding="utf-8", newline="") as f:
        f.write(text)


def _today() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def log(text: str) -> None:
    """lib\\_Common.ahk Log — append 'YYYY-MM-DD HH:mm | <text>' to
    logs\\created.log (BOM only when the file is first created)."""
    logs = REPO_ROOT / "logs"
    logs.mkdir(exist_ok=True)
    line = datetime.now().strftime("%Y-%m-%d %H:%M") + " | " + text + "\n"
    p = logs / "created.log"
    _log_trim_if_over(p)
    if p.exists():
        _append(p, line)
    else:
        _write_new(p, line)


LOG_CAP_BYTES = 1048576     # lib\_Common.ahk LogCapBytes


def _log_trim_if_over(p: Path, cap: int = 0) -> bool:
    """Mirror of lib\\_Common.ahk LogTrimIfOver: past the cap, drop the oldest
    half of the log, cut at a line boundary. The kept half goes to a temp
    file beside the log and replaces it (FileReplaceText's rule — a kill
    mid-write must never empty the log); a refused replace (another writer
    holding the log) leaves it untouched. Never raises."""
    try:
        if not p.exists() or p.stat().st_size <= (cap or LOG_CAP_BYTES):
            return False
        text = _read_text_any(p)
        # AHK's InStr(txt, "`n", , StrLen(txt) // 2) — a 1-based start.
        cut = text.find("\n", max(len(text) // 2 - 1, 0))
        return _replace_text(p, text[cut + 1:] if cut >= 0 else "")
    except OSError:
        return False


def _replace_text(p: Path, content: str) -> bool:
    """Replace a file's content crash-safely (lib\\_Common.ahk FileReplaceText):
    write <name>.tmp-<pid> (UTF-8 BOM, LF, like _write_new) and os.replace it
    over the original. False, with the original untouched and the temp file
    removed, when the replace is refused. Never raises."""
    tmp = p.with_name(f"{p.name}.tmp-{os.getpid()}")
    try:
        _write_new(tmp, content)
        os.replace(tmp, p)
        return True
    except OSError:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        return False


def _read_template(name: str) -> str:
    """Read a template PRESERVING its on-disk line endings. The AHK generators
    fill templates with FileRead + StrReplace + FileAppend (no EOL translation),
    so they emit whatever endings the template has on disk (CRLF on a normal
    .gitattributes checkout). To stay byte-identical the writer must do the same
    — do NOT normalize here. (An earlier version normalized to LF, so MCP-created
    template macros came out LF while GUI-created ones were CRLF.) Content built
    from Python string literals stays LF, matching AHK's backtick-n literals."""
    p = REPO_ROOT / "templates" / name
    with open(p, encoding="utf-8-sig", newline="") as f:
        return f.read()


# ---------------------------------------------------------------------------
# Windows shortcut + AHK process helpers (PowerShell/subprocess, no pywin32)
# ---------------------------------------------------------------------------
def _ps_quote(s: str) -> str:
    return "'" + str(s).replace("'", "''") + "'"


def _run_powershell(command: str, check: bool = True) -> subprocess.CompletedProcess:
    return subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-Command", command],
        capture_output=True, text=True, check=check,
    )


def make_shortcut(link: str, ahk_file: str, workdir: str | None = None,
                  args: str | None = None) -> None:
    """lib\\_Common.ahk MakeAhkShortcut: a .lnk whose Target is the interpreter
    and whose Arguments is the quoted script path (association-proof). Extra
    `args` are appended after the script path (already quoted by the caller)."""
    if workdir is None:
        workdir = os.path.dirname(ahk_file)
    arguments = chr(34) + ahk_file + chr(34)
    if args:
        arguments += " " + args
    Path(link).parent.mkdir(parents=True, exist_ok=True)
    cmd = (
        "$ws=New-Object -ComObject WScript.Shell;"
        f"$s=$ws.CreateShortcut({_ps_quote(link)});"
        f"$s.TargetPath={_ps_quote(AHK_EXE)};"
        f"$s.Arguments={_ps_quote(arguments)};"
        f"$s.WorkingDirectory={_ps_quote(workdir)};"
        "$s.Save()"
    )
    _run_powershell(cmd)


def make_loop_shortcut(base: str, disp: str) -> str:
    """lib\\_Common.ahk MakeLoopShortcut: a 'loop <disp>' entry that runs the
    workflow <base> repeatedly via lib\\LoopRunner.ahk. Voice: 'open loop <disp>'."""
    _ensure_voice_dir()
    loop_runner = REPO_ROOT / "lib" / "LoopRunner.ahk"
    link = VOICE_MACROS / f"loop {disp}.lnk"
    make_shortcut(str(link), str(loop_runner), str(REPO_ROOT), chr(34) + base + chr(34))
    return str(link)


# How long one load check may take. /validate is NOT always dialog-free: a
# #Warn warning still pops its message box under /validate /ErrorStdOut
# (measured), and nothing answers it — so an unbounded check hung forever.
VALIDATE_TIMEOUT_S = 20
# validate_ahk's error text starts with this when the check had to be killed.
# master_preflight keys on it (mirrors AhkValidate's timedOut in _Common.ahk).
VALIDATE_BLOCKED = "The load check of "


def _kill_tree(pid: int) -> None:
    """Kill a process and every child it started (taskkill /T), quietly."""
    try:
        subprocess.run(["taskkill", "/T", "/F", "/PID", str(pid)],
                       capture_output=True, timeout=15)
    except (OSError, subprocess.SubprocessError):
        pass


def validate_ahk(path: str, timeout_s: float | None = None) -> tuple[bool, str]:
    """Load-check a script with AutoHotkey /validate. Returns (ok, error_text).

    Never raises for a slow check: past the timeout the process tree is
    killed and (False, <text starting with VALIDATE_BLOCKED>) comes back, so
    every write-then-validate caller still reaches its rollback. The error
    text is read as UTF-8 (as run_ahk_snippet reads it) — the ANSI default
    turned the quoted source line's non-Latin characters into '?'. Mirrors
    lib\\_Common.ahk AhkValidate."""
    timeout_s = VALIDATE_TIMEOUT_S if timeout_s is None else timeout_s
    try:
        proc = subprocess.Popen(
            [AHK_EXE, "/ErrorStdOut=UTF-8", "/validate", path],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            encoding="utf-8", errors="replace")
    except OSError as e:
        return False, f"Couldn't run the load check for {path}: {e}"
    try:
        out, err = proc.communicate(timeout=timeout_s)
    except subprocess.TimeoutExpired:
        _kill_tree(proc.pid)
        try:
            proc.kill()
        except OSError:
            pass
        try:
            proc.communicate(timeout=5)
        except (subprocess.SubprocessError, OSError):
            pass
        return False, (f"{VALIDATE_BLOCKED}{path} didn't finish within {timeout_s:g} s, so it "
                       f"was stopped. Something in it is waiting on a dialog nobody can see — "
                       f"almost always a #Warn warning, which pops a message box even during a "
                       f"load check. Remove #Warn (or fix what it warns about) and try again.")
    return proc.returncode == 0, (out or err or "").strip()


# ---------------------------------------------------------------------------
# Write, load-check, and roll back to the EXACT previous bytes
#
# Every write-then-validate path goes through here. There used to be five
# hand-written copies with two rollback methods: some restored raw bytes,
# others re-wrote a decoded text snapshot (so a CRLF or BOM-less file came
# back re-encoded while the error said "nothing changed"), none ran when the
# load check itself raised, and create_workflow's "rollback" deleted the
# workflow it was replacing.
# ---------------------------------------------------------------------------
class _FileSnapshot:
    """The exact bytes of files about to be (re)written — None for a file
    that didn't exist, which restore() removes again."""

    def __init__(self, *paths: Path):
        self.saved = [(p, p.read_bytes() if p.exists() else None) for p in paths]

    def prev(self, path: Path) -> bytes | None:
        return next((raw for p, raw in self.saved if p == path), None)

    def restore(self) -> None:
        for p, raw in self.saved:
            try:
                if raw is None:
                    p.unlink(missing_ok=True)
                else:
                    p.write_bytes(raw)
            except OSError:
                pass            # best effort: the original error is what matters


def _write_validated(writes: list, check: list, fail_msg: str,
                     also: tuple = ()) -> _FileSnapshot:
    """Write each (path, text) in `writes` (BOM + LF, _write_new), then
    load-check every entry of `check` — a path, or (path, its own failure
    message). On a failed check — or ANY exception, a timeout or missing
    AutoHotkey included — every written file goes back to its exact previous
    bytes (or away, if it is new) and the error is raised:
    VoiceKitError(message + the load error) for a failed check, the original
    exception otherwise. `also` names extra files to snapshot (ones the
    caller may delete next). Returns the snapshot for callers that bank the
    previous version or need to undo later."""
    snap = _FileSnapshot(*[p for p, _ in writes], *also)
    try:
        for p, text in writes:
            _write_new(p, text)
        for c in check:
            p, msg = c if isinstance(c, tuple) else (c, fail_msg)
            ok, err = validate_ahk(str(p))
            if not ok:
                # Explained HERE, while the failing file is still on disk
                # (the restore below may remove it).
                raise VoiceKitError(f"{msg}\n{_explain_ahk_error(err, script=p)}")
    except BaseException:
        snap.restore()
        raise
    return snap


def _replace_validated(path: Path, text: str, fail_msg: str) -> bytes | None:
    """One file: write, load-check, byte-exact restore on failure. Returns
    the previous bytes (None when the file is new)."""
    return _write_validated([(path, text)], [path], fail_msg).prev(path)



# ---------------------------------------------------------------------------
# Explaining load errors (WP1 — the extraction feedback)
#
# AutoHotkey v2 puts functions, classes and variables in ONE case-insensitive
# namespace, so top-level code that assigns to `LOG` fails with "This Func
# cannot be used as an output variable. Specifically: LOG" — and nothing in
# that text says Log is an AutoHotkey BUILT-IN (it is: the clash in that report was not
# VoiceKit's doing, and a VK_ prefix on our names would not have stopped it),
# or that Notify comes from lib\_Common.ahk. _explain_ahk_error appends one
# plain-English hint naming where the name comes from. It is applied at the
# user-facing raise sites (_write_validated, create_snippet, _clash_error,
# run_ahk_snippet) and deliberately NOT inside validate_ahk: master_preflight
# truncates that text and _error_module parses it. The hint goes at the END,
# so the first "file.ahk (N)" in the text is still AutoHotkey's own.
# ---------------------------------------------------------------------------

# AutoHotkey v2 built-in classes (a clash with one reads "This Class ...").
_AHK_BUILTIN_CLASSES = frozenset(n.lower() for n in """
Any Array BoundFunc Buffer ClipboardAll Class Closure ComObjArray ComObject ComValue
ComValueRef Enumerator Error File Float Func Gui IndexError InputHook Integer Map
MemberError Menu MenuBar MethodError Number Object OSError Primitive PropertyError
RegExMatchInfo String TargetError TimeoutError TypeError UnsetError UnsetItemError
ValueError VarRef ZeroDivisionError
""".split())

# AutoHotkey v2 built-in functions and classes — every name here resolves to a
# Func or Class in a bare v2 process (test_load_error_explains_name_clash
# checks each one against the installed interpreter, so the list can't rot).
# Used to say "an AutoHotkey built-in" without spawning a probe; a name not
# listed here still gets identified by a one-line /validate probe.
AHK_BUILTINS = _AHK_BUILTIN_CLASSES | frozenset(n.lower() for n in """
Abs ACos ASin ATan BlockInput CallbackCreate CallbackFree CaretGetPos Ceil Chr Click
ClipWait ComCall ComObjActive ComObjConnect ComObjFlags ComObjFromPtr ComObjGet
ComObjQuery ComObjType ComObjValue ControlAddItem ControlChooseIndex ControlChooseString
ControlClick ControlDeleteItem ControlFindItem ControlFocus ControlGetChecked
ControlGetChoice ControlGetClassNN ControlGetEnabled ControlGetExStyle ControlGetFocus
ControlGetHwnd ControlGetIndex ControlGetItems ControlGetPos ControlGetStyle
ControlGetText ControlGetVisible ControlHide ControlHideDropDown ControlMove ControlSend
ControlSendText ControlSetChecked ControlSetEnabled ControlSetExStyle ControlSetStyle
ControlSetText ControlShow ControlShowDropDown CoordMode Cos Critical DateAdd DateDiff
DetectHiddenText DetectHiddenWindows DirCopy DirCreate DirDelete DirExist DirMove
DirSelect DllCall Download DriveEject DriveGetCapacity DriveGetFileSystem DriveGetLabel
DriveGetList DriveGetSerial DriveGetSpaceFree DriveGetStatus DriveGetStatusCD
DriveGetType DriveLock DriveRetract DriveSetLabel DriveUnlock Edit EditGetCurrentCol
EditGetCurrentLine EditGetLine EditGetLineCount EditGetSelectedText EditPaste EnvGet
EnvSet Exit ExitApp Exp FileAppend FileCopy FileCreateShortcut FileDelete FileEncoding
FileExist FileGetAttrib FileGetShortcut FileGetSize FileGetTime FileGetVersion
FileInstall FileMove FileOpen FileRead FileRecycle FileRecycleEmpty FileSelect
FileSetAttrib FileSetTime Floor Format FormatTime GetKeyName GetKeySC GetKeyState
GetKeyVK GetMethod GroupActivate GroupAdd GroupClose GroupDeactivate GuiCtrlFromHwnd
GuiFromHwnd HasBase HasMethod HasProp HotIf HotIfWinActive HotIfWinExist
HotIfWinNotActive HotIfWinNotExist Hotkey Hotstring IL_Add IL_Create IL_Destroy
ImageSearch IniDelete IniRead IniWrite InputBox InstallKeybdHook InstallMouseHook InStr
IsAlnum IsAlpha IsDigit IsFloat IsInteger IsLabel IsLower IsNumber IsObject
IsSetRef IsSpace IsTime IsUpper IsXDigit KeyHistory KeyWait ListHotkeys ListLines
ListVars ListViewGetContent Ln LoadPicture Log LTrim Max MenuFromHandle MenuSelect Min
Mod MonitorGet MonitorGetCount MonitorGetName MonitorGetPrimary MonitorGetWorkArea
MouseClick MouseClickDrag MouseGetPos MouseMove MsgBox NumGet NumPut ObjAddRef
ObjBindMethod ObjFromPtr ObjFromPtrAddRef ObjGetBase ObjGetCapacity ObjHasOwnProp
ObjOwnPropCount ObjOwnProps ObjPtr ObjPtrAddRef ObjRelease ObjSetBase ObjSetCapacity
OnClipboardChange OnError OnExit OnMessage Ord OutputDebug Pause Persistent
PixelGetColor PixelSearch PostMessage ProcessClose ProcessExist ProcessGetName
ProcessGetParent ProcessGetPath ProcessSetPriority ProcessWait ProcessWaitClose Random
RegCreateKey RegDelete RegDeleteKey RegExMatch RegExReplace RegRead RegWrite Reload
Round RTrim Run RunAs RunWait Send SendEvent SendInput SendLevel SendMessage SendMode
SendPlay SendText SetCapsLockState SetControlDelay SetDefaultMouseSpeed SetKeyDelay
SetMouseDelay SetNumLockState SetRegView SetScrollLockState SetStoreCapsLockMode
SetTimer SetTitleMatchMode SetWinDelay SetWorkingDir Shutdown Sin Sleep Sort SoundBeep
SoundGetInterface SoundGetMute SoundGetName SoundGetVolume SoundPlay SoundSetMute
SoundSetVolume SplitPath Sqrt StatusBarGetText StatusBarWait StrCompare StrGet StrLen
StrLower StrPtr StrPut StrReplace StrSplit StrTitle StrUpper SubStr Suspend SysGet
SysGetIPAddresses Tan Thread ToolTip TraySetIcon TrayTip Trim Type VarSetStrCapacity
VerCompare WinActivate WinActivateBottom WinActive WinClose WinExist WinGetClass
WinGetClientPos WinGetControls WinGetControlsHwnd WinGetCount WinGetExStyle WinGetID
WinGetIDLast WinGetList WinGetMinMax WinGetPID WinGetPos WinGetProcessName
WinGetProcessPath WinGetStyle WinGetText WinGetTitle WinGetTransColor
WinGetTransparent WinHide WinKill WinMaximize WinMinimize WinMinimizeAll
WinMinimizeAllUndo WinMove WinMoveBottom WinMoveTop WinRedraw WinRestore
WinSetAlwaysOnTop WinSetEnabled WinSetExStyle WinSetRegion WinSetStyle WinSetTitle
WinSetTransColor WinSetTransparent WinShow WinWait WinWaitActive WinWaitClose
WinWaitNotActive
""".split())

# The built-ins most likely to be picked as a variable name — the ones the
# server instructions list by name (the explainer covers every other one).
COMMON_BUILTIN_CLASHES = (
    "Log", "Round", "Max", "Min", "Abs", "Mod", "Exp", "Ln", "Floor", "Ceil",
    "Sort", "Format", "Type", "Random", "Chr", "Ord", "Trim", "Click", "Send",
    "Run", "Sleep", "Edit", "Exit", "Download", "File", "Map", "Array", "Object",
    "Buffer", "Error", "Number", "String", "Integer", "Gui", "Menu",
)

# Column-0 function definitions (`Name(params) {` / `=> expr`) and classes.
# Definitions only — a top-level CALL such as `OnError(LogUncaughtError)`
# has no brace/arrow after its parameter list, so it never counts.
_AHK_DEF_RE = re.compile(r"^([A-Za-z_]\w*)\(.*\)\s*(?:\{|=>)")
_AHK_CLASS_RE = re.compile(r"^class\s+([A-Za-z_]\w*)", re.IGNORECASE)
_AHK_INCLUDE_RE = re.compile(
    r'^[ \t]*#Include(?:Again)?[ \t]+(?:\*i[ \t]+)?"?([^";\r\n]+?)"?[ \t]*(?:[ \t];.*)?$',
    re.IGNORECASE)
_AHK_KEYWORDS = frozenset("if while for loop switch catch return until else try finally "
                          "throw static global local not and or is in contains".split())
_AHK_DEF_CACHE: dict = {}


def _ahk_file_defs(path: Path) -> tuple[dict, list[str]]:
    """({name_lower: (Name, kind, line)}, [raw #Include specs]) for one .ahk
    file. Cached by (path, mtime); an unreadable file gives nothing."""
    try:
        key = (str(path).lower(), path.stat().st_mtime_ns)
    except OSError:
        return {}, []
    if key in _AHK_DEF_CACHE:
        return _AHK_DEF_CACHE[key]
    try:
        lines = _read_text_any(path).splitlines()
    except OSError:
        return {}, []
    defs: dict = {}
    incs: list[str] = []
    for i, line in enumerate(lines, 1):
        m = _AHK_DEF_RE.match(line)
        if m and m.group(1).lower() not in _AHK_KEYWORDS:
            defs.setdefault(m.group(1).lower(), (m.group(1), "function", i))
            continue
        m = _AHK_CLASS_RE.match(line)
        if m:
            defs.setdefault(m.group(1).lower(), (m.group(1), "class", i))
            continue
        m = _AHK_INCLUDE_RE.match(line)
        if m and not m.group(1).startswith("<"):
            incs.append(m.group(1).strip())
    _AHK_DEF_CACHE[key] = (defs, incs)
    return defs, incs


def _resolve_include(spec: str, including: Path, script: Path) -> Path | None:
    """An #Include spec as a file path: %A_LineFile% is the including file,
    %A_ScriptDir% the ROOT script's folder, and a relative spec is taken from
    the including file's folder. A directory include (it only changes the
    include folder) and unknown variables yield None."""
    s = re.sub(r"(?i)%A_LineFile%", lambda _m: str(including), spec)
    s = re.sub(r"(?i)%A_ScriptDir%", lambda _m: str(script.parent), s)
    if "%" in s:
        return None
    p = Path(s)
    if not p.is_absolute():
        p = including.parent / p
    p = Path(os.path.normpath(str(p)))
    return p if p.suffix.lower() == ".ahk" else None


def ahk_defs_in_chain(script: Path) -> dict:
    """Every function/class definition a script sees through its #Include
    chain — itself first, then each include, nested (Browser.ahk pulls in
    UIA.ahk and Clip.ahk): {name_lower: (Name, kind, file, line)}. The first
    definition is listed first; every one is kept, because a "declaration
    conflicts" error is raised AT one of them and names the other."""
    script = Path(script)
    out: dict = {}
    seen: set = set()
    todo = [script]
    while todo and len(seen) < 64:
        f = todo.pop(0)
        k = str(f).lower()
        if k in seen:
            continue
        seen.add(k)
        defs, incs = _ahk_file_defs(f)
        for lk, (name, kind, line) in defs.items():
            out.setdefault(lk, []).append((name, kind, f, line))
        for spec in incs:
            target = _resolve_include(spec, f, script)
            if target is not None:
                todo.append(target)
    return out


def _voicekit_defs() -> dict:
    """Definitions in every VoiceKit library (lib\\*.ahk) and the master —
    the fallback for a clash inside the resident VoiceKit, where the failing
    module doesn't itself include the file that owns the name."""
    out: dict = {}
    for f in sorted((REPO_ROOT / "lib").glob("*.ahk")) + [REPO_ROOT / "VoiceKit.ahk"]:
        for lk, (name, kind, line) in _ahk_file_defs(f)[0].items():
            out.setdefault(lk, (name, kind, f, line))
    return out


def _vk_rel(p: Path) -> str:
    """A path as the user knows it: install-relative when under the install."""
    try:
        return str(Path(p).resolve().relative_to(REPO_ROOT.resolve()))
    except (ValueError, OSError):
        return str(p)


_BUILTIN_PROBE_CACHE: dict = {}
_BUILTIN_PROBE_CLASSES: set = set()   # probe-found built-ins that are classes


def _is_ahk_builtin(name: str) -> bool:
    """Whether `name` is an AutoHotkey built-in function/class: the curated
    list first, then (only for a name it doesn't know) a /validate of
    `<name> := 0` in an otherwise empty script — its error says Func/Class
    exactly when the interpreter already owns the name. Error path only."""
    lk = name.lower()
    if lk in AHK_BUILTINS:
        return True
    if not re.fullmatch(r"[A-Za-z_]\w*", name):
        return False
    if lk not in _BUILTIN_PROBE_CACHE:
        hit = False
        try:
            with tempfile.TemporaryDirectory() as d:
                p = Path(d) / "probe.ahk"
                p.write_text(f"#Requires AutoHotkey v2.0\n{name} := 0\n", encoding="utf-8-sig")
                ok, err = validate_ahk(str(p), 10)
                m = None if ok else re.search(
                    r"This (Func|Class) cannot be used as an output variable", err)
                hit = bool(m)
                if m and m.group(1) == "Class":
                    _BUILTIN_PROBE_CLASSES.add(lk)
        except OSError:
            hit = False
        _BUILTIN_PROBE_CACHE[lk] = hit
    return _BUILTIN_PROBE_CACHE[lk]


def _name_origin(name: str, script: Path | None, prelude_lines: int = 0,
                 snippet: bool = False, skip_line: int = 0) -> str:
    """Where `name` comes from, as a phrase — "an AutoHotkey built-in
    function", "a VoiceKit function defined in lib\\_Common.ahk line 71,
    which this code includes", "a function defined in this same file at line
    3" — or "" when it can't be placed. `skip_line` (a line of `script`)
    is the definition the error itself points at, never the answer."""
    lk = name.lower()
    chain = ahk_defs_in_chain(script) if script is not None and Path(script).exists() else {}
    own = str(script).lower()
    hits = [h for h in chain.get(lk, []) if not (str(h[2]).lower() == own and h[3] == skip_line)]
    builtin = _is_ahk_builtin(name)
    if hits:
        _n, kind, f, line = hits[0]
        if str(f).lower() == own:
            if snippet and line <= prelude_lines:
                if lk == "out":
                    return ("the Out() helper run_ahk_snippet predefines (Out(value) "
                            "reports a line back to you)")
                return "a helper function run_ahk_snippet predefines"
            if snippet:
                return f"a {kind} your own code defines at snippet line {line - prelude_lines}"
            return f"a {kind} defined in this same file at line {line}"
        also = (f" (it replaces AutoHotkey's built-in {name}(), so the name is taken "
                f"either way)" if builtin else "")
        return (f"a VoiceKit {kind} defined in {_vk_rel(f)} line {line}{also}, which this "
                f"code includes")
    if builtin:
        return "an AutoHotkey built-in " + (
            "class" if lk in _AHK_BUILTIN_CLASSES | _BUILTIN_PROBE_CLASSES else "function")
    hit = _voicekit_defs().get(lk)
    if hit is not None:
        _n, kind, f, line = hit
        return (f"a VoiceKit {kind} defined in {_vk_rel(f)} line {line}, which is loaded "
                f"alongside this code")
    return ""


def _rename_example(name: str) -> str:
    return name[:1].lower() + name[1:] + "Text"


_QUOTE_HINT = ('Hint: AutoHotkey v2 escapes a double quote inside "..." as `" '
               '(backtick-quote), not "" as v1 did: "say `"hi`"". Or put the string '
               'in single quotes: \'say "hi"\'.')
_DOUBLED_QUOTE_RE = re.compile(r'\w""|""\w|"""')
_LOAD_ERR_LINE_RE = re.compile(r"(?m)^\s*(\S.*?\.ahk) \((snippet line )?(\d+)\) : ==>")
_RUNTIME_LINE_RE = re.compile(r"UNCAUGHT \w+: (.*?) — snippet line (\d+)")


def _failing_line(err: str, code: str | None) -> str:
    """The source line a load error points at: from `code` for a snippet (its
    errors are already renumbered to the snippet's own lines), else from the
    file the error names (still on disk at every raise site)."""
    m = _LOAD_ERR_LINE_RE.search(err)
    if not m:
        return ""
    n = int(m.group(3))
    if code is not None and m.group(2):
        lines = code.splitlines()
    else:
        try:
            lines = _read_text_any(Path(m.group(1))).splitlines()
        except OSError:
            return ""
    return lines[n - 1] if 0 < n <= len(lines) else ""


def _explain_ahk_error(text: str, code: str | None = None, script: Path | str | None = None,
                       prelude_lines: int = 0) -> str:
    """`text` plus a plain-English hint when it is one of the v2 errors whose
    raw wording never says what went wrong: a name clash (a variable or
    declaration colliding with a built-in, a VoiceKit library function, or
    run_ahk_snippet's Out) or a v1-style "" inside a string. Unchanged when
    nothing matches — a missed pattern costs the hint, never the error.

    `code` is the user's own code when the file is a run_ahk_snippet script
    (errors there are numbered in the snippet's lines, and runtime errors are
    explained too), `script` the file that failed (default: the one the error
    text names), `prelude_lines` the snippet prelude's length."""
    if not text or "\nHint: " in text:
        return text
    hints: list[str] = []
    snippet = code is not None
    m_file = _LOAD_ERR_LINE_RE.search(text)
    if script is None and m_file:
        script = m_file.group(1)
    script = Path(script) if script is not None else None

    # (1) Name clashes at load time.
    m = re.search(r"==> This (Func|Class) cannot be used as an output variable\.\s*"
                  r"Specifically:\s*([A-Za-z_]\w*)", text)
    if m:
        name = m.group(2)
        own = (_name_origin(name, script, prelude_lines, snippet)
               or f"already a {'function' if m.group(1) == 'Func' else 'class'}")
        hints.append(
            f"Hint: '{name}' is {own}. AutoHotkey v2 keeps functions, classes and "
            f"variables in one case-insensitive namespace, so code at the top level of a "
            f"script can't assign to that name (nor can a function that declares it "
            f"global). Rename your variable (e.g. {_rename_example(name)}), or move the "
            f"code inside a function, where plain variables are local.")
    m = re.search(r"==> This (function|class) declaration conflicts with an existing "
                  r"(Func|Class)\.\s*Specifically:\s*([A-Za-z_]\w*)", text)
    if m:
        name = m.group(3)
        # The error sits AT the second definition; the answer is the other one.
        at = (int(m_file.group(3)) + (prelude_lines if m_file.group(2) else 0)) if m_file else 0
        own = _name_origin(name, script, prelude_lines, snippet, skip_line=at) or "already defined"
        hints.append(
            f"Hint: '{name}' is {own}, so this {m.group(1)} can't be declared under the "
            f"same name (AutoHotkey v2 names are case-insensitive). Give yours a "
            f"different name, e.g. My{name}.")

    # (2) Clashes that only show at run time (run_ahk_snippet, exit 3).
    rt = _RUNTIME_LINE_RE.search(text) if snippet else None
    if rt:
        n = int(rt.group(2))
        lines = code.splitlines()
        line = lines[n - 1] if 0 < n <= len(lines) else ""
        msg = rt.group(1)
        # Names inside string literals are text, not references; a name after
        # a dot is a property/method, not the global of that name.
        bare = re.sub(r'"(?:`.|[^"`])*"|\'(?:`.|[^\'`])*\'', '""', line)
        if re.search(r"but got a (Func|Class)\b", msg):
            for w in dict.fromkeys(re.findall(r"(?<![.\w])([A-Za-z_]\w*)\b(?!\s*\()", bare)):
                if w.lower() in _AHK_KEYWORDS:
                    continue
                origin = _name_origin(w, script, prelude_lines, snippet)
                if origin:
                    hints.append(
                        f"Hint: on that line '{w}' is {origin}, not a variable — a variable "
                        f"of that name was never assigned, so AutoHotkey used the function "
                        f"itself. Rename the variable (e.g. {_rename_example(w)}).")
                    break
        elif re.search(r'has no method named "Call"', msg):
            for w in dict.fromkeys(re.findall(r"(?<![.\w])([A-Za-z_]\w*)\s*\(", bare)):
                if w.lower() in _AHK_KEYWORDS:
                    continue
                # Only a name that IS a function can be "hidden" — calling a
                # variable that holds a string fails with the same message.
                if re.search(rf"(?i)\b{re.escape(w)}\b(?!\s*\()", code) and \
                        _name_origin(w, script, prelude_lines, snippet):
                    hints.append(
                        f"Hint: {w}() is called on that line, but a variable or parameter "
                        f"named {w.lower()} in the same scope hides the function (AutoHotkey "
                        f"v2 names are case-insensitive). Rename the variable, e.g. "
                        f"{w.lower()}Val.")
                    break

    # (3) v1-style doubled quotes.
    if m_file:
        spec = re.search(r"Specifically:\s*(.*)", text)
        if _DOUBLED_QUOTE_RE.search(_failing_line(text, code)) or \
                (spec and '"""' in spec.group(1)):
            hints.append(_QUOTE_HINT)
    if not hints:
        return text
    return text.rstrip() + "\n" + "\n".join(hints)


def _ahk_processes() -> list:
    """Every process running OUR interpreter (basename of AHK_EXE — honors
    the VOICEKIT_AHK override, e.g. a UIA build): [{pid, cmd, started}].

    The one Win32_Process query in this module. There used to be three, with
    drifting rules (one hard-coded AutoHotkey64.exe and ignored the override,
    one matched an unanchored pattern); callers now filter these rows in
    Python. `started` is yyyyMMddHHmmss ("" when Windows didn't say)."""
    exe = os.path.basename(AHK_EXE).replace("'", "''")
    cmd = (
        "$list = Get-CimInstance Win32_Process -Filter \"Name = '" + exe + "'\" | "
        "ForEach-Object { @{pid = $_.ProcessId; cmd = $_.CommandLine; "
        "started = $(if ($_.CreationDate) { $_.CreationDate.ToString('yyyyMMddHHmmss') } "
        "else { '' })} }; "
        "ConvertTo-Json -InputObject @($list) -Compress"
    )
    try:
        r = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", cmd],
            capture_output=True, text=True, timeout=30)
        rows = json.loads(r.stdout or "[]")
    except (ValueError, OSError, subprocess.SubprocessError):
        return []
    out = []
    for row in rows if isinstance(rows, list) else [rows]:
        if isinstance(row, dict):     # @($null) serializes as [null] when nothing matched
            out.append({"pid": _as_int(row.get("pid")), "cmd": str(row.get("cmd") or ""),
                        "started": str(row.get("started") or "")})
    return out


def _script_matches(cmd: str, script_name: str, anchored: bool = True) -> bool:
    """Whether a command line runs `script_name` (repo-relative, e.g.
    "lib\\LoopRunner.ahk"). Anchored on REPO_ROOT so another VoiceKit tree's
    copy (the installed one under %LOCALAPPDATA%, a test's temp root) never
    counts as ours; anchored=False matches `\\<script_name>` in any folder.
    Either way the path separator anchors the name, so a macro whose base
    merely ENDS in it (RestartVoicekit.ahk vs \\VoiceKit.ahk) never matches.
    Caseless, like the PowerShell -like it replaces."""
    needle = ((str(REPO_ROOT) if anchored else "") + "\\" + script_name).lower()
    return needle in cmd.lower()


def _ahk_script_running(script_name: str, anchored: bool = True) -> bool:
    """True if our interpreter is running `script_name` (see _script_matches).
    anchored=False is for callers where a false "running" is the SAFE
    mistake (refusing to relaunch a Studio that may hold unsaved work)."""
    return any(_script_matches(p["cmd"], script_name, anchored) for p in _ahk_processes())


# --- Liveness of one pid, without spawning anything -------------------------
# health() rides on EVERY tool response, and asking PowerShell/CIM "is the
# master running" cost ~0.5 s a call (measured) — most of it PowerShell
# starting up. The master already publishes its pid and a 5-second heartbeat
# in logs\master-status.ini; checking that pid through kernel32 takes well
# under a millisecond.
_PROCESS_QUERY_LIMITED_INFORMATION = 0x1000
_SYNCHRONIZE = 0x00100000
_PROCESS_TERMINATE = 0x0001
_STILL_ACTIVE = 259
_ERROR_ACCESS_DENIED = 5


def _kernel32():
    import ctypes
    from ctypes import wintypes
    k = ctypes.WinDLL("kernel32", use_last_error=True)
    k.OpenProcess.restype = wintypes.HANDLE
    k.OpenProcess.argtypes = (wintypes.DWORD, wintypes.BOOL, wintypes.DWORD)
    k.CloseHandle.argtypes = (wintypes.HANDLE,)
    k.GetExitCodeProcess.argtypes = (wintypes.HANDLE, ctypes.POINTER(wintypes.DWORD))
    k.QueryFullProcessImageNameW.argtypes = (wintypes.HANDLE, wintypes.DWORD,
                                             wintypes.LPWSTR, ctypes.POINTER(wintypes.DWORD))
    k.WaitForMultipleObjects.restype = wintypes.DWORD
    k.WaitForMultipleObjects.argtypes = (wintypes.DWORD, ctypes.POINTER(wintypes.HANDLE),
                                         wintypes.BOOL, wintypes.DWORD)
    k.TerminateProcess.argtypes = (wintypes.HANDLE, wintypes.UINT)
    return k


def _pid_image(pid: int) -> str | None:
    """The full image path of a LIVE process, or None when the pid isn't
    running (or can't be opened at all)."""
    import ctypes
    from ctypes import wintypes
    if not pid or os.name != "nt":
        return None
    k = _kernel32()
    h = k.OpenProcess(_PROCESS_QUERY_LIMITED_INFORMATION, False, int(pid))
    if not h:
        return None
    try:
        code = wintypes.DWORD()
        if not k.GetExitCodeProcess(h, ctypes.byref(code)) or code.value != _STILL_ACTIVE:
            return None
        buf = ctypes.create_unicode_buffer(1024)
        size = wintypes.DWORD(len(buf))
        if not k.QueryFullProcessImageNameW(h, 0, buf, ctypes.byref(size)):
            return ""
        return buf.value
    finally:
        k.CloseHandle(h)


def _master_alive() -> bool | None:
    """Fast answer to "is THIS tree's master running", from the pid and
    heartbeat it writes to logs\\master-status.ini (the file sits under
    REPO_ROOT, so another tree's master can never answer for this one).

    True   the pid is alive, runs our interpreter, and the heartbeat is fresh.
    False  the pid is gone and the file settles it: the master exited
           cleanly, or its heartbeat went stale along with it.
    None   can't tell cheaply — no status file; the pid is alive but the
           heartbeat is stale (right after resume from sleep, or a dead
           master's pid reused by some other AutoHotkey process); or the pid
           is gone while its heartbeat is still FRESH and it didn't exit
           cleanly. That last one is a reload in flight: #SingleInstance
           Force has killed the old master and the new one hasn't written
           its pid yet — answering "not running" there would skip a reload
           and tell the caller VoiceKit is down. The caller falls back to
           the process scan for these rare cases."""
    st = read_master_status()
    pid = _as_int(st.get("pid"))
    if not pid:
        return None
    age = _age_seconds(st.get("heartbeat", ""))
    fresh = age is not None and age <= 20
    handover = fresh and st.get("clean_exit") != "1"
    image = _pid_image(pid)
    if image is None:
        return None if handover else False   # handover: let the scan decide
    if image and os.path.basename(image).lower() != os.path.basename(AHK_EXE).lower():
        return False                      # the pid now belongs to something else
    if st.get("clean_exit") == "1" and not fresh:
        return None                       # an exited master's pid, reused by AHK
    return True if fresh else None


def _master_restarting() -> bool:
    """A reload in flight, from the status file alone: the recorded pid is
    gone, but its heartbeat is still fresh and it didn't exit cleanly —
    #SingleInstance Force has replaced the old master and the new one hasn't
    written its pid yet (it does so within its first moments). _master_alive
    answers None for exactly this; this names it, so a caller can wait for
    the new master or say "restarting" instead of "not running"."""
    st = read_master_status()
    pid = _as_int(st.get("pid"))
    if not pid or st.get("clean_exit") == "1":
        return False
    age = _age_seconds(st.get("heartbeat", ""))
    return age is not None and age <= 20 and _pid_image(pid) is None


def _await_master_after_restart(timeout_s: float = 6.0) -> bool | None:
    """If a reload is in flight, wait (briefly) for the new master to report
    in. True once it has; False when the wait ran out mid-restart; None when
    no restart was in flight (the caller's usual check applies)."""
    if not _master_restarting():
        return None
    deadline = time.monotonic() + timeout_s
    while time.monotonic() < deadline:
        if _master_alive() is True:
            return True
        time.sleep(0.25)
    return _master_alive() is True


def voicekit_running() -> bool:
    """True if the resident VoiceKit master (THIS tree's) is running. The
    status file answers almost every time; the process scan only runs when
    it can't (see _master_alive)."""
    alive = _master_alive()
    if alive is not None:
        return alive
    return _ahk_script_running("VoiceKit.ahk")


# ---------------------------------------------------------------------------
# Master health: preflight, quarantine, status file
#
# VoiceKit.ahk pulls hotkeys\_index.ahk in at COMPILE time, so one module that
# stops parsing means the master never starts and every hotkey AND snippet dies
# with it. Mirrors lib\_Common.ahk MasterPreflight / IndexModules — keep the two
# in step, the same way the bridge allocator and the step formats are kept.
# ---------------------------------------------------------------------------
# Derived paths are functions, not constants: tests retarget REPO_ROOT at a
# temp directory, and an import-time constant would keep pointing at the real
# repo (silently testing — or worse, editing — the live install).
def _master_path() -> Path:
    return REPO_ROOT / "VoiceKit.ahk"


def _hotkeys_index() -> Path:
    _seed_live("hotkeys\\_index.ahk")
    return REPO_ROOT / "hotkeys" / "_index.ahk"


def _launcher_path() -> Path:
    return REPO_ROOT / "VoiceKitLauncher.ahk"


def _status_ini() -> Path:
    return REPO_ROOT / "logs" / "master-status.ini"


def _safe_mode_flag() -> Path:
    return REPO_ROOT / "logs" / "safe-mode.flag"


def _runs_ini() -> Path:
    return REPO_ROOT / "logs" / "workflow-runs.ini"


def _runs_log() -> Path:
    return REPO_ROOT / "logs" / "workflow-runs.log"


_INCLUDE_RE = re.compile(r'(?i)#Include\s+"%A_ScriptDir%\\(hotkeys\\[^"]+)"')

# Set when a reload was refused because the files wouldn't load. Reported by
# health() until a reload succeeds — a silently skipped reload is exactly the
# kind of quiet failure this whole change exists to remove.
_LAST_RELOAD_ERROR = ""
_LAST_RELOAD_ERROR_AT = ""      # AHK-style stamp; a master started later clears it
# What the LAST reload's preflight parked: {lowercased rel: the load error}.
# A module can pass its standalone load check and still break the master —
# a function or hotkey that another loaded file also defines only clashes
# when VoiceKit.ahk compiles everything together — so the preflight parks it
# and the reload goes ahead without it. Callers check this to find out
# whether the module they just wrote is one of those (and roll it back).
_LAST_PARKED: dict = {}
# The last VoiceKitLauncher.ahk this module started (a Popen), so the next
# reload can wait for it — see _wait_for_last_launcher.
_LAST_LAUNCH = None
# FastMCP runs sync tools on a thread pool, so two tool calls can overlap.
# Every tool that reloads the master and then reads what that reload parked
# (_LAST_PARKED) holds this for the whole write-reload-check sequence, so one
# call can never judge its write by another call's reload. Re-entrant: the
# tools nest (delete_automation -> _remove_companion_hotkey -> reload).
_RELOAD_LOCK = threading.RLock()


def _serialized(fn):
    """Run `fn` holding _RELOAD_LOCK (see above)."""
    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        with _RELOAD_LOCK:
            return fn(*args, **kwargs)
    return wrapper


def _read_text_any(path: Path) -> str:
    """Read a file whose encoding we don't control — the ONE text reader here.
    AutoHotkey's IniWrite creates logs\\master-status.ini as UTF-16LE with a
    BOM (measured), Notepad may save a macro as UTF-16 or ANSI, and an older
    or hand-edited file may carry no BOM at all. Line endings are left as
    they are on disk."""
    return _decode_any(path.read_bytes())


def _lf(text: str) -> str:
    """Universal newlines (CRLF / CR -> LF) — what Path.read_text gives."""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _decode_any(raw: bytes) -> str:
    """Decode bytes the way _read_text_any reads a file."""
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16")
    if raw[:3] == b"\xef\xbb\xbf":
        return raw.decode("utf-8-sig")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def index_modules() -> dict:
    """Modules hotkeys\\_index.ahk lists, split into the ones that load and the
    ones parked by a preflight. Mirrors lib\\_Common.ahk IndexModules."""
    active: list[str] = []
    parked: list[str] = []
    if _hotkeys_index().exists():
        for line in _read_text_any(_hotkeys_index()).splitlines():
            m = _INCLUDE_RE.search(line)
            if not m:
                continue
            (parked if line.lstrip().startswith(";") else active).append(m.group(1))
    return {"active": active, "quarantined": parked}


# ---------------------------------------------------------------------------
# Per-user files: shipped defaults vs live copies (2026-09-30)
#
# hotkeys\_index.ahk, bridge-map.txt and hotkeys\Snippets.ahk are the user's
# data (every create/delete rewrites them), so git tracks only the shipped
# <name>.default.<ext> and the live copies are made on demand. Mirrors
# lib\_Common.ahk SeedUserFiles / SeedMergeShipped exactly —
# test_seed_user_files_matches_ahk runs both on one sandbox. The path
# accessors (_hotkeys_index, _bridge_map_file, _snippets_file) seed a
# MISSING live file on first use, so an MCP call on a fresh clone works
# before VoiceKit has ever started; the merge of newly shipped lines runs
# with the preflight, like the AHK side.
# ---------------------------------------------------------------------------
USER_FILE_PAIRS = (
    ("hotkeys\\_index.ahk", "hotkeys\\_index.default.ahk"),
    ("bridge-map.txt", "bridge-map.default.txt"),
    ("hotkeys\\Snippets.ahk", "hotkeys\\Snippets.default.ahk"),
)


def _repo_rel(rel: str) -> Path:
    return REPO_ROOT.joinpath(*rel.split("\\"))


def _seed_live(rel: str) -> bool:
    """Create live file `rel` from its shipped default when it is missing
    (a byte copy, like AHK FileCopy). True when it was created — and then
    logged to created.log the way AHK SeedUserFilesLogged does, whichever
    path got here (seed_user_files, or a path getter seeding on first use)."""
    default = next((d for l, d in USER_FILE_PAIRS if l == rel), None)
    if default is None:
        return False
    live, src = _repo_rel(rel), _repo_rel(default)
    if live.exists() or not src.exists():
        return False
    try:
        live.parent.mkdir(parents=True, exist_ok=True)
        live.write_bytes(src.read_bytes())
    except OSError:
        return False
    try:
        log(f"seeded | {rel} | created from its shipped default")
    except OSError:
        pass
    return True


def _seed_map_record(line: str):
    """(combo, file) of a bridge-map record line — a leading ';' is stripped,
    so a commented-out record still parses — or None. Mirrors SeedMapRecord."""
    parts = line.lstrip("; \t").split("|")
    if len(parts) < 3 or not parts[0].strip(" \t") or not parts[2].strip(" \t"):
        return None
    return parts[0].strip(" \t"), parts[2].strip(" \t")


def _seed_append(path: Path, current: str, lines: list[str]) -> bool:
    """Append lines in the file's own line-ending style (SeedAppendLines)."""
    eol = "\r\n" if "\r\n" in current else "\n"
    txt = eol if current and not current.endswith("\n") else ""
    txt += "".join(l + eol for l in lines)
    try:
        if path.exists():
            _append(path, txt)
        else:
            _write_new(path, txt)
        return True
    except OSError:
        return False


def seed_user_files() -> dict:
    """Make the live per-user files: create a missing one from its default,
    then append any SHIPPED _index include / bridge-map record the live file
    lacks — keyed by module file, commented-out copies counting as present,
    skipped when the module file is gone or its key is taken (and then its
    include is held back too). Snippets.ahk is seeded only, never merged.
    Returns {"created": [rel...], "merged": [line...], "index_changed": bool}.
    Mirrors lib\\_Common.ahk SeedUserFiles."""
    out = {"created": [], "merged": [], "index_changed": False}
    for rel, _default in USER_FILE_PAIRS:
        if _seed_live(rel):
            out["created"].append(rel)
            if rel == "hotkeys\\_index.ahk":
                out["index_changed"] = True
    try:
        _seed_merge_shipped(out)
    except OSError:
        pass
    # (A created file was logged by _seed_live itself.)
    for line in out["merged"]:
        log(f"seeded | shipped line added | {line}")
    return out


def _seed_merge_shipped(out: dict) -> None:
    idx, idx_def = _repo_rel("hotkeys\\_index.ahk"), _repo_rel("hotkeys\\_index.default.ahk")
    bmap, bmap_def = _repo_rel("bridge-map.txt"), _repo_rel("bridge-map.default.txt")

    def read_or(f: Path) -> str:
        try:
            return _read_text_any(f)
        except OSError:
            return ""

    def lines(text: str) -> list[str]:     # Loop Parse text, "`n", "`r"
        return [l.replace("\r", "") for l in text.split("\n")]

    live_idx, live_map = read_or(idx), read_or(bmap)
    ship_inc: list[tuple[str, str]] = []
    ship_rec: dict[str, tuple[str, str, str]] = {}   # lower(file) -> (combo, file, line)
    if idx_def.exists():
        for l in lines(read_or(idx_def)):
            m = _INCLUDE_RE.search(l)
            if m and not l.lstrip(" \t").startswith(";"):
                ship_inc.append((m.group(1), l.strip(" \t")))
    if bmap_def.exists():
        for l in lines(read_or(bmap_def)):
            r = _seed_map_record(l)
            if r and not l.lstrip(" \t").startswith(";"):
                ship_rec[r[1].lower()] = (r[0], r[1], l.strip(" \t"))
    have_inc = {m.group(1).lower() for l in lines(live_idx) if (m := _INCLUDE_RE.search(l))}
    have_rec = {r[1].lower() for l in lines(live_map) if (r := _seed_map_record(l))}
    used = _used_bridge_keys()           # once — it reads and scans the whole map
    free = [k for k in BRIDGE_POOL if k not in used]

    def key_of(combo: str) -> str | None:
        if combo[:len(_COMBO_PREFIX)].lower() != _COMBO_PREFIX.lower():
            return None
        return combo[len(_COMBO_PREFIX):].upper()

    add_idx, add_map, blocked = [], [], set()
    for low, (combo, file, line) in sorted(ship_rec.items()):   # AHK Maps enumerate sorted
        if low in have_rec or not _repo_rel(file).exists():
            continue
        k = key_of(combo)
        if k is None or k not in free:
            blocked.add(low)
            continue
        add_map.append(line)
        free.remove(k)
    for rel, line in ship_inc:
        if rel.lower() in have_inc or rel.lower() in blocked or not _repo_rel(rel).exists():
            continue
        add_idx.append(line)
    if add_idx and _seed_append(idx, live_idx, add_idx):
        out["index_changed"] = True
    if add_map:
        _seed_append(bmap, live_map, add_map)
    out["merged"] += add_idx + add_map


def _error_module(err_text: str) -> str:
    """The repo-relative module an AHK load error points at, or "" when the
    failure isn't in a module we're allowed to park (a broken lib\\ or
    VoiceKit.ahk is ours to fix, not to disable). Mirrors MasterErrorModule."""
    if not err_text:
        return ""
    m = re.search(r"(?m)^\s*(\S.*?\.ahk) \(\d+\)", err_text)
    if not m:
        return ""
    # A module whose FILE is gone: AutoHotkey blames the manifest doing the
    # including ('...\hotkeys\_index.ahk (7) : ==> #Include file "...\Gone.ahk"
    # cannot be opened.'), which can't park itself — so name the missing
    # module instead. Only when the manifest is the includer; a missing file
    # included from inside a module already blames that module.
    if re.search(r"(?i)\\hotkeys\\_index\.ahk$", m.group(1)):
        miss = re.search(r'(?i)#Include file "[^"]*\\(hotkeys\\[^"\\]+\.ahk)" cannot be opened',
                         err_text)
        return miss.group(1) if miss else ""
    rel = re.search(r"(?i)\\(hotkeys\\[^\\]+\.ahk)$", m.group(1))
    return rel.group(1) if rel else ""


def _comment_out_include(rel: str, note: str) -> bool:
    """Comment out every uncommented #Include of `rel`, tagging why.
    hotkeys\\_index.ahk's own header documents ";" as the disable mechanism."""
    if not _hotkeys_index().exists():
        return False
    text = _read_text_any(_hotkeys_index())
    out, hit = [], False
    for line in text.splitlines():
        if "\\" + rel in line and not line.lstrip().startswith(";"):
            hit = True
            line = f"; {line}    {note}"
        out.append(line)
    if not hit:
        return False
    # BOM + LF like the AHK mirror (CommentOutLinesContaining). A bare
    # write_text would translate every \n to \r\n on Windows.
    _write_new(_hotkeys_index(), "\n".join(out).rstrip("\n") + "\n")
    return True


def _warn_modules() -> list[str]:
    """Active modules carrying a #Warn directive — the suspects when the
    master's load check times out. Mirrors _Common.ahk MasterWarnModules."""
    out = []
    for rel in index_modules()["active"]:
        try:
            txt = _read_text_any(REPO_ROOT / rel)
        except OSError:
            continue
        if re.search(r"(?im)^[ \t]*#Warn\b", txt):
            out.append(rel)
    return out


def master_preflight(reasons: dict | None = None) -> tuple[bool, list[str], str]:
    """Make sure the master will load, parking broken modules until it does.
    Returns (ok, parked, error_text). ok=False means don't reload. Pass a dict
    as `reasons` to receive {parked rel: the load error that parked it}."""
    parked: list[str] = []
    reasons = {} if reasons is None else reasons
    stamp = _today()
    err = ""
    # The live per-user files first (mirrors MasterPreflight): a fresh clone
    # has none, and an upgrade may ship a hotkey the live manifest lacks.
    seed_user_files()
    for _ in range(6):                    # one broken module can hide the next
        ok, err = validate_ahk(str(_master_path()))
        if not ok and err.startswith(VALIDATE_BLOCKED):
            # Usually a #Warn dialog, but a slow login / AV scan / disk
            # contention can time out too — one retry at twice the bound
            # before parking anything (mirrors MasterPreflight).
            ok, err = validate_ahk(str(_master_path()), VALIDATE_TIMEOUT_S * 2)
        if ok:
            return True, parked, ""
        if err.startswith(VALIDATE_BLOCKED):
            # No error text to read — a dialog blocked the check. Park every
            # active module carrying #Warn (mirrors MasterPreflight).
            hit = False
            for rel in _warn_modules():
                if _comment_out_include(
                        rel, f"quarantined {stamp} — its #Warn warning blocked the load check"):
                    hit = True
                    parked.append(rel)
                    reasons[rel] = err
                    log(f"quarantined | {rel} | #Warn blocked the load check")
            if not hit:
                return False, parked, err
            continue
        rel = _error_module(err)
        why = ("this file is missing" if rel and not (REPO_ROOT / rel).exists()
               else "this file failed to load")
        if not rel or not _comment_out_include(rel, f"quarantined {stamp} — {why}"):
            return False, parked, err
        parked.append(rel)
        reasons[rel] = err
        log(f"quarantined | {rel} | {' '.join(err.split())}")
    ok, err = validate_ahk(str(_master_path()))
    return ok, parked, ("" if ok else err)


def read_master_status() -> dict:
    """logs\\master-status.ini as a flat dict ("" when there is none)."""
    if not _status_ini().exists():
        return {}
    import configparser
    cp = configparser.RawConfigParser()   # Raw: an error message may hold a '%'
    try:
        cp.read_string(_read_text_any(_status_ini()))
    except Exception:
        return {}
    out: dict[str, str] = {}
    for section in cp.sections():
        for k, v in cp.items(section):
            out[k if section == "Master" else f"{section.lower()}_{k}"] = v
    return out


def _age_seconds(stamp: str) -> int | None:
    """Seconds since an AutoHotkey YYYYMMDDHH24MISS timestamp."""
    try:
        return int((datetime.now() - datetime.strptime(stamp, "%Y%m%d%H%M%S")).total_seconds())
    except (ValueError, TypeError):
        return None


# ---------------------------------------------------------------------------
# Server identity: version, install root, stale code (WP10)
#
# "Two sets of VoiceKit tool names" and "a tool the docs mention is missing"
# both come down to WHICH server a client is talking to: the same install
# registered twice (Claude Code's ~/.claude.json and the Desktop config on
# this machine), a dev checkout and an install both registered, or a server
# process that started before its code changed on disk. The installed copy
# has no .git, so the version comes from a VERSION file the build stamps.
# ---------------------------------------------------------------------------
SERVER_STARTED_AT = datetime.now()
_CODE_FILES = (Path(__file__).resolve(), Path(__file__).resolve().with_name("server.py"))


def _mtime_ns(p: Path) -> int:
    try:
        return p.stat().st_mtime_ns
    except OSError:
        return 0


# Stamped at import — the server imports this module when it starts.
_CODE_STAMPS = {p: _mtime_ns(p) for p in _CODE_FILES}


def read_version(root: Path | None = None) -> str:
    """The install's version: the first line of <root>\\VERSION ("1.2.0-dev"
    in the repo; the build stamps "+<git hash>.<date>" into the packaged
    copy). A checkout with an unstamped VERSION gets "+<short hash>" from git
    when git answers within a moment; no VERSION at all is "unknown"."""
    root = Path(root or REPO_ROOT)
    try:
        ver = _read_text_any(root / "VERSION").strip().splitlines()[0].strip()
    except (OSError, IndexError):
        ver = ""
    if not ver:
        return "unknown"
    if "+" not in ver and (root / ".git").exists():
        try:
            r = subprocess.run(["git", "-C", str(root), "rev-parse", "--short", "HEAD"],
                               capture_output=True, text=True, timeout=3)
            h = (r.stdout or "").strip()
            if r.returncode == 0 and re.fullmatch(r"[0-9a-f]{4,40}", h):
                ver += f"+{h}"
        except (OSError, subprocess.SubprocessError):
            pass
    return ver


VERSION = read_version()


def server_stale_files() -> list[str]:
    """server.py / voicekit_writer.py files changed on disk since this server
    process started — two os.stat calls, cheap enough for every response."""
    return [p.name for p, t in _CODE_STAMPS.items() if _mtime_ns(p) != t]


STALE_NOTE = ("This VoiceKit MCP server started before its code changed on disk "
              "({files}), so new or changed tools aren't visible yet. Restart the MCP "
              "server / reconnect it (Claude Code: /mcp; Claude Desktop: restart the app).")


# ---------------------------------------------------------------------------
# Path normalization in every response (WP9, phase 1 — always on)
#
# Tool responses used to carry C:\Users\<name>\... everywhere (snippet errors,
# temp files, list_automations' file paths, client configs), and with it the
# Windows user name. The profile and temp prefixes become %USERPROFILE% and
# %TEMP%: lossless for a reader, and expand_user_paths() turns them back into
# the real paths on the way IN, so an agent can pass back what it was shown.
# ---------------------------------------------------------------------------
def _win_path_variant(path: str, long: bool) -> str:
    """The 8.3 short (long=False) or long form of an existing path, "" when
    Windows can't say (or it equals the input)."""
    try:
        import ctypes
        fn = ctypes.windll.kernel32.GetLongPathNameW if long else \
            ctypes.windll.kernel32.GetShortPathNameW
        buf = ctypes.create_unicode_buffer(1024)
        n = fn(path, buf, 1024)
        out = buf.value if 0 < n < 1024 else ""
    except (OSError, AttributeError, ValueError):
        return ""
    return "" if out.lower() == path.lower() else out


_PATH_RULES_CACHE: list = []


def _path_rules() -> list:
    """[(compiled regex, token)], most specific prefix first (the temp dir
    lives inside the profile). Each prefix matches with either slash style,
    as a doubled-backslash (repr/JSON-escaped) spelling, and in its 8.3 form,
    caselessly, and only on a path boundary — C:\\Users\\bob never eats
    C:\\Users\\bobby."""
    if _PATH_RULES_CACHE:
        return _PATH_RULES_CACHE
    pairs: list[tuple[str, str]] = []
    temps = {tempfile.gettempdir(), os.environ.get("TEMP", ""), os.environ.get("TMP", "")}
    profile = os.environ.get("USERPROFILE", "") or str(Path.home())
    for token, bases in (("%TEMP%", temps), ("%USERPROFILE%", {profile})):
        for b in list(bases):
            b = (b or "").rstrip("\\/")
            if len(b) < 4:                   # never a drive root
                continue
            for v in {b, _win_path_variant(b, True), _win_path_variant(b, False)}:
                if v:
                    pairs.append((v.rstrip("\\/"), token))
    seen: set = set()
    rules = []
    for prefix, token in sorted(pairs, key=lambda pt: -len(pt[0])):
        for spelled in (prefix, prefix.replace("\\", "/"), prefix.replace("\\", "\\\\")):
            if spelled.lower() in seen:
                continue
            seen.add(spelled.lower())
            rules.append((re.compile(re.escape(spelled) + r"(?![\w.\-~$])", re.IGNORECASE),
                          token))
    _PATH_RULES_CACHE.extend(rules)
    return rules


def normalize_paths_text(s: str) -> str:
    """One string with the profile/temp prefixes replaced by their tokens."""
    if not s or not isinstance(s, str):
        return s
    for rx, token in _path_rules():
        if rx.search(s):
            s = rx.sub(lambda _m, t=token: t, s)
    return s


def normalize_paths(obj, exempt: tuple = ()):
    """normalize_paths_text over every string in a response — dict values
    AND keys, lists, tuples — recursively. `exempt` names top-level dict keys
    to leave byte-for-byte (authored source a caller may send back)."""
    if isinstance(obj, str):
        return normalize_paths_text(obj)
    if isinstance(obj, dict):
        return {normalize_paths_text(k) if isinstance(k, str) else k:
                (v if k in exempt else normalize_paths(v)) for k, v in obj.items()}
    if isinstance(obj, list):
        return [normalize_paths(v) for v in obj]
    if isinstance(obj, tuple):
        return tuple(normalize_paths(v) for v in obj)
    return obj


# Tool results whose content is authored source that must round-trip exactly
# (a caller reads it, edits it, and hands it back to update_*).
NORMALIZE_EXEMPT = {
    # reveal's answer IS the unmasked text the caller asked for (audited);
    # masking it again would make the tool pointless. It is normalized
    # (%USERPROFILE%) and, in strict mode, ID-masked by reveal itself.
    "reveal": ("text",),
    "read_macro_source": ("source",),
    "read_hotkey_module": ("body",),
    "read_workflow": ("steps",),
    "read_ai_prompt": ("prompt",),
    "read_reference": ("code",),
    # Its undo path: the old prompt, to hand back to update_ai_prompt.
    "update_ai_prompt": ("previous_prompt",),
}

_PATH_TOKEN_RE = re.compile(r"%(USERPROFILE|TEMP|TMP)%", re.IGNORECASE)


def expand_user_paths(s: str | None) -> str | None:
    """The inverse for tools that ACCEPT a path or a value to act on: the
    privacy tokens a masked response showed (<dir#1c2e>, <file#3a9f>.pdf —
    WP9 phase 2, privacy.py) become the real names, then '%USERPROFILE%\\x'
    becomes the real path. A real path passes unchanged; an unknown or
    one-way (<ssn#..>) token raises."""
    if not s:
        return s
    if "<" in s:
        s = privacy.expand_tokens(s, REPO_ROOT)
    if "%" not in s:
        return s

    def real(m):
        k = m.group(1).upper()
        if k == "USERPROFILE":
            return privacy.profile_dir().rstrip("\\/")
        return privacy.temp_dir().rstrip("\\/")
    return _PATH_TOKEN_RE.sub(real, s)


def _no_tokens(value, what: str) -> None:
    """Authored code, prompts, snippet text and non-path step fields must not
    carry privacy tokens (privacy.refuse_tokens says why and what to do)."""
    if isinstance(value, str) and "<" in value:
        privacy.refuse_tokens(value, what)


def reveal(token: str) -> dict:
    """The real name behind a privacy token (or every token in a masked
    path). Audited in logs\\privacy-audit.log by token, never by value; a
    one-way <ssn#..>/<ein#..> is refused. The answer keeps %USERPROFILE% /
    %TEMP% (phase 1 is always on) and, in strict mode, ID-masks what it
    returns."""
    r = privacy.reveal(token, REPO_ROOT)
    text = privacy.mask_ids_text(normalize_paths_text(r["text"]), REPO_ROOT)
    return {"text": text, "tokens": r["tokens"],
            "note": "Revealed locally and logged (token only) in logs\\privacy-audit.log. "
                    "Pass tokens back as-is where a tool takes a path or value — reveal is "
                    "only for when you must read the name itself."}


# ---------------------------------------------------------------------------
# server_info: which server is this, and which others are out there
# ---------------------------------------------------------------------------
def _python_processes() -> list:
    """Every running python process: [{pid, ppid, cmd, started}] — one CIM
    query, like _ahk_processes."""
    cmd = (
        "$list = Get-CimInstance Win32_Process -Filter \"Name like 'python%'\" | "
        "ForEach-Object { @{pid = $_.ProcessId; ppid = $_.ParentProcessId; "
        "cmd = $_.CommandLine; "
        "started = $(if ($_.CreationDate) { $_.CreationDate.ToString('yyyy-MM-dd HH:mm:ss') } "
        "else { '' })} }; "
        "ConvertTo-Json -InputObject @($list) -Compress"
    )
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", cmd],
                           capture_output=True, text=True, timeout=30)
        rows = json.loads(r.stdout or "[]")
    except (ValueError, OSError, subprocess.SubprocessError):
        return []
    return [{"pid": _as_int(row.get("pid")), "ppid": _as_int(row.get("ppid")),
             "cmd": str(row.get("cmd") or ""), "started": str(row.get("started") or "")}
            for row in (rows if isinstance(rows, list) else [rows]) if isinstance(row, dict)]


def _mcp_server_processes(rows: list | None = None) -> list:
    """Python processes running a VoiceKit-style server.py, one entry per
    SERVER: a venv's python.exe is a launcher that starts the real
    interpreter as its child with the same command line, so a child whose
    parent is also listed folds into the parent's entry."""
    rows = _python_processes() if rows is None else rows

    def script_of(cmd: str) -> str:
        m = re.search(r'(?i)"([^"]*server\.py)"|(\S*server\.py)\b', cmd)
        return (m.group(1) or m.group(2)) if m else ""

    def voicekit_server(cmd: str) -> bool:
        sp = script_of(cmd)
        if not sp:
            return False
        try:     # a VoiceKit server.py sits beside voicekit_writer.py
            if (Path(sp).parent / "voicekit_writer.py").exists():
                return True
        except (OSError, ValueError):
            pass
        return "voicekit" in cmd.lower()
    hits = [r for r in rows if voicekit_server(r["cmd"])]
    pids = {r["pid"] for r in hits}
    me = {os.getpid(), os.getppid()}
    out = []
    for r in hits:
        if r["ppid"] in pids:
            continue                          # the venv launcher's child: folded
        kids = [k["pid"] for k in hits if k["ppid"] == r["pid"]]
        group = {r["pid"], *kids}
        out.append({"pids": sorted(group), "this_server": bool(group & me),
                    "started": r["started"], "command": r["cmd"],
                    "script": script_of(r["cmd"])})
    return out


def _client_config_files() -> list[tuple[str, Path]]:
    """The client configs that can register an MCP server on this machine."""
    home = Path(os.environ.get("USERPROFILE", "") or Path.home())
    appdata = Path(os.environ.get("APPDATA", "") or home / "AppData" / "Roaming")
    return [("Claude Code (user)", home / ".claude.json"),
            ("Claude Desktop", appdata / "Claude" / "claude_desktop_config.json"),
            ("Claude Code (project)", REPO_ROOT / ".mcp.json")]


def _registrations_in(client: str, path: Path) -> list:
    """VoiceKit-looking MCP server entries in one client config. Only the
    command and its arguments are reported — never `env`, which can hold
    secrets."""
    try:
        data = json.loads(_read_text_any(path))
    except (OSError, ValueError):
        return []
    blocks = [("", data.get("mcpServers"))] if isinstance(data, dict) else []
    if isinstance(data, dict) and isinstance(data.get("projects"), dict):
        blocks += [(proj, v.get("mcpServers")) for proj, v in data["projects"].items()
                   if isinstance(v, dict)]
    out = []
    for scope, servers in blocks:
        if not isinstance(servers, dict):
            continue
        for key, spec in servers.items():
            if not isinstance(spec, dict):
                continue
            args = [str(a) for a in spec.get("args") or [] if isinstance(a, (str, int, float))]
            cmd = str(spec.get("command") or "")
            blob = " ".join([key, cmd, *args]).lower()
            if "voicekit" not in blob:
                continue
            script = next((a for a in args if a.lower().endswith("server.py")), "")
            try:
                ours = bool(script) and (Path(script).resolve() ==
                                         (REPO_ROOT / "mcp" / "server.py").resolve())
            except OSError:
                ours = False
            entry = {"client": client, "config": str(path), "name": key,
                     "command": cmd, "args": args, "this_install": ours}
            if scope:
                entry["project"] = scope
            out.append(entry)
    return out


def server_info() -> dict:
    """Which VoiceKit MCP server this is and what else is registered/running
    — the answer to "why do I see two sets of VoiceKit tools?". Read-only;
    paths come back normalized (%USERPROFILE%, %TEMP%)."""
    import sys
    stale = server_stale_files()
    servers = _mcp_server_processes()
    regs = [e for client, p in _client_config_files() if p.exists()
            for e in _registrations_in(client, p)]
    out = {
        "version": VERSION,
        "install_root": str(REPO_ROOT),
        "python": sys.executable,
        "ahk": AHK_EXE,
        "server_pid": os.getpid(),
        "server_started": SERVER_STARTED_AT.strftime("%Y-%m-%d %H:%M:%S"),
        "server_stale": bool(stale),
        "running_servers": servers,
        "registrations": regs,
    }
    notes = []
    if stale:
        notes.append(STALE_NOTE.format(files=", ".join(stale)))
    others = [s for s in servers if not s["this_server"]]
    if others:
        notes.append(f"{len(others)} other server.py process(es) are running besides this "
                     f"one. Each MCP registration starts its own server process with its own "
                     f"copy of the tools, so a client that has VoiceKit registered twice (or "
                     f"two clients) shows two sets of VoiceKit tool names. They act on the "
                     f"same files; a set that looks out of date belongs to an older process "
                     f"— reconnect it.")
    if len(regs) > 1:
        notes.append(f"VoiceKit is registered {len(regs)} times in this machine's client "
                     f"configs (see 'registrations'). Within ONE client, two entries mean "
                     f"two copies of every tool; keep one per client. Entries in different "
                     f"clients (Claude Code vs Claude Desktop) are normal.")
    if not any(r["this_install"] for r in regs) and regs:
        notes.append("None of the registrations point at this install's server.py — this "
                     "server was started some other way (or from another copy).")
    if not notes:
        notes.append("One VoiceKit server, registered once — nothing to untangle.")
    out["note"] = " ".join(notes)
    return normalize_paths(out)


def server_instructions() -> str:
    """The FastMCP `instructions` block, generated at server start from the
    libraries the code a caller writes actually gets — so the list of taken
    names can't drift from lib\\ (test_instructions_names_exist checks every
    name it lists)."""
    return _instructions_text(*_taken_names())


# Prefix families named as a whole instead of name by name.
NAME_FAMILIES = ("Uia", "UIA_", "Browser", "Body", "Hotkey", "Bridge", "Master", "Wf")


def _snippet_scope_defs() -> dict:
    """What run_ahk_snippet code sees: _Common + Browser's chain +
    ExplorerSel + Out."""
    out: dict = {"out": [("Out", "function", Path("<snippet prelude>"), 0)]}
    for lib in ("_Common.ahk", "Browser.ahk", "ExplorerSel.ahk"):
        for lk, hits in ahk_defs_in_chain(REPO_ROOT / "lib" / lib).items():
            out.setdefault(lk, hits)
    return out


def _taken_names(cap: int = 22) -> tuple[list[str], list[str]]:
    """(individual names, family labels like 'Uia*'): the pre-included
    names a caller is likeliest to pick as a variable — the shortest ones
    outside a prefix family — plus the families (fixed ones, and any first
    word shared by three or more definitions)."""
    defs = _snippet_scope_defs()
    names = sorted({hits[0][0] for hits in defs.values()
                    if not hits[0][0].startswith("_")}, key=str.lower)

    def first_word(n: str) -> str:
        m = re.match(r"[A-Z]+(?=[A-Z][a-z])|[A-Z]?[a-z0-9]+|[A-Z]+_?", n)
        return m.group(0) if m else n
    fams = list(NAME_FAMILIES)
    counts: dict = {}
    for n in names:
        w = first_word(n)
        if w != n:
            counts[w] = counts.get(w, 0) + 1
    fams += sorted(w for w, c in counts.items() if c >= 3 and w not in fams)

    def in_family(n: str) -> bool:
        return any(n.startswith(f) and n != f for f in fams)
    solo = sorted((n for n in names if not in_family(n)), key=lambda n: (len(n), n.lower()))
    if "Out" in solo:
        solo.remove("Out")
        solo.insert(0, "Out")
    return solo[:cap], [f + "*" for f in fams]


def _instructions_text(solo: list[str], families: list[str]) -> str:
    return "\n".join([
        "VoiceKit builds and runs AutoHotkey v2 automations (v1 syntax does not load).",
        "",
        "What your AHK code gets pre-included: run_ahk_snippet code has lib\\_Common.ahk, "
        "lib\\Browser.ahk (which pulls in lib\\UIA.ahk and lib\\Clip.ahk) and "
        "lib\\ExplorerSel.ahk (ExplorerSelectedFiles) plus Out(value). "
        "Generated hotkey-module bodies (create_hotkey_module, isolate=True) and launch "
        "macros made with create_launch_macro(ahk_body=...) #Include lib\\_Common.ahk. An "
        "isolate=False module runs inside VoiceKit.ahk itself, beside lib\\_Common.ahk, "
        "lib\\Theme.ahk and every other module. A hand-written macro includes whatever its "
        "own source says (read_macro_source).",
        "",
        "Names you can't assign at the top level of a script: AutoHotkey v2 keeps "
        "functions and variables in ONE case-insensitive namespace, so `log := 1` fails "
        "to load (Log is a built-in). Inside a function, plain variables are local and "
        "any name is fine. A load error that hits a taken name says who owns it.",
        "AutoHotkey built-ins: " + ", ".join(COMMON_BUILTIN_CLASHES) + " (and every other "
        "built-in function and class).",
        "VoiceKit functions: " + ", ".join(solo) + ".",
        "VoiceKit prefix families (every name starting with): " + ", ".join(families) + ".",
        "",
        "Quotes: v2 escapes a double quote inside \"...\" as `\" (backtick-quote); v1's "
        "doubled \"\" does not load. Or use single quotes: 'say \"hi\"'.",
        "",
        "Changing an existing automation: read it first (read_macro_source, "
        "read_hotkey_module, read_workflow, read_ai_prompt), then use edit_macro / "
        "edit_hotkey_module (one exact splice) or update_macro / update_hotkey_module (the "
        "whole file) / update_ai_prompt. create_workflow replaces a workflow of the same "
        "name. Never delete-and-recreate: a hotkey module's key pairing is made by hand.",
        "",
        "Paths: the profile shows as %USERPROFILE% and temp as %TEMP% (in AHK code: "
        "EnvGet(\"USERPROFILE\") / A_Temp). Privacy masking (on by default; set only in "
        "logs\\settings.ini) also shows names under the profile as tokens that keep the "
        "extension: %USERPROFILE%\\Documents\\<dir#1c2e>\\<file#3a9f>.pdf. Pass paths back "
        "exactly as shown wherever a tool takes a path or value (opens, run/focus/capture "
        "steps, run_automation args, batch rows/source; in run_ahk_snippet only inside quoted "
        "strings) — they expand locally — never in saved code or steps. reveal(token) and "
        "run_ahk_snippet(unmask=True) show real names (audited). Masking stops incidental "
        "leakage, not what a snippet deliberately Out()s.",
        "",
        "Files the user selected in File Explorer: a workflow uses {{selected_file}} (exactly "
        "one, else the run refuses to start) or {{selected_files}} (all, each quoted, "
        "space-joined); in a command line (a capture step's command, a run target) both "
        "arrive already quoted, with or without quotes written around them. A capture "
        "step runs a command line hidden from the VoiceKit folder and saves what it "
        "printed as {{name}} — e.g. capture 'Invoice Data' = python \"C:\\tools\\invoice.py\" "
        "{{selected_file}}, then {{Invoice Data}} in later steps. A launch macro takes them as "
        "`#Include \"%A_ScriptDir%\\..\\lib\\ExplorerSel.ahk\"` then "
        "`files := A_Args.Length ? A_Args : ExplorerSelectedFiles()`. "
        "Either way run_automation(args=[...]) passes a fixture file in their place, so test "
        "on a sample under %TEMP%, never a real client document. There is no generic "
        "dry-run.",
        "",
        "Seeing VoiceKit's tools twice, or a tool that should exist is missing? Call "
        "server_info.",
    ])


def health() -> dict:
    """Whether the always-on layer is actually up, and what is turned off.

    Attached to every MCP tool response: a module can take the master down (or
    be parked for taking it down), and until this existed nothing said so — the
    hotkeys just quietly stopped working."""
    global _LAST_RELOAD_ERROR
    running = voicekit_running()
    mods = index_modules()
    st = read_master_status()
    # The flag file alone: the ini's safe_mode is only stamped at master
    # start, so deleting the flag (the documented way out) used to leave
    # safe mode reported until the next restart.
    safe_mode = _safe_mode_flag().exists()
    # A refused reload stops being news once a master has started since —
    # the user fixed it and reloaded by hand (Ctrl+Alt+Shift+R / launcher).
    if _LAST_RELOAD_ERROR and st.get("started", "") > _LAST_RELOAD_ERROR_AT:
        _LAST_RELOAD_ERROR = ""

    stale = server_stale_files()
    out: dict = {
        "master_running": running,
        "modules_loaded": len(mods["active"]),
        "version": VERSION,
        "install_root": str(REPO_ROOT),
        "server_stale": bool(stale),
    }
    if _master_restarting():
        out["master_restarting"] = True     # mid-reload: hotkeys come back in a moment
    if mods["quarantined"]:
        out["quarantined"] = mods["quarantined"]
    if safe_mode:
        out["safe_mode"] = True
    if st:
        if st.get("pid"):
            out["pid"] = st["pid"]
        age = _age_seconds(st.get("heartbeat", ""))
        if age is not None:
            out["heartbeat_age_s"] = age
        # last_error is never cleared by the master, so only report one this
        # master generation hit (last_error_at >= started), with its age —
        # a months-old error attached to every response read as current.
        err_at = st.get("last_error_at", "")
        if st.get("last_error") and err_at and err_at >= st.get("started", ""):
            out["last_error"] = st["last_error"]
            err_age = _age_seconds(err_at)
            if err_age is not None:
                out["last_error_age_s"] = err_age

    notes = []
    if not running:
        notes.append("VoiceKit isn't running, so no hotkey or snippet works. "
                     "Start it with VoiceKitLauncher.ahk.")
    if safe_mode:
        notes.append("Safe mode: VoiceKit kept crashing, so every hotkey module was "
                     "turned off. Re-enable them one at a time in hotkeys\\_index.ahk.")
    elif mods["quarantined"]:
        notes.append(f"{len(mods['quarantined'])} module(s) turned off because they "
                     f"wouldn't load: {', '.join(mods['quarantined'])}. Fix the file, "
                     f"then uncomment its line in hotkeys\\_index.ahk.")
    if _LAST_RELOAD_ERROR:
        notes.append("The last reload was refused (the running copy was left alone): "
                     + _LAST_RELOAD_ERROR)
    if stale:
        notes.append(STALE_NOTE.format(files=", ".join(stale)))
    if notes:
        out["note"] = " ".join(notes)
    return out


def _launch_master(target: Path):
    """Start the master (a separate function so tests can stub the launch).
    Returns the Popen."""
    return subprocess.Popen([AHK_EXE, str(target)])


# How long a reload waits for the previous reload's launcher to finish. A
# launcher normally lives about a second (one /validate of VoiceKit.ahk, then
# it starts the master and exits); the bound covers a slow load check and
# its retry at twice the timeout.
LAUNCHER_WAIT_S = 30


def _wait_for_last_launcher() -> None:
    """Block until the launcher the previous reload started has exited.

    A rolled-back clash reloads twice within a fraction of a second, and the
    second launch used to land while the first launcher was still running
    its preflight: two preflights editing hotkeys\\_index.ahk at once, and —
    because the launcher had no #SingleInstance directive (v2's default is
    Prompt) — the second one sitting on an "already running, replace it?"
    modal on the user's desktop, with the reload the rollback depended on
    held up behind it. The launcher now says #SingleInstance Off too; this
    wait is what keeps the two preflights apart."""
    global _LAST_LAUNCH
    proc, _LAST_LAUNCH = _LAST_LAUNCH, None
    wait = getattr(proc, "wait", None)
    if wait is None:
        return
    try:
        wait(timeout=LAUNCHER_WAIT_S)
    except subprocess.TimeoutExpired:
        # Left running, not killed: a launcher that slow may be showing the
        # user its "VoiceKit couldn't start" message, and killing it would
        # take that away. #SingleInstance Off means the next one won't prompt.
        pass


@_serialized
def reload_voicekit() -> str:
    """Reload the master (so a new snippet/hotkey goes live) by launching it
    again — #SingleInstance Force replaces the running instance. Only acts if
    VoiceKit is already running, to avoid unexpectedly starting the tray app.

    Load-checks FIRST and parks whatever won't compile: reloading into a config
    that doesn't build kills every hotkey and snippet at once, and the master
    stays dead until someone restarts it by hand. What got parked lands in
    _LAST_PARKED.

    Returns 'reloaded', 'not_running' (nothing to reload — the change goes
    live when VoiceKit starts) or 'refused' (running, but the files as they
    stand won't load, so the running copy was left alone)."""
    global _LAST_RELOAD_ERROR, _LAST_RELOAD_ERROR_AT, _LAST_LAUNCH
    _LAST_PARKED.clear()
    # Before even asking whether the master is up: mid-handover the answer
    # can be wrong, and our preflight must not edit _index.ahk while the
    # previous launcher's preflight may still be editing it.
    _wait_for_last_launcher()
    if not voicekit_running():
        return "not_running"
    reasons: dict = {}
    ok, _parked, err = master_preflight(reasons)
    _LAST_PARKED.update({k.lower(): v for k, v in reasons.items()})
    if not ok:
        _LAST_RELOAD_ERROR = " ".join(err.split())[:400]
        _LAST_RELOAD_ERROR_AT = datetime.now().strftime("%Y%m%d%H%M%S")
        return "refused"
    _LAST_RELOAD_ERROR = ""
    if _launcher_path().exists():
        _LAST_LAUNCH = _launch_master(_launcher_path())
    else:
        # The master itself never exits, so there's nothing to wait for next
        # time; its #SingleInstance Force handles a second start.
        _launch_master(_master_path())
    return "reloaded"


def _reload_status() -> str:
    """Reload now and name the outcome. Tolerates a stubbed reload_voicekit
    that returns a bool (True = reloaded; False = refused when a refusal was
    recorded, else not running)."""
    _LAST_PARKED.clear()        # never judge this reload by an earlier one's parking
    st = reload_voicekit()
    if st is True:
        return "reloaded"
    if st is False or st is None:
        return "refused" if _LAST_RELOAD_ERROR else "not_running"
    return st


def _reload_note(status: str, subject: str) -> str:
    """The one wording for "did the reload happen", for every tool that
    reloads. `subject` is what goes live, e.g. "the hotkey"."""
    if status == "reloaded":
        return f"VoiceKit was reloaded, so {subject} is live now."
    if status == "not_running":
        return (f"VoiceKit isn't running, so nothing was reloaded — {subject} goes live "
                f"as soon as it starts (VoiceKitLauncher.ahk).")
    return (f"VoiceKit is running, but the reload was refused because the files as they "
            f"stand won't load — see 'voicekit' in this response. {subject[:1].upper()}"
            f"{subject[1:]} goes live on the next successful reload.")


def _reload_outcome(subject: str) -> dict:
    """Reload and describe it: {'reloaded': bool, 'reload_status', 'reload_note'}."""
    st = _reload_status()
    return {"reloaded": st == "reloaded", "reload_status": st,
            "reload_note": _reload_note(st, subject)}


# For tools whose result never depends on a reload: launch macros, workflows,
# AI actions and isolated bodies all run as their own process, fresh each time.
NO_RELOAD_NOTE = "Live now — it runs as its own process, so nothing needs reloading."


def _parked_reason(rel: str) -> str | None:
    """The load error, if the last reload parked `rel` (else None)."""
    return _LAST_PARKED.get(rel.lower())


_CLASH_HINT = ("Rename the clashing name (or use isolate=True, which runs the code in "
               "its own process) and try again.")


def _clash_error(rel: str, reason: str, what: str, hint: str = _CLASH_HINT) -> VoiceKitError:
    return VoiceKitError(
        f"{rel} loads on its own, but VoiceKit wouldn't load with it — something in "
        f"it clashes with code already loaded into the resident VoiceKit (a function, "
        f"variable, hotkey or hotstring another file also defines: lib\\_Common.ahk, "
        f"lib\\Theme.ahk, VoiceKit.ahk or another module). The reload turned it off, so "
        f"{what}. {hint} The load error:\n{_explain_ahk_error(reason)}")


def _ensure_voice_dir() -> None:
    VOICE_MACROS.mkdir(parents=True, exist_ok=True)


def _require_name(name: str) -> tuple[str, str]:
    phrase = clean_phrase(name)
    base = to_base(phrase)
    if not base or is_reserved(base):
        raise VoiceKitError(
            f"'{name}' isn't a usable name (empty after cleanup, or a reserved "
            f"Windows device name). Pick another.")
    return phrase, base


def _workflow_base(name: str) -> str:
    """A workflow's base from a name that may carry the loop phrase's
    'loop ' prefix ('loop Send Invoice' -> 'SendInvoice') — unless a workflow
    is literally named that way ('Loop Timer' stays 'LoopTimer' when
    workflows\\LoopTimer.steps.txt exists)."""
    n = (name or "").strip()
    base = _require_name(n)[1]
    if n.lower().startswith("loop ") and n[5:].strip() and not _steps_file(base).exists():
        return _require_name(n[5:])[1]
    return base


# Artifact paths. Functions, not constants: tests retarget REPO_ROOT.
def _macro_file(base: str) -> Path:
    return REPO_ROOT / "macros" / f"{base}.ahk"


def _module_file(base: str) -> Path:
    return REPO_ROOT / "hotkeys" / f"{base}.ahk"


def _steps_file(base: str) -> Path:
    return REPO_ROOT / "workflows" / f"{base}.steps.txt"


def _prompt_file(base: str) -> Path:
    return REPO_ROOT / "prompts" / f"{base}.prompt.txt"


def _snippets_file() -> Path:
    _seed_live("hotkeys\\Snippets.ahk")
    return REPO_ROOT / "hotkeys" / "Snippets.ahk"


def _bridge_map_file() -> Path:
    _seed_live("bridge-map.txt")
    return REPO_ROOT / "bridge-map.txt"


# ---------------------------------------------------------------------------
# Content builders
# ---------------------------------------------------------------------------
def _launch_content(phrase: str, ahk_body: str | None) -> str:
    if ahk_body is None:
        tpl = _read_template("launch-template.ahk")
        return tpl.replace("{{PHRASE}}", phrase).replace("{{DATE}}", _today())
    return (
        "#Requires AutoHotkey v2.0\n"
        "#SingleInstance Force\n"
        "; ============================================================\n"
        f";  {phrase}   (created {_today()}, via MCP)\n"
        f';  Trigger by voice:  "open {phrase}"\n'
        ";  Runs top to bottom, then exits.\n"
        "; ============================================================\n"
        '#Include "%A_ScriptDir%\\..\\lib\\_Common.ahk"\n'
        "\n"
        + ahk_body.strip("\n") + "\n"
    )


def _opens_content(phrase: str, target: str) -> str:
    """lib\\_Common.ahk VkOpensMacroContent: the no-code 'Open Something'
    macro. Keep the ';  Opens:  <target>' header line — the Voice Kit home
    window parses it for the row's detail text. test_generators_match_ahk
    byte-compares the two."""
    run_arg = f'"{target}"' if (" " in target and Path(target).exists()) else target
    return (
        "#Requires AutoHotkey v2.0\n"
        "#SingleInstance Force\n"
        "; ============================================================\n"
        f";  {phrase}   (created {_today()})\n"
        f';  Trigger by voice:  "open {phrase}"\n'
        f";  Opens:  {target}\n"
        ";\n"
        ";  Created by New Automation — no code needed. To add steps,\n"
        ";  edit below (building blocks: templates\\launch-template.ahk).\n"
        "; ============================================================\n"
        '#Include "%A_ScriptDir%\\..\\lib\\_Common.ahk"\n'
        f"Run({ahk_str_lit(run_arg)})\n"
    )


def _ai_action_content(phrase: str, base: str) -> str:
    """NewAutomation.ahk NewAIAction: fill templates\\ai-template.ahk."""
    tpl = _read_template("ai-template.ahk")
    return (tpl.replace("{{PHRASE}}", phrase)
               .replace("{{BASE}}", base)
               .replace("{{DATE}}", _today()))


def _indent_body(ahk_body: str) -> str:
    lines = ahk_body.strip("\n").split("\n")
    return "\n".join(("    " + ln) if ln.strip() else "" for ln in lines)


def _hotkey_content(phrase: str, key: str, ahk_body: str | None) -> str:
    if ahk_body is None:
        tpl = _read_template("hotkey-template.ahk")
        return (tpl.replace("{{PHRASE}}", phrase)
                   .replace("{{KEY}}", key)
                   .replace("{{DATE}}", _today()))
    return (
        "#Requires AutoHotkey v2.0\n"
        "; ============================================================\n"
        f";  {phrase}   (created {_today()}, via MCP)\n"
        ";  Loaded by VoiceKit.ahk — do not run this file directly.\n"
        f";  Trigger key:  Ctrl+Alt+Shift+{key}\n"
        f';  Voice pairing: say "show voice shortcuts" -> new shortcut ->\n'
        f";    When I say: {phrase}  ->  Press keys  ->  Ctrl+Alt+Shift+{key}\n"
        "; ============================================================\n"
        "\n"
        f"^!+{key}:: {{\n"
        f"{_indent_body(ahk_body)}\n"
        "}\n"
    )


# ---------------------------------------------------------------------------
# Isolated hotkey modules — mirrors lib\_Common.ahk HotkeyBodyRel /
# HotkeyLauncherContent / HotkeyBodyContent.
#
# VoiceKit is ONE process holding every always-on hotkey and every snippet.
# Custom module code can end that process outright — a bad ComCall offset, an
# unpinned COM vtable (see lib\Acc.ahk) — with no exception to catch, and every
# other hotkey and snippet dies with it. So custom code doesn't live in the
# master: hotkeys\<Base>.ahk is a key binding that launches
# hotkeys\bodies\<Base>.body.ahk in its own process.
# ---------------------------------------------------------------------------
BODY_SUBDIR = "bodies"


def _body_rel(base: str) -> str:
    """Repo-relative body script for <base> (mirrors _Common.ahk HotkeyBodyRel)."""
    return f"hotkeys\\{BODY_SUBDIR}\\{base}.body.ahk"


def _hotkey_launcher_content(phrase: str, base: str, key: str) -> str:
    # A_ScriptDir below is VoiceKit.ahk's folder, not this file's — the module
    # is #Included into the master, like the companion modules.
    return (
        "#Requires AutoHotkey v2.0\n"
        "; ============================================================\n"
        f";  {phrase}   (created {_today()})\n"
        ";  Loaded by VoiceKit.ahk — do not run this file directly.\n"
        ";\n"
        f";  Trigger key:  Ctrl+Alt+Shift+{key}\n"
        ";\n"
        f";  EDIT THE STEPS IN:  {_body_rel(base)}\n"
        ";  They run in their own process on purpose. VoiceKit is one\n"
        ";  process holding every hotkey and every snippet, and code that\n"
        ";  crashes hard would take the whole lot down with it. Out there,\n"
        ";  a crash ends one short-lived process and nothing else.\n"
        ";\n"
        ";  To trigger it by voice (one-time setup, ~30 seconds):\n"
        ';    1. Say: "show voice shortcuts"\n'
        f";    2. Create new shortcut  ->  When I say:  {phrase}\n"
        f";    3. Action: Press keys   ->  Ctrl + Alt + Shift + {key}\n"
        ";  This pairing is recorded in bridge-map.txt.\n"
        "; ============================================================\n"
        "\n"
        f"^!+{key}:: {{\n"
        f'    body := A_ScriptDir "\\{_body_rel(base)}"\n'
        "    if !FileExist(body) {\n"
        f'        TrayTip("The steps for Ctrl+Alt+Shift+{key} are missing:`n" body, "VoiceKit")\n'
        "        return\n"
        "    }\n"
        "    q := Chr(34)                     ; association-proof: run via the interpreter\n"
        "    Run(q A_AhkPath q ' ' q body q)\n"
        "}\n"
    )


def _hotkey_body_content(phrase: str, base: str, key: str, body: str | None) -> str:
    say = phrase.replace("`", "``").replace('"', '`"')
    if not (body or "").strip():
        body = ("; ==== YOUR STEPS BELOW — delete the MsgBox once it works ====\n"
                f'MsgBox("\'{say}\' is wired up! Now edit this file:`n" A_LineFile)')
    return (
        "#Requires AutoHotkey v2.0\n"
        "#SingleInstance Ignore\n"
        "; ============================================================\n"
        f';  Steps for "{say}"  —  Ctrl+Alt+Shift+{key}   (created {_today()})\n'
        ";\n"
        f";  Runs in its own process, started by hotkeys\\{base}.ahk every\n"
        ";  time the key is pressed. A crash in here ends only this\n"
        ";  process — VoiceKit, the other hotkeys and the snippets carry on.\n"
        ";\n"
        ";  #SingleInstance Ignore: pressing the key again while this is\n"
        ";  still running is ignored rather than piling up processes.\n"
        ";\n"
        ";  Long job? Say where you are with BodyStatus — the Voice Kit\n"
        ";  home window and the MCP listing both show the latest line, so\n"
        ";  progress is visible without watching a ToolTip:\n"
        f';    BodyStatus("{base}", "3 / 766  -  Acme Holdings")\n'
        f';    BodyStatusDone("{base}", "766 done")\n'
        ";\n"
        ";  Long LOOP? Check the stop flag each pass - the stop_module\n"
        ";  tool asks nicely (BodyStopRequested), waits, then force-kills:\n"
        f';    if BodyStopRequested("{base}")\n'
        ";        ExitApp()\n"
        "; ============================================================\n"
        '#Include "%A_ScriptDir%\\..\\..\\lib\\_Common.ahk"\n'
        "\n"
        "; One instance per module: #SingleInstance can lose a startup race\n"
        "; (two presses in quick succession); the kernel mutex below cannot.\n"
        f'BodySingleInstance("{base}")\n'
        "\n"
        f"{body}\n"
    )


def workflow_stub(phrase: str, base: str, date: str | None = None) -> str:
    """The generated macros\\<Base>.ahk stub. Mirrors lib\\Workflow.ahk
    WfStubContent (what the Studio's SaveWorkflow writes) byte for byte,
    including the 'Generated by Workflow Studio' marker the overwrite/delete
    guards look for — test_generators_match_ahk runs both."""
    d = date or _today()
    return (
        "#Requires AutoHotkey v2.0\n"
        "#SingleInstance Force\n"
        "; ============================================================\n"
        f";  {phrase}   (workflow, saved {d})\n"
        f';  Trigger by voice:  "open {phrase}"\n'
        ";\n"
        f";  {STUDIO_MARKER} — don't edit steps here.\n"
        f';  Edit by voice:  "open workflow studio"  ->  pick "{phrase}"\n'
        f";  Steps live in:  workflows\\{base}.steps.txt\n"
        "; ============================================================\n"
        '#Include "%A_ScriptDir%\\..\\lib\\_Common.ahk"\n'
        '#Include "%A_ScriptDir%\\..\\lib\\Workflow.ahk"\n'
        f'RunWorkflow(A_ScriptDir "\\..\\workflows\\{base}.steps.txt")\n'
    )


def steps_to_text(phrase: str, steps: list) -> str:
    """Serialize steps to the .steps.txt format (header + type|encA|encB|encC)."""
    out = f"; {phrase} — VoiceKit workflow. Edit it by saying: open workflow studio\n"
    for s in steps:
        t = s[0]
        a = wf_encode(s[1] if len(s) > 1 and s[1] is not None else "")
        b = wf_encode(s[2] if len(s) > 2 and s[2] is not None else "")
        c = wf_encode(s[3] if len(s) > 3 and s[3] is not None else "")
        out += f"{t}|{a}|{b}|{c}\n"
    return out


# ---------------------------------------------------------------------------
# Create
# ---------------------------------------------------------------------------
def _require_free_lnk(disp: str) -> Path:
    """Refuse a voice phrase whose Start Menu entry already exists — creating it
    would silently hijack another automation's phrase (NewAutomation checks the
    same thing since the production review)."""
    _ensure_voice_dir()
    link = VOICE_MACROS / f"{disp}.lnk"
    if link.exists():
        raise VoiceKitError(
            f'The voice phrase "open {disp}" is already taken by another Start Menu '
            f"entry. Pick another name.")
    return link


def create_launch_macro(name: str, ahk_body: str | None = None,
                        opens: str | None = None) -> dict:
    """A standalone macro + Start Menu entry. Exactly one content source:
    `opens` (the no-code 'Open Something' generator — app/file/folder/URL) or
    `ahk_body` (custom AHK v2 code), or neither (placeholder template).
    The .lnk is named SpaceOut(base) — the same name listing, Delete, and
    first-run reinstall reconstruct — so they never drift (phrases with digits
    don't round-trip otherwise)."""
    _no_tokens(ahk_body, "ahk_body")
    if ahk_body and opens:
        raise VoiceKitError("Give either 'opens' or 'ahk_body', not both.")
    if opens is not None:
        # '%USERPROFILE%\...' (the form responses show) means the real path.
        opens = expand_user_paths(opens.strip())
        if not opens:
            raise VoiceKitError("'opens' is empty — give an app, file, folder, or https:// URL "
                                "(or use 'ahk_body' for custom code).")
        # 'opens' is a single target, advertised as no-code data — a line break
        # would escape the ';  Opens:' comment header and inject executable AHK.
        if "\n" in opens or "\r" in opens:
            raise VoiceKitError("'opens' can't contain line breaks — it names one thing to open. "
                                "For multi-line code use 'ahk_body'.")
    phrase, base = _require_name(name)
    macro = _macro_file(base)
    if macro.exists():
        raise VoiceKitError(f"A macro named '{base}' already exists. Pick another name.")
    disp = space_out(base)
    link = _require_free_lnk(disp)
    content = _opens_content(phrase, opens) if opens else _launch_content(phrase, ahk_body)
    _replace_validated(macro, content,
                       "The generated macro failed to load, so nothing was kept:")
    make_shortcut(str(link), str(macro))
    log(f"launch | {disp} | macros\\{base}.ahk" + (f" | opens {opens}" if opens else ""))
    return {"type": "launch_macro", "phrase": disp, "file": str(macro),
            "shortcut": str(link), "voice_phrase": f"open {disp}",
            "note": f'{NO_RELOAD_NOTE} Say "open {disp}".'}


def create_ai_action(name: str, prompt: str) -> dict:
    """An AI text action (mirrors NewAutomation.ahk NewAIAction): the prompt is
    saved to prompts\\<Base>.prompt.txt and the macro is templates\\ai-template.ahk
    filled in. Select text anywhere, say 'open <name>' — the AI's answer replaces
    it. Needs the user's OpenRouter key (AI Settings); creating the action never
    calls the network."""
    _no_tokens(prompt, "The prompt")
    phrase, base = _require_name(name)
    prompt = prompt.strip()
    if not prompt:
        raise VoiceKitError("The AI action needs a prompt — what should the AI do "
                            "with the selected text?")
    macro = _macro_file(base)
    if macro.exists():
        raise VoiceKitError(f"A macro named '{base}' already exists. Pick another name.")
    disp = space_out(base)
    link = _require_free_lnk(disp)
    prompt_file = _prompt_file(base)
    _write_validated([(prompt_file, prompt + "\n"), (macro, _ai_action_content(phrase, base))],
                     [macro], "The generated AI action failed to load, so nothing was kept:")
    make_shortcut(str(link), str(macro))
    log(f"ai-action | {disp} | macros\\{base}.ahk")
    return {"type": "ai_action", "phrase": disp, "file": str(macro),
            "prompt_file": str(prompt_file), "shortcut": str(link),
            "voice_phrase": f"open {disp}",
            "note": NO_RELOAD_NOTE + " Runs on the user's selected text; needs their "
                    "OpenRouter key (home window -> AI Settings) the first time."}


def read_ai_prompt(name: str) -> dict:
    """The full current prompt of an AI action (list_automations only previews
    the first 120 characters — read before rewriting)."""
    base = _require_name(name)[1]
    prompt_file = _prompt_file(base)
    if not prompt_file.exists():
        raise VoiceKitError(f"No AI action named '{base}' (looked for "
                            f"{prompt_file.name}).{_near_miss(base)}")
    return {"type": "ai_action", "base": base, "name": space_out(base),
            "prompt_file": str(prompt_file),
            "prompt": _lf(_read_text_any(prompt_file)).strip()}


def update_ai_prompt(name: str, prompt: str) -> dict:
    """Rewrite an existing AI action's prompt (the macro itself is untouched).
    Returns the previous prompt so an unwanted overwrite can be undone."""
    _no_tokens(prompt, "The prompt")
    prev = read_ai_prompt(name)          # same resolution + not-found error
    prompt = prompt.strip()
    if not prompt:
        raise VoiceKitError("The new prompt can't be empty.")
    _write_new(Path(prev["prompt_file"]), prompt + "\n")
    log(f"ai-prompt updated | {prev['base']}")
    return {"type": "ai_action", "base": prev["base"], "prompt_file": prev["prompt_file"],
            "previous_prompt": prev["prompt"],
            "note": "Applies the next time the action runs."}


# Names hotkeys\ already owns. Overwriting Snippets.ahk would wipe EVERY
# snippet the user has — reachable, because clean_phrase("Snippets") is
# "Snippets". _index.ahk needs no entry: clean_phrase strips the underscore,
# so the manifest's name can't come out of the naming pipeline at all.
RESERVED_MODULE_BASES = {"snippets"}


def normalize_bridge_key(requested: str) -> str:
    """A caller-requested bridge key -> the single pool character it means.
    Tolerates a full 'Ctrl+Alt+Shift+K' and lowercase letters. Raises when it
    isn't an allocatable key."""
    k = (requested or "").strip()
    low = k.lower()
    if low.startswith("ctrl+alt+shift+"):
        k = k[len("ctrl+alt+shift+"):].strip()
    if len(k) == 1 and k.isalpha():
        k = k.upper()
    if len(k) != 1 or k not in BRIDGE_POOL:
        raise VoiceKitError(
            f"'{requested}' isn't a usable hotkey key. Pick one character from: "
            f"{BRIDGE_POOL}  ({', '.join(RESERVED_BRIDGE_KEYS[:-1])} and "
            f"{RESERVED_BRIDGE_KEYS[-1]} are reserved by VoiceKit itself, and "
            f"Space/Enter can't be triggered programmatically).")
    return k


def _module_bridge_key(base: str) -> str:
    """The Ctrl+Alt+Shift key already registered to hotkeys\\<base>.ahk, or ""."""
    rel = f"hotkeys\\{base}.ahk".lower()
    for e in get_bridge_map()["entries"]:
        if e["file"].lower() == rel:
            return e["combo"].rsplit("+", 1)[-1]
    return ""


def _uncomment_include(rel: str) -> bool:
    """Undo a quarantine: switch a commented-out #Include of `rel` back on.
    The inverse of _comment_out_include — re-creating a module that got parked
    for not loading is exactly the "I fixed it" signal."""
    index = _hotkeys_index()
    if not index.exists():
        return False
    out, hit = [], False
    for line in _read_text_any(index).splitlines():
        if (line.lstrip().startswith(";") and "#Include" in line
                and "\\" + rel in line):
            # Rebuilt canonically rather than un-prefixed, which also drops the
            # "quarantined <date> — ..." tag we appended when parking it.
            out.append(f'#Include "%A_ScriptDir%\\{rel}"')
            hit = True
        else:
            out.append(line)
    if hit:
        _write_new(index, "\n".join(out).rstrip("\n") + "\n")   # BOM + LF, like AHK
    return hit


# ---- The overwritten version, banked one step deep -------------------------
# A replace/update keeps the file it overwrote in logs\<kind>\<Base>.prev.ahk
# (kind: module-backups | macro-backups) so an unwanted overwrite stays
# undoable, while the response carries only a hash + first lines — echoing
# the full old body cost a real session its length on every one of five
# rewrites. The backup is the RAW previous bytes (encoding and line endings
# as they were), so what comes back is exactly what was there.
_MODULE_BACKUPS = "module-backups"
_MACRO_BACKUPS = "macro-backups"


def _backup_file(kind: str, base: str) -> Path:
    return REPO_ROOT / "logs" / kind / f"{base}.prev.ahk"


def _module_backup_file(base: str) -> Path:
    """logs\\module-backups\\<Base>.prev.ahk — the code a replace overwrote."""
    return _backup_file(_MODULE_BACKUPS, base)


def _bank_previous(kind: str, base: str, raw: bytes, prefix: str) -> dict:
    """Bank `raw` as <base>'s previous version; returns the response fields
    (<prefix>_sha256 / _length / _first_lines / _backup). The hash and length
    describe the decoded text — what read_*(previous=True) gives back."""
    backup = _backup_file(kind, base)
    backup.parent.mkdir(parents=True, exist_ok=True)
    backup.write_bytes(raw)
    text = _decode_any(raw)
    return {f"{prefix}_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
            f"{prefix}_length": len(text),
            f"{prefix}_first_lines": text.splitlines()[:3],
            f"{prefix}_backup": str(backup)}


def _read_backup(kind: str, base: str, never: str) -> tuple[Path, str]:
    """(path, text) of <base>'s banked previous version, or raise saying why
    there is none (`never`: when a backup gets written)."""
    backup = _backup_file(kind, base)
    if not backup.exists():
        raise VoiceKitError(f"No saved previous version of '{base}' — {never}")
    return backup, _read_text_any(backup)


def _require_module(name: str, verb: str) -> tuple[str, str, Path]:
    """(phrase, base, module file) of an EXISTING hotkey module, refusing
    VoiceKit's own files in hotkeys\\ and naming what the name is instead
    when there's no such module."""
    phrase, base = _require_name(name)
    if base.lower() in RESERVED_MODULE_BASES:
        raise VoiceKitError(
            f"'{base}' is one of VoiceKit's own files in hotkeys\\ (Snippets.ahk holds "
            f"every snippet you have), not a module, and can't be {verb} this way.")
    module = _module_file(base)
    if not module.exists():
        raise VoiceKitError(f"No hotkey module '{base}' — list_automations shows "
                            f"what exists.{_near_miss(base)}")
    return phrase, base, module


def _was_parked(rel: str) -> bool:
    return rel.lower() in [m.lower() for m in index_modules()["quarantined"]]


def _undo_if_parked(rel: str, snap: _FileSnapshot, was_parked: bool,
                    added_index: bool = False, added_map: bool = False,
                    hint: str = _CLASH_HINT) -> None:
    """If the reload just parked `rel` — it loads alone but not inside the
    master — put everything back and raise. The files get their exact
    previous bytes (a brand-new one is removed), registry lines this call
    added are removed again, and a module that was live before is switched
    back on and reloaded — so the user ends up where they started instead of
    with a module that silently stopped working."""
    reason = _parked_reason(rel)
    if reason is None:
        return
    snap.restore()
    if added_map:
        _remove_matching_lines(_bridge_map_file(), lambda ln: _map_line_targets(ln, rel))
    if added_index:
        _remove_matching_lines(_hotkeys_index(), lambda ln: _include_targets(ln, rel))
        what = ("nothing was kept" if not (REPO_ROOT / rel).exists()
                else "the previous version was put back")
    elif not was_parked:
        _uncomment_include(rel)
        st = _reload_status()
        if st == "reloaded" and _parked_reason(rel) is not None:
            # The restored version was parked as well, so the clash is with
            # some OTHER file — the module stays off until that is fixed.
            what = ("the previous version was put back, but VoiceKit wouldn't load "
                    "with that either, so it is turned off for now (the clash is "
                    "with another file)")
        elif st == "reloaded":
            what = "the previous version was put back and VoiceKit reloaded with it"
        else:
            what = "the previous version was put back"
    else:
        what = "the previous version was put back (it stays turned off, as it was before)"
    raise _clash_error(rel, reason, what, hint)


def _also_parked_note(rel: str) -> str:
    """Other modules the last reload turned off — possibly because they clash
    with the one just written (AutoHotkey blames whichever of two clashing
    definitions it meets SECOND, which need not be the new one)."""
    others = [r for r in _LAST_PARKED if r != rel.lower()]
    if not others:
        return ""
    return (f" Note: that reload also turned off {', '.join(others)} because VoiceKit "
            f"wouldn't load with it — if that started with this change, the two may "
            f"define the same function or hotkey.")


def _isolated_live_note(base: str) -> str:
    """What an isolated body change means right now: the body runs fresh on
    every press, but only while its key binding is loaded."""
    rel = f"hotkeys\\{base}.ahk"
    if _was_parked(rel):
        return (f"The steps are saved, but the key binding ({rel}) is turned off — "
                f"it was quarantined after a failed load, so the key does nothing until "
                f"it is switched back on in hotkeys\\_index.ahk.")
    if not voicekit_running():
        return ("The steps are saved; VoiceKit isn't running, so the key works once it "
                "starts (VoiceKitLauncher.ahk). Nothing needs reloading.")
    return "Live now — the body runs fresh on each press, so nothing needs reloading."


def read_hotkey_module(name: str, previous: bool = False) -> dict:
    """The current code of a hotkey module, so it can be edited rather than
    rewritten from memory. Returns the body for an isolated module (the half
    that holds the steps) and the module file itself otherwise.

    previous=True returns the version the LAST REPLACE overwrote instead
    (create_hotkey_module and update_hotkey_module bank it in
    logs\\module-backups\\) — the undo path for an unwanted replace, now that
    the replace response carries only a hash of the old body rather than the
    whole text."""
    phrase, base, module = _require_module(name, "read")
    if previous:
        backup, text = _read_backup(
            _MODULE_BACKUPS, base, "a backup is written when a replace or update "
                                   "overwrites it, and it hasn't been replaced.")
        return {
            "type": "hotkey_module", "phrase": phrase, "base": base,
            "previous": True, "file": str(backup), "body": text,
            "note": ("The full file as it stood before the last replace. It is the "
                     "complete file, so to restore its content, pass this whole text "
                     "to update_hotkey_module (it is saved back as UTF-8 with LF "
                     "line endings)."),
        }
    body_file = REPO_ROOT / _body_rel(base)
    isolated = body_file.exists()
    src = body_file if isolated else module
    return {
        "type": "hotkey_module", "phrase": phrase, "base": base,
        "isolated": isolated, "file": str(src),
        "bridge_key": (lambda k: f"Ctrl+Alt+Shift+{k}" if k else "")(_module_bridge_key(base)),
        "body": _read_text_any(src),
        "note": ("This is the body that runs in its own process; hotkeys\\"
                 f"{base}.ahk is just the key binding. Change it with "
                 "edit_hotkey_module (one splice) or update_hotkey_module (the whole "
                 "file, verbatim)." if isolated else
                 "This module runs in VoiceKit's own process (isolate=False). Change "
                 "it with edit_hotkey_module or update_hotkey_module."),
    }


@_serialized
def create_hotkey_module(name: str, ahk_body: str | None = None,
                         isolate: bool = True, key: str | None = None) -> dict:
    """Create — or REPLACE — an always-on hotkey module.

    Replacing in place mirrors create_workflow, deliberately: a module used to
    be un-overwritable, so changing one meant delete-then-recreate. That cost
    more than inventing a new name, which is how a session ends up with five
    near-identical throwaway modules.

    A replace KEEPS the module's existing Ctrl+Alt+Shift key. That key is
    paired to a Voice Access shortcut by hand and there is no API to re-pair
    it, so reallocating would silently break the user's voice trigger.

    `key` asks for a specific one instead of taking whatever comes next off
    the pool; it must be free, and it is only honoured when creating."""
    _no_tokens(ahk_body, "ahk_body")
    phrase, base = _require_name(name)
    if base.lower() in RESERVED_MODULE_BASES:
        raise VoiceKitError(
            f"'{base}' is one of VoiceKit's own files in hotkeys\\ and can't be used as "
            f"a module name — hotkeys\\Snippets.ahk holds every snippet you have, and "
            f"_index.ahk is the module manifest. Pick another name.")
    module = _module_file(base)
    body_file = REPO_ROOT / _body_rel(base)
    replacing = module.exists()

    wanted = normalize_bridge_key(key) if key else ""
    # Reuse the registered key on a replace; only a genuinely new module draws
    # from the pool.
    key = _module_bridge_key(base) if replacing else ""
    registered = bool(key)
    if registered and wanted and wanted != key:
        # Moving an existing module to another key would strand the Voice
        # Access shortcut that was paired to the old one by hand.
        raise VoiceKitError(
            f"'{base}' already answers to Ctrl+Alt+Shift+{key}, and replacing it keeps "
            f"that key so the Voice Access pairing keeps working. To move it to "
            f"Ctrl+Alt+Shift+{wanted}, delete the module and create it again — then "
            f"re-pair the phrase in Voice Access.")
    if not key:
        if wanted:
            if wanted in _used_bridge_keys():
                taken = next((e["phrase"] for e in get_bridge_map()["entries"]
                              if e["combo"].rsplit("+", 1)[-1].upper() == wanted),
                             "a commented-out line in bridge-map.txt")
                free = get_bridge_map()["free_keys"]
                raise VoiceKitError(
                    f"Ctrl+Alt+Shift+{wanted} is already taken (by '{taken}'). "
                    f"Free right now: {', '.join(free) if free else '(none)'}.")
            key = wanted
        else:
            key = allocate_bridge_key()
            if key is None:
                raise VoiceKitError(f"No free bridge keys left (all {len(BRIDGE_POOL)} used). "
                                    "Retire one in bridge-map.txt first.")

    rel = f"hotkeys\\{base}.ahk"
    was_parked = _was_parked(rel)
    # Written, load-checked, and on any failure put back to the exact previous
    # bytes (or removed, if new). The launcher is what the MASTER loads, so its
    # check protects every hotkey. The body runs in its own process, where a
    # mistake costs that process and nothing else — its check is there to TELL
    # the caller now, instead of a load-error dialog on the first key press in
    # a process nobody is watching (and before _Common's error logging exists).
    if isolate:
        body_file.parent.mkdir(parents=True, exist_ok=True)
        snap = _write_validated(
            [(module, _hotkey_launcher_content(phrase, base, key)),
             (body_file, _hotkey_body_content(phrase, base, key, ahk_body))],
            [module, (body_file, "The steps failed to load, so nothing was changed "
                                 "(they run in their own process, so this could never "
                                 "have hurt VoiceKit — this check just tells you now "
                                 "rather than on the first key press):")],
            "The generated module failed to load, so nothing was changed:")
    else:
        snap = _write_validated(
            [(module, _hotkey_content(phrase, key, ahk_body))], [module],
            "The generated module failed to load, so nothing was changed:",
            also=(body_file,))
        # An isolate=False rewrite of a previously isolated module orphans its body.
        if body_file.exists():
            body_file.unlink()

    # Wire it in only where it isn't already — a replace must never append a
    # second #Include or a second bridge-map line.
    listed = [m.lower() for m in index_modules()["active"]]
    unparked = added_index = False
    if was_parked:
        unparked = _uncomment_include(rel)
    elif rel.lower() not in listed:
        _append(_hotkeys_index(), f'\n#Include "%A_ScriptDir%\\{rel}"')
        added_index = True
    if not registered:
        _append(_bridge_map_file(), f"Ctrl+Alt+Shift+{key}|{phrase}|{rel}|{_today()}\n")

    rl = _reload_outcome("the hotkey")
    _undo_if_parked(rel, snap, was_parked, added_index=added_index, added_map=not registered)
    log(f"hotkey | {phrase} | Ctrl+Alt+Shift+{key} | {rel}"
        + (" | isolated" if isolate else "") + (" | replaced" if replacing else ""))
    note = rl["reload_note"]
    if replacing:
        note = (f"Replaced in place, keeping Ctrl+Alt+Shift+{key} — the Voice Access "
                f"pairing still works. ") + note
    if unparked:
        note += (" Its #Include had been commented out (quarantined after a failed "
                 "load) and is switched back on.")
    note += _also_parked_note(rel)
    out = {
        "type": "hotkey_module", "phrase": phrase, "file": str(module),
        "bridge_key": f"Ctrl+Alt+Shift+{key}", **rl,
        "isolated": isolate, "replaced": replacing,
        "voice_pairing": ([] if replacing else [
            'Say: "show voice shortcuts"',
            "Create a new shortcut -> When I say: " + phrase,
            f"Action: Press keys -> Ctrl + Alt + Shift + {key}",
        ]),
        "note": note,
    }
    if replacing:
        prev = snap.prev(body_file)
        out.update(_bank_previous(_MODULE_BACKUPS, base,
                                  prev if prev is not None else snap.prev(module),
                                  "previous_body"))
    if isolate:
        out["body_file"] = str(body_file)
        out["note"] += (f" The steps live in {_body_rel(base)} and run in their own "
                        f"process, so a crash there can't take VoiceKit down; "
                        f"{module.name} is just the key binding.")
    return out


def _splice_once(src: str, old_string: str, new_string: str, read_hint: str) -> str:
    """File-Edit-tool semantics shared by edit_hotkey_module and edit_macro:
    old_string must appear in `src` exactly once; distinct errors for empty /
    identical / missing / ambiguous. `read_hint` names the tool that shows
    the current code, so each error points at the right place to look."""
    if old_string == "":
        raise VoiceKitError(f"old_string can't be empty — {read_hint} shows "
                            f"the current code to pick a splice point from.")
    if old_string == new_string:
        raise VoiceKitError("old_string and new_string are identical — nothing to change.")
    n = src.count(old_string)
    if n == 0:
        raise VoiceKitError(
            f"old_string doesn't appear in the current code — {read_hint} "
            f"shows what's there now (whitespace must match exactly).")
    if n > 1:
        raise VoiceKitError(
            f"old_string appears {n} times — include more surrounding context so "
            f"it matches exactly once.")
    return src.replace(old_string, new_string, 1)


@_serialized
def edit_hotkey_module(name: str, old_string: str, new_string: str) -> dict:
    """Exact-match splice into a hotkey module's code — file-Edit-tool
    semantics: old_string must appear in the current code exactly once.

    This is the cheap iteration path item 2 of the scrape feedback left
    open: most fixes change one function, and a full create_hotkey_module
    replace costs the whole body in tokens. For an isolated module (the
    default kind) the BODY is edited — the same file read_hotkey_module
    returns — and because a body runs fresh on every press, the edit is live
    with NO master reload. An in-process module (isolate=False) is edited in
    its module file and the master reloaded.

    The edited file must still load: it is /validate'd and put back to its
    exact previous bytes on failure, so a typo'd splice can never brick the
    module (or, for an in-process one, the master). An unwanted edit is
    undone by calling again with the two strings swapped — a splice is its
    own inverse, which is why there is no backup file here like the replace
    path keeps."""
    _no_tokens(new_string, "new_string")
    phrase, base, module = _require_module(name, "edited")
    body_file = REPO_ROOT / _body_rel(base)
    isolated = body_file.exists()
    target = body_file if isolated else module
    # Normalize to LF before matching: generated files are LF, but a body the
    # user touched in Notepad is CRLF, and an agent's old_string (written with
    # plain \n) would then never match. AHK reads either; LF is canonical.
    src = _read_text_any(target).replace("\r\n", "\n")
    old_string = old_string.replace("\r\n", "\n")
    new_string = new_string.replace("\r\n", "\n")
    edited = _splice_once(src, old_string, new_string, "read_hotkey_module")
    rel = f"hotkeys\\{base}.ahk"
    was_parked = _was_parked(rel)
    snap = _write_validated(
        [(target, edited)], [target],
        "That edit would stop the module loading, so it was rolled back and nothing "
        "changed:")
    out = {"type": "hotkey_module", "phrase": phrase, "base": base,
           "isolated": isolated, "file": str(target), "edited": True,
           "body_sha256": hashlib.sha256(edited.encode("utf-8")).hexdigest(),
           "body_length": len(edited),
           "undo": "Call edit_hotkey_module again with old_string and new_string swapped."}
    if isolated:
        # A body runs fresh on every press — nothing to reload.
        out["note"] = _isolated_live_note(base)
    else:
        rl = _reload_outcome("the change")
        _undo_if_parked(rel, snap, was_parked)
        out.update(rl)
        out["note"] = ("The module runs inside the resident VoiceKit. " + rl["reload_note"]
                       + _also_parked_note(rel))
    log(f"edit hotkey | {phrase} | {target.name}")
    return out


def _defines_combo(code: str, key: str) -> bool:
    """Whether AHK code DEFINES the Ctrl+Alt+Shift+<key> hotkey — as a
    '^!+k::' line (any modifier order, optional ~ * $ prefixes) or a
    Hotkey("^!+k", ...) call — on a line that isn't a comment. A bare
    substring test was satisfied by Send("^!+b"), a comment mentioning the
    old binding, or ^!+backspace::, and a false pass silently strands the
    hand-made Voice Access pairing this check exists to protect."""
    perms = "|".join(re.escape("".join(p)) for p in itertools.permutations("^!+"))
    k = re.escape(key)
    as_label = re.compile(rf"^\s*[~*$]*(?:{perms}){k}\s*::", re.I)
    as_call = re.compile(rf"\bHotkey\(\s*([\"'])[~*$]*(?:{perms}){k}\1", re.I)
    in_block = False
    for line in code.splitlines():
        s = line.strip()
        if in_block:
            if s.startswith("*/") or s.endswith("*/"):
                in_block = False
            continue
        if s.startswith("/*"):
            in_block = not s.endswith("*/") or len(s) < 4
            continue
        if not s or s.startswith(";"):
            continue
        if as_label.match(line):
            return True
        code_part = re.split(r"\s;", line, maxsplit=1)[0]   # drop a trailing comment
        if as_call.search(code_part):
            return True
    return False


@_serialized
def update_hotkey_module(name: str, code: str) -> dict:
    """Replace a hotkey module's code in full — the write half of the
    read_hotkey_module round trip, for rewrites too broad for
    edit_hotkey_module's single splice.

    It rewrites exactly the file read_hotkey_module returns: the BODY for an
    isolated module (live on the next press, no reload), the module file
    itself for an in-process one (master reloaded). Unlike
    create_hotkey_module — which wraps what you pass in a fresh generated
    header — nothing is re-wrapped here: `code` IS the new file (saved as
    UTF-8 with LF line endings), so a read → tweak → update loop can never
    nest one generated header inside another, and restoring a banked backup
    is one verbatim call.

    The new code is load-checked and the file put back to its exact previous
    bytes on failure; the overwritten file is banked one step deep in
    logs\\module-backups\\ (the same slot a create replace uses) and comes
    back via read_hotkey_module(name, previous=True).

    An in-process module file IS its key binding, so the new code must still
    define the registered combo — dropping ^!+<key> would leave
    bridge-map.txt, press_hotkey and the user's hand-made Voice Access
    pairing pointing at a dead key. Changing a combo stays
    delete-then-recreate, exactly as it is for a create replace."""
    _no_tokens(code, "code")
    phrase, base, module = _require_module(name, "changed")
    if not (code or "").strip():
        raise VoiceKitError("The new code is empty — to remove a module, use "
                            "delete_automation.")

    body_file = REPO_ROOT / _body_rel(base)
    isolated = body_file.exists()
    target = body_file if isolated else module

    key = _module_bridge_key(base)
    if not isolated and key and not _defines_combo(code, key):
        raise VoiceKitError(
            f"Rejected: the new code never defines this module's registered "
            f"combo Ctrl+Alt+Shift+{key} (no '^!+{key}::' hotkey line or "
            f"Hotkey(\"^!+{key}\", ...) call outside a comment), which would "
            f"leave bridge-map.txt, press_hotkey and the user's hand-made "
            f"Voice Access pairing pointing at a dead key. Keep the ^!+{key} "
            f"definition; to change a combo, delete the module and create it "
            f"again, then re-pair the phrase in Voice Access.")

    rel = f"hotkeys\\{base}.ahk"
    was_parked = _was_parked(rel)
    snap = _write_validated(
        [(target, code.replace("\r\n", "\n").strip("\n") + "\n")], [target],
        "The new code failed to load, so the module was left exactly as it was:")

    unparked = False
    rl: dict = {}
    if not isolated:
        # An update that makes a parked module load again is the "I fixed it"
        # signal, same as a create replace — and here the file just proven to
        # load IS the file that was parked, so un-parking is safe.
        if was_parked:
            unparked = _uncomment_include(rel)
        rl = _reload_outcome("the change")
        _undo_if_parked(rel, snap, was_parked)
    banked = _bank_previous(_MODULE_BACKUPS, base, snap.prev(target), "previous_body")
    log(f"hotkey updated | {phrase} | {target.name}")
    if isolated:
        note = _isolated_live_note(base)
    else:
        note = "The module runs inside the resident VoiceKit. " + rl["reload_note"]
    if unparked:
        note += (" Its #Include had been commented out (quarantined after a "
                 "failed load) and is switched back on.")
    if not isolated:
        note += _also_parked_note(rel)
    note += (" read_hotkey_module(name, previous=True) returns the overwritten "
             "version if this needs undoing.")
    return {
        "type": "hotkey_module", "phrase": phrase, "base": base,
        "isolated": isolated, "file": str(target), "updated": True,
        "bridge_key": f"Ctrl+Alt+Shift+{key}" if key else "",
        **banked, **rl,
        "note": note,
    }


# ---------------------------------------------------------------------------
# Macro source access + near-miss identification
#
# 2026-08-05 feedback (the Split Pages debugging run): the #1 gap was that no
# tool could READ a macro's .ahk source — the session cost a human manually
# pasting 250 lines into chat, and the patch that came back had to be
# hand-applied. read_ai_prompt existed for AI actions; this is the same loop
# closed for every macros\ script, read (read_macro_source) and write
# (update_macro full replace with backup; edit_macro splice) both.
# ---------------------------------------------------------------------------
def _macro_backup_file(base: str) -> Path:
    """logs\\macro-backups\\<Base>.prev.ahk — the source an update overwrote.
    One file per macro, one step deep, same shape as the module backups."""
    return _backup_file(_MACRO_BACKUPS, base)


def _macro_kind(base: str, txt: str) -> str:
    """What a macros\\<base>.ahk actually is. The listing used to leave a
    hand-written application distinguishable from a no-code opener only by a
    MISSING field — a signal the feedback's author nearly missed."""
    if base.lower() in _BUILTINS_LOWER:
        return "tool"
    if STUDIO_MARKER in txt:
        return "workflow_stub"
    if _prompt_file(base).exists():
        return "ai_action"
    if _OPENS_RE.search(txt):
        return "opens"
    return "raw_script"


# The no-code opener's header line, which the home window parses too.
_OPENS_RE = re.compile(r"^;\s+Opens:\s+(\S.*?)\s*$", re.MULTILINE)


def _macro_info(base: str) -> tuple[Path, str, str]:
    """(path, text, kind) of macros\\<base>.ahk — THE classifier every
    caller branches on (the listing, the near-miss errors, the update/edit
    guards, create_workflow's overwrite guard and the delete cross-guards
    used to re-derive it five slightly different ways). kind is "" when
    there is no such file."""
    path = _macro_file(base)
    if not path.exists():
        return path, "", ""
    txt = _read_text_any(path)
    return path, txt, _macro_kind(base, txt)


def _macro_description(txt: str, cap: int = 400) -> str:
    """The leading banner-comment block of a hand-written macro, as prose.

    Split Pages carries an excellent self-describing header that the MCP
    used to throw away — surfacing it answers "what even is this thing" at
    the index level. Border lines (all =/-) are dropped, the ';' prefixes
    stripped, and the result capped for listing use."""
    lines: list[str] = []
    for raw in txt.splitlines():
        s = raw.strip()
        if s.startswith(";"):
            body = s.lstrip(";").strip()
            if body and set(body) <= set("=-—─ "):
                continue                      # banner border, not content
            if body:
                lines.append(body)
            continue                          # blank comment line: stay in block
        if lines:
            break              # first non-comment line after the block ends it
        if not s or s.startswith("#"):
            continue           # directives/blanks before the block
        break                  # code before any comment: no header to surface
    desc = "\n".join(lines)
    return desc if len(desc) <= cap else desc[:cap].rstrip() + " …"


def _near_miss(base: str) -> str:
    """What `base` DOES exist as, for a tool that looked for something it
    isn't. The index knows — an error that stops at "no <X> named that" while
    a same-named script sits in macros\\ costs the caller a whole reasoning
    loop rediscovering it (2026-08-05 feedback, item 4). "" when the name
    matches nothing at all."""
    disp = space_out(base)
    hits: list[str] = []
    if _steps_file(base).exists():
        hits.append(f"a workflow (workflows\\{base}.steps.txt) — "
                    f"read_workflow('{disp}') shows its steps")
    _macro, _txt, kind = _macro_info(base)
    if kind:
        if kind == "tool":
            hits.append(f"one of VoiceKit's own tools (macros\\{base}.ahk) — "
                        f"run it with run_automation('{disp}')")
        elif kind == "workflow_stub" and not hits:
            hits.append(f"a generated workflow stub (macros\\{base}.ahk) whose "
                        f"steps file is missing — recreate it with create_workflow")
        elif kind == "ai_action":
            hits.append(f"an AI action — read_ai_prompt('{disp}') has its prompt")
        elif kind == "opens":
            hits.append(f"a no-code launch macro (macros\\{base}.ahk) — "
                        f"read_macro_source('{disp}') has its code")
        elif kind == "raw_script":
            hits.append(f"a hand-written script (macros\\{base}.ahk, no steps "
                        f"file) — read_macro_source('{disp}') returns its source")
    if _module_file(base).exists():
        hits.append(f"a hotkey module — read_hotkey_module('{disp}') has its "
                    f"code, press_hotkey('{disp}') triggers it")
    if not hits:
        return ""
    return f" '{base}' does exist, as " + "; and as ".join(hits) + "."


_MACRO_KIND_NOTES = {
    "tool": ("Part of VoiceKit itself — readable for reference, but update_macro "
             "and edit_macro refuse it."),
    "workflow_stub": ("A GENERATED workflow stub — the steps file is canonical and "
                      "this file is regenerated on every save, so don't edit it: "
                      "read_workflow / create_workflow are the right tools."),
    "ai_action": ("An AI action's generated macro — its behavior lives in the "
                  "prompt file (read_ai_prompt / update_ai_prompt), which is "
                  "usually what you want to change."),
    "opens": ("A no-code 'Open Something' macro. update_macro replaces it; "
              "edit_macro splices one change."),
    "raw_script": ("A hand-written script. update_macro replaces it (backed up, "
                   "load-checked); edit_macro splices one change."),
}


def read_macro_source(name: str, previous: bool = False) -> dict:
    """The full source of macros\\<Base>.ahk, so a debugging session never
    again needs a human to open the file and paste 250 lines into chat.

    previous=True returns the version the last update_macro overwrote (banked
    in logs\\macro-backups\\) — the undo path, mirroring
    read_hotkey_module(previous=True)."""
    phrase, base = _require_name(name)
    macro, txt, kind = _macro_info(base)
    if not kind:
        raise VoiceKitError(f"No macro named '{base}' (looked for "
                            f"macros\\{base}.ahk).{_near_miss(base)}")
    if previous:
        backup, text = _read_backup(
            _MACRO_BACKUPS, base, "a backup is written when update_macro overwrites "
                                  "it, and it hasn't been updated.")
        return {"type": "macro", "phrase": phrase, "base": base, "previous": True,
                "file": str(backup), "source": text,
                "note": "The full file as it stood before the last update_macro."}
    out = {"type": "macro", "phrase": phrase, "base": base,
           "name": space_out(base), "kind": kind, "file": str(macro),
           "source": txt, "note": _MACRO_KIND_NOTES[kind]}
    desc = _macro_description(txt)
    if desc:
        out["description"] = desc
    return out


def _updatable_macro(base: str, verb: str) -> Path:
    """The macros\\ file `verb` may change, with the guards both write paths
    share: builtins are VoiceKit's own machinery, and a workflow stub is
    regenerated on every save so editing it only creates a lie."""
    macro, _txt, kind = _macro_info(base)
    if not kind:
        raise VoiceKitError(f"No macro named '{base}' to {verb} — create a new one "
                            f"with create_launch_macro.{_near_miss(base)}")
    if kind == "tool":
        raise VoiceKitError(f"'{base}' is part of VoiceKit itself and can't be "
                            f"changed this way.")
    if kind == "workflow_stub":
        raise VoiceKitError(
            f"'{base}' is a workflow's GENERATED stub — it is regenerated on every "
            f"save, so an edit here would silently vanish. The steps file is "
            f"canonical: read_workflow('{space_out(base)}') shows it, "
            f"create_workflow replaces it.")
    return macro


def update_macro(name: str, ahk_source: str) -> dict:
    """Replace a macro's source in full — the write half of the read/patch
    loop (an agent used to produce a patch a human then had to hand-apply).

    The new source (saved as UTF-8 with LF line endings) is load-checked and
    the file put back to its exact previous bytes on failure, reusing the
    create-tools' discard rule: a broken write must never be what's left
    behind. The overwritten file is banked one step deep (logs\\macro-backups\\,
    raw bytes) and retrievable via read_macro_source(name, previous=True), so
    an unwanted overwrite stays undoable without the response echoing it."""
    _no_tokens(ahk_source, "ahk_source")
    phrase, base = _require_name(name)
    if not (ahk_source or "").strip():
        raise VoiceKitError("The new source is empty — to remove a macro, use "
                            "delete_automation.")
    macro = _updatable_macro(base, "update")
    prev = _replace_validated(
        macro, ahk_source.replace("\r\n", "\n").strip("\n") + "\n",
        "The new source failed to load, so the macro was left exactly as it was:")
    banked = _bank_previous(_MACRO_BACKUPS, base, prev, "previous_source")
    log(f"macro updated | {base}")
    return {
        "type": "macro", "phrase": phrase, "base": base, "file": str(macro),
        "updated": True, **banked,
        "note": ("Live on its next run — macros run fresh each time, so nothing needs "
                 "reloading. read_macro_source(name, previous=True) returns the "
                 "overwritten version if this needs undoing."),
    }


def edit_macro(name: str, old_string: str, new_string: str) -> dict:
    """Exact-match splice into a macro's source — the patch-shaped tool.
    Same semantics as edit_hotkey_module: old_string must match exactly once,
    the edited file is load-checked and rolled back on failure, and undo is
    the same call with the strings swapped (which is why there's no backup
    file here, unlike update_macro)."""
    _no_tokens(new_string, "new_string")
    phrase, base = _require_name(name)
    macro = _updatable_macro(base, "edit")
    src = _read_text_any(macro).replace("\r\n", "\n")
    edited = _splice_once(src, old_string.replace("\r\n", "\n"),
                          new_string.replace("\r\n", "\n"), "read_macro_source")
    _replace_validated(macro, edited,
                       "That edit would stop the macro loading, so it was rolled back "
                       "and nothing changed:")
    log(f"macro edited | {base}")
    return {
        "type": "macro", "phrase": phrase, "base": base, "file": str(macro),
        "edited": True,
        "source_sha256": hashlib.sha256(edited.encode("utf-8")).hexdigest(),
        "source_length": len(edited),
        "note": "Live on its next run — macros run fresh each time, so nothing needs reloading.",
        "undo": "Call edit_macro again with old_string and new_string swapped.",
    }


# ---------------------------------------------------------------------------
# Logs — the observability the 2026-08-05 feedback found missing: Split Pages
# had written "saved of pageCount | src" on every run, which would have shown
# the varying-failure-point signature of its race BEFORE anyone read code —
# but no tool could read it, so "it keeps crashing" stayed an interview.
# ---------------------------------------------------------------------------
LOG_FILES = {
    "activity": ("created.log",
                 "one line per notable event — automations created/deleted, "
                 "quarantines, and per-run lines tools write via Log() "
                 "(Split Pages: 'split-pages | <saved> of <pages> | <file>')"),
    "errors": ("errors.log",
               "uncaught AutoHotkey errors from the resident VoiceKit AND every "
               "macro/body process — script name, message, file+line (the text "
               "of the error popup, captured by _Common.ahk's OnError)"),
    "runs": ("workflow-runs.log",
             "the workflow engine's per-step run log (run_automation and "
             "run_workflow_batch already attach the relevant tail as 'trace')"),
}


def read_log(log: str = "activity", tail: int = 50, filter: str | None = None) -> dict:
    """The tail of one of VoiceKit's logs, newest last. `filter` (a
    case-insensitive substring) is applied BEFORE the tail, so
    read_log('activity', 20, 'split-pages') is the last 20 Split Pages runs
    however much else happened in between."""
    key = (log or "activity").strip().lower()
    if key not in LOG_FILES:
        raise VoiceKitError(
            f"No log called '{log}'. Available: "
            + "; ".join(f"'{k}' ({v[0]})" for k, v in LOG_FILES.items()) + ".")
    fname, desc = LOG_FILES[key]
    p = REPO_ROOT / "logs" / fname
    tail = max(1, min(_as_int(tail, 50), 500))
    out: dict = {"log": key, "file": str(p), "description": desc}
    if not p.exists():
        out["lines"] = []
        out["note"] = (f"logs\\{fname} doesn't exist yet — nothing has been "
                       f"written to this log on this machine.")
        return out
    lines = [l for l in _read_text_any(p).splitlines() if l.strip()]
    out["total_lines"] = len(lines)
    if filter and str(filter).strip():
        needle = str(filter).strip().lower()
        lines = [l for l in lines if needle in l.lower()]
        out["matched_lines"] = len(lines)
    out["lines"] = lines[-tail:]
    if len(lines) > tail:
        out["note"] = (f"Showing the newest {tail} of {len(lines)} "
                       f"{'matching ' if filter else ''}lines — raise 'tail' "
                       f"(max 500) for more.")
    return out


@_serialized
def create_snippet(abbrev: str, expansion: str) -> dict:
    _no_tokens(expansion, "The expansion")
    abbrev = re.sub(r"\s", "", abbrev.strip())
    # Same rules as lib\_Common.ahk SnippetValidate: ':' ends the abbreviation
    # early and '`' is AHK's escape character (':*:ab`::hi' is "Invalid
    # hotkey") — either one stops Snippets.ahk loading, every snippet with it.
    if not abbrev or ":" in abbrev or "`" in abbrev:
        raise VoiceKitError("The abbreviation can't be empty or contain a colon (:) or a "
                            "backtick (`) — either one breaks the hotstring format and would "
                            "disable every snippet.")
    if expansion.strip() == "{":
        raise VoiceKitError("The expansion can't be just '{' — that's code-block syntax "
                            "in the snippets file (it would break loading entirely).")
    snip = _snippets_file()
    existing = _read_text_any(snip) if snip.exists() else ""
    # Case-insensitive like the GUI (and like AHK hotstring firing itself):
    # a case-differing duplicate would load fine but never fire (shadowed).
    if re.search(rf"^:[^:]*:{re.escape(abbrev)}::", existing, re.MULTILINE | re.IGNORECASE):
        raise VoiceKitError(f"A snippet for '{abbrev}' already exists. Delete it first to change it.")
    exp = snip_encode(expansion)                # one line: ` ; and newlines escaped
    rel = "hotkeys\\Snippets.ahk"
    was_parked = _was_parked(rel)
    # Appended (not rewritten — the rest of the file keeps its bytes), and on
    # a failed load check or any exception put back exactly as it was.
    snap = _FileSnapshot(snip)
    try:
        _append(snip, f"\n:*:{abbrev}::{exp}")
        ok, err = validate_ahk(str(snip))
        if not ok:
            raise VoiceKitError(f"That snippet would break Snippets.ahk, so it was not kept:\n"
                                f"{_explain_ahk_error(err, script=snip)}")
    except BaseException:
        snap.restore()
        raise
    rl = _reload_outcome("the snippet")
    _undo_if_parked(rel, snap, was_parked,
                    hint="An abbreviation some other file already defines is the usual cause.")
    log(f"snippet | {abbrev}")
    return {"type": "snippet", "abbrev": abbrev, "file": str(snip), **rl,
            "note": rl["reload_note"] + _also_parked_note(rel)}


# The step fields that take a path to ACT on (index into type|a|b|c): a run
# target, a focus launch command, a capture command line. Privacy tokens and
# %USERPROFILE% / %TEMP% expand there; a token anywhere else in a step is
# refused (it would be typed or matched literally, and read_workflow returns
# steps verbatim).
_STEP_PATH_FIELD = {"run": 1, "focus": 2, "capture": 2}


def _expand_step_paths(s):
    """A step with its path fields passed through expand_user_paths; privacy
    tokens refused in every other field; everything else untouched."""
    if not s:
        return s
    out = list(s)
    i = _STEP_PATH_FIELD.get(out[0])
    for j in range(1, len(out)):
        if j != i:
            _no_tokens(out[j], f"A '{out[0]}' step")
    if i is not None and len(out) > i and isinstance(out[i], str):
        out[i] = expand_user_paths(out[i])
    return tuple(out) if isinstance(s, tuple) else out


def create_workflow(name: str, steps: list) -> dict:
    phrase, base = _require_name(name)
    if not steps:
        raise VoiceKitError("A workflow needs at least one step.")
    # Path-taking fields accept what a response showed (%USERPROFILE%\...,
    # %TEMP%\...): the engine's Run() expands no variables, so store the real
    # path — exactly what passing the real path would have stored.
    steps = [_expand_step_paths(s) for s in steps]
    depth = 0
    for i, s in enumerate(steps):
        if not s or s[0] not in STEP_TYPES:
            raise VoiceKitError(f"Step {i + 1}: unknown type '{s[0] if s else ''}'. "
                                f"Valid types: {', '.join(STEP_TYPES)}.")
        t = s[0]
        a = str(s[1]).strip() if len(s) >= 2 and s[1] is not None else ""
        b = str(s[2]).strip() if len(s) >= 3 and s[2] is not None else ""
        c = str(s[3]).strip() if len(s) >= 4 and s[3] is not None else ""
        # Every semantic check lives HERE, not in the MCP schema's to_abc: this
        # module is also imported directly (the fallback surface when the MCP
        # connection is down), and a guard that only the server applied let a
        # malformed drag or an empty click through that door.
        if t == "run" and not a:
            raise VoiceKitError(f"Step {i + 1}: a 'run' step needs a target — a program, "
                                f"file, folder or https:// URL.")
        if t in ("focus", "waitwin", "close", "move", "drag",
                 "click", "dblclick", "rclick", "hover") and not a:
            raise VoiceKitError(f"Step {i + 1}: a '{t}' step needs a window (e.g. "
                                f"'ahk_exe notepad.exe' or part of its title).")
        if t in ("text", "keys") and not (len(s) >= 2 and str(s[1] or "")):
            raise VoiceKitError(f"Step {i + 1}: a '{t}' step needs something to "
                                f"{'type' if t == 'text' else 'send'}.")
        if t in ("click", "dblclick", "rclick", "hover"):
            if not b and not c:
                raise VoiceKitError(f"Step {i + 1}: a '{t}' step needs an element name "
                                    f"(preferred) or a window-relative 'x,y' position.")
            if c and not _XY_RE.match(c):
                raise VoiceKitError(f"Step {i + 1}: a '{t}' position must be window-relative "
                                    f"'x,y' in whole pixels (got '{c}').")
        if t == "drag" and not _DRAG_RE.match(c):
            raise VoiceKitError(f"Step {i + 1}: a 'drag' step needs 'x1,y1,x2,y2' — the "
                                f"window-relative press and release points (got '{c}').")
        if t == "move" and b not in MOVE_POSITIONS:
            raise VoiceKitError(f"Step {i + 1}: a 'move' step needs a position: "
                                f"{', '.join(MOVE_POSITIONS)} (got '{b}').")
        # ask|label|suggestion| — a blank label would collapse every input
        # into the engine's "Input" fallback; require one at create time.
        if s[0] == "ask" and not (len(s) >= 2 and str(s[1]).strip()):
            raise VoiceKitError(f"Step {i + 1}: an 'ask' step needs a label (what to ask the user for).")
        # collect|label|element| — same rule: a blank label would collapse
        # every collected value into the engine's "Collected" fallback column.
        if s[0] == "collect" and not (len(s) >= 2 and str(s[1]).strip()):
            raise VoiceKitError(f"Step {i + 1}: a 'collect' step needs a label "
                                f"(the sheet column the value is saved under).")
        # wait|ms| — whole milliseconds, or a RANGE ("600-1400") that pauses a
        # random amount in between. lib\Workflow.ahk WfWaitMs is the engine-side
        # parser; keep this in step with it.
        if s[0] == "wait":
            ms = str(s[1]).strip() if len(s) >= 2 else ""
            if not WAIT_MS_RE.match(ms):
                raise VoiceKitError(f"Step {i + 1}: a 'wait' step needs whole milliseconds "
                                    f"(e.g. 1500) or a range (e.g. 600-1400), got '{ms}'.")
        # set|name|value| — a nameless value can never be referenced, and
        # braces in the name mean the author confused it with a {{reference}}.
        if s[0] == "set":
            nm = str(s[1]).strip() if len(s) >= 2 else ""
            if not nm:
                raise VoiceKitError(f"Step {i + 1}: a 'set' step needs a name "
                                    f"(what later steps write as {{{{name}}}}).")
            if "{" in nm or "}" in nm:
                raise VoiceKitError(f"Step {i + 1}: a 'set' name goes in without braces "
                                    f"— use 'Greeting', then {{{{Greeting}}}} in later steps.")
        # capture|name|command|seconds — same naming rule as set (a name is
        # what later steps write as {{name}}); a command to run; and a
        # timeout that is blank (the engine's 30 s) or a positive number of
        # seconds. lib\Workflow.ahk WfCaptureSecs would quietly read junk as
        # 30, so junk is refused here, where the author can still fix it.
        if s[0] == "capture":
            if not a:
                raise VoiceKitError(f"Step {i + 1}: a 'capture' step needs a name "
                                    f"(what later steps write as {{{{name}}}}).")
            if "{" in a or "}" in a:
                raise VoiceKitError(f"Step {i + 1}: a 'capture' name goes in without braces "
                                    f"— use 'Invoice Data', then {{{{Invoice Data}}}} in later steps.")
            if not b:
                raise VoiceKitError(f"Step {i + 1}: a 'capture' step needs a command to run.")
            if c and not (CAPTURE_SECS_RE.match(c) and float(c) > 0):
                raise VoiceKitError(f"Step {i + 1}: a 'capture' timeout must be a positive "
                                    f"number of seconds, or blank for 30 (got '{c}').")
        # fill|window|label|value — a window and a label to find the box by;
        # the value must be GIVEN ("" clears the box, so a forgotten value
        # must not silently become a clear) and be one line (SendText turns a
        # line break into Enter, which submits a web form). The value is NOT
        # stripped: spaces may matter to the box, and the engine types it as is.
        if t == "fill":
            if not a:
                raise VoiceKitError(f"Step {i + 1}: a 'fill' step needs a window (e.g. "
                                    f"'ahk_exe chrome.exe' or part of its title).")
            lab = fill_label(str(s[2]) if len(s) >= 3 and s[2] is not None else "")
            if lab["err"]:
                raise VoiceKitError(f"Step {i + 1}: a 'fill' step {lab['err']}.")
            if len(s) < 4 or s[3] is None:
                raise VoiceKitError(f"Step {i + 1}: a 'fill' step needs a value — '' clears "
                                    f"the box, so an empty value has to be given on purpose.")
            if "\n" in str(s[3]) or "\r" in str(s[3]):
                raise VoiceKitError(f"Step {i + 1}: a 'fill' value can't contain a line break "
                                    f"(it would press Enter and could submit the form) — use a "
                                    f"'text' step for multi-line text.")
        # if|window|element|condType — reject a blank window or bad condType at
        # create time (the engine would otherwise mis-branch or stop mid-run).
        if s[0] == "if":
            if not (len(s) >= 2 and str(s[1]).strip()):
                raise VoiceKitError(f"Step {i + 1}: an 'if' step needs a window to check.")
            cond = s[3] if len(s) >= 4 else ""
            if cond == "clipboardchanged":
                raise VoiceKitError(f"Step {i + 1}: 'clipboardchanged' only works on a waitfor "
                                    f"step, not an if — 'changed' needs a before-value that an "
                                    f"if has no way to supply.")
            if cond not in IF_CONDS:
                raise VoiceKitError(f"Step {i + 1}: 'if' condition must be one of "
                                    f"{', '.join(IF_CONDS)} (got '{cond}').")
            if cond in CONDS_NEEDING_ELEMENT and not (len(s) >= 3 and str(s[2]).strip()):
                raise VoiceKitError(f"Step {i + 1}: '{cond}' needs something to look for.")
            depth += 1
        # waitfor|window|element|condType[,seconds] — a wrong condType would
        # stop the run, and a missing window would make WinExist("") match
        # whatever the previous step touched.
        elif s[0] == "waitfor":
            raw = s[3] if len(s) >= 4 else ""
            cond, _, secs = str(raw).partition(",")
            cond = cond.strip()
            if cond not in WAIT_CONDS:
                raise VoiceKitError(f"Step {i + 1}: 'waitfor' condition must be one of "
                                    f"{', '.join(WAIT_CONDS)} (got '{cond}').")
            if secs.strip():
                try:
                    if float(secs) <= 0:
                        raise ValueError
                except ValueError:
                    raise VoiceKitError(f"Step {i + 1}: 'waitfor' seconds must be a "
                                        f"positive number (got '{secs.strip()}').")
            if cond != "clipboardchanged" and not (len(s) >= 2 and str(s[1]).strip()):
                raise VoiceKitError(f"Step {i + 1}: a 'waitfor' step needs a window to watch.")
            if cond in CONDS_NEEDING_ELEMENT and not (len(s) >= 3 and str(s[2]).strip()):
                raise VoiceKitError(f"Step {i + 1}: '{cond}' needs something to wait for.")
        elif s[0] == "else":
            if depth == 0:
                raise VoiceKitError(f"Step {i + 1}: 'else' has no matching 'if' above it.")
        elif s[0] == "endif":
            if depth == 0:
                raise VoiceKitError(f"Step {i + 1}: 'endif' has no matching 'if' above it.")
            depth -= 1
    if depth > 0:
        raise VoiceKitError(f"{depth} 'if' step(s) are missing a matching 'endif'.")
    macro, _txt, kind = _macro_info(base)
    if kind == "tool":
        raise VoiceKitError(f"'{space_out(base)}' is one of VoiceKit's own tools, so it "
                            f"can't be a workflow name. Pick another name.")
    if kind and kind != "workflow_stub":
        raise VoiceKitError(f"A hand-written macro named '{base}' already exists. "
                            f"Pick another name.")
    steps_file = _steps_file(base)
    # A replace snapshots both files first: a stub that fails its load check
    # (a lib file mid-edit, say) puts the OLD workflow back byte for byte —
    # this used to delete the very workflow it was replacing.
    replacing = steps_file.exists()
    _write_validated([(steps_file, steps_to_text(phrase, steps)),
                      (macro, workflow_stub(phrase, base))], [macro],
                     "The generated workflow stub failed to load, so nothing was changed:")
    _ensure_voice_dir()
    disp = space_out(base)
    link = VOICE_MACROS / f"{disp}.lnk"
    if not link.exists():
        make_shortcut(str(link), str(macro))
    loop_link = make_loop_shortcut(base, disp)   # companion 'loop <disp>' entry (repeats until stopped)
    log(f"workflow | {phrase} | workflows\\{base}.steps.txt")
    return {"type": "workflow", "phrase": disp, "steps_file": str(steps_file),
            "stub": str(macro), "shortcut": str(link), "loop_shortcut": loop_link,
            "step_count": len(steps), "voice_phrase": f"open {disp}",
            "loop_voice_phrase": f"open loop {disp}", "replaced": replacing,
            "note": NO_RELOAD_NOTE + " Run it by saying \"open " + disp + "\"."}


# ---------------------------------------------------------------------------
# Read / list
# ---------------------------------------------------------------------------
# VoiceKit's own tools: listed separately, never deletable through MCP.
_BUILTINS = {"NewAutomation", "WorkflowStudio", "RecordMySteps", "VoiceKitHelp", "VoiceKitHome", "AskAI"}
# Deletion guard compares case-insensitively: spoken names round-trip through
# CleanPhrase's Title Case ("Ask AI" -> base "AskAi"), and NTFS would happily
# match AskAi.ahk to AskAI.ahk — an exact-case check would not protect it.
_BUILTINS_LOWER = {b.lower() for b in _BUILTINS}
# Shipped tools that are delete-protected but stay EDITABLE through
# update_macro / edit_macro (a hand-written tool someone may well want to fix
# in place — Split Pages is the case that prompted this).
_DELETE_PROTECTED_LOWER = _BUILTINS_LOWER | {"splitpages"}


def _body_status_file(base: str) -> Path:
    """logs\\body-status-<Base>.txt — mirrors lib\\_Common.ahk BodyStatusFile."""
    # Sanitized out here, not inline: a backslash inside an f-string
    # expression is a syntax error before Python 3.12.
    safe = re.sub(r"[^\w\-]", "", base)
    return REPO_ROOT / "logs" / f"body-status-{safe}.txt"


def body_status(base: str) -> dict | None:
    """The latest line a long-running hotkey body published, or None.

    A body runs in its own process and may run for hours; BodyStatus() in
    _Common.ahk lets it say where it got to, and this is the read half. `age`
    is seconds since the line was written — a body that died leaves a line
    that simply stops getting newer, which is the distinction a caller needs.
    """
    f = _body_status_file(base)
    if not f.exists():
        return None
    try:
        raw = _read_text_any(f).strip()
    except OSError:
        return None
    if not raw:
        return None
    stamp, _, text = raw.partition("|")
    out = {"text": text or raw, "when": stamp}
    age = _age_seconds(stamp)
    if age is not None:
        out["age_seconds"] = age
    return out


def _body_stop_file(base: str) -> Path:
    """logs\\body-stop-<Base>.flag — mirrors lib\\_Common.ahk BodyStopFile.
    Same sanitize as _body_status_file: the two sides must derive the same
    name or a graceful stop can never happen."""
    safe = re.sub(r"[^\w\-]", "", base)
    return REPO_ROOT / "logs" / f"body-stop-{safe}.flag"


def _body_processes() -> list:
    """Every running isolated-module body: [{base, pid, started}].

    Matched on OUR interpreter's basename (honors VOICEKIT_AHK) running a
    script under a hotkeys\\bodies\\ folder — the same shape the feedback's
    hand-rolled PowerShell hack matched, but anchored on the path segments
    and the .body.ahk suffix so a coincidental substring can't be caught —
    and on REPO_ROOT, so a same-named body from another VoiceKit tree is
    never listed (or stopped) as this one's."""
    prefix = re.escape(str(REPO_ROOT))
    out = []
    for row in _ahk_processes():
        m = re.search(prefix + r"\\hotkeys\\bodies\\([^\\/:*?\"<>|]+)\.body\.ahk",
                      row["cmd"], re.I)
        if m:
            out.append({"base": m.group(1), "pid": row["pid"], "started": row["started"]})
    return out


# Worked examples shipped as readable code. Not generator templates —
# nothing fills these in; they are copied and edited. They live here rather
# than in a tool description because an MCP-only client (Claude Desktop) has
# no filesystem access, so "it's in templates\" would be no answer at all.
REFERENCES = {
    "web-scrape": (
        "templates/web-scrape-body.ahk",
        "Worked scraper body: search term -> walk results by keyboard -> "
        "capture each in a background tab -> dedupe -> resumable queue. "
        "Copy into hotkeys\\bodies\\<Base>.body.ahk and edit the CONFIGURE ME "
        "block. Distilled from a real harvest of a few hundred listings.",
    ),
}


def read_reference(name: str = "") -> dict:
    """A worked example to copy, or the catalogue when asked for nothing."""
    if not (name or "").strip():
        return {"references": [{"name": k, "description": v[1]}
                               for k, v in REFERENCES.items()],
                "note": "Call again with a name to get the code."}
    key = name.strip().lower()
    if key not in REFERENCES:
        raise VoiceKitError(
            f"No reference called '{name}'. Available: "
            f"{', '.join(sorted(REFERENCES))} (call with no name for the catalogue).")
    rel, desc = REFERENCES[key]
    path = REPO_ROOT / rel
    if not path.exists():
        raise VoiceKitError(f"The '{key}' reference is missing from this install ({rel}).")
    return {"name": key, "description": desc, "file": str(path),
            "code": _read_text_any(path),
            "note": ("Copy this into hotkeys\\bodies\\<Base>.body.ahk (create the "
                     "module first with create_hotkey_module, then edit_hotkey_module "
                     "or a full replace), and change the CONFIGURE ME block. Its "
                     "#Include paths are already right for that folder.")}


def read_module_status(name: str) -> dict:
    """One module's answer to "is it still going, and how far in?".

    The two halves that answer it live apart — the status LINE a body
    published (body_status) and whether its process is alive (list_running) —
    and either alone misleads: a fresh-looking line from a body that died at
    3 a.m. reads as progress, and a running body with no line reads as
    nothing happening. This joins them and says which case it is."""
    _phrase, base, _module = _require_module(name, "checked")
    procs = [p for p in _body_processes() if p["base"].lower() == base.lower()]
    st = body_status(base)
    out = {"name": space_out(base), "base": base, "running": bool(procs),
           "pids": [p["pid"] for p in procs], "status": st}
    if procs:
        out["started_at"] = procs[0]["started"]
        age = _age_seconds(procs[0]["started"])
        if age is not None:
            out["running_seconds"] = age
    # The four states, named — the whole point of joining the two halves.
    if procs and st:
        out["note"] = (f"Running for {out.get('running_seconds', '?')} s; last reported "
                       f"{st.get('age_seconds', '?')} s ago. A status line that stops "
                       f"getting newer while the process lives means it is stuck, not done.")
    elif procs:
        out["note"] = ("Running, but it has never published a status line — add "
                       f'BodyStatus("{base}", "...") to see progress.')
    elif st:
        out["note"] = (f"Not running. Its last line was written {st.get('age_seconds', '?')} s "
                       f"ago — that is where it got to, whether it finished or died.")
    else:
        out["note"] = "Not running, and it has never published a status line."
    return out


def list_running() -> dict:
    """The lifecycle view: which isolated-module bodies are executing NOW.

    press_hotkey starts a body; this shows it running (pid, start time, how
    long, and the latest BodyStatus line it published); stop_module ends it.
    Before this, 'is it still going?' meant hand-rolled Win32_Process
    pattern-matching from a shell."""
    running = []
    for p in _body_processes():
        e = {"name": space_out(p["base"]), "base": p["base"], "pid": p["pid"],
             "started_at": p["started"]}
        age = _age_seconds(p["started"])
        if age is not None:
            e["running_seconds"] = age
        st = body_status(p["base"])
        if st:
            e["last_status"] = st
        running.append(e)
    return {"running": running,
            "note": ("Nothing running. Only isolated module bodies appear here — "
                     "isolate=False modules run inside the resident VoiceKit."
                     if not running else
                     "stop_module stops one; a body that checks BodyStopRequested() "
                     "in its loop exits cleanly, anything else is force-killed "
                     "after the grace period.")}


def stop_module(name: str, grace_s: int = 5) -> dict:
    """Stop a module's running body: cooperative first, force-kill fallback.

    Drops logs\\body-stop-<Base>.flag (which BodyStopRequested() in
    _Common.ahk consumes — generated bodies document the check), waits up to
    grace_s for the process(es) to exit, then TerminateProcess whatever
    remains — scoped to the module's own PIDs, never a name pattern. The
    flag is removed afterwards in every path, so a stop request can never
    leak into the module's next run. A process that is still there at the
    end (one this server may not end, e.g. running elevated) is reported
    with stopped=False and still_running, never as stopped."""
    phrase, base = _require_name(name)
    procs = [p for p in _body_processes() if p["base"].lower() == base.lower()]
    if not procs:
        raise VoiceKitError(
            f"'{base}' has no body process running — list_running shows what does. "
            f"(Only isolated modules run as their own process.)")
    grace_s = max(0, min(_as_int(grace_s, 5), 60))
    pids = [p["pid"] for p in procs]
    flag = _body_stop_file(base)
    how = "killed"
    # Handles are opened BEFORE the flag is written: a body that exits the
    # instant it sees the flag can't have its pid recycled out from under us,
    # and waiting on a handle reports the exit the moment it happens (the old
    # loop spawned a PowerShell every 0.4 s to ask).
    import ctypes
    from ctypes import wintypes
    k = _kernel32()
    handles: list = []           # [(pid, handle)]
    unopened: list = []          # alive, but we may not wait on or end it
    blind: list = []             # refused even a look (access denied)
    for pid in map(int, pids):
        h = k.OpenProcess(_SYNCHRONIZE | _PROCESS_TERMINATE
                          | _PROCESS_QUERY_LIMITED_INFORMATION, False, pid)
        denied = not h and ctypes.get_last_error() == _ERROR_ACCESS_DENIED
        if h:
            handles.append((pid, h))
        elif _pid_image(pid) is not None:
            # A pid that won't open for SYNCHRONIZE|TERMINATE but still
            # answers QUERY_LIMITED (which works across integrity levels) is
            # alive and out of reach — typically a body running elevated.
            # It used to count as "already gone": the flag was deleted before
            # it could see it and the call reported a stop that never happened.
            unopened.append(pid)
        elif denied:
            # Refused even QUERY_LIMITED (a protected process, or one this
            # server may not look at). Not "gone" — a gone pid fails with
            # ERROR_INVALID_PARAMETER, not ACCESS_DENIED: the scan listed it a
            # moment ago, and only the scan can say whether it still runs.
            blind.append(pid)

    def _alive_unopened() -> list:
        return [p for p in unopened if _pid_image(p) is not None]

    def _alive_blind() -> list:
        if not blind:
            return []
        listed = {int(p["pid"]) for p in _body_processes()}
        return [p for p in blind if p in listed]

    def _handle_alive(h) -> bool:
        code = wintypes.DWORD()
        return bool(k.GetExitCodeProcess(h, ctypes.byref(code))) and code.value == _STILL_ACTIVE

    def _wait_all(ms: int) -> bool:
        """Everything exited within `ms`? Handles are waited on; a pid with
        no handle can only be polled, so the wait goes in 200 ms slices."""
        deadline = time.monotonic() + max(0, ms) / 1000
        arr = (wintypes.HANDLE * len(handles))(*[h for _, h in handles]) if handles else None
        while True:
            left = max(0, int((deadline - time.monotonic()) * 1000))
            step = min(200, left) if (unopened or blind) else left
            done = True
            if handles:
                done = k.WaitForMultipleObjects(len(handles), arr, True, step) < len(handles)
            # A blind pid can't be polled cheaply (only the process scan sees
            # it), so it gets the whole wait and is checked once at the end.
            pending = bool(_alive_unopened()) or (bool(blind) and left > 0)
            if done and not pending:
                return not _alive_blind()
            if left == 0:
                return False
            if done:                     # only the unopened ones left: poll
                time.sleep(step / 1000)

    refused: list = []           # TerminateProcess said no
    try:
        flag.parent.mkdir(parents=True, exist_ok=True)
        flag.write_text("stop\n", encoding="utf-8")
        if _wait_all(grace_s * 1000):
            how = "graceful"
        else:
            ended = False
            for pid, h in handles:       # scoped to this module's own processes
                if not _handle_alive(h):
                    continue
                if k.TerminateProcess(h, 1):
                    ended = True
                else:
                    refused.append(pid)
            if ended:                    # nothing was ended: nothing new to wait for
                _wait_all(5000)
        survivors = sorted({pid for pid, h in handles if _handle_alive(h)}
                           | set(_alive_unopened()) | set(_alive_blind()))
    finally:
        for _pid, h in handles:
            k.CloseHandle(h)
        try:
            flag.unlink(missing_ok=True)
        except OSError:
            pass
    unreachable = sorted(set(unopened) | set(blind))
    if survivors:
        out_of_reach = bool(refused) or bool(set(survivors) & set(unreachable))
        why = ("Windows refused to end them" if refused
               else "could not open process (access denied?)" if set(survivors) <= set(blind)
               else "this server isn't allowed to end them" if out_of_reach
               else "they were ended but still hadn't exited 5 s later")
        if out_of_reach:
            why += (" — usually because they run elevated (as administrator) and this "
                    "server doesn't. End them from an elevated Task Manager")
        log(f"stop module | {phrase} | FAILED | still running {survivors} | pids {pids}")
        out = {"stopped": False, "phrase": phrase, "base": base, "pids": pids,
               "still_running": survivors, "how": "failed",
               "note": f"Couldn't stop pid(s) {', '.join(map(str, survivors))}: {why}."}
        if unreachable:
            out["could_not_open"] = unreachable
        return out
    if how == "graceful" and unreachable:
        # Gone — but a process this server couldn't open is one it couldn't
        # watch either: whether it saw the flag or ended some other way can't
        # be told, so it is never reported as a graceful stop.
        how = "exited"
    log(f"stop module | {phrase} | {how} | pids {pids}")
    if how == "graceful":
        note = "It saw the stop flag and exited on its own."
    elif how == "exited":
        note = (f"pid(s) {', '.join(map(str, unreachable))}: could not open process (access "
                f"denied?), so this server couldn't watch it. It is gone now, but whether it "
                f"saw the stop flag can't be told.")
    else:
        note = (f"It didn't stop within {grace_s} s (no BodyStopRequested check, or a "
                f"blocking call), so it was force-killed.")
    out = {"stopped": True, "phrase": phrase, "base": base, "pids": pids, "how": how,
           "note": note}
    if unreachable:
        out["could_not_open"] = unreachable
    return out


def _uia_probe() -> Path:
    """mcp\\uia_probe.ahk — the read-only out-of-process UIA inspector."""
    return REPO_ROOT / "mcp" / "uia_probe.ahk"


def _run_uia_probe(args: list, timeout: int = 30) -> str:
    """Run the probe and return its UTF-8 report text.

    The probe writes to a temp FILE rather than stdout: AutoHotkey is a
    GUI-subsystem exe and its stdout is only trustworthy through a real
    redirect (the same trap AhkValidate works around). An "error=..." line in
    the report is an answer; a nonzero exit or a missing report is a failure.
    """
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "probe.txt"
        r = subprocess.run(
            [AHK_EXE, "/ErrorStdOut", str(_uia_probe()), str(out), *map(str, args)],
            capture_output=True, text=True, timeout=timeout,
        )
        if r.returncode != 0 or not out.exists():
            detail = (r.stdout or r.stderr or "").strip()
            raise VoiceKitError(
                f"UIA probe failed (exit {r.returncode})"
                + (f": {detail}" if detail else ""))
        return out.read_text(encoding="utf-8-sig")


def _parse_kv(text: str) -> dict:
    """key=value lines -> dict (first '=' splits; later ones stay in the value)."""
    rec = {}
    for line in text.splitlines():
        k, sep, v = line.partition("=")
        if sep:
            rec[k] = v
    return rec


def inspect_focus() -> dict:
    """What currently has keyboard focus, read over UI Automation.

    Read-only: no clicks, no keystrokes, no focus changes. The report is the
    element as UIA sees it — name, control type, value, rect — plus up to
    three ancestors and the window it lives in. The name a FOCUSED element
    reports often differs from the text you'd copy off the page (measured on
    listing tiles: copied text said '<title> in <City>, ST', the focused
    name was '<title>, $price, <City>, ST, listing <id>'), which is exactly
    the mismatch this exists to show before a matching predicate is written.
    """
    rec = _parse_kv(_run_uia_probe(["focus"], timeout=20))
    if "error" in rec:
        return {"focused": False, "note": rec["error"]}
    out = {"focused": True,
           "name": rec.get("name", ""),
           "control_type": rec.get("control_type", ""),
           "value": rec.get("value", ""),
           "automation_id": rec.get("automation_id", ""),
           "enabled": rec.get("enabled") == "1"}
    try:
        out["control_type_id"] = int(rec.get("control_type_id", "0"))
    except ValueError:
        out["control_type_id"] = 0
    parts = rec.get("rect", "").split(",")
    if len(parts) == 4 and all(p.lstrip("-").isdigit() for p in parts):
        x, y, w, h = (int(p) for p in parts)
        out["rect"] = {"x": x, "y": y, "w": w, "h": h}
    else:
        out["rect"] = None
    out["ancestry"] = [rec[k] for k in ("ancestor_1", "ancestor_2", "ancestor_3")
                       if k in rec]
    for k in ("window_title", "window_class", "window_exe"):
        if k in rec:
            out[k] = rec[k]
    return out


def dump_uia_tree(window: str | None = None, max_depth: int = 8,
                  name_filter: str | None = None, max_lines: int = 300) -> dict:
    """A window's UIA tree as an indented text outline — 'what am I looking at'.

    Read-only, same probe as inspect_focus. `window` is an AutoHotkey WinTitle
    (empty = the active window); the probe matches titles as substrings. The
    walk is bounded (depth, output lines, visited-node budget), so a browser's
    enormous tree costs a capped report, never a stall.
    """
    max_depth = max(1, min(_as_int(max_depth, 8), 25))
    max_lines = max(1, min(_as_int(max_lines, 300), 2000))
    txt = _run_uia_probe(
        ["tree", window or "", max_depth, max_lines, name_filter or ""])
    header, _, tree = txt.partition("\n\n")
    rec = _parse_kv(header)
    if "error" in rec:
        raise VoiceKitError(rec["error"])
    out = {"window_title": rec.get("window_title", ""),
           "window_class": rec.get("window_class", ""),
           "window_exe": rec.get("window_exe", ""),
           "lines": _as_int(rec.get("lines"), 0),
           "tree": tree.rstrip("\n")}
    if out["lines"] >= max_lines:
        out["note"] = ("Output hit the max_lines cap — raise max_lines, lower "
                       "max_depth, or set name_filter to narrow it.")
    return out


def list_automations() -> dict:
    """Everything sayable/typable, in the same categories the Voice Kit home
    window shows: workflows (+ loop phrase), launch macros (+ what they open),
    AI actions (+ prompt preview), hotkey modules (+ combo), snippets
    (+ expansion), and VoiceKit's own tools. An automation the user gave a
    companion hotkey (the home window's Hotkey button) carries it as
    "hotkey" on its own entry — the companion module isn't listed twice."""
    combos = {e["file"]: e for e in get_bridge_map()["entries"]}
    # Companion hotkeys fold into their automation's entry (mirrors the
    # home window) instead of listing as modules of their own.
    companion_combo = {}
    for file, e in combos.items():
        parent = _companion_parent(file)
        if parent:
            companion_combo[parent.lower()] = e["combo"]

    def _with_hotkey(entry: dict) -> dict:
        combo = companion_combo.get(entry["base"].lower())
        if combo:
            entry["hotkey"] = combo
        return entry

    macros_dir = REPO_ROOT / "macros"
    launch, workflows, ai_actions, tools = [], [], [], []
    for f in sorted(macros_dir.glob("*.ahk")):
        base = f.stem
        disp = space_out(base)
        _path, txt, kind = _macro_info(base)
        if kind == "tool":
            say = "open voice kit" if base.lower() == "voicekithome" else f"open {disp}"
            tools.append(_with_hotkey(
                {"name": disp, "base": base, "voice_phrase": say, "file": str(f)}))
            continue
        if kind == "workflow_stub":
            workflows.append(_with_hotkey(
                {"name": disp, "base": base, "file": str(f),
                 "voice_phrase": f"open {disp}",
                 "loop_voice_phrase": f"open loop {disp}"}))
            continue
        if kind == "ai_action":
            preview = _lf(_read_text_any(_prompt_file(base))).strip()
            ai_actions.append(_with_hotkey(
                {"name": disp, "base": base, "file": str(f),
                 "voice_phrase": f"open {disp}",
                 "prompt_preview": preview[:120]}))
            continue
        # The kind is explicit now. It used to be inferable only from a
        # MISSING 'opens' field, which left a full interactive application
        # (Split Pages) reading as just another one-line opener.
        entry = {"name": disp, "base": base, "file": str(f),
                 "voice_phrase": f"open {disp}", "kind": kind}
        if kind == "opens":
            entry["opens"] = _OPENS_RE.search(txt).group(1).strip()
        else:
            desc = _macro_description(txt)
            if desc:
                entry["description"] = desc
        launch.append(_with_hotkey(entry))

    hotkeys = []
    combos_lower = {k.lower(): e for k, e in combos.items()}
    for f in sorted((REPO_ROOT / "hotkeys").glob("*.ahk")):
        if f.stem.lower() in ("_index", "snippets") or f.stem.lower().endswith(".default"):
            continue    # VoiceKit's own files (and their shipped defaults)
        if f.name.lower().endswith(COMPANION_SUFFIX):
            continue    # a companion — already folded into its automation above
        entry = {"name": space_out(f.stem), "base": f.stem, "file": str(f)}
        bm = combos_lower.get(f"hotkeys\\{f.stem}.ahk".lower())
        if bm:
            entry["combo"] = bm["combo"]
            entry["voice_phrase"] = bm["phrase"]
        st = body_status(f.stem)
        if st:
            entry["status"] = st        # what a long-running body last reported
        hotkeys.append(entry)

    snippets = []
    snip = _snippets_file()
    if snip.exists():
        for line in _read_text_any(snip).splitlines():
            m = re.match(r"^:\*?[^:]*:(.+?)::(.*)$", line.strip())
            if m:
                exp = m.group(2).strip()
                snippets.append({"abbrev": m.group(1),
                                 "dynamic": exp in ("{", ""),
                                 "expansion_preview": "" if exp in ("{", "") else exp[:80]})

    return {"workflows": workflows, "launch_macros": launch, "ai_actions": ai_actions,
            "hotkey_modules": hotkeys, "snippets": snippets, "voicekit_tools": tools}


def read_workflow(name: str) -> dict:
    base = _workflow_base(name)
    steps_file = _steps_file(base)
    if not steps_file.exists():
        raise VoiceKitError(f"No workflow named '{base}' (looked for "
                            f"{steps_file.name}).{_near_miss(base)}")
    steps = []
    for line in _read_text_any(steps_file).splitlines():
        line = line.strip()
        if not line or line.startswith(";"):
            continue
        parts = line.split("|")
        while len(parts) < 4:
            parts.append("")
        steps.append({"type": parts[0],
                      "a": wf_decode(parts[1]),
                      "b": wf_decode(parts[2]),
                      "c": wf_decode(parts[3])})
    return {"base": base, "name": space_out(base), "steps": steps}


def get_bridge_map() -> dict:
    mapfile = _bridge_map_file()
    entries = []
    if mapfile.exists():
        for line in _read_text_any(mapfile).splitlines():
            line = line.strip()
            if not line or line.startswith(";"):
                continue
            # Trim each field like the GUI does (VoiceKitHome.ahk) — the file is
            # the user's hand-maintained recreate list, so padded fields happen.
            parts = [p.strip() for p in line.split("|")]
            if len(parts) >= 3:
                entries.append({"combo": parts[0], "phrase": parts[1], "file": parts[2],
                                "created": parts[3] if len(parts) > 3 else ""})
    # Free = the same rule the allocator and the GUI's key dropdown use, so a
    # commented-out registration still holds its key (see _used_bridge_keys).
    used = _used_bridge_keys()
    free = [k for k in BRIDGE_POOL if k not in used]
    return {"entries": entries, "free_keys": free, "reserved_keys": list(RESERVED_BRIDGE_KEYS)}


# ---------------------------------------------------------------------------
# Run / trigger / delete
# ---------------------------------------------------------------------------
def _run_args(args) -> list[str]:
    """run_automation's `args`, checked: a list of plain strings (each becomes
    one argv entry — no shell is involved, so spaces, & and quotes arrive
    exactly as given). Line breaks and NULs are refused: a Windows command
    line can't carry them faithfully."""
    if args is None:
        return []
    if isinstance(args, str) or not isinstance(args, (list, tuple)):
        raise VoiceKitError("'args' is a list of strings, one per argument — e.g. "
                            "[\"%TEMP%\\\\sample.pdf\"].")
    out = []
    for a in args:
        if not isinstance(a, str):
            raise VoiceKitError(f"Every argument must be a string (got {type(a).__name__}).")
        if any(ch in a for ch in "\r\n\0"):
            raise VoiceKitError("An argument can't contain a line break.")
        out.append(a)
    return out


def _workflow_file_args(args: list[str]) -> list[str]:
    """A workflow's arguments stand in for the File Explorer selection
    ({{selected_file}} / {{selected_files}}), so each must be an existing file
    or folder. Returned as full paths (%USERPROFILE% / %TEMP% forms expanded)."""
    out = []
    for a in args:
        p = expand_user_paths(a.strip())
        path = Path(p) if p else None
        if not path or not path.is_absolute() or not (path.is_file() or path.is_dir()):
            raise VoiceKitError(
                f"'{a}' isn't an existing file or folder (give a full path). A workflow's "
                f"arguments are files standing in for the File Explorer selection — its "
                f"{{{{selected_file}}}} / {{{{selected_files}}}} values. Try it on a fixture: "
                f"create a sample file under %TEMP% (run_ahk_snippet can write one) rather "
                f"than pointing at a real client document.")
        out.append(str(path.resolve()))
    return out


# Fields where the engine resolves {{Name}} (lib\Workflow.ahk WfStepSubstText):
# "a" = paramA, "b" = paramB (read_workflow's keys).
_SUBST_FIELDS = {**{t: ("a",) for t in ("text", "keys", "run", "waitwin", "move", "close", "drag")},
                 **{t: ("a", "b") for t in ("focus", "click", "dblclick", "rclick", "hover",
                                             "if", "waitfor")},
                 **{t: ("b",) for t in ("collect", "set", "capture")},
                 "fill": ("a", "b", "c")}     # fill's value is the one substituted paramC
_VAR_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def _workflow_uses_selected_file(base: str) -> bool:
    """True when a workflow references the SINGULAR {{selected_file}} and no
    step of its own defines that name (mirrors WfSelectionNeeds = 2). Such a
    run refuses to start unless exactly one file is selected — with a modal
    popup on the user's desktop — so run_automation checks the args count
    first and refuses instead."""
    steps = read_workflow(base)["steps"]
    defined = set()
    for s in steps:
        if s["type"] in ("ask", "collect", "set", "capture"):
            defined.add(s["a"].strip().lower())
    for s in steps:
        for f in _SUBST_FIELDS.get(s["type"], ()):
            for m in _VAR_RE.finditer(s[f]):
                nm = m.group(1).lower()
                if nm == "selected_file" and nm not in defined:
                    return True
    return False


def run_automation(name: str, wait_seconds: int = 0, args: list | None = None) -> dict:
    """Trigger any spoken automation (workflow, launch macro, AI action, or
    VoiceKit tool) by launching its macro — the same thing Voice Access does
    when the user says 'open <name>', including the working directory (the
    macro's own folder, like its Start Menu shortcut). A 'loop <name>' phrase
    starts a workflow's loop companion (repeats until stopped — the floating
    Stop Looping button, Ctrl+Alt+Shift+X, or a failing step). With
    wait_seconds > 0, waits that long for the script to finish and reports how
    it went; otherwise fire-and-forget.

    `args` are passed to the script as its command-line arguments (A_Args), as
    an argv list — no shell. For a workflow they must be existing files or
    folders: they become its selection ({{selected_file}} / {{selected_files}})
    instead of whatever File Explorer has selected. A raw script decides what
    its arguments mean (a %USERPROFILE% / %TEMP% form or a privacy token in one
    arrives as the real path — WP9). Refused for Workflow Studio (its argument means 'load
    this file' / '/record') and for 'loop <name>' phrases."""
    phrase, base = _require_name(name)
    argv = _run_args(args)
    macro = _macro_file(base)
    if macro.exists():
        if argv and base.lower() == "workflowstudio":
            raise VoiceKitError("Workflow Studio doesn't take arguments here — its argument "
                                "means 'load this steps file' or '/record'. Run it without args.")
        # Workflow Studio is #SingleInstance Force: relaunching it would kill an
        # open session and lose unsaved recorded steps (the GUI's OpenStudioSafely
        # guards this; do the same headlessly by refusing rather than clobbering).
        # Deliberately NOT anchored on REPO_ROOT: here a false "running" only
        # refuses a launch, while a false "not running" destroys a session.
        if base.lower() == "workflowstudio" and _ahk_script_running("WorkflowStudio.ahk",
                                                                    anchored=False):
            raise VoiceKitError("Workflow Studio is already open — launching it again would discard "
                                "any unsaved recorded steps. Ask the user to save or close it first.")
        is_workflow = _steps_file(base).exists()
        if argv and not is_workflow:
            # What a masked response showed (a privacy token, %USERPROFILE%)
            # is what an agent passes back; the script gets the real path.
            argv = [expand_user_paths(a) for a in argv]
        if argv and is_workflow:
            argv = _workflow_file_args(argv)
            # The engine would refuse too — but with a modal popup on the
            # user's desktop that nobody here can dismiss. Refuse first.
            if len(argv) != 1 and _workflow_uses_selected_file(base):
                raise VoiceKitError(
                    f"This workflow uses {{{{selected_file}}}} — exactly one file — but "
                    f"{len(argv)} arguments were given. Pass one file, or use "
                    f"{{{{selected_files}}}} in the workflow to take several.")
        since = datetime.now().strftime("%Y%m%d%H%M%S")
        # cwd = the macro's folder: what its Start Menu shortcut sets, so a
        # script using relative paths behaves the same here as by voice.
        r = _launch_and_report([AHK_EXE, str(macro), *argv], space_out(base), str(macro),
                               wait_seconds, cwd=str(macro.parent))
        if argv:
            r["args"] = argv
        # A workflow's generated stub always exits 0, so its exit code says
        # nothing — but the engine records what the run did. Non-workflow
        # macros (AI actions, tools) leave no record and report as before.
        if is_workflow:
            r = _attach_outcome(r, base, since)
        return r
    # 'loop <name>': the workflow's loop companion (same target as its
    # 'loop <name>.lnk' — lib\LoopRunner.ahk resolves the steps file itself).
    if phrase.lower().startswith("loop ") and phrase[5:].strip():
        wf_base = to_base(phrase[5:])
        if _steps_file(wf_base).exists():
            if argv:
                raise VoiceKitError("A loop doesn't take arguments (LoopRunner's own arguments "
                                    "are the workflow and a batch file). Run the workflow once "
                                    "with args, or use run_workflow_batch for rows of inputs.")
            _refuse_if_loop_running()
            since = datetime.now().strftime("%Y%m%d%H%M%S")
            r = _launch_and_report(
                [AHK_EXE, str(REPO_ROOT / "lib" / "LoopRunner.ahk"), wf_base],
                f"loop {space_out(wf_base)}", str(REPO_ROOT / "lib" / "LoopRunner.ahk"),
                wait_seconds, cwd=str(REPO_ROOT))
            if not r.get("finished"):
                r["note"] = ("Looping until stopped — the floating Stop Looping button, "
                             "saying 'click stop looping', or Ctrl+Alt+Shift+X ends it; "
                             "it also stops itself if a step fails.")
            return _attach_outcome(r, wf_base, since, exit_codes=True)
        raise VoiceKitError(f"No workflow named '{wf_base}' to loop — "
                            f"list_automations shows what exists.{_near_miss(wf_base)}")
    raise VoiceKitError(f"No automation named '{base}' to run — "
                        f"list_automations shows what exists.{_near_miss(base)}")


# Hard cap on wait_seconds. Claude Desktop cancels a tool call at ~240 s;
# a blocking wait that outlives the cancel answers a dead request, which
# crashes the MCP session ("Request already responded to") and disconnects
# the server. Stay well under the cancel window — for longer runs, callers
# should fire-and-forget and check back (read_workflow_sheet / list state).
WAIT_CAP_SECONDS = 120


def _launch_and_report(cmd: list, launched: str, file: str, wait_seconds: int,
                       cwd: str | None = None) -> dict:
    # cwd: what the voice shortcut would use (a macro's own folder; the repo
    # root for a loop) — without it a launch ran in the MCP server's cwd.
    proc = subprocess.Popen(cmd, cwd=cwd)
    result = {"launched": launched, "file": file}
    capped = min(int(wait_seconds), WAIT_CAP_SECONDS)
    if capped < wait_seconds:
        result["wait_capped"] = (f"wait_seconds {wait_seconds} was capped to {WAIT_CAP_SECONDS} "
                                 f"(longer blocking waits get cancelled by the client and can "
                                 f"wedge the connection — poll instead).")
    wait_seconds = capped
    if wait_seconds > 0:
        try:
            rc = proc.wait(timeout=wait_seconds)
            result["finished"] = True
            result["exit_code"] = rc
            result["note"] = ("Finished cleanly." if rc == 0 else
                              f"Exited with code {rc} — something in it failed.")
        except subprocess.TimeoutExpired:
            result["finished"] = False
            result["note"] = (f"Still running after {wait_seconds}s — a long automation, "
                              "a window it's waiting for, or a failing step showing its popup.")
    else:
        result["finished"] = False
        result["note"] = "Running on the desktop now; a failing step shows a popup naming it."
    return result


# ---------------------------------------------------------------------------
# What a run actually did
#
# The exit code used to be the only signal, and lib\LoopRunner.ahk exited 0
# unconditionally — so every batch reported "Finished cleanly", including the
# ones that died at step 7 and collected nothing. The engine now records each
# run (logs\workflow-runs.ini, plus a step-by-step logs\workflow-runs.log) and
# these read it back.
# ---------------------------------------------------------------------------
_RUN_OUTCOMES = {
    "ok":        "Every step ran.",
    "failed":    "A step failed and the run stopped there.",
    "stopped":   "Stopped early — the Stop Looping button, Ctrl+Alt+Shift+X, or a cancelled dialog.",
    "cancelled": "Cancelled before any step ran.",
    "error":     "It never got as far as running.",
    "running":   "Still running.",
}

# lib\LoopRunner.ahk's exit codes — the fallback when no record was written.
_LOOP_EXIT = {
    0: ("ok", "Every planned pass finished."),
    1: ("failed", "A step failed and the run stopped there."),
    2: ("stopped", "Stopped early — the Stop Looping button, Ctrl+Alt+Shift+X, or a cancelled dialog."),
    3: ("error", "It never got as far as running."),
}


def _as_int(v, default: int = 0) -> int:
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return default


def read_run_record(base: str) -> dict:
    """A workflow's last recorded run, from logs\\workflow-runs.ini (written by
    lib\\Workflow.ahk — UTF-16LE with a BOM, like everything AutoHotkey's
    IniWrite creates). {} when there is no record for it."""
    if not _runs_ini().exists():
        return {}
    import configparser
    cp = configparser.RawConfigParser()   # Raw: a reason string may hold a '%'
    try:
        cp.read_string(_read_text_any(_runs_ini()))
    except Exception:
        return {}
    for section in cp.sections():
        if section.lower() == base.lower():
            return dict(cp.items(section))
    return {}


def _run_trace(base: str, limit: int = 60) -> list[str]:
    """The tail of logs\\workflow-runs.log covering this workflow's last run —
    from its final '---- run started' marker on. This is the per-step evidence
    that otherwise has to be reconstructed from clipboard contents and window
    titles after the fact."""
    if not _runs_log().exists():
        return []
    try:
        lines = [l for l in _read_text_any(_runs_log()).splitlines() if f"  {base}  " in l]
    except Exception:
        return []
    start = 0
    for i, line in enumerate(lines):
        if "---- run started" in line:
            start = i
    return lines[start:][-limit:]


def _run_report(base: str, since: str = "") -> dict:
    """The last run of `base`, as a reportable block.

    `since` is an AutoHotkey YYYYMMDDHH24MISS stamp taken just before launching:
    a record older than that belongs to an EARLIER run, and reporting a stale
    outcome as this run's would be worse than reporting none at all."""
    rec = read_run_record(base)
    if not rec:
        return {}
    if since and rec.get("started", "") < since:
        return {}
    outcome = rec.get("outcome", "")
    out = {
        "outcome": outcome,
        "ok": outcome == "ok",
        "what_happened": _RUN_OUTCOMES.get(outcome, outcome),
        "steps_total": _as_int(rec.get("steps_total")),
        "started": rec.get("loop_started_text") or rec.get("started_text", ""),
        "log": str(_runs_log()),
    }
    if _as_int(rec.get("failed_step")):
        out["failed_step"] = _as_int(rec.get("failed_step"))
        out["step"] = rec.get("step", "")
    if rec.get("reason"):
        out["reason"] = rec["reason"]
    if _as_int(rec.get("passes_total")) or _as_int(rec.get("passes_done")):
        out["passes_done"] = _as_int(rec.get("passes_done"))
        out["passes_total"] = _as_int(rec.get("passes_total"))
    if outcome != "ok":
        trace = _run_trace(base)
        if trace:
            out["trace"] = trace
    return out


def _run_note(run: dict) -> str:
    """One sentence saying how a run went, for the tool response's note."""
    bits = [run["what_happened"]]
    if run.get("failed_step"):
        bits.append(f"Step {run['failed_step']} of {run['steps_total']}: "
                    f"{run.get('step', '')} — {run.get('reason', '')}")
    elif run.get("reason"):
        bits.append(run["reason"])
    if "passes_total" in run:
        bits.append(f"{run['passes_done']} of {run['passes_total']} pass(es) completed.")
    return " ".join(b for b in bits if b)


def _attach_outcome(r: dict, base: str, since: str, exit_codes: bool = False) -> dict:
    """Replace exit-code guesswork with what the run recorded about itself.
    `exit_codes` says the launched process reports one meaningfully (only
    lib\\LoopRunner.ahk does — a generated workflow stub always exits 0)."""
    run = _run_report(base, since)
    if run and run.get("outcome") != "running":
        r["run"] = run
        r["ok"] = run["ok"]
        r["note"] = _run_note(run)
        return r
    if exit_codes and r.get("finished") and isinstance(r.get("exit_code"), int):
        outcome, note = _LOOP_EXIT.get(r["exit_code"], ("error", f"Exited with code {r['exit_code']}."))
        r["ok"] = outcome == "ok"
        r["outcome"] = outcome
        r["note"] = f"{note} (No run record was written — check {_runs_log()}.)"
    elif not r.get("finished"):
        r["note"] = (r.get("note", "") + f" Nothing recorded yet; when it ends, "
                     f"its outcome lands in {_runs_ini()}.").strip()
    return r


def _workflow_labels(base: str, step_type: str) -> list[str]:
    """Unique labels of a workflow's ask or collect steps, in step order
    (mirrors lib\\Workflow.ahk WfAskLabels / WfCollectLabels, including the
    blank-label fallbacks)."""
    steps = read_workflow(base)["steps"]
    if step_type == "ask":
        return _input_labels(steps)
    labels: list[str] = []
    seen: set[str] = set()
    for s in steps:
        if s["type"] != step_type:
            continue
        lab = s["a"].strip() or ("Input" if step_type == "ask" else "Collected")
        if lab.lower() not in seen:
            seen.add(lab.lower())
            labels.append(lab)
    return labels


_BUILTIN_VALUE_NAMES = {"clipboard", "date", "time", "datetime", "selected_file", "selected_files"}


def _input_labels(steps: list[dict]) -> list[str]:
    """A workflow's INPUTS in step order (mirrors lib\\Workflow.ahk WfAskLabels):
    every unique ask label, then every {{Name}} a fill VALUE uses that no ask /
    collect / set / capture step defines and no built-in covers
    (WfFillInputNames) — a batch column that a fill step types into its box,
    with no ask step typing it anywhere else. Steps are read_workflow dicts."""
    labels: list[str] = []
    seen: set[str] = set()
    for s in steps:
        if s["type"] == "ask":
            lab = s["a"].strip(" \t") or "Input"
            if lab.lower() not in seen:
                seen.add(lab.lower())
                labels.append(lab)
    defined = set()
    for s in steps:
        if s["type"] in ("ask", "collect", "set", "capture"):
            nm = s["a"].strip(" \t") or {"ask": "Input", "collect": "Collected"}.get(s["type"], "")
            if nm:
                defined.add(nm.lower())
    for s in steps:
        if s["type"] != "fill":
            continue
        for m in _VAR_RE.finditer(s.get("c", "")):
            nm = m.group(1)
            k = nm.lower()
            if k in defined or k in seen or k in _BUILTIN_VALUE_NAMES:
                continue
            seen.add(k)
            labels.append(nm)
    return labels


# ---------------------------------------------------------------------------
# Batch rows from a CSV or Excel range (roadmap WP7)
#
# Everything happens HERE, in Python: any source (a CSV in any of the
# encodings Excel saves, or a range of an .xlsx tab) is normalized into the
# same logs\mcp-batch-<stamp>-<pid>.csv that lib\LoopRunner.ahk already reads
# through WfLoopCsvRows — so the AHK side, its CSV dialect and every on-disk
# format stay byte-for-byte unchanged (no lockstep). Python also does ALL the
# row skipping, and writes a metadata column holding each row's source row
# number: a record with a non-blank cell is never dropped by WfLoopCsvRows, so
# pass N is always the Nth row written, and that row names where it came from.
#
# Cell VALUES are client data (invoice amounts, names): nothing here puts one in
# a response or an error, except the dry_run preview the caller asks for.
# ---------------------------------------------------------------------------
BATCH_CSV_EXTS = (".csv", ".tsv", ".txt")
BATCH_XLSX_EXTS = (".xlsx", ".xlsm")
# The extra column every source-mode batch file carries (WfLoopCsvRows ignores
# columns that aren't ask labels). A label with this name is refused.
BATCH_META_COLUMN = "VoiceKit source row"
# What Excel shows (and saves into a CSV) for a formula that failed. Typing
# "#N/A" into a submitted form is worse than stopping, so these are refused.
EXCEL_ERRORS = frozenset({
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
    "#GETTING_DATA", "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#CONNECT!",
    "#BUSY!", "#UNKNOWN!",
})
_CELL_RANGE_RE = re.compile(
    r"^\$?([A-Z]{1,3})?\$?(\d+)?(?::\$?([A-Z]{1,3})?\$?(\d+)?)?$", re.IGNORECASE)
_MAX_COL = 16384                       # XFD — Excel's last column


def _col_num(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _col_letter(n: int) -> str:
    s = ""
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _parse_cell_range(text: str | None) -> tuple:
    """(sheet or "", first col, first row, last col or None, last row or None).
    'A2:C40', 'A:C' (every row), 'A2:C' (open-ended: to the last used row),
    optionally prefixed 'Sheet!' / "'Sheet name'!". Empty = the whole used
    area. None for a last col/row means "as far as the data goes"."""
    text = (text or "").strip()
    sheet = ""
    if "!" in text:
        sheet, text = text.rsplit("!", 1)
        sheet = sheet.strip()
        if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
            sheet = sheet[1:-1].replace("''", "'")
        text = text.strip()
    if not text:
        return sheet, 1, 1, None, None
    m = _CELL_RANGE_RE.match(text)
    if not m or ":" not in text or not (m.group(1) and m.group(3)):
        raise VoiceKitError(
            f"cell_range '{text}' isn't a range VoiceKit reads. Use two corners with column "
            f"letters: 'A2:C40', 'A2:C' (from row 2 to the last used row) or 'A:C' (every row).")
    c1, c2 = _col_num(m.group(1)), _col_num(m.group(3))
    r1 = int(m.group(2)) if m.group(2) else 1
    r2 = int(m.group(4)) if m.group(4) else None
    if r1 < 1 or (r2 is not None and r2 < 1):
        raise VoiceKitError(f"cell_range '{text}': rows start at 1.")
    if max(c1, c2) > _MAX_COL:
        raise VoiceKitError(f"cell_range '{text}' goes past Excel's last column (XFD).")
    c1, c2 = min(c1, c2), max(c1, c2)          # 'C2:A40' means what it obviously means
    if r2 is not None and r2 < r1:
        r1, r2 = r2, r1
    return sheet, c1, r1, c2, r2


def _range_text(c1: int, r1: int, c2: int, r2: int | None) -> str:
    return f"{_col_letter(c1)}{r1}:{_col_letter(c2)}{'' if r2 is None else r2}"


def _is_excel_error(text: str) -> bool:
    return text.strip().upper() in EXCEL_ERRORS


def _num_text(v: float) -> str:
    """A float as a person would type it: integers without '.0', no exponent,
    and at Excel's own precision (15 significant digits) so float noise
    (1234.4999999999998, 0.30000000000000004) reads as Excel displays it."""
    if v != v or v in (float("inf"), float("-inf")):
        return repr(v)
    if v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    from decimal import Decimal
    s = format(Decimal(f"{v:.15g}"), "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return "0" if s in ("-0", "") else s


def _xlsx_cell_text(cell, date_format: str) -> tuple[str, bool]:
    """(text, is_error) for one openpyxl cell (read_only, data_only)."""
    from datetime import date as _date, time as _time, timedelta as _td
    v = getattr(cell, "value", None)
    if v is None:
        return "", False
    if getattr(cell, "data_type", "") == "e" or (isinstance(v, str) and _is_excel_error(v)):
        return str(v).strip(), True
    if isinstance(v, bool):
        return ("TRUE" if v else "FALSE"), False
    if isinstance(v, datetime):
        if date_format:
            return v.strftime(date_format), False
        return v.strftime("%Y-%m-%d" if v.time() == _time(0) else "%Y-%m-%d %H:%M:%S"), False
    if isinstance(v, _date):
        return v.strftime(date_format or "%Y-%m-%d"), False
    if isinstance(v, _time):
        return v.strftime("%H:%M:%S"), False
    if isinstance(v, _td):
        return str(v), False
    if isinstance(v, (int, float)):
        s = str(v) if isinstance(v, int) else _num_text(v)
        # An ID or ZIP formatted "00000" keeps the zeros Excel shows — the one
        # number format that changes WHAT the value is, not how it looks.
        fmt = str(getattr(cell, "number_format", "") or "")
        if re.fullmatch(r"0+", fmt) and re.fullmatch(r"\d+", s):
            s = s.zfill(len(fmt))
        return s, False
    return str(v), False


def _openpyxl():
    try:
        import openpyxl
    except ImportError:
        raise VoiceKitError(
            "Reading an .xlsx needs the openpyxl package, which this MCP server's Python "
            "doesn't have. Install it into mcp\\.venv — pip install openpyxl, i.e. "
            f"\"{REPO_ROOT / 'mcp' / '.venv' / 'Scripts' / 'python.exe'}\" -m pip install "
            "openpyxl (or re-run mcp\\Setup-MCP.bat) — then reconnect the server. A CSV "
            "works without it: in Excel, File > Save As > 'CSV UTF-8'.") from None
    return openpyxl


def _xlsx_grid(path: Path, sheet: str, c1: int, r1: int, c2, r2, date_format: str) -> dict:
    openpyxl = _openpyxl()
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise VoiceKitError(f"Couldn't read {path.name} — another program has it locked. "
                            f"Close it (or save a copy) and try again.") from None
    except Exception as e:  # noqa: BLE001 — a corrupt/encrypted workbook, a renamed .xls
        raise VoiceKitError(f"Couldn't open {path.name} as an Excel workbook "
                            f"({type(e).__name__}). Only .xlsx / .xlsm are read; save an "
                            f"old .xls as .xlsx first.") from None
    try:
        names = list(wb.sheetnames)
        if sheet:
            hit = next((n for n in names if n.lower() == sheet.strip().lower()), None)
            if hit is None:
                raise VoiceKitError(f"{path.name} has no tab named '{sheet}'. "
                                    f"Its tabs: {', '.join(names)}.")
            ws = wb[hit]
        else:
            ws = wb.active
            if ws is None or not hasattr(ws, "iter_rows"):
                ws = wb[names[0]]
        if not hasattr(ws, "iter_rows"):
            raise VoiceKitError(f"'{ws.title}' in {path.name} is a chart, not a sheet of cells.")
        if r2 is None or c2 is None:
            # Read-only mode trusts the file's recorded <dimension>, and some
            # writers record just "A1" — an open-ended range would then stop
            # at row/column 1. Forget it and read to where the rows really end.
            try:
                ws.reset_dimensions()
            except AttributeError:
                pass
        rows = []
        for i, row in enumerate(ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2)):
            rows.append((r1 + i, [_xlsx_cell_text(c, date_format) for c in row]))
        return {"kind": "xlsx", "sheet": ws.title, "rows": rows}
    finally:
        wb.close()


def _sniff_delimiter(text: str, suffix: str) -> str:
    """Conservative: .tsv, or a tab in the deciding line, is tab-separated;
    else a comma there means comma; semicolons and no comma is European-locale
    Excel's ';'. Quoted parts are ignored."""
    if suffix == ".tsv":
        return "\t"
    # The FIRST LINE THAT HAS ANY of them decides; lines with none are
    # skipped, so a title line ('Invoice detail') above a ';' table can't make it
    # read as one column. No delimiter anywhere = one column (comma).
    for line in text.split("\n", 50)[:50]:
        line = re.sub(r'"[^"]*"', "", line)
        if "\t" in line:
            return "\t"
        if "," in line:
            return ","
        if ";" in line:
            return ";"
    return ","


def _csv_decode(raw: bytes) -> tuple[str, str]:
    """(text, encoding). UTF-16 by BOM (Excel's 'Unicode Text'), UTF-8 with
    or without BOM, else Windows-1252 — what Excel's plain 'CSV (Comma
    delimited)' writes on a US machine."""
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16"), "utf-16"
    try:
        return raw.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        return raw.decode("cp1252"), "cp1252"
    except UnicodeDecodeError:
        # A byte Windows-1252 leaves undefined: not text in either encoding.
        # Refuse rather than type a replacement character into a return.
        raise VoiceKitError("That file isn't UTF-8 or Windows-1252 text. In Excel: File > "
                            "Save As > 'CSV UTF-8 (Comma delimited)', then try again.") from None


def _csv_grid(path: Path, c1: int, r1: int, c2, r2) -> dict:
    try:
        raw = path.read_bytes()
    except PermissionError:
        raise VoiceKitError(f"Couldn't read {path.name} — another program has it locked. "
                            f"Close it and try again.") from None
    text, enc = _csv_decode(raw)
    delim = _sniff_delimiter(text, path.suffix.lower())
    recs = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    rows = []
    for n, rec in enumerate(recs, start=1):     # row N = CSV record N (a quoted
        if n < r1 or (r2 is not None and n > r2):   # newline doesn't start one)
            continue
        cells = rec[c1 - 1:] if c2 is None else rec[c1 - 1:c2]
        if c2 is not None:
            cells = cells + [""] * (c2 - c1 + 1 - len(cells))
        rows.append((n, [(c, _is_excel_error(c)) for c in cells]))
    # A closed range reaching past the file's end stops where the file does —
    # as an xlsx range does (openpyxl yields no rows past the data), so both
    # report the same skips and the same next_cell_range, and 'A2:C1048576'
    # doesn't build a million phantom blank rows.
    return {"kind": "csv", "encoding": enc,
            "delimiter": {"\t": "tab", ",": "comma", ";": "semicolon"}[delim], "rows": rows}


def _read_batch_source(path: Path, sheet: str, cell_range: str, date_format: str) -> dict:
    """The source's cells inside the range: {'rows': [(row number, [(text,
    is_error), ...]), ...], 'c1', 'c2', 'r1', 'r2' (None = open-ended), plus
    kind/sheet/encoding/delimiter}. Every row is the same width. An open end
    is trimmed to the last row / column holding something."""
    rsheet, c1, r1, c2, r2 = _parse_cell_range(cell_range)
    if rsheet and sheet and rsheet.lower() != sheet.lower():
        raise VoiceKitError(f"sheet='{sheet}' and cell_range names '{rsheet}' — pick one.")
    sheet = sheet or rsheet
    ext = path.suffix.lower()
    if ext in BATCH_XLSX_EXTS:
        g = _xlsx_grid(path, sheet, c1, r1, c2, r2, date_format)
    else:
        if sheet:
            raise VoiceKitError(f"sheet= is for an Excel workbook; {path.name} is a text file "
                                f"with one table.")
        g = _csv_grid(path, c1, r1, c2, r2)

    def filled(cells) -> bool:
        return any(t.strip() or err for t, err in cells)
    rows = g["rows"]
    if r2 is None:
        while rows and not filled(rows[-1][1]):
            rows.pop()
    if c2 is None:
        width = 0
        for _, cells in rows:
            for j in range(len(cells) - 1, -1, -1):
                if cells[j][0].strip() or cells[j][1]:
                    width = max(width, j + 1)
                    break
        width = max(width, 1)
        c2 = c1 + width - 1
    width = c2 - c1 + 1
    g["rows"] = [(n, (cells + [("", False)] * width)[:width]) for n, cells in rows]
    g.update(c1=c1, c2=c2, r1=r1, r2=r2)
    return g


def _value_like(t: str) -> bool:
    """A cell that reads like DATA, not a column title: a number or amount,
    a date, an ID-ish digit run, an e-mail address, or long text."""
    t = (t or "").strip()
    if not t:
        return False
    if len(t) > 40 or "@" in t or re.search(r"\d{3}", t):
        return True
    if re.fullmatch(r"\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}", t):
        return True
    num = re.sub(r"[\s$,%()]", "", t).lstrip("-+")
    return bool(num) and bool(re.fullmatch(r"\d*\.?\d+", num))


def _header_row_looks_real(hcells: list, data: list, labels: list, columns: dict | None) -> bool:
    """Whether the row read as the header can be QUOTED in an error. Cell
    values are client data, and when a range starts one row too low its
    'header' is the first client's row — so header names are only listed when
    the row looks like titles: a cell names an ask label or a `columns`
    header, or (no cell reads like data) some column has data-like values
    below its title. An all-text grid can't be told apart — not listed."""
    texts = [t.strip() for t, _e in hcells]
    wanted = {l.casefold() for l in labels}
    wanted |= {str(v).strip().casefold() for v in (columns or {}).values()
               if isinstance(v, str) and not re.match(r"(?i)col\s*:", v.strip())}
    if any(t and t.casefold() in wanted for t in texts):
        return True
    if any(_value_like(t) for t in texts):
        return False
    for j, t in enumerate(texts):
        if t and any(j < len(cells) and _value_like(cells[j][0]) for _n, cells in data[:5]):
            return True
    return False


def _resolve_source_col(spec, headers: dict, g: dict, what: str, header: bool,
                        headers_ok: bool = True) -> int:
    """An absolute column number for a header name, 'col:C' or a 1-based
    column number (A = 1). `headers` maps lower-cased header -> column.
    headers_ok=False: the header row may be data, so errors name column
    letters only, never its cells."""
    c1, c2 = g["c1"], g["c2"]
    span = f"{_col_letter(c1)}-{_col_letter(c2)}" if c1 != c2 else _col_letter(c1)

    def available() -> str:
        if header and headers and headers_ok:
            names = [f"{h} ({_col_letter(c)})" for h, c in headers.values()]
            return f"Headers in this range: {', '.join(names)}."
        if header and headers:
            return (f"Row {g.get('_header_row', '?')} — read as the header — doesn't look like "
                    f"a header row, so its cells aren't listed (they may be client data): set "
                    f"header=False or give columns= by letter, e.g. columns={{'Amount': "
                    f"'col:A'}}. This range spans {span}.")
        return (f"header=False, so name columns as 'col:B' or a number (A = 1); this "
                f"range spans {span}.")
    col = None
    if isinstance(spec, bool):
        spec = str(spec)
    if isinstance(spec, int):
        col = spec
    else:
        s = str(spec).strip()
        m = re.fullmatch(r"(?i)col\s*:\s*([A-Z]{1,3}|\d+)", s)
        if m:
            v = m.group(1)
            col = int(v) if v.isdigit() else _col_num(v)
        elif header and s.lower() in headers:
            col = headers[s.lower()][1]
        else:
            raise VoiceKitError(f"No column for {what} '{s}'. {available()}")
    if not (c1 <= col <= c2):
        raise VoiceKitError(f"{what} is mapped to column {_col_letter(col) or col}, outside "
                            f"the range ({span}). {available()}")
    return col


def _compress_rows(nums: list[int]) -> str:
    """[3,4,5,9] -> '3-5, 9' — row lists stay short in a response."""
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(str(nums[i]) if i == j else f"{nums[i]}-{nums[j]}")
        i = j + 1
    return ", ".join(out)


def _map_batch_source(g: dict, labels: list[str], header: bool, columns: dict | None,
                      require: list | None, skip_if_filled: list | None,
                      skip_blank: bool) -> dict:
    """Map a source grid onto the ask labels and filter it. Returns
    {'runs': [(row number, [values in label order])], 'columns', 'done_columns',
    'header_row', 'data_rows', 'skipped': {reason: [row numbers]}}. Raises on
    an unresolvable label or an Excel error cell (by row + column, never value)."""
    rows = list(g["rows"])
    headers: dict = {}
    header_row = None
    if header:
        if not rows:
            raise VoiceKitError("The range is empty — there's no header row to read.")
        header_row, hcells = rows.pop(0)
        for j, (t, _err) in enumerate(hcells):
            h = t.strip()
            if h and h.lower() not in headers:          # first of a duplicate wins (as in AHK)
                headers[h.lower()] = (h, g["c1"] + j)
        if not headers:
            raise VoiceKitError(
                f"Row {header_row} — the first row of the range, read as the header — is "
                f"empty. Start cell_range at the header row (e.g. 'A4:C'), or pass "
                f"header=False and map columns as 'col:B'.")
    headers_ok = bool(header) and _header_row_looks_real(hcells, rows, labels, columns)
    g["_header_row"] = header_row
    cmap = {str(k).strip().lower(): v for k, v in (columns or {}).items()}
    unknown = [k for k in (columns or {}) if str(k).strip().lower() not in
               {l.lower() for l in labels}]
    if unknown:
        raise VoiceKitError(f"columns names {', '.join(repr(str(k)) for k in unknown)}, which "
                            f"isn't an ask label of this workflow. Its labels: "
                            f"{', '.join(labels)}.")
    resolved: dict = {}
    col_of: dict = {}
    for lab in labels:
        spec = cmap.get(lab.lower(), lab)
        if lab.lower() not in cmap and not header:
            raise VoiceKitError(
                f"header=False, so every ask label needs a column in `columns` — '{lab}' has "
                f"none. e.g. columns={{'{lab}': 'col:A'}}.")
        col = _resolve_source_col(spec, headers, g, f"the label '{lab}'", header, headers_ok)
        col_of[lab] = col
        hdr = g["rows"][0][1][col - g["c1"]][0].strip() if header else ""
        # The header cell is quoted back only when the row reads as titles
        # (a data row's cell is client data).
        resolved[lab] = {"column": _col_letter(col),
                         **({"header": hdr} if header and headers_ok else {})}
    done_cols = []
    for spec in (skip_if_filled or []):
        col = _resolve_source_col(spec, headers, g, "skip_if_filled column", header,
                                  headers_ok)
        done_cols.append(col)
    req = []
    for r in (require or []):
        hit = next((l for l in labels if l.lower() == str(r).strip().lower()), None)
        if hit is None:
            raise VoiceKitError(f"require names '{r}', which isn't an ask label. Its labels: "
                                f"{', '.join(labels)}.")
        req.append(hit)

    skipped: dict = {"already_done": [], "blank": [], "required_blank": []}
    runs, bad = [], []
    c1 = g["c1"]
    for n, cells in rows:
        done = [cells[c - c1] for c in done_cols]
        for c, (t, err) in zip(done_cols, done):
            if err:
                bad.append(f"row {n} (column {_col_letter(c)}: {t})")
        if any(t.strip() for t, err in done if not err):
            skipped["already_done"].append(n)
            continue
        vals = {lab: cells[col_of[lab] - c1] for lab in labels}
        if skip_blank and not any(t.strip() or err for t, err in vals.values()):
            skipped["blank"].append(n)
            continue
        if any(not vals[lab][0].strip() and not vals[lab][1] for lab in req):
            skipped["required_blank"].append(n)
            continue
        for lab in labels:
            t, err = vals[lab]
            if err:
                bad.append(f"row {n} (column {_col_letter(col_of[lab])}: {t.strip()})")
        runs.append((n, [vals[lab][0] for lab in labels]))
    if bad:
        more = f", and {len(bad) - 10} more" if len(bad) > 10 else ""
        raise VoiceKitError(
            f"The source has Excel error values where the workflow would read them: "
            f"{'; '.join(bad[:10])}{more}. Nothing was run — fix those cells (or skip the "
            f"rows with require / skip_if_filled / cell_range) and try again.")
    return {"runs": runs, "columns": resolved, "done_columns": done_cols, "require": req,
            "header_row": header_row, "data_rows": len(rows),
            "skipped": {k: v for k, v in skipped.items() if v}}


def _own_sheet_rows(sheet: Path, labels: list[str]) -> dict:
    """The workflow's own inputs sheet, read the way WfLoopCsvRows reads it
    (UTF-8, comma, header names every label, a record skipped only when EVERY
    cell is blank after Trim) — so pass N here is pass N there."""
    try:
        raw = sheet.read_bytes()
    except PermissionError:
        raise VoiceKitError(f"Couldn't read {sheet.name} — another program has it locked. "
                            f"Close it and try again.") from None
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise VoiceKitError(
            f"{sheet.name} isn't saved as UTF-8 (Excel's plain 'CSV' is ANSI), and the loop "
            f"reads its own sheet as UTF-8 — accented text would arrive garbled. In Excel: "
            f"File > Save As > 'CSV UTF-8 (Comma delimited)', then try again.") from None
    recs = list(csv.reader(io.StringIO(text, newline="")))
    if len(recs) < 2:
        raise VoiceKitError(f"{sheet.name} needs a header row plus at least one row.")
    cols: dict = {}
    for j, h in enumerate(recs[0]):
        if h.strip(" \t") and h.strip(" \t").lower() not in cols:
            cols[h.strip(" \t").lower()] = (h.strip(" \t"), j)
    missing = [l for l in labels if l.lower() not in cols]
    if missing:
        raise VoiceKitError(f"{sheet.name} is missing a column for: {', '.join(missing)}.")
    runs, bad = [], []
    for n, rec in enumerate(recs[1:], start=2):
        if not any(c.strip(" \t") for c in rec):
            continue
        vals = []
        for lab in labels:
            j = cols[lab.lower()][1]
            v = rec[j] if j < len(rec) else ""
            if _is_excel_error(v):
                bad.append(f"row {n} (column {_col_letter(j + 1)}: {v.strip()})")
            vals.append(v)
        runs.append((n, vals))
    if bad:
        raise VoiceKitError(f"{sheet.name} has Excel error values: {'; '.join(bad[:10])}. "
                            f"Nothing was run.")
    if not runs:
        raise VoiceKitError(f"{sheet.name} has a header row but no rows of answers.")
    return {"runs": runs, "data_rows": len(recs) - 1,
            "columns": {l: {"column": _col_letter(cols[l.lower()][1] + 1),
                            "header": cols[l.lower()][0]} for l in labels}}


def _write_batch_csv(header: list[str], rows: list[list[str]]) -> Path:
    """The per-run batch file LoopRunner reads (earlier ones are cleared —
    callers have already checked no loop is running). utf-8-sig, RFC-4180,
    CRLF: the dialect WfCsvParse reads."""
    logs = REPO_ROOT / "logs"
    logs.mkdir(parents=True, exist_ok=True)
    for old in list(logs.glob("mcp-batch*.csv")):
        try:
            old.unlink()
        except OSError:
            pass
    batch = logs / f"mcp-batch-{datetime.now():%Y%m%d%H%M%S}-{os.getpid()}.csv"
    with open(batch, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    return batch


def _batch_target(name: str) -> tuple[str, list[str]]:
    """(base, ask labels) of a workflow a batch can feed, or a clear refusal."""
    base = _workflow_base(name)
    steps_file = _steps_file(base)
    if not steps_file.exists():
        raise VoiceKitError(f"No workflow named '{base}' (looked for "
                            f"{steps_file.name}).{_near_miss(base)}")
    labels = _workflow_labels(base, "ask")
    if not labels:
        raise VoiceKitError(
            f"'{space_out(base)}' has no inputs (no ask steps, and no {{{{name}}}} in a "
            f"fill value), so there's nothing to feed "
            f"rows into — run it with run_automation('loop {space_out(base)}') instead.")
    return base, labels


def _launch_batch(base: str, batch: Path, n: int, wait_seconds: int) -> dict:
    since = datetime.now().strftime("%Y%m%d%H%M%S")
    r = _launch_and_report(
        [AHK_EXE, str(REPO_ROOT / "lib" / "LoopRunner.ahk"), base, str(batch)],
        f"loop {space_out(base)} — batch of {n} row(s)",
        str(REPO_ROOT / "lib" / "LoopRunner.ahk"), wait_seconds, cwd=str(REPO_ROOT))
    r["rows"] = n
    collects = _workflow_labels(base, "collect")
    if collects:
        r["collects"] = collects
        r["results_note"] = (f"This workflow collects: {', '.join(collects)}. When the loop "
                             f"finishes, call read_workflow_sheet('{space_out(base)}') to read "
                             f"the values back.")
    if not r.get("finished"):
        r["note"] = ("Running one pass per row; it stops by itself after the last row. "
                     "The user can end it early with the Stop Looping button or Ctrl+Alt+Shift+X.")
    return _attach_outcome(r, base, since, exit_codes=True)


def run_workflow_batch(name: str, rows: list | None = None, wait_seconds: int = 0, *,
                       source: str | None = None, sheet: str | None = None,
                       cell_range: str | None = None, header: bool = True,
                       columns: dict | None = None, require: list | None = None,
                       skip_if_filled: list | None = None, skip_blank: bool = True,
                       date_format: str | None = None, dry_run: bool = False,
                       preview_rows: int = 10) -> dict:
    """Run a workflow's loop once per row, headlessly — no chooser dialog, no
    per-pass questions. The rows come from EXACTLY ONE of:

    - `rows`: a list of dicts answering every ask label (case-insensitive
      keys; extras are ignored). An all-blank row is refused.
    - `source`: a .csv/.tsv/.txt or .xlsx/.xlsm file, optionally narrowed to a
      `sheet` and `cell_range`, with `columns` mapping labels to headers /
      'col:C' / column numbers, and rows skipped by `skip_blank`, `require`
      and `skip_if_filled` — all in Python, so pass N is the Nth row kept and
      the response names its source row. `dry_run` maps and filters and
      returns a preview WITHOUT launching anything.

    The rows are written to a per-run logs\\mcp-batch-<stamp>.csv and
    lib\\LoopRunner.ahk is launched with it as its second argument; the
    floating Stop Looping bar still shows, so the user keeps control. Where
    collected values land: the workflow's own sheet as `source` fills them in
    beside each row; every other batch appends them to that sheet as new rows
    (inputs used + values collected). The source file is never written.

    Refused while any loop is running: LoopRunner is #SingleInstance Force,
    so a second launch used to kill the running batch mid-run (an agent
    retrying after finished=False was enough) — and a shared batch file was
    rewritten underneath it."""
    source_args = {"sheet": sheet, "cell_range": cell_range, "columns": columns,
                   "require": require, "skip_if_filled": skip_if_filled,
                   "date_format": date_format}
    if (rows is None) == (source is None or not str(source).strip()):
        raise VoiceKitError("Give exactly one of rows (a list of label -> value objects) or "
                            "source (a CSV / Excel file path).")
    if rows is not None:
        given = [k for k, v in source_args.items() if v] + \
            [k for k, v in (("header", header is not True), ("skip_blank", skip_blank is not True),
                            ("dry_run", bool(dry_run))) if v]
        if given:
            raise VoiceKitError(f"{', '.join(given)} only apply with source=; with rows=, "
                                f"filter the rows yourself.")
        return _run_batch_rows(name, rows, wait_seconds)
    return _run_batch_source(name, str(source), wait_seconds, sheet or "", cell_range or "",
                             bool(header), columns, require, skip_if_filled, bool(skip_blank),
                             date_format or "", bool(dry_run), preview_rows)


def _run_batch_rows(name: str, rows: list, wait_seconds: int) -> dict:
    base, labels = _batch_target(name)
    if not isinstance(rows, list) or not rows:
        raise VoiceKitError("Give at least one row of inputs (a list of objects, "
                            f"each answering: {', '.join(labels)}).")
    clean_rows = []
    for i, row in enumerate(rows):
        if not isinstance(row, dict):
            raise VoiceKitError(f"Row {i + 1} isn't an object of label -> value.")
        got = {str(k).strip().lower(): ("" if v is None else str(v)) for k, v in row.items()}
        missing = [l for l in labels if l.lower() not in got]
        if missing:
            raise VoiceKitError(f"Row {i + 1} is missing: {', '.join(missing)}. "
                                f"Every row must answer all inputs: {', '.join(labels)}.")
        # read_workflow_sheet shows these values path-normalized; a row passed
        # back as shown must type the real path, not "%USERPROFILE%\...".
        vals = [expand_user_paths(got[l.lower()]) for l in labels]
        if not any(v.strip() for v in vals):
            # The loop reader skips blank rows, so pass N would no longer be
            # row N — and an all-blank row is almost always a caller mistake.
            raise VoiceKitError(f"Row {i + 1} has no values — every input "
                                f"({', '.join(labels)}) is blank. Drop the row or fill it in.")
        clean_rows.append(vals)
    _refuse_if_loop_running()
    batch = _write_batch_csv(labels, clean_rows)
    return _launch_batch(base, batch, len(clean_rows), wait_seconds)


def _run_batch_source(name, source, wait_seconds, sheet, cell_range, header, columns,
                      require, skip_if_filled, skip_blank, date_format, dry_run,
                      preview_rows) -> dict:
    base, labels = _batch_target(name)
    if any(l.lower() == BATCH_META_COLUMN.lower() for l in labels):
        raise VoiceKitError(f"An ask label is named '{BATCH_META_COLUMN}', which a source "
                            f"batch uses for its own bookkeeping column. Rename that step.")
    path = Path(expand_user_paths(source.strip().strip('"')))
    if not path.is_absolute():
        raise VoiceKitError(f"source must be a full path (got '{source}'); "
                            f"%USERPROFILE%\\... and %TEMP%\\... are fine.")
    if not path.is_file():
        raise VoiceKitError(f"No file at {path}.")
    ext = path.suffix.lower()
    if ext not in BATCH_CSV_EXTS + BATCH_XLSX_EXTS:
        raise VoiceKitError(f"{path.name}: only .csv / .tsv / .txt and .xlsx / .xlsm are read "
                            f"(an old .xls: save it as .xlsx first).")
    if date_format:
        try:
            probe = datetime(2026, 3, 15, 14, 30).strftime(date_format)
        except (ValueError, TypeError):
            probe = ""
        if "%" not in date_format or not probe or probe == date_format:
            raise VoiceKitError(f"date_format '{date_format}' isn't a strftime pattern — "
                                f"e.g. '%m/%d/%Y' (03/15/2026) or '%Y-%m-%d'.")
    try:
        preview_rows = max(0, min(int(preview_rows), 50))
    except (TypeError, ValueError):
        preview_rows = 10
    try:
        mtime = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
    except OSError:
        mtime = ""
    out: dict = {"base": base, "name": space_out(base), "source": str(path),
                 "source_modified": mtime}
    lock = path.with_name("~$" + path.name)
    if lock.exists():
        out["open_in_excel"] = True
        out["open_note"] = ("It's open in Excel: this read the last SAVED version. Save it "
                            "first if it has edits you want run.")
    own = Path(os.path.abspath(REPO_ROOT / "workflows" / f"{base}.inputs.csv"))
    notes: list[str] = []
    try:                        # samefile also catches a short-name / mapped-drive spelling
        is_own = own.is_file() and os.path.samefile(path, own)
    except OSError:
        is_own = False
    if is_own or os.path.normcase(os.path.abspath(path)) == os.path.normcase(str(own)):
        extra = [k for k, v in (("sheet", sheet), ("cell_range", cell_range),
                                ("columns", columns), ("require", require),
                                ("skip_if_filled", skip_if_filled), ("date_format", date_format),
                                ("header", not header), ("skip_blank", not skip_blank)) if v]
        if extra:
            raise VoiceKitError(
                f"That's this workflow's own inputs sheet, which runs as-is so collected "
                f"values are written back beside each row — {', '.join(extra)} can't apply "
                f"to it. To run part of it, copy those rows to another file and pass that "
                f"(results are then appended to the sheet as new rows).")
        m = _own_sheet_rows(path, labels)
        runs = m["runs"]
        out.update(source_kind="workflow_sheet", columns=m["columns"],
                   source_rows=m["data_rows"])
        skipped = {}
        blanks = m["data_rows"] - len(runs)
        if blanks:
            skipped["blank"] = {"count": blanks}
        batch_header, batch_rows = None, None
        out["results_land"] = (f"Collected values are written into {path.name} beside each row "
                               f"(into {base}.results.csv instead if the sheet is open in Excel).")
    else:
        g = _read_batch_source(path, sheet, cell_range, date_format)
        m = _map_batch_source(g, labels, header, columns, require, skip_if_filled, skip_blank)
        runs = m["runs"]
        out["source_kind"] = g["kind"]
        if g["kind"] == "xlsx":
            out["sheet"] = g["sheet"]
        else:
            out["encoding"], out["delimiter"] = g["encoding"], g["delimiter"]
            if date_format:
                notes.append("date_format applies to Excel date cells; a CSV's text is "
                             "passed exactly as written.")
        last_read = (g["rows"][-1][0] if g["rows"] else g["r1"])
        # The range as actually read: an open end shows where the data stopped.
        out["cell_range"] = _range_text(g["c1"], g["r1"], g["c2"],
                                        g["r2"] if g["r2"] is not None else last_read)
        if m["header_row"]:
            out["header_row"] = m["header_row"]
        out["columns"] = m["columns"]
        if m["done_columns"]:
            out["skip_if_filled"] = [_col_letter(c) for c in m["done_columns"]]
        out["source_rows"] = m["data_rows"]
        skipped = {k: {"count": len(v), "rows": _compress_rows(v)}
                   for k, v in m["skipped"].items()}
        batch_header = labels + [BATCH_META_COLUMN]
        batch_rows = [vals + [str(n)] for n, vals in runs]
        out["_resume"] = (g, m, last_read)
        out["results_land"] = (
            f"Collected values are appended to workflows\\{base}.inputs.csv as new rows "
            f"(inputs used + values collected), or to {base}.results.csv if that sheet is "
            f"open in Excel. {path.name} is never written.")
    out["passes"] = len(runs)
    out["pass_rows"] = _compress_rows([n for n, _ in runs])
    out["skipped"] = skipped
    if not _workflow_labels(base, "collect"):
        out.pop("results_land", None)
    if notes:
        out["notes"] = notes
    resume = out.pop("_resume", None)

    def set_resume(start_row: int):
        if not resume:
            return
        g, m, _ = resume
        end = g["r2"] if (g["r2"] is not None and start_row <= g["r2"]) else None
        rng = _range_text(g["c1"], start_row, g["c2"], end)
        out["next_cell_range"] = rng
        # A resumed range starts below the header, so it maps by column letter.
        out["resume_with"] = {
            "cell_range": rng, "header": False,
            "columns": {lab: "col:" + v["column"] for lab, v in m["columns"].items()},
            **({"sheet": out["sheet"]} if out.get("sheet") else {}),
            **({"skip_if_filled": ["col:" + _col_letter(c) for c in m["done_columns"]]}
               if m["done_columns"] else {}),
            **({"require": m["require"]} if m["require"] else {}),
            **({"skip_blank": False} if not skip_blank else {}),
            **({"date_format": date_format} if date_format else {}),
        }
    if resume:
        set_resume(resume[2] + 1)

    if dry_run:
        out["dry_run"] = True
        out["launched"] = False
        out["preview"] = [{"source_row": n, **dict(zip(labels, vals))}
                          for n, vals in runs[:preview_rows]]
        out["note"] = (f"Dry run — nothing launched. {len(runs)} pass(es) would run"
                       f"{' (rows ' + out['pass_rows'] + ')' if runs else ''}. The preview "
                       f"shows cell values; a real run's response never does.")
        if _loop_running():
            out["note"] += " A loop is running right now, so a real run would be refused."
        return out
    if not runs:
        raise VoiceKitError(
            "No rows left to run after skipping ("
            + (", ".join(f"{k}: {v['count']}" for k, v in skipped.items()) or "the range is empty")
            + "). Check cell_range, columns and the skip options with dry_run=True.")
    _refuse_if_loop_running()
    # The own sheet goes over in its canonical spelling (the loop compares
    # full paths to decide it's the sheet and write results beside rows).
    batch = own if batch_rows is None else _write_batch_csv(batch_header, batch_rows)
    r = _launch_batch(base, batch, len(runs), wait_seconds)
    r.pop("rows", None)
    out.update(r)
    run = r.get("run") or {}
    done = run.get("passes_done")
    if resume and isinstance(done, int) and 0 <= done < len(runs) and run.get("outcome") != "ok":
        set_resume(runs[done][0])
        out["resume_note"] = (f"{done} of {len(runs)} pass(es) finished; next_cell_range "
                              f"starts at source row {runs[done][0]}, the first that didn't.")
    return out


def _running_loops() -> list[str]:
    """Workflow bases of THIS tree's running loops / batches
    (lib\\LoopRunner.ahk <Base> [batch.csv]); "?" when the base can't be read."""
    out = []
    for p in _ahk_processes():
        if _script_matches(p["cmd"], "lib\\LoopRunner.ahk"):
            m = re.search(r'(?i)\\lib\\LoopRunner\.ahk"?\s+"?([^"\s]+)', p["cmd"])
            out.append(m.group(1) if m else "?")
    return out


def _loop_running() -> bool:
    """True while a loop / batch (lib\\LoopRunner.ahk) is running — the poll
    signal for batches too long to block on (see WAIT_CAP_SECONDS)."""
    return bool(_running_loops())


def _refuse_if_loop_running() -> None:
    """LoopRunner is #SingleInstance Force: launching a loop while one runs
    REPLACES it, cutting that run short (its completed rows are journaled,
    but the rest of the batch never runs). So refuse, naming the loop."""
    loops = _running_loops()
    if loops:
        names = ", ".join(f"'{space_out(b)}'" if b != "?" else "one" for b in loops)
        raise VoiceKitError(
            f"A loop is already running ({names}). Only one loop runs at a time, and "
            f"starting another would cut that one short. Wait for it to finish — "
            f"read_workflow_sheet reports loop_running — or have the user stop it (the "
            f"Stop Looping button or Ctrl+Alt+Shift+X), then try again.")


def read_workflow_sheet(name: str) -> dict:
    """A workflow's data files, parsed: the inputs sheet
    (workflows\\<Base>.inputs.csv — ask columns plus any collect columns runs
    have filled in) and, if present, <Base>.results.csv (the overflow file
    written when the sheet itself was locked, e.g. open in Excel). Also
    reports loop_running — poll this after a fire-and-forget batch: results
    are written when the loop ends, so loop_running=False means the data
    here is final."""
    base = _workflow_base(name)
    if not _steps_file(base).exists():
        raise VoiceKitError(f"No workflow named '{base}'.{_near_miss(base)}")
    out: dict = {"base": base, "name": space_out(base), "loop_running": _loop_running()}
    for key, fname in (("sheet", f"{base}.inputs.csv"), ("results", f"{base}.results.csv")):
        p = REPO_ROOT / "workflows" / fname
        if not p.exists():
            out[key] = None
            continue
        # _read_text_any, not a strict utf-8 open: Excel saves a CSV as ANSI
        # unless told otherwise, and that used to raise here.
        recs = list(csv.reader(io.StringIO(_read_text_any(p), newline="")))
        header = [h.strip() for h in recs[0]] if recs else []
        rows = [dict(zip(header, rec + [""] * (len(header) - len(rec))))
                for rec in recs[1:] if any(c.strip() for c in rec)]
        out[key] = {"file": str(p), "columns": header, "rows": rows}
    # What the last run did, when this workflow has a run record — "the sheet
    # is empty" and "the sheet is stale" need the same evidence to tell apart.
    run = _run_report(base)
    if run:
        out["last_run"] = run
    # No sheet is three different situations, and they used to read identically.
    if out["sheet"] is None and out["results"] is None:
        if not run:
            out["note"] = ("No data, and no record of this workflow ever running on this machine. "
                           "The sheet appears once inputs are batched or a run collects values.")
        elif not run["ok"]:
            out["note"] = (f"No data because the last run didn't finish. {_run_note(run)} "
                           f"Step-by-step trace: {_runs_log()}.")
        elif not _workflow_labels(base, "collect"):
            out["note"] = ("The last run finished, but this workflow has no collect steps, "
                           "so it never produces data. Add one to capture a value.")
        else:
            out["note"] = ("The last run finished but saved nothing — its collect steps came "
                           "back empty. Check that they point at the right selection or box.")
    return out


# The prelude every snippet runs behind. Placeholders (__OUT__ / __ERR__ /
# __PRELUDE_LINES__) are substituted per run — none of them changes the line
# count, so __PRELUDE_LINES__ can be computed from this very text.
#
# _SnipFmt: Out() used to be a bare FileAppend, and handing it an OBJECT — the
# natural move with UiaRect's {x,y,w,h} — THREW inside Out, which fed the
# uncaught-error hole below. Now any value prints readably (arrays, Maps,
# object props one level deep; COM wrappers fall back to their type name).
#
# _SnipUncaught: an uncaught AHK error in a GUI-subsystem process pops a MODAL
# dialog (measured 2026-08-03 — /ErrorStdOut covers load errors only), so the
# snippet hung the full timeout and reported nothing; a real session lost 20
# minutes binary-searching its own code with Out() markers. OnError turns that
# into an immediate exit 3 carrying the message, the failing line (numbered in
# the SNIPPET's own lines), and everything Out() wrote before it.
_SNIPPET_PRELUDE = (
    "#Requires AutoHotkey v2.0\n"
    # v2's load-time warnings (VarUnset, Unreachable) are MODAL MsgBoxes shown
    # before any code runs: `Out(zz)` or a trailing `return` sat the whole
    # timeout with nothing in output or errors. Off — a real unset read still
    # throws at run time and exits 3 through OnError with the variable named.
    "#Warn All, Off\n"
    "; VoiceKit run_ahk_snippet — runs once, in its own process, then exits.\n"
    '#Include "__COMMON__"\n'
    '#Include "__BROWSER__"\n'    # Browser.ahk includes UIA.ahk itself
    '#Include "__EXPLORERSEL__"\n'   # ExplorerSelectedFiles — self-contained
    "_SnipFmt(v, depth := 0) {\n"
    "    if !IsObject(v)\n"
    "        return String(v)\n"
    "    if (depth >= 2)\n"
    "        return Type(v)\n"
    '    parts := ""\n'
    "    try {\n"
    "        if (v is Array) {\n"
    "            for item in v {\n"
    '                parts .= (parts = "" ? "" : ", ") _SnipFmt(item, depth + 1)\n'
    "                if (A_Index >= 20) {\n"
    '                    parts .= ", ..."\n'
    "                    break\n"
    "                }\n"
    "            }\n"
    '            return "[" parts "]"\n'
    "        }\n"
    "        if (v is Map) {\n"
    "            for k, item in v {\n"
    '                parts .= (parts = "" ? "" : ", ") _SnipFmt(k, depth + 1) ": " _SnipFmt(item, depth + 1)\n'
    "                if (A_Index >= 20) {\n"
    '                    parts .= ", ..."\n'
    "                    break\n"
    "                }\n"
    "            }\n"
    '            return "Map{" parts "}"\n'
    "        }\n"
    "        for k, item in v.OwnProps() {\n"
    '            parts .= (parts = "" ? "" : ", ") k ": " _SnipFmt(item, depth + 1)\n'
    "            if (A_Index >= 20) {\n"
    '                parts .= ", ..."\n'
    "                break\n"
    "            }\n"
    "        }\n"
    '        return (Type(v) = "Object" ? "" : Type(v) " ") "{" parts "}"\n'
    "    }\n"
    "    return Type(v)\n"
    "}\n"
    "Out(text) {\n"
    '    try FileAppend(_SnipFmt(text) "`n", "__OUT__", "UTF-8")\n'
    "}\n"
    "_SnipUncaught(err, mode) {\n"
    '    msg := "UNCAUGHT " Type(err)\n'
    "    if (err is Error) {\n"
    '        try msg .= ": " err.Message\n'
    "        try {\n"
    '            if (err.Extra != "")\n'
    '                msg .= " (" err.Extra ")"\n'
    "        }\n"
    '        try msg .= (err.File = A_ScriptFullPath && err.Line > __PRELUDE_LINES__) ? " — snippet line " (err.Line - __PRELUDE_LINES__) : " — " err.File " line " err.Line\n'
    "        try {\n"
    '            if (err.What != "")\n'
    '                msg .= " [in " err.What "]"\n'
    "        }\n"
    "    } else {\n"
    '        try msg .= ": " _SnipFmt(err)\n'
    "    }\n"
    '    try FileAppend(msg "`n", "__OUT__", "UTF-8")\n'
    '    try FileAppend(msg "`n", "__ERR__", "UTF-8")\n'
    "    ExitApp(3)\n"
    "}\n"
    # -1: run BEFORE _Common.ahk's LogUncaughtError (registered during its
    # include). A snippet error is already fully reported through this
    # channel — without the -1 it would ALSO land in the user's real
    # logs\errors.log, turning every failed probe (and every deliberate
    # conformance-test error) into a line in their error history.
    "OnError(_SnipUncaught, -1)\n"
)


def run_ahk_snippet(code: str, timeout_s: int = 15) -> dict:
    """Run a one-off AutoHotkey v2 snippet in its own throwaway process.

    The Tab Probe case from the scrape feedback: a 25-second diagnostic
    used to cost a real bridge key, registry lines, Voice Access pairing
    notes and a delete_automation afterwards. This costs nothing: the code is
    written to a temp folder, run once, and gone.

    What the snippet gets, so a probe needs zero path knowledge: lib\\
    _Common.ahk, lib\\Browser.ahk and lib\\ExplorerSel.ahk pre-included (Browser
    pulls in lib\\UIA.ahk, so UiaFocused / UiaFind / UiaDumpTree / BodyStatus /
    BrowserGrabPage / BrowserTypeVerified / ExplorerSelectedFiles are all in
    scope), and an Out(text)
    helper that appends a line to the report this call returns — the reliable channel, since
    AutoHotkey is a GUI-subsystem exe whose stdout only works through a real
    redirect. Out() takes any value: objects (UiaRect's {x,y,w,h} among them)
    print readably instead of throwing. An ExitApp after the code keeps 'runs
    once' true even when the snippet declares hotkeys or GUIs, and the temp
    folder is the working directory so relative file writes vanish with it.

    Errors never vanish (2026-08-03 feedback: they used to): a snippet that
    won't LOAD returns exit 2 with the syntax error in 'errors'; an UNCAUGHT
    runtime error exits 3 immediately — no modal dialog, no timeout wait —
    with the message and failing snippet line in both 'output' (in sequence
    with the Out() lines that preceded it) and 'errors'.

    On timeout the process is killed and the report still carries whatever
    Out() wrote — for a wedged probe, the partial log IS the diagnostic.

    Privacy tokens (<dir#..>/<file#..>, WP9) in the code are expanded to the
    real names locally, but only inside quoted string literals (escaped for
    AHK; that literal's %USERPROFILE% / %TEMP% expand too) — a token
    anywhere else is refused. The code stays in the temp folder only."""
    if not (code or "").strip():
        raise VoiceKitError("No code to run.")
    code = privacy.expand_ahk_code(code, REPO_ROOT)
    timeout_s = max(1, min(_as_int(timeout_s, 15), 120))
    with tempfile.TemporaryDirectory() as d:
        out_file = Path(d) / "snippet-out.txt"
        err_file = Path(d) / "snippet-err.txt"
        script = Path(d) / "snippet.ahk"
        n_prelude = _SNIPPET_PRELUDE.count("\n")
        prelude = (_SNIPPET_PRELUDE
                   .replace("__COMMON__", str(REPO_ROOT / "lib" / "_Common.ahk"))
                   .replace("__BROWSER__", str(REPO_ROOT / "lib" / "Browser.ahk"))
                   .replace("__EXPLORERSEL__", str(REPO_ROOT / "lib" / "ExplorerSel.ahk"))
                   .replace("__OUT__", str(out_file))
                   .replace("__ERR__", str(err_file))
                   .replace("__PRELUDE_LINES__", str(n_prelude)))
        script.write_text(prelude + f"{code}\n" + "ExitApp(0)\n",
                          encoding="utf-8-sig", newline="")
        timed_out = False
        exit_code = None
        errors = ""
        try:
            r = subprocess.run([AHK_EXE, "/ErrorStdOut=UTF-8", str(script)],
                               capture_output=True, text=True, encoding="utf-8",
                               errors="replace", timeout=timeout_s, cwd=d)
            exit_code = r.returncode
            errors = (r.stdout or "").strip() or (r.stderr or "").strip()
        except subprocess.TimeoutExpired:
            timed_out = True     # subprocess.run kills and reaps on timeout
        # A load error names the line in the GENERATED file — "line 67" of a
        # two-line snippet. Re-number past the prelude so it points at the
        # snippet's own line. Anchored to the script filename so a "(42)"
        # inside the quoted code excerpt is never touched.
        if errors:
            errors = re.sub(
                re.escape(script.name) + r" \((\d+)\)",
                lambda m: (f"{script.name} (snippet line {int(m.group(1)) - n_prelude})"
                           if int(m.group(1)) > n_prelude else m.group(0)),
                errors)
        # Runtime errors report through a FILE like Out() does — the pipe is
        # only trustworthy for load errors, which are written before the GUI
        # subsystem matters.
        if err_file.exists():
            runtime_err = _read_text_any(err_file).strip()
            if runtime_err:
                errors = f"{errors}\n{runtime_err}".strip() if errors else runtime_err
        output = ""
        if out_file.exists():
            output = _read_text_any(out_file).rstrip("\n")
        if errors:
            # Name clashes / v1 quotes get a plain-English hint (WP1) — load
            # errors (exit 2) and runtime ones (exit 3) alike. Explained while
            # the script is still on disk: its #Include chain is the index.
            errors = _explain_ahk_error(errors, code=code, script=script,
                                        prelude_lines=n_prelude)
            # The throwaway file's full temp path is noise (it is gone when
            # this returns) and carries the Windows user name.
            errors = errors.replace(str(script), script.name)
        result = {"exit_code": exit_code, "output": output, "errors": errors,
                  "timed_out": timed_out}
        if timed_out:
            result["note"] = (f"Killed after {timeout_s} s. 'output' holds what "
                              f"Out() wrote before that — for a wedged probe the "
                              f"partial log is the diagnostic. An uncaught error "
                              f"can't be the cause (those exit immediately with "
                              f"code 3), so the code was still busy or blocked — "
                              f"an endless loop, a wait that never returned, or "
                              f"a dialog it opened itself.")
        elif exit_code == 3 and errors:
            result["note"] = ("The snippet hit an uncaught AHK error — the "
                              "UNCAUGHT line in 'output'/'errors' names it and "
                              "the failing snippet line.")
        elif exit_code and errors:
            result["note"] = "The snippet failed to load or run — see 'errors'."
        return result


def _run_inline_ahk(script: str, timeout: int = 15) -> None:
    """Run a short throwaway AHK v2 script via stdin (AutoHotkey's `*` mode) —
    no temp files."""
    subprocess.run([AHK_EXE, "/ErrorStdOut", "*"], input=script,
                   capture_output=True, text=True, timeout=timeout, check=True)


def resolve_bridge(name_or_key: str) -> dict:
    """Find a bridge-map entry by voice phrase, module name, or bare key
    (e.g. 'Toggle Timer', 'ToggleTimer', or 'A'). Raises if there's no match.

    Matches in PRECEDENCE order — exact phrase, then module base, then bare key
    letter — across all entries, so a coincidental single-letter key match never
    beats an entry whose actual phrase/name equals the input."""
    entries = get_bridge_map()["entries"]
    want = name_or_key.strip()
    want_base = to_base(clean_phrase(want))

    def rows():
        for e in entries:
            key = e["combo"].rsplit("+", 1)[-1]
            # A companion (hotkeys\<Base>.hotkey.ahk) answers to its PARENT's
            # name — the name list_automations shows it under.
            mod_base = (_companion_parent(e["file"])
                        or re.sub(r"(?i)^hotkeys\\|\.ahk$", "", e["file"]))
            yield e, key, mod_base

    for match in (
        lambda e, key, mb: want.lower() == e["phrase"].lower(),
        lambda e, key, mb: want_base and want_base.lower() == mb.lower(),
        lambda e, key, mb: len(want) == 1 and want.upper() == key.upper(),
    ):
        for e, key, mb in rows():
            if match(e, key, mb):
                return {"combo": e["combo"], "phrase": e["phrase"], "key": key, "module": e["file"]}
    raise VoiceKitError(f"No hotkey module matches '{name_or_key}' — "
                        f"get_bridge_map shows what exists.")


def press_hotkey(name_or_key: str) -> dict:
    """Trigger an always-on hotkey module by synthesizing its Ctrl+Alt+Shift
    combo (SendLevel 1, so VoiceKit's hook hotkeys hear it — the same way the
    home window's Run button does it). Only combos registered in bridge-map.txt
    can be pressed. Requires the resident VoiceKit master to be running."""
    hit = resolve_bridge(name_or_key)
    # The key comes from bridge-map.txt (rsplit on '+'). Even though the file is
    # ours, constrain it to ONE character of the bridge pool (letters, digits
    # and the measured-safe punctuation) before it enters an executed Send
    # string — a tampered/garbled combo field must never inject keystrokes.
    # Membership in the pool, not a character class: it is the tighter test
    # (one known character, never a sequence) and it stays correct as the pool
    # grows, which a hand-written regex would not.
    if len(hit["key"]) != 1 or hit["key"].upper() not in BRIDGE_POOL:
        raise VoiceKitError(f"Refusing to press a malformed bridge key '{hit['key']}' "
                            f"(from {hit['combo']}). Check bridge-map.txt.")
    # Mid-reload the old master is gone and the new one may not have its
    # hotkeys yet: a press now would land nowhere and still report success.
    # Wait briefly for the new master, and say "restarting" if it's slow —
    # never "not running", which it isn't.
    restarted = _await_master_after_restart()
    if restarted is False:
        raise VoiceKitError(
            f"VoiceKit is restarting (a reload is in progress), so {hit['combo']} isn't "
            f"registered yet — try again in a few seconds.")
    if restarted is None and not voicekit_running():
        raise VoiceKitError(
            "VoiceKit isn't running, so its hotkeys aren't registered — pressing "
            f"{hit['combo']} would do nothing. Start it first (VoiceKitLauncher.ahk).")
    _run_inline_ahk(
        "#Requires AutoHotkey v2.0\n"
        "SendLevel 1\n"
        f'Send "^!+{hit["key"].lower()}"\n'
        "Sleep 150\n"
        "ExitApp\n")
    return {"pressed": hit["combo"], "phrase": hit["phrase"], "module": hit["module"],
            "note": ("The combo was sent to the resident VoiceKit. An isolated module's "
                     "steps run in their own process — list_running / read_module_status "
                     "show it.")}


def _remove_matching_lines(path: Path, predicate) -> int:
    """Rewrite a file (UTF-8 BOM, LF) dropping lines for which predicate(line)
    is True. Returns the number removed; the file is left untouched when
    nothing matched (no churn on unrelated deletes).

    The BOM is kept because the AHK side of this operation keeps it:
    _Common.ahk RemoveLinesContaining rewrites through FileOpen(.., "w",
    "UTF-8"), which writes EF BB BF (measured). Stripping it here meant an
    MCP delete silently changed bridge-map.txt / _index.ahk in a way the GUI
    doing the same delete did not."""
    if not path.exists():
        return 0
    # Universal newlines, since the rewrite is LF: a CRLF file (saved by an
    # older Notepad) would otherwise keep a \r on every surviving line.
    text = _lf(_read_text_any(path))
    kept, removed = [], 0
    for line in text.split("\n"):
        if predicate(line):
            removed += 1
        else:
            kept.append(line)
    if removed:
        with open(path, "w", encoding="utf-8-sig", newline="") as f:
            f.write("\n".join(kept))
    return removed


# ---- Matching an EXISTING artifact's lines and files -----------------------
# clean_phrase is lossy for a CamelCase base ('WebScrapeDemo' -> 'Webscrapedemo')
# and NTFS is case-insensitive, so a delete used to unlink the file (the
# filesystem ignores case) while case-sensitive substring tests left its
# #Include and bridge-map line behind — and the dangling #Include then refused
# every reload. Lines are therefore matched on the exact FIELD, caselessly,
# and names resolve to the real on-disk stem before anything is derived
# from them (SpaceOut for .lnk names needs the real capitals).
def _include_targets(line: str, rel: str) -> bool:
    """True if `line` is an #Include (live or parked) of repo-relative `rel`
    ('hotkeys\\Foo.ahk'): the include's path is `rel` or ends in '\\'+rel,
    compared caselessly. A longer name that merely contains it never counts."""
    m = re.search(r'(?i)#Include\s+(?:"([^"]+)"|(\S+))', line)
    if not m:
        return False
    target = (m.group(1) or m.group(2)).lower()
    rel = rel.lower()
    return target == rel or target.endswith("\\" + rel)


def _map_line_targets(line: str, rel: str) -> bool:
    """True if a bridge-map line (live or commented) registers file `rel`:
    the trimmed FILE field (third) equals it, caselessly."""
    parts = [p.strip() for p in line.split("|")]
    return len(parts) >= 3 and parts[2].lower() == rel.lower()


def _name_keys(name: str) -> list[str]:
    """Lower-case stems a typed name can mean, most canonical first: the
    CleanPhrase base ('Web Scrape Demo' -> 'webscrapedemo'), then the name with
    only whitespace removed (a name holding punctuation clean_phrase strips)."""
    keys = []
    for k in (to_base(clean_phrase(name)).lower(), re.sub(r"\s", "", name).lower()):
        if k and k not in keys:
            keys.append(k)
    return keys


def _resolve_existing_base(name: str, type: str) -> str:
    """The real on-disk stem an existing automation of `type` goes by, or the
    canonical to_base(clean_phrase(name)) when nothing matches (so a new name
    still gets CleanPhrase). Compared caselessly against the stems of the
    files that type owns — and, for a hotkey module, against the registry
    lines too, so a module whose file is already gone still resolves."""
    stems: list[str] = []
    if type == "hotkey_module":
        hk = REPO_ROOT / "hotkeys"
        if hk.is_dir():
            stems += [p.stem for p in hk.glob("*.ahk")]
        mods = index_modules()
        for rel in mods["active"] + mods["quarantined"]:
            stems.append(re.sub(r"(?i)^hotkeys\\|\.ahk$", "", rel))
        stems += [re.sub(r"(?i)^hotkeys\\|\.ahk$", "", e["file"])
                  for e in get_bridge_map()["entries"]]
        # Never a companion (Foo.hotkey) or a nested path.
        stems = [s for s in stems if "." not in s and "\\" not in s]
    else:
        for folder, suffix in (("macros", ".ahk"), ("workflows", ".steps.txt"),
                               ("prompts", ".prompt.txt")):
            d = REPO_ROOT / folder
            if d.is_dir():
                stems += [p.name[:-len(suffix)] for p in d.iterdir()
                          if p.name.lower().endswith(suffix)]
    # A leading underscore is VoiceKit's own (hotkeys\_index.ahk, the module
    # manifest): clean_phrase can never produce one, and the whitespace key
    # must not make it reachable either.
    stems = [s for s in stems if not s.startswith("_")]
    for key in _name_keys(name):
        for st in stems:
            if st.lower() == key:
                return st
    return to_base(clean_phrase(name))


def _remove_companion_hotkey(base: str, result: dict) -> None:
    """The home window's Hotkey button can give any automation a companion
    Ctrl+Alt+Shift key: hotkeys\\<Base>.hotkey.ahk, registered in _index.ahk and
    bridge-map.txt. Deleting the automation must retire all three (mirroring the
    GUI's delete parity), or a live key keeps pointing at a ghost and its bridge
    key stays allocated forever."""
    module = REPO_ROOT / _companion_rel(base)
    existed = module.exists()
    if existed:
        module.unlink()
        result["removed"].append(str(module))
    rel = _companion_rel(base)
    n = _remove_matching_lines(_hotkeys_index(), lambda ln: _include_targets(ln, rel))
    n += _remove_matching_lines(_bridge_map_file(), lambda ln: _map_line_targets(ln, rel))
    if existed or n:
        # So the companion key stops working now.
        result.update(_reload_outcome("the companion hotkey's removal"))


@_serialized
def delete_automation(name: str, type: str) -> dict:
    """Remove an automation and its artifacts (including a companion hotkey
    assigned in the home window). `type` is one of
    launch_macro | workflow | hotkey_module | snippet | ai_action."""
    result = {"type": type, "removed": []}

    if type == "snippet":
        abbrev = re.sub(r"\s", "", name.strip())
        snip = _snippets_file()
        text = _read_text_any(snip) if snip.exists() else ""
        # Case-insensitive like the GUI (hotstrings themselves fire caselessly).
        pat = re.compile(rf"^:[^:]*:{re.escape(abbrev)}::(.*)$", re.IGNORECASE)
        matches = [m for m in (pat.match(ln.strip()) for ln in text.split("\n")) if m]
        if not matches:
            raise VoiceKitError(f"No snippet '{abbrev}' found.")
        # A code-block snippet's line is ':*:abbrev:: {' — removing just that
        # line orphans the block, Snippets.ahk stops loading, and the reload
        # below would take down EVERY hotkey and snippet. Refuse, like Home does.
        if any(m.group(1).strip() in ("{", "") for m in matches):
            raise VoiceKitError(
                f"'{abbrev}' is a code-block (dynamic) snippet — deleting its line would break "
                f"Snippets.ahk and every hotkey with it. Edit hotkeys\\Snippets.ahk by hand "
                f"and remove the whole block instead.")
        snap = _FileSnapshot(snip)              # roll back — never reload a broken file
        try:
            _remove_matching_lines(snip, lambda ln: pat.match(ln.strip()) is not None)
            ok, err = validate_ahk(str(snip))
            if not ok:
                raise VoiceKitError(f"Removing '{abbrev}' would break Snippets.ahk, so it "
                                    f"was rolled back:\n{err}")
        except BaseException:
            snap.restore()
            raise
        result["removed"].append(f"snippet {abbrev}")
        result.update(_reload_outcome("the removal"))
        log(f"deleted snippet | {abbrev}")
        return result

    # The real on-disk stem, not clean_phrase's: 'WebScrapeDemo' must stay
    # 'WebScrapeDemo' (not 'Webscrapedemo') for .lnk names and line matches.
    base = _resolve_existing_base(name, type)
    if base.lower() in _DELETE_PROTECTED_LOWER:
        raise VoiceKitError(f"'{base}' is part of VoiceKit itself and can't be deleted"
                            + (" (it can still be changed with update_macro / edit_macro)."
                               if base.lower() not in _BUILTINS_LOWER else "."))

    if type == "hotkey_module":
        # Checked before anything is touched: hotkeys\Snippets.ahk holds every
        # snippet the user has, and deleting it as a "module" would wipe them all.
        if base.lower() in RESERVED_MODULE_BASES:
            raise VoiceKitError(
                f"'{base}' is one of VoiceKit's own files in hotkeys\\ and can't be deleted "
                f"as a module — hotkeys\\Snippets.ahk holds every snippet you have. Delete "
                f"a single snippet with type='snippet'.")
        module = _module_file(base)
        existed = module.exists()
        if existed:
            module.unlink()
            result["removed"].append(str(module))
        body = REPO_ROOT / _body_rel(base)      # isolated modules keep their steps here
        if body.exists():
            body.unlink()
            result["removed"].append(str(body))
        backup = _module_backup_file(base)      # the replace-undo snapshot
        if backup.exists():
            backup.unlink()
            result["removed"].append(str(backup))
        # Its last status line and any pending stop flag too (the AHK side,
        # DeleteHotkeyModuleArtifacts, does the same): a module later given
        # this name must not show the deleted one's progress as its own.
        for leftover in (_body_status_file(base), _body_stop_file(base)):
            try:
                if leftover.exists():
                    leftover.unlink()
                    result["removed"].append(str(leftover))
            except OSError:
                pass
        rel = f"hotkeys\\{base}.ahk"
        n = _remove_matching_lines(_hotkeys_index(), lambda ln: _include_targets(ln, rel))
        n += _remove_matching_lines(_bridge_map_file(), lambda ln: _map_line_targets(ln, rel))
        if not existed and n == 0:
            raise VoiceKitError(f"No hotkey module '{base}' found — nothing was deleted.")
        result.update(_reload_outcome("the removal"))
        log(f"deleted hotkey | {base}")
        return result

    if type == "workflow":
        steps_file = _steps_file(base)
        macro, _txt, kind = _macro_info(base)
        is_stub = kind == "workflow_stub"
        # Cross-guard BEFORE touching anything: a same-named launch macro / AI
        # action must not lose its Start Menu shortcut to a mistyped `type`.
        if kind and not is_stub:
            hint = "ai_action" if kind == "ai_action" else "launch_macro"
            raise VoiceKitError(f"'{base}' isn't a workflow — delete it with type='{hint}'.")
        if not steps_file.exists() and not is_stub:
            raise VoiceKitError(f"No workflow named '{base}' found — nothing was deleted.")
        if steps_file.exists():
            steps_file.unlink()
            result["removed"].append(str(steps_file))
        for data in (f"{base}.inputs.csv", f"{base}.results.csv"):   # sheet + results overflow
            p = REPO_ROOT / "workflows" / data
            if p.exists():
                p.unlink()
                result["removed"].append(str(p))
        # Unsaved loop results (lib\WorkflowLoop.ahk's journal): left behind they
        # would be merged into the sheet of the NEXT workflow given this name.
        jdir = REPO_ROOT / "logs" / "loop-journal"
        if jdir.is_dir():
            for p in jdir.glob(f"{base}.*.jnl"):
                if re.fullmatch(re.escape(base) + r"\.\d{14}-\d+\.jnl", p.name, re.I):
                    try:                          # a live loop may be appending right now;
                        p.unlink()                # never let that abort the delete midway
                        result["removed"].append(str(p))
                    except OSError:
                        pass
        if is_stub:
            macro.unlink()
            result["removed"].append(str(macro))
        disp = space_out(base)
        _delete_lnk(disp, result)
        _delete_lnk(f"loop {disp}", result)      # companion loop entry
        _remove_companion_hotkey(base, result)   # companion hotkey, if assigned
        log(f"deleted workflow | {disp}")
        return result

    if type == "ai_action":
        macro = _macro_file(base)
        prompt_file = _prompt_file(base)
        if not macro.exists() and not prompt_file.exists():
            raise VoiceKitError(f"No AI action '{base}' found.")
        if macro.exists():
            macro.unlink()
            result["removed"].append(str(macro))
        if prompt_file.exists():
            prompt_file.unlink()
            result["removed"].append(str(prompt_file))
        _remove_macro_backup(base, result)
        for nm in {clean_phrase(name), space_out(base)}:
            _delete_lnk(nm, result)
        _remove_companion_hotkey(base, result)
        log(f"deleted ai-action | {base}")
        return result

    if type == "launch_macro":
        macro, _txt, kind = _macro_info(base)
        if not kind:
            raise VoiceKitError(f"No launch macro '{base}' found.")
        if kind == "workflow_stub":
            raise VoiceKitError(f"'{base}' is a workflow, not a launch macro — "
                                f"delete it with type='workflow'.")
        if kind == "ai_action":
            raise VoiceKitError(f"'{base}' is an AI action — delete it with "
                                f"type='ai_action' so its prompt file goes too.")
        macro.unlink()
        result["removed"].append(str(macro))
        _remove_macro_backup(base, result)   # the update_macro undo snapshot
        # Launch shortcut may be named by phrase or SpaceOut(base); try both.
        for nm in {clean_phrase(name), space_out(base)}:
            _delete_lnk(nm, result)
        _remove_companion_hotkey(base, result)
        log(f"deleted launch | {base}")
        return result

    raise VoiceKitError(f"Unknown automation type '{type}'.")


def _delete_lnk(display_name: str, result: dict) -> None:
    lnk = VOICE_MACROS / f"{display_name}.lnk"
    if lnk.exists():
        lnk.unlink()
        result["removed"].append(str(lnk))


def _remove_macro_backup(base: str, result: dict) -> None:
    """Deleting a macro retires its update_macro undo snapshot too — the same
    parity rule the module delete keeps for logs\\module-backups\\."""
    backup = _macro_backup_file(base)
    if backup.exists():
        backup.unlink()
        result["removed"].append(str(backup))
