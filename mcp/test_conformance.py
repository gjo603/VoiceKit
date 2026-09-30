"""Conformance test: prove voicekit_writer's output is byte-compatible with the
real AutoHotkey engine, and that the naming/encoding ports match AHK exactly.

Non-invasive: it only writes temp files and reads existing committed artifacts —
it never creates automations in macros\\, hotkeys\\, or the Start Menu. Every
test that writes goes through _sandbox(), which retargets REPO_ROOT *and*
VOICE_MACROS at a temp tree; and VOICE_MACROS is pointed away from the real
Start Menu for the whole run anyway, so even a guard test that slips past its
refusal can't reach the user's 'Voice Macros' folder.

Run:  python test_conformance.py     (prints PASS/FAIL, exit 1 on failure)
Also collectable by pytest (test_* functions).
"""

from __future__ import annotations

import contextlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

import voicekit_writer as vk

REAL_ROOT = vk.REPO_ROOT
REAL_COMMON = REAL_ROOT / "lib" / "_Common.ahk"
HARNESS = REAL_ROOT / "mcp" / "_conformance" / "roundtrip.ahk"
RUN_HARNESS = REAL_ROOT / "mcp" / "_conformance" / "runrecord.ahk"
BODY_HARNESS = REAL_ROOT / "mcp" / "_conformance" / "bodystatus.ahk"
GEN_HARNESS = REAL_ROOT / "mcp" / "_conformance" / "generators.ahk"
BATCH_HARNESS = REAL_ROOT / "mcp" / "_conformance" / "batchrows.ahk"

# Never the real Start Menu, not even by accident (see the module docstring).
vk.VOICE_MACROS = Path(tempfile.gettempdir()) / "vk-conformance-startmenu"


@contextlib.contextmanager
def _sandbox(reload=lambda: "not_running", libs=("_Common.ahk",)):
    """A throwaway VoiceKit tree: REPO_ROOT and VOICE_MACROS (the Start Menu
    folder) retargeted at a temp dir, reload_voicekit stubbed (pass
    reload=None to keep the real one), and the named lib\\ files copied in —
    a generated body #Includes ..\\..\\lib\\_Common.ahk, so creating or
    editing an isolated module load-checks against it. Yields the root."""
    tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
    d = Path(tmp.name)
    orig = (vk.REPO_ROOT, vk.VOICE_MACROS, vk.reload_voicekit)
    vk.REPO_ROOT, vk.VOICE_MACROS = d, d / "startmenu"
    if reload is not None:
        vk.reload_voicekit = reload
    try:
        for sub in ("macros", "hotkeys", "workflows", "prompts", "logs", "startmenu"):
            (d / sub).mkdir()
        if libs:
            (d / "lib").mkdir()
            for lib in libs:
                shutil.copyfile(REAL_ROOT / "lib" / lib, d / "lib" / lib)
        yield d
    finally:
        vk.REPO_ROOT, vk.VOICE_MACROS, vk.reload_voicekit = orig
        tmp.cleanup()


@contextlib.contextmanager
def _patched(**attrs):
    """Temporarily replace voicekit_writer attributes (restored on exit)."""
    orig = {k: getattr(vk, k) for k in attrs}
    for k, v in attrs.items():
        setattr(vk, k, v)
    try:
        yield
    finally:
        for k, v in orig.items():
            setattr(vk, k, v)

# Adversarial params: pipe, percent, newline, tab, quotes, backslashes, GUID path.
STEPS = [
    ("focus", "Notepad ahk_exe notepad.exe", "notepad.exe", ""),
    ("text", "a|b%c\nsecond\tline & 100%", "", ""),
    ("click", "Untitled - Notepad", "File | Save As...", "120,44"),
    ("hover", "Untitled - Notepad", "Format", ""),          # hover by element name
    ("hover", "ahk_exe notepad.exe", "", "50,60"),          # hover by position only
    ("drag", "ahk_exe notepad.exe", "", "10,20,300,220"),   # drag: press/travel/release path
    ("run", 'explorer.exe "::{20D04FE0-3AEA-1069-A2D8-08002B30309D}"', "", ""),
    ("run", r'C:\Program Files\App "x".exe', "", ""),
    ("keys", "^s", "", ""),
    ("wait", "800", "", ""),
    # A wait may also be a RANGE, which the engine turns into Random(lo, hi) —
    # so a looped workflow doesn't pause identically every pass. The writer
    # must let it through untouched, and the engine must read it back.
    ("wait", "600-1400", "", ""),
    ("ask", "Customer name", "Acme, Inc", ""),              # input, with a suggestion
    ("collect", "Order number", "", ""),                    # collect the selection
    ("collect", "Total | price", "Amount box", ""),         # collect a named box's value
    # set: names a value; the value may reference other names, and carries the
    # same adversarial characters as any other param.
    ("set", "Greeting", "Hi {{Customer name}} | 100% done\nline two", ""),
    # capture: name | cmd.exe command line (quotes, pipe, percent, {{...}}) |
    # timeout seconds — and the blank-timeout form (the engine's 30 s).
    ("capture", "Invoice Data", r'python "C:\x y\invoice.py" {{selected_files}} | findstr 100%', "90"),
    ("capture", "Stamp", "echo %DATE%", ""),
    # fill: window | label (with an occurrence suffix and a literal ##) |
    # VALUE in paramC — the one paramC the engine substitutes — plus the
    # empty value that clears a box.
    ("fill", "ahk_exe chrome.exe", "Amount | Line ##1#2", "{{Amount}} 1,234.50 | 100%"),
    ("fill", "ahk_exe chrome.exe", "Memo", ""),
    # waitfor: condType alone, condType + seconds, and the windowless clipboard
    # wait — paramC carries both halves and must survive the round trip.
    ("waitfor", "ahk_exe notepad.exe", "Save | As...", "elementexists"),
    ("waitfor", "ahk_exe notepad.exe", "No results found", "textnotvisible,30"),
    ("waitfor", "", "", "clipboardchanged,5"),
    ("text", "{{Greeting}} for {{Order number}}", "", ""),  # {{...}} survives encoding
    # if/else/endif branching (condType in paramC): must round-trip too.
    ("if", "ahk_exe notepad.exe", "", "winexists"),
    ("text", "then branch", "", ""),
    ("else", "", "", ""),
    ("text", "else branch", "", ""),
    ("endif", "", "", ""),
    ("if", "Untitled - Notepad", "File | Save", "elementexists"),
    ("endif", "", "", ""),
]


def _expected_line(step) -> str:
    t, a, b, c = (list(step) + ["", "", ""])[:4]
    return f"{t}|{vk.wf_encode(a)}|{vk.wf_encode(b)}|{vk.wf_encode(c)}\n"


def test_workflow_roundtrip_through_real_engine():
    """Python emits .steps.txt -> real WorkflowLoad parses it -> re-emit -> match."""
    with tempfile.TemporaryDirectory() as d:
        in_file = Path(d) / "in.steps.txt"
        out_file = Path(d) / "out.txt"
        vk._write_new(in_file, vk.steps_to_text("Round Trip", STEPS))  # BOM+LF, as create_workflow writes

        r = subprocess.run(
            [vk.AHK_EXE, "/ErrorStdOut", str(HARNESS), str(in_file), str(out_file)],
            capture_output=True, text=True, timeout=30,
        )
        assert r.returncode == 0, f"harness failed: {r.stdout}{r.stderr}"

        got = out_file.read_text(encoding="utf-8-sig").replace("\r\n", "\n")
        want = "".join(_expected_line(s) for s in STEPS)
        assert got == want, f"round-trip mismatch:\n--- got ---\n{got}\n--- want ---\n{want}"


def test_run_record_reads_back_what_the_engine_wrote():
    """The engine records a run; the writer reads it back the same way.

    This is a two-implementation seam like every other one here: AutoHotkey's
    IniWrite creates the file as UTF-16LE with a BOM, and the reason string it
    stores is what run_workflow_batch reports instead of the exit code (which
    is always 0 and always said "Finished cleanly")."""
    with tempfile.TemporaryDirectory() as d:
        name = "ConformanceRunRecord"
        r = subprocess.run(
            [vk.AHK_EXE, "/ErrorStdOut", str(RUN_HARNESS), d, name],
            capture_output=True, text=True, timeout=60,
        )
        assert r.returncode == 0, f"harness failed: {r.stdout}{r.stderr}"

        ini = Path(d) / "workflow-runs.ini"
        log = Path(d) / "workflow-runs.log"
        assert ini.exists(), "engine wrote no run record"
        assert ini.read_bytes()[:2] == b"\xff\xfe", "expected AHK's UTF-16LE BOM"
        assert log.exists(), "engine wrote no run log"

        orig_ini, orig_log = vk._runs_ini, vk._runs_log
        try:
            vk._runs_ini = lambda: ini
            vk._runs_log = lambda: log
            rec = vk.read_run_record(name)
            assert rec, "writer read no record back"
            rep = vk._run_report(name)
        finally:
            vk._runs_ini, vk._runs_log = orig_ini, orig_log

        assert rep["outcome"] == "failed" and rep["ok"] is False, rep
        assert rep["failed_step"] == 2, rep          # not 1, not 3 — the `if` step
        assert rep["steps_total"] == 3, rep
        assert "notacondition" in rep["reason"], rep  # the string the popup used to eat
        assert rep["passes_done"] == 2 and rep["passes_total"] == 5, rep
        assert any("FAILED at step 2/3" in l for l in rep["trace"]), rep["trace"]
        assert not any("step 3/3" in l for l in rep["trace"]), "step after the failure ran"

        # A record older than `since` belongs to an earlier run and must not be
        # reported as this one's.
        try:
            vk._runs_ini, vk._runs_log = lambda: ini, lambda: log
            assert vk._run_report(name, since="29990101000000") == {}
            assert vk._run_report(name, since=rec["started"]) != {}

            # A plain run after a loop replaces the loop's record WHOLE: no
            # stale pass counts or loop start time reported as this run's.
            stale = vk.read_run_record(name + "Stale")
            assert stale.get("outcome") == "ok", stale
            for k in ("passes_done", "passes_total", "collected_rows",
                      "loop_started", "loop_started_text"):
                assert k not in stale, f"stale loop key {k!r} survived: {stale}"
            srep = vk._run_report(name + "Stale")
            assert "passes_total" not in srep, srep
            assert srep["started"] == stale["started_text"], srep
            # ...and writing one section never disturbs another's.
            assert vk.read_run_record(name).get("outcome") == "failed"
        finally:
            vk._runs_ini, vk._runs_log = orig_ini, orig_log


def test_body_status_reads_back_what_ahk_wrote():
    """A long-running hotkey body publishes progress; the writer reads it back.

    Another AHK-writes / Python-reads seam: BodyStatus() in _Common.ahk writes
    "<yyyyMMddHHmmss>|<text>" as UTF-8, and body_status() parses it here. Both
    sides implement the format independently, which is exactly the arrangement
    that drifts unless a real round trip pins it."""
    with tempfile.TemporaryDirectory() as d:
        base, text = "ConformanceBodyZz", "3 / 766 - Acme Holdings"
        r = subprocess.run(
            [vk.AHK_EXE, "/ErrorStdOut", str(BODY_HARNESS), d, base, text],
            capture_output=True, text=True, timeout=60,
        )
        assert r.returncode == 0, f"harness failed: {r.stdout}{r.stderr}"

        # The AHK side chose this path from the base alone; the Python side
        # must choose the same NAME, or the read half looks at the wrong file.
        written = Path(d) / "logs" / f"body-status-{base}.txt"
        assert written.exists(), f"the AHK helper wrote no status file: {list(Path(d).rglob('*'))}"
        assert written.name == vk._body_status_file(base).name, "the two sides name it differently"

        orig = vk._body_status_file
        try:
            vk._body_status_file = lambda b: Path(d) / "logs" / f"body-status-{b}.txt"
            st = vk.body_status(base)
        finally:
            vk._body_status_file = orig

        assert st, "writer read no status back"
        assert st["text"] == text, st
        assert re.fullmatch(r"\d{14}", st["when"]), st       # AHK's A_Now stamp
        assert 0 <= st["age_seconds"] < 120, st              # just written
        assert vk.body_status("NoSuchBodyZz") is None

    # The names each side derives must match, not just the contents.
    assert vk._body_status_file("MyModule").name == "body-status-MyModule.txt"


def test_uia_probe_tree_reads_a_real_window():
    """dump_uia_tree end to end: the committed probe (mcp\\uia_probe.ahk on top
    of lib\\UIA.ahk) walks a real window in ANOTHER process and the writer
    parses the report back. One more AHK-writes / Python-reads seam — the
    header / blank line / tree layout and its UTF-8 encoding are implemented
    independently on both sides, which is the arrangement that drifts."""
    with tempfile.TemporaryDirectory() as d:
        target = Path(d) / "target.ahk"
        target.write_text(
            '#Requires AutoHotkey v2.0\n'
            '#SingleInstance Off\n'
            'g := Gui("+AlwaysOnTop", "VKConformanceUia_ZZ")\n'
            'g.AddText("w220", "Probe Label Zz")\n'
            'g.AddButton("w220", "Probe Button Zz")\n'
            'g.Show("x0 y0 NoActivate")\n'          # do not steal the user's focus
            'SetTimer(() => ExitApp(0), -60000)\n',
            encoding="utf-8")
        p = subprocess.Popen([vk.AHK_EXE, "/ErrorStdOut", str(target)])
        try:
            got = vk.dump_uia_tree("VKConformanceUia_ZZ", max_depth=6, max_lines=100)
            assert got["window_title"] == "VKConformanceUia_ZZ", got
            assert 'Button "Probe Button Zz"' in got["tree"], got["tree"]
            assert got["lines"] >= 3, got
            assert "note" not in got, got            # nowhere near the cap

            # The filter narrows the report but keeps the walk (and the indent).
            filt = vk.dump_uia_tree("VKConformanceUia_ZZ", name_filter="Probe Button")
            assert 'Button "Probe Button Zz"' in filt["tree"], filt["tree"]
            assert "Probe Label Zz" not in filt["tree"], filt["tree"]

            # A window that doesn't exist is an error, not an empty dump.
            try:
                vk.dump_uia_tree("NoSuchWindowZzQ_13579")
                assert False, "a missing window must raise"
            except vk.VoiceKitError as e:
                assert "not found" in str(e).lower()
        finally:
            p.terminate()
            p.wait(timeout=10)      # the temp dir can't be removed while the
                                    # target still holds its script file


def test_uia_probe_focus_report_parses():
    """inspect_focus returns a parseable report about the REAL desktop (it is
    read-only, so running it here is safe). Whatever happens to hold focus
    while the suite runs, the shape must hold: focused=True with the element
    fields, or focused=False with a note saying why."""
    got = vk.inspect_focus()
    assert isinstance(got, dict) and "focused" in got, got
    if got["focused"]:
        for k in ("name", "control_type", "control_type_id", "value",
                  "automation_id", "enabled", "rect", "ancestry"):
            assert k in got, f"missing {k}: {got}"
        assert isinstance(got["ancestry"], list), got
    else:
        assert got.get("note"), got


def test_created_files_have_bom_and_lf():
    """AHK writes newly-created files with a UTF-8 BOM and LF endings."""
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "x.txt"
        vk._write_new(p, "line1\nline2\n")
        raw = p.read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf"), "missing UTF-8 BOM on created file"
        assert b"\r\n" not in raw, "created file must use LF, not CRLF"


_GEN_DATE = "2031-02-03"            # pinned on both sides; never today's date
_GEN_BODY = 'MsgBox("x")\nSleep(10)   ; a comment\nx := "a `;b"'
_GEN_CASES = [
    ("launcher", "Morning Tabs", "MorningTabs", "Q", _GEN_DATE),
    ("launcher", 'Say "hi" `now', "SayHiNow", "[", _GEN_DATE),
    ("launcher", "Plan 9", "Plan9", "7", _GEN_DATE),
    ("body", "Morning Tabs", "MorningTabs", "Q", "", _GEN_DATE),
    ("body", 'Say "hi" `now', "SayHiNow", "[", "", _GEN_DATE),
    ("body", "Scrape Queue", "ScrapeQueue", "K", _GEN_BODY, _GEN_DATE),
    ("opens", "Team Board", "https://example.com/x?a=1&b=2", _GEN_DATE),
    ("opens", "Semi Colon", "https://example.com/?q=a ;b", _GEN_DATE),
    ("opens", "Progs", r"C:\Program Files", _GEN_DATE),         # exists + space: quoted
    ("opens", "Win Dir", r"C:\Windows", _GEN_DATE),
    ("opens", "Missing Dir", r"C:\No Such Dir Zz 7f3a", _GEN_DATE),
    ("opens", "Quote Tick", 'notepad.exe "C:\\a b\\c.txt" `x', _GEN_DATE),
    ("stub", "Open Public User Folder", "OpenPublicUserFolder", _GEN_DATE),
    ("stub", "Plan 9 Tabs", "Plan9Tabs", _GEN_DATE),
    ("clean", "meeting notes"), ("clean", "MEETING NOTES"), ("clean", "open2tabs"),
    ("clean", "meeting-notes!!"), ("clean", "  multiple   spaces\tand tabs  "),
    ("clean", "Caf\u00e9 na\u00efve \u2014 plan"), ("clean", " \n foo bar \n "), ("clean", "! foo !"), ("clean", ""),
    ("spaceout", "OpenPublicUserFolder"), ("spaceout", "Plan9Tabs"),
    ("spaceout", "VoiceKitHome"), ("spaceout", "ABCDef"), ("spaceout", "x"),
    ("snipenc", "a;b"), ("snipenc", "``"), ("snipenc", "line1\r\nline2\rline3\nline4"),
    ("snipenc", "tab\there ; not a comment"), ("snipenc", "caf\u00e9 \u2014 \u2713"),
    ("snipenc", "`n stays literal"), ("snipenc", "{"),
    ("errmod", 'C:\\Automations\\VoiceKit\\hotkeys\\Foo.ahk (12) : ==> Missing "}"\n     Specifically: blah'),
    ("errmod", 'C:\\My Voice Kit\\hotkeys\\Bar Baz.ahk (3) : ==> Missing ")"'),
    ("errmod", 'C:\\Automations\\VoiceKit\\lib\\Acc.ahk (40) : ==> Missing "}"'),
    ("errmod", ""),
    ("errmod", 'C:\\VK\\hotkeys\\_index.ahk (3) : ==> #Include file "C:\\VK\\hotkeys\\Gone Module.ahk" cannot be opened.'),
    ("errmod", 'C:\\VK\\hotkeys\\Foo.ahk (2) : ==> #Include file "C:\\VK\\hotkeys\\helper.ahk" cannot be opened.'),
    ("errmod", 'C:\\VK\\hotkeys\\_index.ahk (4) : ==> Missing "}"'),
]


def _gen_python(case):
    kind, args = case[0], case[1:]
    if kind == "stub":
        return vk.workflow_stub(*args)
    if kind in ("clean", "spaceout", "snipenc", "errmod"):
        return {"clean": vk.clean_phrase, "spaceout": vk.space_out,
                "snipenc": vk.snip_encode, "errmod": vk._error_module}[kind](args[0])
    date = args[-1]
    with _patched(_today=lambda: date):
        if kind == "launcher":
            return vk._hotkey_launcher_content(*args[:-1])
        if kind == "body":
            return vk._hotkey_body_content(*args[:-1])
        if kind == "opens":
            return vk._opens_content(*args[:-1])
    raise AssertionError(f"no Python generator for {kind}")


def test_generators_match_ahk():
    """Every generator the MCP mirrors, run on BOTH sides and byte-compared:
    the isolated-module launcher and body, the Open Something macro, the
    workflow stub, and the CleanPhrase / SpaceOut / SnipEncode /
    MasterErrorModule ports. The AHK half runs the real functions through
    mcp\\_conformance\\generators.ahk with the same inputs and a pinned date
    (the stub is WfStubContent, the function the Studio's save calls). Then
    every generated script is load-checked in a sandbox tree laid out like
    the real one, since matching bytes are no use if the bytes don't load."""
    with _sandbox(libs=tuple(p.name for p in (REAL_ROOT / "lib").glob("*.ahk"))) as d:
        src, out = d / "cases.txt", d / "out.txt"
        src.write_bytes("\x1d".join("\x1e".join(c) for c in _GEN_CASES).encode("utf-8"))
        r = subprocess.run([vk.AHK_EXE, "/ErrorStdOut", str(GEN_HARNESS), str(src), str(out)],
                           capture_output=True, text=True, timeout=60)
        assert r.returncode == 0, f"generators.ahk failed ({r.returncode}): {r.stdout}{r.stderr}"
        got = out.read_bytes().decode("utf-8").split("\x1d")
        assert len(got) == len(_GEN_CASES), (len(got), len(_GEN_CASES))
        for case, ahk in zip(_GEN_CASES, got):
            py = _gen_python(case)
            assert ahk == py, (f"{case[0]} drifted for {case[1:]!r}:\n--- AHK ---\n{ahk}\n"
                               f"--- Python ---\n{py}")
        # The generated scripts load, each where it would really live.
        (d / "hotkeys" / "bodies").mkdir()
        where = {"launcher": d / "hotkeys", "body": d / "hotkeys" / "bodies",
                 "opens": d / "macros", "stub": d / "macros"}
        for i, (case, text) in enumerate(zip(_GEN_CASES, got)):
            if case[0] not in where:
                continue
            f = where[case[0]] / f"gen{i}.ahk"
            f.write_bytes(text.encode("utf-8"))
            ok, err = vk.validate_ahk(str(f))
            assert ok, f"generated {case[0]} {case[1:]!r} doesn't load:\n{err}"
        assert vk.STUDIO_MARKER in got[[c[0] for c in _GEN_CASES].index("stub")]


def test_naming_and_encoding_parity():
    assert vk.clean_phrase("meeting notes") == "Meeting Notes"
    assert vk.clean_phrase("MEETING NOTES") == "Meeting Notes"
    assert vk.clean_phrase("open2tabs") == "Open2tabs"          # StrTitle: no cap after digit
    assert vk.clean_phrase("meeting-notes!!") == "Meetingnotes"  # punctuation deleted, not spaced
    assert vk.to_base("Meeting Notes") == "MeetingNotes"
    assert vk.space_out("OpenPublicUserFolder") == "Open Public User Folder"
    for r in ("CON", "prn", "Nul", "COM1", "LPT9"):
        assert vk.is_reserved(r), r
    for ok in ("Notes", "Com", "Con1", ""):
        assert not vk.is_reserved(ok) or ok == "", ok
    # WfEncode/WfDecode inverse on nasty input
    for s in ("a|b%c\r\nx", "100% | done", "plain", "%25%7C%0A"):
        assert vk.wf_decode(vk.wf_encode(s)) == s


def test_allocate_bridge_key_is_free_and_in_pool():
    """Against a sandbox map (never the user's live one): taken keys —
    active or commented out — are skipped, and the answer is in the pool."""
    with _sandbox() as d:
        (d / "bridge-map.txt").write_text(
            "; map\nCtrl+Alt+Shift+A|a|hotkeys\\A.ahk|x\n"
            ";Ctrl+Alt+Shift+B|b|hotkeys\\B.ahk|x\n", encoding="utf-8")
        k = vk.allocate_bridge_key()
        assert k in vk.BRIDGE_POOL
        used = (d / "bridge-map.txt").read_text(encoding="utf-8-sig")
        assert f"Ctrl+Alt+Shift+{k}|" not in used, "allocator returned an in-use key"
        assert k not in "ABENRXHIC", f"allocator handed out a taken/reserved key: {k}"


def test_ai_template_validates_and_fills():
    """The committed AI-action template must load-check as-is (placeholders live
    only in comments/strings), and the fill must leave no {{...}} behind."""
    tpl = vk.REPO_ROOT / "templates" / "ai-template.ahk"
    ok, err = vk.validate_ahk(str(tpl))
    assert ok, f"ai-template.ahk failed /validate: {err}"
    filled = vk._ai_action_content("Fix Grammar", "FixGrammar")
    assert "{{" not in filled and "}}" not in filled, "unfilled placeholder left in AI action"
    assert '"open Fix Grammar"' in filled
    assert "FixGrammar.prompt.txt" in filled


def test_template_fill_preserves_line_endings():
    """Template-filled macros (AI action + placeholder launch/hotkey) must carry
    the SAME line endings the template has on disk — the AHK generators fill via
    FileRead+FileAppend with no EOL translation, so byte parity requires it."""
    def tpl_is_crlf(name):
        return b"\r\n" in (vk.REPO_ROOT / "templates" / name).read_bytes()

    ai = vk._ai_action_content("Fix Grammar", "FixGrammar")
    assert ("\r\n" in ai) == tpl_is_crlf("ai-template.ahk"), \
        "AI-action macro endings diverge from the template AHK reads"
    launch = vk._launch_content("Test Ping", None)          # None => template path
    assert ("\r\n" in launch) == tpl_is_crlf("launch-template.ahk"), \
        "launch placeholder endings diverge from its template"
    hotkey = vk._hotkey_content("Test Bridge", "A", None)   # None => template path
    assert ("\r\n" in hotkey) == tpl_is_crlf("hotkey-template.ahk"), \
        "hotkey placeholder endings diverge from its template"


def test_opens_guards():
    """opens rejects line breaks / empties (injection + dead-macro guards).
    (press_hotkey's malformed-key guard is exercised for real in
    test_fix_regressions_via_monkeypatch.)"""
    for bad in ("", "   ", "https://x\nMsgBox(1)", "a\rb"):
        try:
            vk.create_launch_macro("Guard Probe", opens=bad)
            assert False, f"opens={bad!r} should be rejected"
        except vk.VoiceKitError:
            pass


def test_ahk_str_lit_and_opens_content():
    """ahk_str_lit mirrors NewAutomation.ahk AhkStrLit; _opens_content keeps the
    ';  Opens:' header line the home window parses."""
    assert vk.ahk_str_lit('say "hi"') == '"say `"hi`""'
    assert vk.ahk_str_lit("tick`mark") == '"tick``mark"'
    assert vk.ahk_str_lit("a\r\nb") == '"a`nb"'
    assert vk.ahk_str_lit("a ;b") == '"a `;b"', "a space-preceded ';' starts a comment"
    c = vk._opens_content("Team Board", "https://example.com/x")
    assert ";  Opens:  https://example.com/x\n" in c
    assert 'Run("https://example.com/x")\n' in c
    # An existing path containing a space gets embedded quotes (Run needs them).
    spaced = r"C:\Program Files"
    if Path(spaced).exists():
        c2 = vk._opens_content("Progs", spaced)
        assert f'Run("`"{spaced}`"")' in c2


_STRLIT_CASES = [
    "plain", 'say "hi"', "tick`mark", "a\r\nb", "SELECT 1 ; SELECT 2",
    "https://example.com/?q=a ;b", "semi;colon", "`;", "``", '"', "tab\there",
    "C:\\Program Files\\x.exe", "multi\nline ; text `n literal",
]


def _run_common_probe(d: Path, body: str, *args: str) -> None:
    """Run a throwaway script that #Includes the SANDBOX copy of _Common.ahk
    (its uncaught-error logger then writes beside it, never into the user's
    errors.log)."""
    script = d / "probe.ahk"
    script.write_text("#Requires AutoHotkey v2.0\n#SingleInstance Off\n"
                      f'#Include "{d / "lib" / "_Common.ahk"}"\n' + body + "\nExitApp(0)\n",
                      encoding="utf-8")
    r = subprocess.run([vk.AHK_EXE, "/ErrorStdOut", str(script), *args],
                       capture_output=True, text=True, timeout=30)
    assert r.returncode == 0, r.stdout + r.stderr


# One table for both implementations of the fill label rule (the AHK suite
# tests\engine-fill-selftest.ahk checks the same rows against the engine).
FILL_LABELS = ["Amount", "Amount#2", "Amount #2", "Line ##4", "Line ###4", "Box #A",
               "A##B", "  Name  ", "#3", "Amount#0", "Tax # 12", "x#1234567", "##", "a#b#3"]


def test_fill_label_matches_the_engine():
    """fill_label is a second implementation of lib\\Workflow.ahk WfFillLabel
    (the writer's guard refuses what the engine would refuse at run time).
    Run the REAL AHK function on the same labels and compare field by field."""
    sep = "\x1e"
    with tempfile.TemporaryDirectory() as d:
        src, out, script = Path(d) / "in.txt", Path(d) / "out.txt", Path(d) / "probe.ahk"
        src.write_text(sep.join(FILL_LABELS), encoding="utf-8")
        script.write_text(
            "#Requires AutoHotkey v2.0\n#SingleInstance Off\n#NoTrayIcon\n"
            f'#Include "{REAL_ROOT / "lib" / "Workflow.ahk"}"\n'
            "o := ''\n"
            f'for s in StrSplit(FileRead(A_Args[1], "UTF-8"), Chr({ord(sep)})) {{\n'
            "    L := WfFillLabel(s)\n"
            "    o .= L.name '|' L.occ '|' (L.explicit ? 1 : 0) '|' (L.err != '' ? 1 : 0) '`n'\n"
            "}\n"
            'FileAppend(o, A_Args[2], "UTF-8")\nExitApp(0)\n', encoding="utf-8")
        r = subprocess.run([vk.AHK_EXE, "/ErrorStdOut", str(script), str(src), str(out)],
                           capture_output=True, text=True, timeout=30)
        assert r.returncode == 0, r.stdout + r.stderr
        got = out.read_text(encoding="utf-8-sig").split("\n")[:len(FILL_LABELS)]
    for raw, line in zip(FILL_LABELS, got):
        py = vk.fill_label(raw)
        want = f"{py['name']}|{py['occ']}|{1 if py['explicit'] else 0}|{1 if py['err'] else 0}"
        assert line == want, (raw, line, want)


def test_fill_inputs_match_the_engine():
    """A {{Name}} a fill VALUE uses that nothing defines is an INPUT of the
    workflow (the batch path feeds only inputs, and an ask step would type its
    answer somewhere). _input_labels decides what a batch must answer; the
    engine's WfAskLabels decides what the loop reads — run both on one
    workflow, through the real WorkflowLoad, and compare."""
    steps = [("ask", "Client", "", ""), ("set", "Later", "x", ""),
             ("collect", "Seen", "", ""),
             ("fill", "w", "{{Label Only}}", "{{Client}} {{Box 1}} {{Later}} {{date}} {{box 1}}"),
             ("fill", "w", "Other", "{{ Box 2 }}|{{Seen}}"), ("text", "{{Typo}}", "", ""),
             ("ask", "", "", ""), ("fill", "w", "Z", "{{selected_file}}{{Input}}")]
    with tempfile.TemporaryDirectory() as d:
        sf, out, script = Path(d) / "x.steps.txt", Path(d) / "out.txt", Path(d) / "probe.ahk"
        vk._write_new(sf, vk.steps_to_text("Probe", steps))
        script.write_text(
            "#Requires AutoHotkey v2.0\n#SingleInstance Off\n#NoTrayIcon\n"
            f'#Include "{REAL_ROOT / "lib" / "Workflow.ahk"}"\n'
            "o := ''\n"
            "for l in WfAskLabels(WorkflowLoad(A_Args[1]))\n"
            "    o .= l '`n'\n"
            'FileAppend(o, A_Args[2], "UTF-8")\nExitApp(0)\n', encoding="utf-8")
        r = subprocess.run([vk.AHK_EXE, "/ErrorStdOut", str(script), str(sf), str(out)],
                           capture_output=True, text=True, timeout=30)
        assert r.returncode == 0, r.stdout + r.stderr
        ahk = out.read_text(encoding="utf-8-sig").split("\n")[:-1]
    py = vk._input_labels([{"type": t, "a": a, "b": b, "c": c} for t, a, b, c in steps])
    assert py == ["Client", "Input", "Box 1", "Box 2"], py
    assert ahk == py, (ahk, py)


def test_ahk_str_lit_matches_ahk():
    """ahk_str_lit is a second implementation of lib\\_Common.ahk AhkStrLit
    (every generated Run()/SendText() literal on both sides goes through
    one of them). Run the REAL AHK function on the same inputs and compare —
    and prove every result actually loads, ';' included: an unescaped ' ;'
    starts a comment even inside a string, which is how a Type Text hotkey
    for 'SELECT 1 ; SELECT 2' used to be wired in and then parked."""
    sep = "\x1e"
    with _sandbox() as d:
        src, out = d / "in.txt", d / "out.txt"
        src.write_text(sep.join(_STRLIT_CASES), encoding="utf-8", newline="")
        _run_common_probe(
            d,
            's := ""\n'
            'for v in StrSplit(FileRead(A_Args[1], "UTF-8"), Chr(30))\n'
            '    s .= (A_Index > 1 ? Chr(30) : "") AhkStrLit(v)\n'
            'f := FileOpen(A_Args[2], "w", "UTF-8-RAW"), f.Write(s), f.Close()',
            str(src), str(out))
        got = out.read_text(encoding="utf-8").split(sep)
        want = [vk.ahk_str_lit(v) for v in _STRLIT_CASES]
        for case, a, p in zip(_STRLIT_CASES, got, want):
            assert a == p, f"AhkStrLit drifted for {case!r}:\n  AHK   : {a}\n  Python: {p}"
        assert len(got) == len(want), (len(got), len(want))
        # Every literal loads, as an assignment and inside a Run() call.
        lits = d / "lits.ahk"
        lits.write_text("#Requires AutoHotkey v2.0\n" + "".join(
            f"v{i} := {lit}\nif false\n    Run({lit})\n" for i, lit in enumerate(want)),
            encoding="utf-8")
        ok, err = vk.validate_ahk(str(lits))
        assert ok, f"a generated literal doesn't load:\n{err}"
        # ...and the opener that used to be unreachable now gets created.
        got = vk.create_launch_macro("Semi Probe Zz", opens="https://example.com/?q=a ;b")
        assert Path(got["file"]).exists(), got


def test_builtin_tools_match_ahk():
    """VoiceKit's own tools are listed twice: lib\\_Common.ahk VkBuiltinTools /
    VkEditableTools (Home's tool rows, the Studio's refusal to save a workflow
    over a tool, Home's delete guard) and voicekit_writer._BUILTINS /
    _DELETE_PROTECTED_LOWER (list_automations' kind, the delete and edit
    guards). Nothing pinned them together, and both missed Split Pages. Run
    the AHK lists and compare."""
    with _sandbox() as d:
        out = d / "tools.txt"
        _run_common_probe(
            d,
            's := ""\n'
            'for k in VkBuiltinTools()\n    s .= "B|" k "`n"\n'
            'for k in VkEditableTools()\n    s .= "E|" k "`n"\n'
            's .= "caseless|" VkDeleteProtected("askai") VkDeleteProtected("SPLITPAGES")'
            ' VkDeleteProtected("MorningTabs") "`n"\n'
            'f := FileOpen(A_Args[1], "w", "UTF-8-RAW"), f.Write(s), f.Close()',
            str(out))
        rows = [ln.split("|", 1) for ln in out.read_text(encoding="utf-8").splitlines() if ln]
        builtins = {v for k, v in rows if k == "B"}
        editable = {v for k, v in rows if k == "E"}
        assert builtins == vk._BUILTINS, f"builtin tools drifted:\n  AHK   : {sorted(builtins)}\n  Python: {sorted(vk._BUILTINS)}"
        both = {b.lower() for b in builtins | editable}
        assert both == vk._DELETE_PROTECTED_LOWER, (sorted(both), sorted(vk._DELETE_PROTECTED_LOWER))
        assert "splitpages" in both and "splitpages" not in vk._BUILTINS_LOWER, \
            "Split Pages is delete-protected but stays editable"
        assert dict(rows)["caseless"] == "110", dict(rows)["caseless"]


def test_snippet_abbrev_backtick_is_refused():
    """A backtick is AHK's escape character: ':*:ab`::hi' is 'Invalid hotkey',
    and the preflight would then park Snippets.ahk — every snippet the user
    has. Refused up front, like lib\\_Common.ahk SnippetValidate does."""
    with _sandbox() as d:
        snip = d / "hotkeys" / "Snippets.ahk"
        snip.write_text("#Requires AutoHotkey v2.0\n:*:/keep::kept\n", encoding="utf-8")
        before = snip.read_bytes()
        for bad in ("ab`c", "`x", "a:b"):
            try:
                vk.create_snippet(bad, "hello")
                assert False, f"abbreviation {bad!r} must be refused"
            except vk.VoiceKitError as e:
                assert "backtick" in str(e) or "colon" in str(e), e
        assert snip.read_bytes() == before, "a refused snippet must not touch the file"


def test_delete_module_takes_its_status_and_stop_files():
    """Deleting a hotkey module takes its body-status line and any pending stop
    flag with it (lib\\_Common.ahk DeleteHotkeyModuleArtifacts does the same):
    a module later given the same name must not show the old one's progress."""
    with _sandbox() as d:
        hk = d / "hotkeys"
        (hk / "StatusProbeZz.ahk").write_text("statusProbe := 1\n", encoding="utf-8")
        (hk / "_index.ahk").write_text(
            '#Requires AutoHotkey v2.0\n#Include "%A_ScriptDir%\\hotkeys\\StatusProbeZz.ahk"\n',
            encoding="utf-8")
        (d / "bridge-map.txt").write_text(
            "Ctrl+Alt+Shift+Q|status probe zz|hotkeys\\StatusProbeZz.ahk|2026-09-30\n",
            encoding="utf-8")
        status = d / "logs" / "body-status-StatusProbeZz.txt"
        stop = d / "logs" / "body-stop-StatusProbeZz.flag"
        other = d / "logs" / "body-status-StatusProbeZzOther.txt"
        for f in (status, stop, other):
            f.write_text("20260930000000|766 done\n", encoding="utf-8")
        got = vk.delete_automation("StatusProbeZz", "hotkey_module")
        assert not status.exists() and not stop.exists(), got
        assert other.exists(), "another module's status line must stay"
        assert str(status) in got["removed"], got["removed"]


def test_snippet_guards_reject_before_writing():
    """The brace-block and duplicate guards fire before any file is touched
    (so these calls are safe in a conformance run)."""
    snip = vk.REPO_ROOT / "hotkeys" / "Snippets.ahk"
    before = snip.read_bytes() if snip.exists() else b""
    try:
        vk.create_snippet("/conformance_probe", "{")
        assert False, "a lone '{' expansion must be rejected"
    except vk.VoiceKitError:
        pass
    existing = vk.list_automations()["snippets"]
    if existing:
        for variant in (existing[0]["abbrev"], existing[0]["abbrev"].upper()):
            try:
                vk.create_snippet(variant, "dup")
                assert False, f"duplicate abbreviation must be rejected ({variant})"
            except vk.VoiceKitError:
                pass  # caseless like the GUI — hotstrings fire caselessly
    after = snip.read_bytes() if snip.exists() else b""
    assert before == after, "guards must not modify Snippets.ahk"


def test_wait_step_guard_matches_the_engine():
    """create_workflow accepts exactly the waits WfWaitMs accepts.

    A `wait` may be whole milliseconds or a range ("600-1400") the engine
    turns into Random(lo, hi). Both sides parse it independently, so the
    accepted set has to be pinned on the Python side too — a wait the writer
    lets through but the engine rejects fails halfway into a run instead of
    at create time. (The guard fires before any file is written, so these
    calls are safe in a conformance run.)"""
    for good in ("0", "1500", "600-1400", " 250 ", "600 - 1400"):
        assert vk.WAIT_MS_RE.match(good.strip()), f"should be accepted: {good!r}"
    for bad in ("", "soon", "1.5", "600-", "-600", "600-1400-2000", "1,500"):
        assert not vk.WAIT_MS_RE.match(bad.strip()), f"should be rejected: {bad!r}"
    for bad in ("soon", "", "1.5"):
        try:
            vk.create_workflow("Conformance Wait Probe Zz", [("wait", bad, "", "")])
            assert False, f"a wait of {bad!r} must be rejected"
        except vk.VoiceKitError as e:
            assert "wait" in str(e), e


def test_list_and_bridge_resolution():
    """list_automations categories exist; VoiceKit's tools are separated out;
    bridge entries resolve by phrase and by key letter."""
    listing = vk.list_automations()
    for k in ("workflows", "launch_macros", "ai_actions", "hotkey_modules",
              "snippets", "voicekit_tools"):
        assert k in listing, f"missing category {k}"
    tool_bases = {t["base"] for t in listing["voicekit_tools"]}
    if (vk.REPO_ROOT / "macros" / "VoiceKitHome.ahk").exists():
        assert "VoiceKitHome" in tool_bases
        home = next(t for t in listing["voicekit_tools"] if t["base"] == "VoiceKitHome")
        assert home["voice_phrase"] == "open voice kit"
    entries = vk.get_bridge_map()["entries"]
    if entries:
        e = entries[0]
        key = e["combo"].rsplit("+", 1)[-1]
        assert vk.resolve_bridge(e["phrase"])["combo"] == e["combo"]
        assert vk.resolve_bridge(key)["combo"] == e["combo"]
    try:
        vk.resolve_bridge("No Such Hotkey Exists Xyz")
        assert False, "unknown hotkey must raise"
    except vk.VoiceKitError:
        pass


def test_builtins_protected_from_delete():
    try:
        vk.delete_automation("Voice Kit Home", "launch_macro")
        assert False, "deleting a builtin must be refused"
    except vk.VoiceKitError as e:
        assert "part of VoiceKit" in str(e)


def test_tool_named_workflow_delete_keeps_the_tools_companion():
    """A legacy workflow saved under a tool's name shares the tool's base, so
    retiring 'its' companion hotkey would take the TOOL's shipped one
    (RecordMySteps.hotkey.ahk, ^!+W). AHK DeleteWorkflowArtifacts skips the
    retire for a tool; the writer refuses the whole delete for a tool name
    before touching anything — either way the companion must survive."""
    with _sandbox() as d:
        comp = d / "hotkeys" / "RecordMySteps.hotkey.ahk"
        comp.write_text("; companion\n", encoding="utf-8")
        line = "Ctrl+Alt+Shift+W|open record my steps|hotkeys\\RecordMySteps.hotkey.ahk|2026-07-23"
        (d / "bridge-map.txt").write_text(f"; map\n{line}\n", encoding="utf-8")
        (d / "hotkeys" / "_index.ahk").write_text(
            '#Include "%A_ScriptDir%\\hotkeys\\RecordMySteps.hotkey.ahk"\n', encoding="utf-8")
        steps = d / "workflows" / "RecordMySteps.steps.txt"
        steps.write_text("wait|100||\n", encoding="utf-8")
        try:
            vk.delete_automation("Record My Steps", "workflow")
            assert False, "a tool-named workflow delete must be refused"
        except vk.VoiceKitError as e:
            assert "part of VoiceKit" in str(e), e
        assert comp.exists() and line in (d / "bridge-map.txt").read_text(encoding="utf-8")
        assert "RecordMySteps.hotkey.ahk" in (d / "hotkeys" / "_index.ahk").read_text(
            encoding="utf-8"), "the companion's include must stay wired"


def test_delete_guards_are_nondestructive():
    """Every guard that protects against a bad delete fires BEFORE any file is
    touched (safe to exercise against the real repo)."""
    # Dynamic code-block snippet (/date ships in the repo): refuse, don't corrupt.
    snip = vk.REPO_ROOT / "hotkeys" / "Snippets.ahk"
    if snip.exists() and ":/date::" in snip.read_text(encoding="utf-8-sig"):
        before = snip.read_bytes()
        try:
            vk.delete_automation("/date", "snippet")
            assert False, "dynamic snippet delete must be refused"
        except vk.VoiceKitError as e:
            assert "code-block" in str(e)
        assert snip.read_bytes() == before, "guard must not modify Snippets.ahk"
    # Nonexistent names raise instead of reporting success.
    for typ in ("workflow", "hotkey_module"):
        try:
            vk.delete_automation("Never Existed Xyz Q", typ)
            assert False, f"{typ} delete of a nonexistent name must raise"
        except vk.VoiceKitError as e:
            assert "No " in str(e)
    # A non-workflow macro must not lose its .lnk to type='workflow'.
    listing = vk.list_automations()
    guardable = [e for e in listing["launch_macros"] + listing["ai_actions"]
                 if e["base"].lower() not in vk._DELETE_PROTECTED_LOWER]
    if guardable:
        try:
            vk.delete_automation(guardable[0]["name"], "workflow")
            assert False, "workflow delete must cross-guard non-workflow macros"
        except vk.VoiceKitError as e:
            assert "isn't a workflow" in str(e)


def test_fix_regressions_via_monkeypatch():
    """Regression tests for the reviewed fixes, isolated by monkeypatching."""
    # 1. run_automation refuses to relaunch a running Workflow Studio.
    orig_running = vk._ahk_script_running
    vk._ahk_script_running = lambda *a, **k: True
    try:
        vk.run_automation("Workflow Studio")
        assert False, "must refuse to clobber a running Studio"
    except vk.VoiceKitError as e:
        assert "already open" in str(e)
    finally:
        vk._ahk_script_running = orig_running
    # 2. resolve_bridge precedence: exact phrase beats an earlier key-letter hit.
    orig_map = vk.get_bridge_map
    vk.get_bridge_map = lambda: {"entries": [
        {"combo": "Ctrl+Alt+Shift+A", "phrase": "Toggle Timer",
         "file": "hotkeys\\ToggleTimer.ahk", "created": ""},
        {"combo": "Ctrl+Alt+Shift+B", "phrase": "A", "file": "hotkeys\\A.ahk", "created": ""},
    ], "free_keys": [], "reserved_keys": []}
    try:
        assert vk.resolve_bridge("A")["combo"] == "Ctrl+Alt+Shift+B", \
            "exact phrase must beat a coincidental key-letter match"
        assert vk.resolve_bridge("Toggle Timer")["combo"] == "Ctrl+Alt+Shift+A"
    finally:
        vk.get_bridge_map = orig_map
    # 3. press_hotkey rejects a malformed key before anything runs.
    orig_resolve = vk.resolve_bridge
    vk.resolve_bridge = lambda n: {"combo": 'Ctrl+Alt+Shift+a" "^!+c', "phrase": "x",
                                   "key": 'a" "^!+c', "module": "hotkeys\\X.ahk"}
    try:
        vk.press_hotkey("x")
        assert False, "malformed bridge key must be refused"
    except vk.VoiceKitError as e:
        assert "malformed" in str(e)
    finally:
        vk.resolve_bridge = orig_resolve
    # 4. 'loop <name>' resolves through run_automation (unknown -> clean error).
    try:
        vk.run_automation("loop no such workflow zz")
        assert False, "unknown loop target must raise"
    except vk.VoiceKitError as e:
        assert "to loop" in str(e)


def test_companion_hotkey_folds_into_listing_and_delete_cleans_up():
    """A companion hotkey (hotkeys\\<Base>.hotkey.ahk + _index + bridge-map
    lines, written by the home window's Hotkey button) must (a) appear as the
    automation's own "hotkey" field — never as a separate hotkey module, (b)
    count its key as used, and (c) be retired when the automation is deleted."""
    with _sandbox(reload=lambda: False) as d:
        base = "CompanionProbeZz"
        (Path(d) / "workflows" / f"{base}.steps.txt").write_text(
            "; probe\nwait|100||\n", encoding="utf-8")
        (Path(d) / "macros" / f"{base}.ahk").write_text(
            f";  {vk.STUDIO_MARKER}\n", encoding="utf-8")
        (Path(d) / "hotkeys" / f"{base}.hotkey.ahk").write_text(
            "#Requires AutoHotkey v2.0\n^!+K:: {\n}\n", encoding="utf-8")
        (Path(d) / "hotkeys" / "_index.ahk").write_text(
            "#Requires AutoHotkey v2.0\n"
            f'#Include "%A_ScriptDir%\\hotkeys\\{base}.hotkey.ahk"',
            encoding="utf-8")
        (Path(d) / "bridge-map.txt").write_text(
            f"Ctrl+Alt+Shift+K|open companion probe zz|hotkeys\\{base}.hotkey.ahk|2026-07-15\n",
            encoding="utf-8")

        listing = vk.list_automations()
        wf = next(w for w in listing["workflows"] if w["base"] == base)
        assert wf.get("hotkey") == "Ctrl+Alt+Shift+K", "companion must fold into its automation"
        assert all(not h["base"].lower().endswith(".hotkey")
                   for h in listing["hotkey_modules"]), \
            "companion must not be listed as a hotkey module of its own"
        assert "K" not in vk.get_bridge_map()["free_keys"], "companion key must count as used"

        vk.delete_automation("Companion Probe Zz", "workflow")
        assert not (Path(d) / "hotkeys" / f"{base}.hotkey.ahk").exists(), \
            "delete must remove the companion module"
        assert f"{base}.hotkey.ahk" not in (Path(d) / "hotkeys" / "_index.ahk").read_text(
            encoding="utf-8-sig"), "delete must unwire the companion from _index.ahk"
        assert f"{base}.hotkey.ahk" not in (Path(d) / "bridge-map.txt").read_text(
            encoding="utf-8-sig"), "delete must free the companion's bridge key"


def test_hotkey_module_replace_keeps_key_and_guards_reserved_names():
    """Re-creating a module REPLACES it (like create_workflow) instead of
    erroring — the asymmetry that used to force delete-then-recreate. The
    replace must keep the registered key (its Voice Access pairing is manual
    and can't be recreated), must not duplicate registry lines, and must
    un-quarantine a parked #Include. Reserved names stay refused."""
    with _sandbox(reload=lambda: False) as d:
        (Path(d) / "hotkeys" / "_index.ahk").write_text(
            "#Requires AutoHotkey v2.0\n", encoding="utf-8")
        (Path(d) / "bridge-map.txt").write_text("; map\n", encoding="utf-8")

        first = vk.create_hotkey_module("Replace Probe Zz", 'MsgBox("one")')
        key = first["bridge_key"]
        assert first["replaced"] is False
        assert first["voice_pairing"], "a new module must return pairing steps"

        second = vk.create_hotkey_module("Replace Probe Zz", 'MsgBox("two")')
        assert second["replaced"] is True, "re-creating must replace, not raise"
        assert second["bridge_key"] == key, \
            "a replace must keep the key — the Voice Access pairing is manual"
        assert 'MsgBox("one")' in vk.read_hotkey_module(
            "Replace Probe Zz", previous=True)["body"], \
            "the overwritten body must stay retrievable so a replace can be undone"
        assert 'MsgBox("two")' in (Path(d) / vk._body_rel("ReplaceProbeZz")).read_text(
            encoding="utf-8-sig"), "the new body must be on disk"
        assert not second["voice_pairing"], "a replace needs no re-pairing"

        idx = (Path(d) / "hotkeys" / "_index.ahk").read_text(encoding="utf-8-sig")
        bmap = (Path(d) / "bridge-map.txt").read_text(encoding="utf-8-sig")
        assert idx.count("ReplaceProbeZz.ahk") == 1, "replace duplicated an #Include"
        assert bmap.count("ReplaceProbeZz.ahk") == 1, "replace duplicated a bridge-map line"

        # Deleting must leave the registry files byte-compatible with what
        # the AHK side writes: RemoveLinesContaining keeps the UTF-8 BOM.
        (Path(d) / "bridge-map.txt").write_bytes(
            b"\xef\xbb\xbf; map\nCtrl+Alt+Shift+Z|Other|hotkeys\\Other.ahk|2026-01-01\n")
        vk._remove_matching_lines(Path(d) / "bridge-map.txt",
                                  lambda ln: "Other.ahk" in ln)
        assert (Path(d) / "bridge-map.txt").read_bytes().startswith(b"\xef\xbb\xbf"), \
            "a delete must not strip the BOM the AHK side preserves"

        # A module parked by the load-failure quarantine comes back on when
        # it is re-created — that IS the "I fixed it" signal.
        vk._comment_out_include("hotkeys\\ReplaceProbeZz.ahk", "quarantined test")
        assert vk.index_modules()["quarantined"] == ["hotkeys\\ReplaceProbeZz.ahk"]
        third = vk.create_hotkey_module("Replace Probe Zz", 'MsgBox("three")')
        assert vk.index_modules()["active"] == ["hotkeys\\ReplaceProbeZz.ahk"], \
            "re-creating a quarantined module must switch its #Include back on"
        assert third["bridge_key"] == key
        assert (Path(d) / "hotkeys" / "_index.ahk").read_text(
            encoding="utf-8-sig").count("ReplaceProbeZz.ahk") == 1

        # "Snippets" is refused before anything is written — replacing it
        # would wipe every snippet the user has. (_index needs no guard:
        # clean_phrase strips the underscore, so that name is unreachable.)
        (Path(d) / "hotkeys" / "Snippets.ahk").write_text(
            ":*:/keepme::still here\n", encoding="utf-8")
        try:
            vk.create_hotkey_module("Snippets", 'MsgBox("clobber")')
            assert False, "'Snippets' must be refused as a module name"
        except vk.VoiceKitError as e:
            assert "VoiceKit's own files" in str(e)
        assert "/keepme" in (Path(d) / "hotkeys" / "Snippets.ahk").read_text(
            encoding="utf-8-sig"), "the guard must not touch Snippets.ahk"
        assert vk._require_name("_index")[1] == "Index", \
            "clean_phrase must keep '_index' unreachable as a module base"


def test_bridge_pool_matches_ahk_and_honours_a_requested_key():
    """BRIDGE_POOL is a second implementation of lib\\_Common.ahk BridgeKeyPool
    and must not drift from it. Also: a caller can ask for a specific key
    (the pool has symbols now), a taken key is refused, and a replace keeps
    the key it already had rather than moving it."""
    ahk_src = (vk.REPO_ROOT / "lib" / "_Common.ahk").read_text(encoding="utf-8-sig")
    m = re.search(r"BridgeKeyPool\(\)\s*\{\s*return\s+\"([^\"]*)\"", ahk_src)
    assert m, "couldn't find BridgeKeyPool() in lib\\_Common.ahk"
    assert m.group(1) == vk.BRIDGE_POOL, (
        f"bridge pool drifted:\n  AHK   : {m.group(1)}\n  Python: {vk.BRIDGE_POOL}")
    for reserved in "ENRXHIC":
        assert reserved not in vk.BRIDGE_POOL, f"{reserved} is reserved by VoiceKit"
    assert "|" not in vk.BRIDGE_POOL, "'|' would break the bridge-map delimiter"
    assert '"' not in vk.BRIDGE_POOL and "`" not in vk.BRIDGE_POOL, \
        "quote/backtick would break the generated AHK string literals"

    # normalize_bridge_key: tolerant input, strict output.
    assert vk.normalize_bridge_key("j") == "J"
    assert vk.normalize_bridge_key("Ctrl+Alt+Shift+[") == "["
    assert vk.normalize_bridge_key(" ] ") == "]"
    for bad in ("E", "|", "F13", "", "ab"):
        try:
            vk.normalize_bridge_key(bad)
            assert False, f"'{bad}' must not be accepted as a bridge key"
        except vk.VoiceKitError:
            pass

    with _sandbox(reload=lambda: False) as d:
        (Path(d) / "hotkeys" / "_index.ahk").write_text(
            "#Requires AutoHotkey v2.0\n", encoding="utf-8")
        (Path(d) / "bridge-map.txt").write_text("; map\n", encoding="utf-8")

        got = vk.create_hotkey_module("Bracket Probe Zz", 'TrayTip("x")', key="[")
        assert got["bridge_key"] == "Ctrl+Alt+Shift+[", got["bridge_key"]
        assert "[" not in vk.get_bridge_map()["free_keys"], "the key must now read as used"

        try:
            vk.create_hotkey_module("Other Probe Zz", 'TrayTip("y")', key="[")
            assert False, "a taken key must be refused"
        except vk.VoiceKitError as e:
            assert "already taken" in str(e)

        # A replace keeps its key; asking for a different one is refused
        # rather than silently stranding the Voice Access pairing.
        same = vk.create_hotkey_module("Bracket Probe Zz", 'TrayTip("z")')
        assert same["bridge_key"] == "Ctrl+Alt+Shift+["
        try:
            vk.create_hotkey_module("Bracket Probe Zz", 'TrayTip("z")', key="J")
            assert False, "re-keying via replace must be refused"
        except vk.VoiceKitError as e:
            assert "keeps that key" in str(e)


def test_hotkey_replace_returns_hash_not_body():
    """A replace reports the overwritten body as sha256 + length + first lines
    and banks the full text in logs\\module-backups\\, instead of echoing it
    back — the echo was the single largest token sink of a real session (a
    ~300-line body replaced five times paid for the OLD body every time).
    read_hotkey_module(previous=True) is the undo path, so the promise the
    response used to carry still holds; delete removes the backup with the
    module."""
    import hashlib
    with _sandbox(reload=lambda: True):
        vk.create_hotkey_module("Hash Probe Zz", 'TrayTip("one")')
        try:
            vk.read_hotkey_module("Hash Probe Zz", previous=True)
            assert False, "no backup may exist before the first replace"
        except vk.VoiceKitError as e:
            assert "replaced" in str(e).lower(), e
        prev = vk.read_hotkey_module("Hash Probe Zz")["body"]

        got = vk.create_hotkey_module("Hash Probe Zz", 'TrayTip("two")')
        assert got["replaced"] is True, got
        assert "previous_body" not in got, "the full body must not be echoed"
        assert got["previous_body_sha256"] == hashlib.sha256(prev.encode("utf-8")).hexdigest()
        assert got["previous_body_length"] == len(prev)
        assert got["previous_body_first_lines"] == prev.splitlines()[:3]

        back = vk.read_hotkey_module("Hash Probe Zz", previous=True)
        assert back["body"] == prev, "the backup must hold the exact overwritten text"
        assert Path(got["previous_body_backup"]).exists()

        vk.delete_automation("Hash Probe Zz", "hotkey_module")
        assert not Path(got["previous_body_backup"]).exists(), \
            "delete must remove the backup with the module"


def test_edit_hotkey_module_splices_and_rolls_back():
    """edit_hotkey_module: exact-match splice with file-Edit-tool semantics.
    The uniqueness guards, the load-check-and-roll-back on a breaking edit,
    the swap-the-strings undo, and the no-reload rule for isolated bodies all
    get proven here — against the real /validate, so the rollback path runs
    for real."""
    import hashlib
    reloads = []
    with _sandbox(reload=lambda: reloads.append(1) or True):

        vk.create_hotkey_module("Edit Probe Zz",
                                'TrayTip("one")\nTrayTip("keep me")')
        before = vk.read_hotkey_module("Edit Probe Zz")["body"]
        reloads.clear()

        got = vk.edit_hotkey_module("Edit Probe Zz",
                                    'TrayTip("one")', 'TrayTip("two")')
        after = vk.read_hotkey_module("Edit Probe Zz")["body"]
        assert 'TrayTip("two")' in after and 'TrayTip("keep me")' in after
        assert 'TrayTip("one")' not in after
        assert got["isolated"] is True and "reloaded" not in got
        assert not reloads, "a body edit must not reload the master"
        assert got["body_sha256"] == hashlib.sha256(after.encode("utf-8")).hexdigest()

        # Guards: no match / ambiguous / identical / reserved / missing.
        for bad, why in [
            (('TrayTip("gone")', "x"), "doesn't appear"),
            (("TrayTip", "x"), "times"),
            (('TrayTip("two")', 'TrayTip("two")'), "identical"),
        ]:
            try:
                vk.edit_hotkey_module("Edit Probe Zz", *bad)
                assert False, f"must refuse: {why}"
            except vk.VoiceKitError as e:
                assert why in str(e), (why, str(e))
        for name in ("Snippets", "No Such Module Zz"):
            try:
                vk.edit_hotkey_module(name, "a", "b")
                assert False, f"must refuse editing '{name}'"
            except vk.VoiceKitError:
                pass

        # A breaking edit is rolled back by the real load check.
        try:
            vk.edit_hotkey_module("Edit Probe Zz",
                                  'TrayTip("two")', 'TrayTip("two"')
            assert False, "a splice that kills the parse must be refused"
        except vk.VoiceKitError as e:
            assert "rolled back" in str(e)
        assert vk.read_hotkey_module("Edit Probe Zz")["body"] == after, \
            "the file must be exactly as it was before the bad edit"

        # Undo is the same call with the strings swapped.
        vk.edit_hotkey_module("Edit Probe Zz", 'TrayTip("two")', 'TrayTip("one")')
        assert vk.read_hotkey_module("Edit Probe Zz")["body"] == before

        # In-process module: the module file is the target, and the master
        # IS reloaded (that's where isolate=False code runs).
        vk.create_hotkey_module("Inline Probe Zz", 'MsgBox("aa")', isolate=False)
        reloads.clear()
        got = vk.edit_hotkey_module("Inline Probe Zz", 'MsgBox("aa")', 'MsgBox("bb")')
        assert got["isolated"] is False and got["reloaded"] is True
        assert reloads, "an in-process edit must reload the master"
        assert 'MsgBox("bb")' in vk.read_hotkey_module("Inline Probe Zz")["body"]


def test_update_hotkey_module_replaces_verbatim_and_guards():
    """update_hotkey_module: the full-file counterpart of edit_hotkey_module.
    The file read_hotkey_module returns is written VERBATIM (no re-wrapping,
    unlike a create replace — so read -> tweak -> update can never nest one
    generated header inside another), load-checked with byte-exact rollback,
    and the overwritten text banked in the same module-backups slot a create
    replace uses. An in-process module must keep defining its registered
    combo (any modifier order), or the write is refused before touching disk."""
    import hashlib
    reloads = []
    with _sandbox(reload=lambda: reloads.append(1) or True):

        # --- isolated module: the BODY is replaced verbatim, no reload ---
        vk.create_hotkey_module("Upd Probe Zz", 'TrayTip("one")')
        before = vk.read_hotkey_module("Upd Probe Zz")["body"]
        new_body = before.replace('TrayTip("one")', 'TrayTip("two")')
        reloads.clear()
        got = vk.update_hotkey_module("Upd Probe Zz", new_body)
        after = vk.read_hotkey_module("Upd Probe Zz")["body"]
        assert after == new_body, "the file must be written verbatim (no re-wrap)"
        assert got["updated"] is True and got["isolated"] is True
        assert "reloaded" not in got and not reloads, \
            "a body update must not reload the master"
        assert got["bridge_key"].startswith("Ctrl+Alt+Shift+")
        assert "previous_body" not in got, "the full old body must not be echoed"
        assert got["previous_body_sha256"] == \
            hashlib.sha256(before.encode("utf-8")).hexdigest()
        assert vk.read_hotkey_module("Upd Probe Zz", previous=True)["body"] == before, \
            "the bank must hold the exact overwritten text"

        # --- a breaking rewrite is rolled back by the real load check,
        #     and the bank still holds the last GOOD previous ---
        vk.update_hotkey_module("Upd Probe Zz", after)   # banks `after`
        try:
            vk.update_hotkey_module(
                "Upd Probe Zz", after.replace('TrayTip("two")', 'TrayTip("two"'))
            assert False, "a rewrite that kills the parse must be refused"
        except vk.VoiceKitError as e:
            assert "left exactly as it was" in str(e), e
        assert vk.read_hotkey_module("Upd Probe Zz")["body"] == after, \
            "the file must be byte-identical after the failed update"
        assert vk.read_hotkey_module("Upd Probe Zz", previous=True)["body"] == after, \
            "a failed update must not touch the bank"

        # --- guards: empty code / reserved name / unknown name ---
        for name, code, why in [
            ("Upd Probe Zz", "   \n", "empty"),
            ("Snippets", 'TrayTip("x")', "VoiceKit's own"),
            ("No Such Module Zz", 'TrayTip("x")', "No hotkey module"),
        ]:
            try:
                vk.update_hotkey_module(name, code)
                assert False, f"must refuse: {why}"
            except vk.VoiceKitError as e:
                assert why in str(e), (why, str(e))

        # --- in-process: the combo must survive, and the master reloads ---
        made = vk.create_hotkey_module("Inline Upd Zz", 'MsgBox("aa")',
                                       isolate=False)
        key = made["bridge_key"].rsplit("+", 1)[-1]
        src = vk.read_hotkey_module("Inline Upd Zz")["body"]
        try:
            vk.update_hotkey_module(
                "Inline Upd Zz",
                '#Requires AutoHotkey v2.0\nMsgBox("no binding")')
            assert False, "must refuse in-process code that drops the combo"
        except vk.VoiceKitError as e:
            assert f"Ctrl+Alt+Shift+{key}" in str(e), e
        assert vk.read_hotkey_module("Inline Upd Zz")["body"] == src, \
            "a combo rejection must happen before anything is written"
        reloads.clear()
        got = vk.update_hotkey_module(
            "Inline Upd Zz", src.replace('MsgBox("aa")', 'MsgBox("bb")'))
        assert got["isolated"] is False and got["reloaded"] is True
        assert reloads, "an in-process update must reload the master"
        assert 'MsgBox("bb")' in vk.read_hotkey_module("Inline Upd Zz")["body"]
        # A different modifier ORDER still counts as defining the combo.
        vk.update_hotkey_module(
            "Inline Upd Zz",
            src.replace('MsgBox("aa")', 'MsgBox("cc")')
               .replace(f"^!+{key}::", f"!+^{key}::"))
        assert 'MsgBox("cc")' in vk.read_hotkey_module("Inline Upd Zz")["body"]


def test_run_ahk_snippet_runs_and_reports():
    """run_ahk_snippet: one-off code in a throwaway process, reporting through
    the Out() file channel (GUI-subsystem stdout is not trustworthy). Uses the
    REAL repo root so the pre-included lib files resolve; the snippet itself
    only writes inside its temp folder, so this stays non-invasive."""
    got = vk.run_ahk_snippet('Out("hello " (1+1))\nOut("second line")')
    assert got["exit_code"] == 0 and got["timed_out"] is False, got
    assert got["output"] == "hello 2\nsecond line", got

    # ExitApp code comes back as-is.
    assert vk.run_ahk_snippet("ExitApp(7)")["exit_code"] == 7

    # The pre-included libs really are there (UiaControlTypeName needs no COM).
    got = vk.run_ahk_snippet('Out(UiaControlTypeName(50000))')
    assert got["output"] == "Button", got

    # A snippet that won't load reports the error instead of pretending.
    got = vk.run_ahk_snippet("this is not ahk(")
    assert got["exit_code"] not in (0, None), got
    assert got["errors"], "the load error text must come back"

    # A wedged snippet is killed at the timeout, and the partial log survives.
    got = vk.run_ahk_snippet('Out("started")\nSleep(60000)', timeout_s=2)
    assert got["timed_out"] is True, got
    assert got["output"] == "started", got

    try:
        vk.run_ahk_snippet("   ")
        assert False, "empty code must be refused"
    except vk.VoiceKitError:
        pass


def test_run_ahk_snippet_surfaces_uncaught_errors():
    """2026-08-03 feedback: an uncaught AHK error left NOTHING behind — a
    GUI-subsystem process pops a modal error dialog (measured: /ErrorStdOut
    covers load errors only), so the snippet sat the whole timeout and the
    report was empty; a real session lost 20 minutes binary-searching its own
    code with Out() markers. The prelude's OnError must turn that into an
    immediate exit 3 that names the error and the SNIPPET's own line, keeps
    everything Out() wrote before it, and never waits for the timeout."""
    t0 = time.time()
    got = vk.run_ahk_snippet('Out("step1")\nv := ""\nv.NoSuchMethod()\nOut("never")',
                             timeout_s=30)
    assert time.time() - t0 < 20, "an uncaught error must exit at once, not wait out the timeout"
    assert got["timed_out"] is False and got["exit_code"] == 3, got
    assert got["output"].startswith("step1\n"), "Out() lines before the error must survive"
    assert "never" not in got["output"], got
    assert "NoSuchMethod" in got["output"] and "snippet line 3" in got["output"], got
    assert "NoSuchMethod" in got["errors"], "the error must also land in 'errors'"
    assert "uncaught" in got.get("note", "").lower(), got

    # The reported case verbatim: StrSplit on an object (what UiaRect returns).
    got = vk.run_ahk_snippet('StrSplit({x: 1}, ",")')
    assert got["exit_code"] == 3 and "StrSplit" in got["errors"], got
    assert "snippet line 1" in got["errors"], got

    # Load errors point at the snippet's own line too, not line ~67 of the
    # generated wrapper file.
    got = vk.run_ahk_snippet('Out("fine")\nif (')
    assert got["exit_code"] == 2, got
    assert "snippet line 2" in got["errors"], got


def test_run_ahk_snippet_out_serializes_objects():
    """Out() takes any value now. It used to be a bare FileAppend, so the
    natural move — Out(UiaRect(el)) — THREW inside Out and fed the silent
    uncaught-error hole above. Objects print their properties (alphabetical:
    OwnProps enumeration order), arrays and Maps their items, and COM
    wrappers at least their type name instead of a crash."""
    got = vk.run_ahk_snippet(
        'Out({x: 12, y: 34, w: 100, h: 40})\n'
        'Out([1, 2, "three"])\n'
        'Out(Map("a", 1))\n'
        'Out(3.5)')
    assert got["exit_code"] == 0 and got["timed_out"] is False, got
    lines = got["output"].split("\n")
    assert lines[0].startswith("{") and lines[0].endswith("}"), got
    for prop in ("x: 12", "y: 34", "w: 100", "h: 40"):
        assert prop in lines[0], got
    assert lines[1] == "[1, 2, three]", got
    assert lines[2] == "Map{a: 1}", got
    assert lines[3] == "3.5", got


def test_mcp_cancellation_is_not_connection_fatal():
    """2026-08-03 feedback, the costliest failure: ONE client-side cancel of a
    slow dump_uia_tree killed the MCP connection for the rest of the session.
    The SDK answered the cancelled id with an error reply ("Request
    cancelled"); the client had already dropped the id, took the orphaned
    reply as a protocol fault, and closed the pipes — every later call was
    -32000 Connection closed. server.py now suppresses that reply (the spec's
    position: after a cancellation the receiver SHOULD NOT respond), and this
    proves it ON THE WIRE against the real server over real stdio: cancel an
    in-flight tools/call and assert (a) no response for that id is ever
    written — not even once the abandoned worker finishes — and (b) the same
    connection still answers afterwards."""
    import json
    import sys
    import threading

    mcp_dir = vk.REPO_ROOT / "mcp"
    py = mcp_dir / ".venv" / "Scripts" / "python.exe"
    if not py.exists():
        py = Path(sys.executable)
    if subprocess.run([str(py), "-c", "import fastmcp"],
                      capture_output=True).returncode != 0:
        print("      (skipped: fastmcp not importable — the MCP add-on isn't installed)")
        return

    proc = subprocess.Popen(
        [str(py), str(mcp_dir / "server.py")],
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        cwd=str(mcp_dir))
    replies: list = []
    stderr_lines: list = []
    lock = threading.Lock()

    def _read():
        for raw in proc.stdout:
            line = raw.decode("utf-8", "replace").strip()
            if not line:
                continue
            try:
                msg = json.loads(line)
            except ValueError:
                continue
            with lock:
                replies.append(msg)

    def _read_err():
        # Drained so the server can't block on a full stderr pipe; kept so a
        # failure can SHOW the server's own traceback instead of a bare
        # "no reply" (this test failed exactly once, under full-suite load,
        # with nothing to say for itself — never again).
        for raw in proc.stderr:
            stderr_lines.append(raw.decode("utf-8", "replace").rstrip())

    threading.Thread(target=_read, daemon=True).start()
    threading.Thread(target=_read_err, daemon=True).start()

    def send(obj):
        proc.stdin.write(json.dumps(obj).encode("utf-8") + b"\n")
        proc.stdin.flush()

    def wait_for_id(rid, deadline_s):
        deadline = time.time() + deadline_s
        while time.time() < deadline:
            with lock:
                for m in replies:
                    if m.get("id") == rid:
                        return m
            time.sleep(0.1)
        return None

    def diagnostics():
        tail = [ln for ln in stderr_lines if ln.strip()][-8:]
        with lock:
            seen = [m.get("id") for m in replies]
        return f"reply ids seen: {seen}; server stderr tail: {tail}"

    try:
        send({"jsonrpc": "2.0", "id": 1, "method": "initialize",
              "params": {"protocolVersion": "2024-11-05", "capabilities": {},
                         "clientInfo": {"name": "vk-conformance", "version": "0"}}})
        assert wait_for_id(1, 30), f"server never answered initialize — {diagnostics()}"
        send({"jsonrpc": "2.0", "method": "notifications/initialized"})

        # Slow enough that the cancellation always lands while the call is in
        # flight (the notification goes out within milliseconds; the snippet
        # costs an AHK process spawn plus 1.5 s) — but SHORT enough that the
        # abandoned worker is finished well inside the orphan window below,
        # so its teardown can never coincide with the follow-up request.
        send({"jsonrpc": "2.0", "id": 2, "method": "tools/call",
              "params": {"name": "run_ahk_snippet",
                         "arguments": {"code": "Sleep 1500", "timeout_s": 20}}})
        send({"jsonrpc": "2.0", "method": "notifications/cancelled",
              "params": {"requestId": 2, "reason": "client cancelled"}})

        # The 8 s window outlives the worker — a reply arriving after the
        # thread finishes would be the same connection-killer.
        late = wait_for_id(2, 8)
        assert late is None, (
            f"a cancelled request must never be answered — this orphaned reply "
            f"is what killed real sessions: {late}")

        # And the connection is still a connection.
        send({"jsonrpc": "2.0", "id": 3, "method": "tools/list"})
        listing = wait_for_id(3, 30)
        assert listing, f"no reply to tools/list after a cancellation — {diagnostics()}"
        assert listing.get("result", {}).get("tools"), (
            f"tools/list answered but not with tools: {listing} — {diagnostics()}")
        # The one flake this test ever had was the server tripping over its
        # own cancel ("Request already responded to", a real race — see
        # test_quiet_cancel_survives_a_late_respond): a traceback in stderr
        # is a failure even when the connection happened to survive it.
        bad = [ln for ln in stderr_lines if "already responded" in ln or "Traceback" in ln]
        assert not bad, f"the server raised while handling the cancel — {diagnostics()}"
    finally:
        try:
            proc.stdin.close()
        except OSError:
            pass
        try:
            proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            proc.kill()


def test_quiet_cancel_survives_a_late_respond():
    """The race behind the wire test's one flake ("Request already responded
    to" in the server's stderr): a sync tool runs in a worker thread that a
    cancel can't interrupt, so the handler can come back from it and reach
    respond() with no checkpoint in between — AFTER cancel() marked the
    request completed. Stock respond() asserts on that, and the assertion
    escaped into the server's task group. Driven deterministically here on
    a real RequestResponder: cancel, then respond — nothing may raise, no
    reply may be written, and the request leaves the in-flight table once."""
    import asyncio
    import importlib
    import sys
    try:
        import anyio  # noqa: F401
        from mcp.shared.session import RequestResponder
    except Exception:
        print("      (skipped: the mcp SDK isn't importable here)")
        return
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    try:
        server = importlib.import_module("server")
    except Exception as e:  # noqa: BLE001
        print(f"      (skipped: server.py not importable here: {e})")
        return
    # Importing server.py installed the patch; a False here would mean the
    # SDK's shape moved and the stock (reply-sending, asserting) path is live.
    assert (getattr(RequestResponder.respond, "__name__", "") == "_respond_unless_cancelled"
            and getattr(RequestResponder.cancel, "__name__", "") == "_cancel_without_reply"), \
        "the SDK shape changed — server.py's quiet-cancel patch no longer applies"

    sent, completed = [], []

    class FakeSession:
        async def _send_response(self, request_id, response):
            sent.append((request_id, response))

    async def scenario():
        r = RequestResponder(request_id=7, request_meta=None, request=None,
                             session=FakeSession(), on_complete=completed.append)
        with r:
            await r.cancel()                 # the client's notifications/cancelled
            await r.respond({"late": True})  # the worker finished anyway
        return r

    r = asyncio.run(scenario())
    assert sent == [], f"a cancelled request must never be answered: {sent}"
    assert completed == [r], "the request must leave the in-flight table exactly once"

    # An uncancelled request still answers normally.
    async def normal():
        r2 = RequestResponder(request_id=8, request_meta=None, request=None,
                              session=FakeSession(), on_complete=completed.append)
        with r2:
            await r2.respond({"ok": True})
    asyncio.run(normal())
    assert sent == [(8, {"ok": True})], sent


def test_list_running_and_stop_module():
    """The lifecycle pair, against REAL processes. The graceful path is the
    AHK-reads/Python-writes seam for the stop flag: stop_module (Python,
    _body_stop_file) drops the flag, BodyStopRequested (_Common.ahk,
    BodyStopFile) must derive the SAME name, consume it, and the body exits
    0 — if the two sides ever disagree, 'graceful' becomes 'killed' and this
    fails. The kill path proves a body that never checks still stops, scoped
    to its own PID."""
    with _sandbox(reload=None) as d:
        bodies = Path(d) / "hotkeys" / "bodies"
        bodies.mkdir(parents=True)
        # Name parity first — cheaper to diagnose than a failed stop.
        assert vk._body_stop_file("My Module!").name == "body-stop-MyModule.flag"

        polite = bodies / "StopProbeZz.body.ahk"
        polite.write_text(
            "#Requires AutoHotkey v2.0\n"
            f'#Include "{REAL_COMMON}"\n'
            "root := A_Args[1]\n"
            "Loop 300 {\n"
            '    if BodyStopRequested("StopProbeZz", root)\n'
            "        ExitApp(0)\n"
            "    Sleep(100)\n"
            "}\n"
            "ExitApp(9)\n", encoding="utf-8")
        p = subprocess.Popen([vk.AHK_EXE, "/ErrorStdOut", str(polite), d])
        try:
            deadline = time.monotonic() + 15
            seen = None
            while time.monotonic() < deadline:
                seen = next((e for e in vk.list_running()["running"]
                             if e["base"] == "StopProbeZz"), None)
                if seen:
                    break
            assert seen, "the launched body never showed up in list_running"
            assert seen["pid"] == p.pid, seen
            assert seen["name"] == "Stop Probe Zz", seen

            got = vk.stop_module("Stop Probe Zz", grace_s=20)
            assert got["how"] == "graceful", got
            assert p.wait(timeout=10) == 0, "a graceful stop must exit via ExitApp(0)"
            assert not vk._body_stop_file("StopProbeZz").exists(), \
                "the flag must not leak into the next run"
            assert not any(e["base"] == "StopProbeZz"
                           for e in vk.list_running()["running"])
            try:
                vk.stop_module("Stop Probe Zz")
                assert False, "stopping a stopped module must say so"
            except vk.VoiceKitError as e:
                assert "no body process" in str(e)
        finally:
            if p.poll() is None:
                p.kill()
                p.wait(timeout=10)

        # A body that never checks the flag: killed, exit code nonzero.
        rude = bodies / "KillProbeZz.body.ahk"
        rude.write_text("#Requires AutoHotkey v2.0\nSleep(120000)\nExitApp(0)\n",
                        encoding="utf-8")
        p2 = subprocess.Popen([vk.AHK_EXE, "/ErrorStdOut", str(rude)])
        try:
            deadline = time.monotonic() + 15
            while time.monotonic() < deadline:
                if any(e["base"] == "KillProbeZz"
                       for e in vk.list_running()["running"]):
                    break
            got = vk.stop_module("Kill Probe Zz", grace_s=1)
            assert got["how"] == "killed", got
            assert p2.wait(timeout=10) != 0, "a force-kill can't exit 0"
            assert not vk._body_stop_file("KillProbeZz").exists()
        finally:
            if p2.poll() is None:
                p2.kill()
                p2.wait(timeout=10)


def test_body_single_instance_mutex():
    """BodySingleInstance: the duplicate-instance fix, against real processes.

    #SingleInstance compares hidden-window titles and a starting script only
    creates that window once loading finishes, so two quick launches both
    pass the check (measured — that race is how a session ended up with two
    scrapers walking the same grid). The kernel mutex is atomic: with the
    first instance holding it, a second launch of the same base must exit 0
    quietly BEFORE reaching the body's code, while a different base sails
    through."""
    with tempfile.TemporaryDirectory() as d:
        # #SingleInstance Off on purpose: the directive must not help, so a
        # second exit-0 here proves the MUTEX alone stops the duplicate. (Off
        # also avoids v2's DEFAULT — Prompt — whose modal dialog would wedge
        # the second instance before the mutex line ever ran; measured.)
        marker = Path(d) / "ran.txt"
        script = Path(d) / "MutexProbeZz.ahk"
        script.write_text(
            "#Requires AutoHotkey v2.0\n"
            "#SingleInstance Off\n"
            f'#Include "{REAL_COMMON}"\n'
            'BodySingleInstance("MutexProbeZz")\n'
            f'FileAppend("ran`n", "{marker}", "UTF-8")\n'
            'DllCall("Sleep", "UInt", 30000)\n', encoding="utf-8")
        p1 = subprocess.Popen([vk.AHK_EXE, "/ErrorStdOut", str(script)])
        p2 = None
        try:
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline and not marker.exists():
                time.sleep(0.1)
            assert marker.exists(), "the first instance never claimed the mutex"

            p2 = subprocess.Popen([vk.AHK_EXE, "/ErrorStdOut", str(script)])
            assert p2.wait(timeout=15) == 0, \
                "the losing instance must exit 0 quietly, like #SingleInstance Ignore"
            assert p1.poll() is None, "the holder must keep running"
            assert marker.read_text(encoding="utf-8-sig") == "ran\n", \
                "the losing instance must exit BEFORE the body's code runs"

            # A different base is a different mutex — it must not be blocked.
            other_marker = Path(d) / "other.txt"
            other = Path(d) / "OtherProbeZz.ahk"
            other.write_text(
                "#Requires AutoHotkey v2.0\n"
                "#SingleInstance Off\n"
                f'#Include "{REAL_COMMON}"\n'
                'BodySingleInstance("OtherProbeZz")\n'
                f'FileAppend("ran`n", "{other_marker}", "UTF-8")\n'
                "ExitApp(0)\n", encoding="utf-8")
            p3 = subprocess.Popen([vk.AHK_EXE, "/ErrorStdOut", str(other)])
            assert p3.wait(timeout=15) == 0
            assert other_marker.exists(), "a different base must not be blocked"
        finally:
            for p in (p1, p2):
                if p is not None and p.poll() is None:
                    p.kill()
                    p.wait(timeout=10)


def test_read_module_status_joins_line_and_liveness():
    """read_module_status names the four states. The one that matters is
    "line but no process": a body that died leaves a status line behind, and
    reporting that line without the liveness check reads as progress."""
    with _sandbox(reload=None) as d:
        bodies = Path(d) / "hotkeys" / "bodies"
        bodies.mkdir(parents=True)
        (Path(d) / "hotkeys" / "StatProbeZz.ahk").write_text("; launcher\n", encoding="utf-8")
        # Never ran: no process, no line.
        got = vk.read_module_status("Stat Probe Zz")
        assert got["running"] is False and got["status"] is None, got
        assert "never published" in got["note"], got

        # A body that publishes, then exits: line WITHOUT a process.
        body = bodies / "StatProbeZz.body.ahk"
        body.write_text(
            "#Requires AutoHotkey v2.0\n"
            "#SingleInstance Off\n"
            f'#Include "{REAL_COMMON}"\n'
            'BodyStatus("StatProbeZz", "12 / 766 - Acme", A_Args[1])\n'
            "ExitApp(0)\n", encoding="utf-8")
        r = subprocess.run([vk.AHK_EXE, "/ErrorStdOut", str(body), d],
                           capture_output=True, text=True, timeout=30)
        assert r.returncode == 0, r.stdout + r.stderr
        got = vk.read_module_status("Stat Probe Zz")
        assert got["running"] is False, got
        assert got["status"]["text"] == "12 / 766 - Acme", got
        assert "Not running" in got["note"], got

        # Running WITH a line: both halves, and the stuck-vs-working hint.
        live = bodies / "LiveProbeZz.body.ahk"
        (Path(d) / "hotkeys" / "LiveProbeZz.ahk").write_text("; launcher\n", encoding="utf-8")
        live.write_text(
            "#Requires AutoHotkey v2.0\n"
            "#SingleInstance Off\n"
            f'#Include "{REAL_COMMON}"\n'
            'BodyStatus("LiveProbeZz", "3 / 9 - working", A_Args[1])\n'
            'DllCall("Sleep", "UInt", 30000)\n', encoding="utf-8")
        p = subprocess.Popen([vk.AHK_EXE, "/ErrorStdOut", str(live), d])
        try:
            deadline = time.monotonic() + 15
            got = {}
            while time.monotonic() < deadline:
                got = vk.read_module_status("Live Probe Zz")
                if got["running"] and got["status"]:
                    break
                time.sleep(0.3)
            assert got["running"] is True and got["pids"] == [p.pid], got
            assert got["status"]["text"] == "3 / 9 - working", got
            assert "running_seconds" in got and "stuck" in got["note"], got
        finally:
            if p.poll() is None:
                p.kill()
                p.wait(timeout=10)

        try:
            vk.read_module_status("No Such Module Zz")
            assert False, "an unknown module must raise"
        except vk.VoiceKitError as e:
            assert "No hotkey module" in str(e)


def test_reference_bodies_are_readable_and_load():
    """Every shipped reference must be readable through the MCP surface AND
    actually load. A reference body carries a BODY's ..\\..\\lib\\ include
    paths, so it only load-checks from a hotkeys\\bodies\\ folder — which is
    exactly where it gets copied, and the reason this test builds that layout
    instead of validating it in place. A worked example that doesn't parse
    would be worse than none."""
    cat = vk.read_reference()
    assert any(r["name"] == "web-scrape" for r in cat["references"]), cat
    try:
        vk.read_reference("no-such-reference-zz")
        assert False, "an unknown reference must raise"
    except vk.VoiceKitError as e:
        assert "Available" in str(e)

    libs = ("_Common.ahk", "UIA.ahk", "Clip.ahk", "Browser.ahk")   # Browser.ahk includes UIA + Clip
    for entry in cat["references"]:
        got = vk.read_reference(entry["name"])
        assert got["code"].startswith("#Requires AutoHotkey v2.0"), entry
        with tempfile.TemporaryDirectory() as d:
            bodies = Path(d) / "hotkeys" / "bodies"
            bodies.mkdir(parents=True)
            (Path(d) / "lib").mkdir()
            for lib in libs:
                (Path(d) / "lib" / lib).write_bytes((vk.REPO_ROOT / "lib" / lib).read_bytes())
            copy = bodies / "RefProbeZz.body.ahk"
            copy.write_text(got["code"], encoding="utf-8-sig", newline="")
            ok, err = vk.validate_ahk(str(copy))
            assert ok, f"reference '{entry['name']}' does not load as a body: {err}"


def test_get_bridge_map_strips_padded_fields():
    """Hand-padded bridge-map lines (the file is the user's recreate list) must
    parse like the GUI, which Trim()s every field."""
    with _sandbox(reload=None) as d:
        (Path(d) / "bridge-map.txt").write_text(
            "Ctrl+Alt+Shift+B | Mute Mic | hotkeys\\MuteMic.ahk | 2026-01-02\n",
            encoding="utf-8")
        bm = vk.get_bridge_map()
        assert bm["entries"][0]["combo"] == "Ctrl+Alt+Shift+B"
        assert bm["entries"][0]["phrase"] == "Mute Mic"
        assert "B" not in bm["free_keys"], "padded key must count as used"
        assert vk.resolve_bridge("Mute Mic")["key"] == "B"


def test_macro_source_read_update_edit_and_kinds():
    """2026-08-05 feedback (the Split Pages debugging run), items 1 + 3: the
    MCP could not read a macro's .ahk source (a human pasted 250 lines into
    chat) nor write a fix back (the patch was hand-applied), and the listing
    showed a full interactive application as just another launch macro,
    distinguishable only by a MISSING 'opens' field. read_macro_source /
    update_macro / edit_macro close the loop; list_automations now carries an
    explicit kind and surfaces a raw script's header comment."""
    import hashlib
    with _sandbox(reload=lambda: False) as d:
        raw = Path(d) / "macros" / "RawProbeZz.ahk"
        raw.write_text(
            "#Requires AutoHotkey v2.0\n"
            "#SingleInstance Force\n"
            "; ============================================================\n"
            ";  Raw Probe Zz — a hand-written interactive tool.\n"
            ";  It cycles pages and asks questions, like Split Pages.\n"
            "; ============================================================\n"
            'TrayTip("one")\n', encoding="utf-8")
        opens = Path(d) / "macros" / "OpensProbeZz.ahk"
        opens.write_text(vk._opens_content("Opens Probe Zz", "https://example.com/x"),
                         encoding="utf-8")
        (Path(d) / "workflows" / "StubProbeZz.steps.txt").write_text(
            "; probe\nwait|100||\n", encoding="utf-8")
        stub = Path(d) / "macros" / "StubProbeZz.ahk"
        stub.write_text(vk.workflow_stub("Stub Probe Zz", "StubProbeZz"),
                        encoding="utf-8")
        (Path(d) / "macros" / "AskAI.ahk").write_text("; builtin stand-in\n",
                                                      encoding="utf-8")

        # The taxonomy tells the truth now: explicit kind, header surfaced.
        listing = vk.list_automations()
        raw_e = next(e for e in listing["launch_macros"] if e["base"] == "RawProbeZz")
        assert raw_e["kind"] == "raw_script", raw_e
        assert "hand-written interactive tool" in raw_e.get("description", ""), raw_e
        assert "====" not in raw_e.get("description", ""), "banner borders must be dropped"
        opens_e = next(e for e in listing["launch_macros"] if e["base"] == "OpensProbeZz")
        assert opens_e["kind"] == "opens", opens_e
        assert opens_e["opens"] == "https://example.com/x", opens_e

        # Read: full source plus what the thing IS.
        got = vk.read_macro_source("Raw Probe Zz")
        assert got["kind"] == "raw_script" and 'TrayTip("one")' in got["source"], got
        assert "hand-written interactive tool" in got.get("description", ""), got
        got = vk.read_macro_source("Stub Probe Zz")
        assert got["kind"] == "workflow_stub", got
        assert "steps file is canonical" in got["note"], got

        # No backup before the first update.
        try:
            vk.read_macro_source("Raw Probe Zz", previous=True)
            assert False, "no backup may exist before the first update"
        except vk.VoiceKitError as e:
            assert "hasn't been updated" in str(e), e

        # Update: banked, hashed, undoable; live source replaced. (Two
        # TrayTip lines so the ambiguous-splice guard below has a real
        # ambiguity to refuse.)
        prev = vk.read_macro_source("Raw Probe Zz")["source"]
        got = vk.update_macro("Raw Probe Zz",
                              '#Requires AutoHotkey v2.0\nTrayTip("two")\n'
                              'TrayTip("keep me")\n')
        assert got["updated"] is True, got
        assert "previous_source" not in got, "the old source must not be echoed"
        assert got["previous_source_sha256"] == hashlib.sha256(
            prev.encode("utf-8")).hexdigest(), got
        assert got["previous_source_length"] == len(prev), got
        assert 'TrayTip("two")' in vk.read_macro_source("Raw Probe Zz")["source"]
        back = vk.read_macro_source("Raw Probe Zz", previous=True)
        assert back["source"] == prev, "the backup must hold the exact overwritten text"

        # A breaking update is rolled back by the real load check.
        before_bytes = raw.read_bytes()
        try:
            vk.update_macro("Raw Probe Zz", 'TrayTip("unterminated\n')
            assert False, "a source that fails /validate must be refused"
        except vk.VoiceKitError as e:
            assert "left exactly as it was" in str(e), e
        assert raw.read_bytes() == before_bytes, "rollback must restore exact bytes"

        # Guards: builtins, generated stubs, nonexistent, empty.
        for bad_name, why in [("Ask AI", "part of VoiceKit"),
                              ("Stub Probe Zz", "GENERATED stub"),
                              ("Never Existed Zz", "No macro named")]:
            try:
                vk.update_macro(bad_name, "#Requires AutoHotkey v2.0\n")
                assert False, f"update of '{bad_name}' must be refused"
            except vk.VoiceKitError as e:
                assert why in str(e), (bad_name, str(e))
        try:
            vk.update_macro("Raw Probe Zz", "   ")
            assert False, "an empty source must be refused"
        except vk.VoiceKitError as e:
            assert "empty" in str(e), e

        # Splice: applied, validated, rolled back, undone by swapping.
        got = vk.edit_macro("Raw Probe Zz", 'TrayTip("two")', 'TrayTip("three")')
        assert got["edited"] is True, got
        after = vk.read_macro_source("Raw Probe Zz")["source"]
        assert 'TrayTip("three")' in after and 'TrayTip("keep me")' in after
        assert 'TrayTip("two")' not in after
        try:
            vk.edit_macro("Raw Probe Zz", 'TrayTip("three")', 'TrayTip("three"')
            assert False, "a splice that kills the parse must be refused"
        except vk.VoiceKitError as e:
            assert "rolled back" in str(e), e
        assert vk.read_macro_source("Raw Probe Zz")["source"] == after
        for bad, why in [(('TrayTip("gone")', "x"), "doesn't appear"),
                         (("TrayTip", "x"), "times"),
                         (("", "x"), "empty"),
                         (('TrayTip("three")', 'TrayTip("three")'), "identical")]:
            try:
                vk.edit_macro("Raw Probe Zz", *bad)
                assert False, f"must refuse: {why}"
            except vk.VoiceKitError as e:
                assert why in str(e), (why, str(e))

        # Delete parity: the undo snapshot goes with the macro.
        assert Path(got_backup := vk._macro_backup_file("RawProbeZz")).exists()
        vk.delete_automation("Raw Probe Zz", "launch_macro")
        assert not Path(got_backup).exists(), \
            "delete must remove the update_macro backup with the macro"


def test_read_log_tails_and_filters():
    """2026-08-05 feedback item 2: Split Pages logged 'saved of pageCount'
    every run, and no tool could read it — the varying-failure-point
    signature of its race sat unread in created.log while the bug was
    reconstructed by interview. read_log: filter BEFORE tail (so a tool's
    last N runs survive unrelated traffic), clamped tail, and a missing file
    is an answer, not an error."""
    with _sandbox(reload=None) as d:
        logs = Path(d) / "logs"
        logs.mkdir(exist_ok=True)
        lines = [f"2026-08-05 10:{i:02d} | noise | line {i}" for i in range(60)]
        lines.insert(10, "2026-08-05 09:00 | split-pages | 3 of 12 | a.pdf")
        lines.insert(30, "2026-08-05 09:30 | split-pages | 7 of 12 | a.pdf")
        (logs / "created.log").write_text("\n".join(lines) + "\n", encoding="utf-8")

        got = vk.read_log()
        assert got["log"] == "activity" and len(got["lines"]) == 50, got
        assert got["total_lines"] == 62, got
        assert "note" in got, "a truncated read must say so"

        got = vk.read_log("activity", tail=10, filter="SPLIT-PAGES")
        assert got["matched_lines"] == 2, got
        assert [l for l in got["lines"]] == [
            "2026-08-05 09:00 | split-pages | 3 of 12 | a.pdf",
            "2026-08-05 09:30 | split-pages | 7 of 12 | a.pdf"], got

        got = vk.read_log("errors")
        assert got["lines"] == [] and "doesn't exist yet" in got["note"], got

        got = vk.read_log("activity", tail=99999)
        assert len(got["lines"]) == 62, "tail clamps to 500, not below the file size"

        try:
            vk.read_log("no-such-log")
            assert False, "an unknown log name must raise"
        except vk.VoiceKitError as e:
            assert "created.log" in str(e), "the error must list what exists"


def test_near_miss_errors_name_what_the_name_is():
    """2026-08-05 feedback item 4: "Looked for SplitPages.steps.txt" told the
    mechanism but not that the index KNEW SplitPages existed as a raw script.
    Errors that know about near-miss entities save entire reasoning loops:
    every not-found path now says what the name IS and which tool reads it."""
    with _sandbox(reload=None) as d:
        (Path(d) / "macros" / "RawProbeZz.ahk").write_text(
            "#Requires AutoHotkey v2.0\n; a tool\nTrayTip(1)\n", encoding="utf-8")
        (Path(d) / "hotkeys" / "HkProbeZz.ahk").write_text("; module\n", encoding="utf-8")
        (Path(d) / "workflows" / "WfProbeZz.steps.txt").write_text(
            "; wf\nwait|100||\n", encoding="utf-8")

        # The SplitPages case verbatim: a workflow tool pointed at a script.
        try:
            vk.read_workflow("Raw Probe Zz")
            assert False
        except vk.VoiceKitError as e:
            assert "hand-written script" in str(e) and "read_macro_source" in str(e), e
        for fn in (vk.read_workflow_sheet,
                   lambda n: vk.run_workflow_batch(n, [{"x": "1"}])):
            try:
                fn("Raw Probe Zz")
                assert False
            except vk.VoiceKitError as e:
                assert "read_macro_source" in str(e), e

        # A macro tool pointed at a hotkey module, and the reverse.
        try:
            vk.read_macro_source("Hk Probe Zz")
            assert False
        except vk.VoiceKitError as e:
            assert "hotkey module" in str(e) and "read_hotkey_module" in str(e), e
        try:
            vk.read_hotkey_module("Raw Probe Zz")
            assert False
        except vk.VoiceKitError as e:
            assert "read_macro_source" in str(e), e

        # run_automation on a module: the trigger tool is press_hotkey.
        try:
            vk.run_automation("Hk Probe Zz")
            assert False
        except vk.VoiceKitError as e:
            assert "press_hotkey" in str(e), e

        # A workflow reachable only by its steps file still gets named.
        try:
            vk.read_macro_source("Wf Probe Zz")
            assert False
        except vk.VoiceKitError as e:
            assert "read_workflow" in str(e), e

        # A name that matches nothing keeps the plain error.
        try:
            vk.read_workflow("Truly Nothing Zz")
            assert False
        except vk.VoiceKitError as e:
            assert "does exist" not in str(e), e


def _ahk_procs_mentioning(text: str) -> list:
    """PIDs of AutoHotkey / cmd processes whose command line contains `text`."""
    # A literal for a single-quoted PowerShell -like pattern: wildcards get a
    # backtick, a single quote is doubled.
    lit = re.sub(r"([`*?\[\]])", r"`\1", str(text)).replace("'", "''")
    r = vk._run_powershell(
        "Get-CimInstance Win32_Process -Filter \"Name='AutoHotkey64.exe' or Name='cmd.exe'\" | "
        "Where-Object { $_.CommandLine -like '*" + lit + "*' } | "
        "ForEach-Object { $_.ProcessId }", check=False)
    return [x for x in (r.stdout or "").split() if x.strip().isdigit()]


def test_delete_camelcase_module_leaves_no_dangling_lines():
    """clean_phrase('WebScrapeDemo') is 'Webscrapedemo'. NTFS ignores case, so
    the file used to be deleted while the case-sensitive line matchers kept its
    #Include and bridge-map line — and the dangling #Include then made every
    later reload fail its load check. Lines must match caselessly on the exact
    field (a longer name that merely CONTAINS it stays), and names resolve to
    the real on-disk stem (so the .lnk name keeps its capitals)."""
    with _sandbox() as d:
        hk = d / "hotkeys"
        (hk / "bodies").mkdir()
        (d / "VoiceKit.ahk").write_text(
            '#Requires AutoHotkey v2.0\n#Include "%A_ScriptDir%\\hotkeys\\_index.ahk"\n',
            encoding="utf-8")
        (hk / "WebScrapeDemo.ahk").write_text("demoProbe := 1\n", encoding="utf-8")
        (hk / "bodies" / "WebScrapeDemo.body.ahk").write_text("x := 1\n", encoding="utf-8")
        (hk / "WebScrapeDemoExtra.ahk").write_text("demoExtra := 1\n", encoding="utf-8")
        (hk / "_index.ahk").write_text(
            "#Requires AutoHotkey v2.0\n"
            '#Include "%A_ScriptDir%\\hotkeys\\WebScrapeDemo.ahk"\n'
            '#Include "%A_ScriptDir%\\hotkeys\\WebScrapeDemoExtra.ahk"\n', encoding="utf-8")
        (d / "bridge-map.txt").write_text(
            "; map\n"
            "Ctrl+Alt+Shift+Q | web scrape demo | hotkeys\\WebScrapeDemo.ahk | 2026-09-01\n"
            "Ctrl+Alt+Shift+V|web scrape demo extra|hotkeys\\WebScrapeDemoExtra.ahk|2026-09-01\n",
            encoding="utf-8")

        got = vk.delete_automation("WebScrapeDemo", "hotkey_module")
        idx = (hk / "_index.ahk").read_text(encoding="utf-8-sig")
        bm = (d / "bridge-map.txt").read_text(encoding="utf-8-sig")
        assert not (hk / "WebScrapeDemo.ahk").exists(), got
        assert not (hk / "bodies" / "WebScrapeDemo.body.ahk").exists(), got
        assert "WebScrapeDemo.ahk\"" not in idx, f"dangling #Include left behind:\n{idx}"
        assert "WebScrapeDemoExtra.ahk" in idx, "a longer name must not be caught"
        assert "hotkeys\\WebScrapeDemo.ahk" not in bm, f"bridge-map line left behind:\n{bm}"
        assert "WebScrapeDemoExtra.ahk" in bm
        ok, err = vk.validate_ahk(str(d / "VoiceKit.ahk"))
        assert ok, f"the sandbox master must still load after the delete:\n{err}"

        # A CamelCase WORKFLOW: its shortcuts are named by SpaceOut of the REAL
        # stem ('Demo Flow Zz'), which clean_phrase's 'Demoflowzz' can't recover.
        (d / "workflows" / "DemoFlowZz.steps.txt").write_text("wait|100||\n", encoding="utf-8")
        (d / "macros" / "DemoFlowZz.ahk").write_text(
            f";  {vk.STUDIO_MARKER}\n", encoding="utf-8")
        for lnk in ("Demo Flow Zz.lnk", "loop Demo Flow Zz.lnk"):
            (d / "startmenu" / lnk).write_bytes(b"")
        # An unsaved loop-results journal (lib\WorkflowLoop.ahk) goes with the
        # workflow — a same-named workflow created later must not inherit its
        # rows — while another workflow's journal is left alone.
        jdir = d / "logs" / "loop-journal"
        jdir.mkdir(parents=True, exist_ok=True)
        (jdir / "DemoFlowZz.20260930000000-1234.jnl").write_text("0|0|0\n", encoding="utf-8")
        (jdir / "DemoFlowZzOther.20260930000000-1234.jnl").write_text("0|0|0\n", encoding="utf-8")
        vk.delete_automation("DemoFlowZz", "workflow")
        assert not any((d / "startmenu").iterdir()), \
            f"shortcuts left behind: {list((d / 'startmenu').iterdir())}"
        assert not (d / "workflows" / "DemoFlowZz.steps.txt").exists()
        assert not (jdir / "DemoFlowZz.20260930000000-1234.jnl").exists(), "journal left behind"
        assert (jdir / "DemoFlowZzOther.20260930000000-1234.jnl").exists(), "wrong journal removed"


def test_delete_snippets_module_is_refused_untouched():
    """hotkeys\\Snippets.ahk holds every snippet the user has; deleting it as a
    'hotkey_module' must be refused before any file or line is touched."""
    with _sandbox() as d:
        snip = d / "hotkeys" / "Snippets.ahk"
        snip.write_text(":*:/zz::probe\n", encoding="utf-8")
        idx = d / "hotkeys" / "_index.ahk"
        idx.write_text('#Include "%A_ScriptDir%\\hotkeys\\Snippets.ahk"\n', encoding="utf-8")
        before = (snip.read_bytes(), idx.read_bytes())
        for name in ("Snippets", "snippets", " SNIPPETS "):
            try:
                vk.delete_automation(name, "hotkey_module")
                assert False, f"deleting {name!r} as a module must be refused"
            except vk.VoiceKitError as e:
                assert "snippet" in str(e).lower(), e
            assert (snip.read_bytes(), idx.read_bytes()) == before, "nothing may be touched"
        # The manifest itself stays unreachable too (no leading-underscore stems).
        try:
            vk.delete_automation("_index", "hotkey_module")
            assert False, "_index must never resolve to the manifest"
        except vk.VoiceKitError:
            pass
        assert idx.read_bytes() == before[1]


def test_preflight_parks_missing_and_warn_modules():
    """Python mirror of MasterPreflight: (1) a module whose FILE is gone makes
    AutoHotkey blame _index.ahk ('#Include file ... cannot be opened'), which
    can't park itself — the missing module's line must be parked instead;
    (2) a #Warn warning pops a (hidden) dialog even under /validate, so the
    check is bounded, the tree is killed, and the #Warn module is parked;
    (3) the rewritten manifest keeps LF endings like the AHK side."""
    orig_timeout = vk.VALIDATE_TIMEOUT_S
    with _sandbox() as d:
        vk.VALIDATE_TIMEOUT_S = 3
        try:
            hk = d / "hotkeys"
            (d / "VoiceKit.ahk").write_text(
                '#Requires AutoHotkey v2.0\n#Include "%A_ScriptDir%\\hotkeys\\_index.ahk"\n',
                encoding="utf-8")
            (hk / "Keep.ahk").write_text("keepProbe := 1\n", encoding="utf-8")
            (hk / "_index.ahk").write_text(
                "#Requires AutoHotkey v2.0\n"
                '#Include "%A_ScriptDir%\\hotkeys\\Keep.ahk"\n'
                '#Include "%A_ScriptDir%\\hotkeys\\Gone Module.ahk"\n', encoding="utf-8")
            miss = ('C:\\VK\\hotkeys\\_index.ahk (3) : ==> #Include file '
                    '"C:\\VK\\hotkeys\\Gone Module.ahk" cannot be opened.')
            assert vk._error_module(miss) == "hotkeys\\Gone Module.ahk"
            assert vk._error_module('C:\\VK\\hotkeys\\_index.ahk (4) : ==> Missing "}"') == ""
            inner = ('C:\\VK\\hotkeys\\Foo.ahk (2) : ==> #Include file '
                     '"C:\\VK\\hotkeys\\helper.ahk" cannot be opened.')
            assert vk._error_module(inner) == "hotkeys\\Foo.ahk"

            ok, parked, err = vk.master_preflight()
            raw = (hk / "_index.ahk").read_bytes()
            assert ok and parked == ["hotkeys\\Gone Module.ahk"], (ok, parked, err)
            assert b"\r" not in raw, "the parked manifest must stay LF, like the AHK side"
            txt = raw.decode("utf-8-sig")
            assert '; #Include "%A_ScriptDir%\\hotkeys\\Gone Module.ahk"' in txt, txt
            assert "missing" in txt and '\n#Include "%A_ScriptDir%\\hotkeys\\Keep.ahk"' in txt

            warn = d / "fx_warn_py.ahk"
            warn.write_text("#Requires AutoHotkey v2.0\n#Warn\nMsgBox(neverAssignedPy)\n",
                            encoding="utf-8")
            t0 = time.monotonic()
            ok, err = vk.validate_ahk(str(warn))
            took = time.monotonic() - t0
            assert not ok and err.startswith(vk.VALIDATE_BLOCKED) and "#Warn" in err, err
            assert took < 12, f"the bound didn't hold: {took:.1f}s"
            time.sleep(0.3)
            assert not _ahk_procs_mentioning("fx_warn_py.ahk"), "a blocked check left a process"
            t0 = time.monotonic()
            assert vk.validate_ahk(str(hk / "Keep.ahk"))[0] and time.monotonic() - t0 < 5

            (hk / "Warny.ahk").write_text("#Warn\nwarnProbe := neverAssignedPy2\n",
                                          encoding="utf-8")
            with open(hk / "_index.ahk", "a", encoding="utf-8", newline="") as f:
                f.write('#Include "%A_ScriptDir%\\hotkeys\\Warny.ahk"\n')
            ok, parked, err = vk.master_preflight()
            assert ok and parked == ["hotkeys\\Warny.ahk"], (ok, parked, err)
            assert '; #Include "%A_ScriptDir%\\hotkeys\\Warny.ahk"' in \
                (hk / "_index.ahk").read_text(encoding="utf-8-sig")
            assert not _ahk_procs_mentioning(str(d)), "preflight left a process behind"
        finally:
            vk.VALIDATE_TIMEOUT_S = orig_timeout


def test_validate_reads_error_text_as_utf8():
    """validate_ahk used bare /ErrorStdOut (ANSI) while run_ahk_snippet reads
    UTF-8: the 'Specifically:' excerpt of non-Latin code came back as '?'."""
    with tempfile.TemporaryDirectory() as d:
        bad = Path(d) / "fx_utf8.ahk"
        bad.write_text('#Requires AutoHotkey v2.0\nx := "日本語 Ωmega\n', encoding="utf-8-sig")
        ok, err = vk.validate_ahk(str(bad))
        assert not ok and "日本語" in err, err


def test_health_reports_only_current_state():
    """health() rides on every response, so stale state there is noise that
    reads as current: last_error only when this master generation hit it (with
    its age), safe mode from the flag file alone, and a refused reload
    forgotten once a master has started since."""
    orig = (vk.voicekit_running, vk._LAST_RELOAD_ERROR, vk._LAST_RELOAD_ERROR_AT)
    with _sandbox() as d:
        vk.voicekit_running = lambda: True
        try:
            ini = d / "logs" / "master-status.ini"

            def status(**kv):
                body = "[Master]\r\n" + "".join(f"{k}={v}\r\n" for k, v in kv.items())
                ini.write_bytes(b"\xff\xfe" + body.encode("utf-16-le"))   # IniWrite's form

            status(pid=1, started="20260901120000", heartbeat="20260901120000",
                   last_error="old boom", last_error_at="20260801000000", safe_mode=1)
            h = vk.health()
            assert "last_error" not in h, "an error from an earlier master is not news"
            assert "safe_mode" not in h, "safe mode is the flag file, not a stale ini value"

            status(pid=1, started="20260901120000", heartbeat="20260901120000",
                   last_error="new boom", last_error_at="20260901130000")
            (d / "logs" / "safe-mode.flag").write_text("x", encoding="utf-8")
            h = vk.health()
            assert h.get("last_error") == "new boom" and "last_error_age_s" in h, h
            assert h.get("safe_mode") is True

            vk._LAST_RELOAD_ERROR, vk._LAST_RELOAD_ERROR_AT = "refused x", "20260901110000"
            assert "last reload was refused" not in vk.health().get("note", ""), \
                "a master started after the refusal clears it"
            vk._LAST_RELOAD_ERROR, vk._LAST_RELOAD_ERROR_AT = "refused y", "20260901130000"
            assert "last reload was refused" in vk.health().get("note", "")
        finally:
            vk.voicekit_running, vk._LAST_RELOAD_ERROR, vk._LAST_RELOAD_ERROR_AT = orig


def test_process_matching_is_scoped_to_this_tree():
    """A body running under ANOTHER VoiceKit tree (the installed copy, a test
    root) must not be listed — or stopped — as this tree's."""
    with tempfile.TemporaryDirectory() as a, tempfile.TemporaryDirectory() as b:
        bodies = Path(a) / "hotkeys" / "bodies"
        bodies.mkdir(parents=True)
        probe = bodies / "ScopeProbeZz.body.ahk"
        probe.write_text("#Requires AutoHotkey v2.0\n#SingleInstance Off\nSleep(30000)\nExitApp\n",
                         encoding="utf-8")
        orig_root = vk.REPO_ROOT
        p = subprocess.Popen([vk.AHK_EXE, "/ErrorStdOut", str(probe)])
        try:
            vk.REPO_ROOT = Path(a)
            deadline = time.monotonic() + 15
            seen = False
            while time.monotonic() < deadline and not seen:
                seen = any(e["base"] == "ScopeProbeZz" for e in vk._body_processes())
            assert seen, "its own tree must see it"
            vk.REPO_ROOT = Path(b)
            assert not any(e["base"] == "ScopeProbeZz" for e in vk._body_processes()), \
                "another tree must not"
            assert not vk._ahk_script_running("VoiceKit.ahk"), \
                "a master is only 'running' when it is THIS tree's"
        finally:
            vk.REPO_ROOT = orig_root
            p.kill()
            p.wait(timeout=10)


def test_used_bridge_keys_match_ahk():
    """ONE rule for "is this bridge key taken", implemented twice: Python
    _used_bridge_keys (allocator, get_bridge_map's free list, a key= request)
    and AHK BridgeFreeKeys (New Automation's key dropdown, Home's Hotkey
    button). The two used to disagree — the allocator missed a hand-padded
    'Ctrl+Alt+Shift+B | ...' line, get_bridge_map ignored commented-out
    lines — so one key could be handed out twice. Both run on the same
    tricky map here and must return the same free list."""
    fixture = "\r\n".join([
        "; VoiceKit bridge map - mentions Ctrl+Alt+Shift+E (reserved anyway)",
        "Ctrl+Alt+Shift+A|Alpha|hotkeys\\Alpha.ahk|2026-01-01",
        "Ctrl+Alt+Shift+B | Padded | hotkeys\\Padded.ahk | 2026-01-02",
        "; Ctrl+Alt+Shift+D|Parked|hotkeys\\Parked.ahk|2026-01-03",
        ";Ctrl+Alt+Shift+F  |  Parked padded | hotkeys\\PP.ahk",
        "\t  Ctrl+Alt+Shift+G|Indented|hotkeys\\G.ahk|",
        "ctrl+alt+shift+j|lower|hotkeys\\Lower.ahk|",
        "Ctrl+Alt+Shift+[|Bracket|hotkeys\\Bracket.ahk|",
        "Ctrl+Alt+Shift+;|Semi|hotkeys\\Semi.ahk|",
        "Ctrl+Alt+Shift+\\|Backslash|hotkeys\\Bs.ahk|",
        "note: the old combo was Ctrl+Alt+Shift+K|, see above",
        "Ctrl+Alt+Shift+LongName|Not a key|hotkeys\\Long.ahk|",
        "Ctrl+Alt+Shift+M",
        "Ctrl+Alt+Shift+ 5 |space before the key|hotkeys\\Five.ahk|",
        ""])
    with _sandbox() as d:
        mapfile = d / "bridge-map.txt"
        mapfile.write_bytes(b"\xef\xbb\xbf" + fixture.encode("utf-8"))
        used = vk._used_bridge_keys()
        assert used == set("ABDFGJ[;\\KM5"), sorted(used)
        py_free = "".join(vk.get_bridge_map()["free_keys"])
        assert vk.allocate_bridge_key() == py_free[0] == "L", \
            "the allocator must skip every used key, padded and commented ones included"

        script, out = d / "freekeys.ahk", d / "freekeys.txt"
        # The SANDBOX copy of _Common: its uncaught-error logger writes beside
        # itself, so a failure here can never land in the user's errors.log.
        script.write_text(
            "#Requires AutoHotkey v2.0\n#SingleInstance Off\n"
            f'#Include "{d / "lib" / "_Common.ahk"}"\n'
            's := ""\nfor k in BridgeFreeKeys(A_Args[1])\n    s .= k\n'
            'FileAppend(s, A_Args[2], "UTF-8")\nExitApp(0)\n', encoding="utf-8")
        r = subprocess.run([vk.AHK_EXE, "/ErrorStdOut", str(script), str(mapfile), str(out)],
                           capture_output=True, text=True, timeout=30)
        assert r.returncode == 0, r.stdout + r.stderr
        ahk_free = out.read_text(encoding="utf-8-sig")
        assert ahk_free == py_free, f"free keys drifted:\n  AHK   : {ahk_free}\n  Python: {py_free}"

        # A requested key held only by a COMMENTED line is taken too.
        try:
            vk.create_hotkey_module("Dee Probe Zz", 'TrayTip("x")', key="D")
            assert False, "a key held by a commented-out line must be refused"
        except vk.VoiceKitError as e:
            assert "already taken" in str(e), e
    assert vk.get_bridge_map()["reserved_keys"] == list(vk.RESERVED_BRIDGE_KEYS)
    assert not set(vk.RESERVED_BRIDGE_KEYS) & set(vk.BRIDGE_POOL)


_SEED_INC_S = '#Include "%A_ScriptDir%\\hotkeys\\Snippets.ahk"'
_SEED_INC_W = '#Include "%A_ScriptDir%\\hotkeys\\Ship.hotkey.ahk"'
_SEED_REC_W = "Ctrl+Alt+Shift+W|open ship|hotkeys\\Ship.hotkey.ahk|2026-01-01"
# Each case: the live files before the run (None = missing). The defaults and
# the shipped module are the same for every case.
_SEED_CASES = {
    "fresh": (None, None, None),
    "upgrade": ("; manifest\n" + _SEED_INC_S + "\n" + '#Include "%A_ScriptDir%\\hotkeys\\Mine.ahk"',
                "; map\nCtrl+Alt+Shift+A|mine|hotkeys\\Mine.ahk|2026-02-02\n",
                "; my snippets only\n"),
    "commented": ("; manifest\n" + _SEED_INC_S + "\n; " + _SEED_INC_W + "    parked by me\n",
                  "; map\n; " + _SEED_REC_W + "\n", None),
    "rekeyed": ("; manifest\n" + _SEED_INC_S + "\n" + _SEED_INC_W + "\n",
                "; map\nCtrl+Alt+Shift+Q|open ship|hotkeys\\Ship.hotkey.ahk|2026-03-03\n", None),
    "key_taken": ("; manifest\n" + _SEED_INC_S + "\n",
                  "; map\nCtrl+Alt+Shift+W|other|hotkeys\\Other.ahk|2026-04-04\n", None),
    "crlf": ("; manifest\r\n" + _SEED_INC_S + "\r\n", "; map\r\n", None),
}
_SEED_LIVE = ("hotkeys/_index.ahk", "bridge-map.txt", "hotkeys/Snippets.ahk")


def _seed_fixture(d: Path, live: tuple) -> None:
    (d / "hotkeys").mkdir(exist_ok=True)
    put = lambda rel, text: (d / rel).write_bytes(b"\xef\xbb\xbf" + text.encode("utf-8"))
    put("hotkeys/_index.default.ahk", "; manifest\n" + _SEED_INC_S + "\n\n" + _SEED_INC_W + "\n")
    put("bridge-map.default.txt", "; map\n" + _SEED_REC_W + "\n")
    put("hotkeys/Snippets.default.ahk", "; snippets\n:*:/d::x\n")
    put("hotkeys/Ship.hotkey.ahk", "; shipped companion\n")
    for rel, text in zip(_SEED_LIVE, live):
        if text is not None:
            put(rel, text)


def test_seed_user_files_matches_ahk():
    """The live per-user files (hotkeys\\_index.ahk, bridge-map.txt,
    hotkeys\\Snippets.ahk) are gitignored and made from their shipped
    *.default copies, with newly shipped lines merged in — by AHK
    SeedUserFiles on every master start and by Python seed_user_files before
    an MCP reload. Two implementations of one rule: run both on identical
    sandboxes, case by case, and require identical files and reports."""
    for name, live in _SEED_CASES.items():
        with _sandbox() as d:
            ahk_root = d / "ahk"
            ahk_root.mkdir()
            _seed_fixture(ahk_root, live)
            py_root = d / "py"
            py_root.mkdir()
            _seed_fixture(py_root, live)
            out = d / "seed.txt"
            _run_common_probe(
                d,
                "r := SeedUserFiles(A_Args[1])\n"
                's := JoinList(r.created, ",") "`n" JoinList(r.merged, ",") "`n" (r.indexChanged ? 1 : 0)\n'
                'FileAppend(s, A_Args[2], "UTF-8-RAW")',
                str(ahk_root), str(out))
            ahk_created, ahk_merged, ahk_changed = out.read_text(encoding="utf-8").split("\n")
            with _patched(REPO_ROOT=py_root):
                r = vk.seed_user_files()
            assert ahk_created == ",".join(r["created"]), (name, ahk_created, r["created"])
            assert ahk_merged == ",".join(r["merged"]), (name, ahk_merged, r["merged"])
            assert ahk_changed == ("1" if r["index_changed"] else "0"), (name, ahk_changed, r)
            for rel in _SEED_LIVE:
                a, p = (ahk_root / rel).read_bytes(), (py_root / rel).read_bytes()
                assert a == p, f"{name}: {rel} drifted:\n  AHK   : {a!r}\n  Python: {p!r}"
            idx = (py_root / "hotkeys/_index.ahk").read_bytes().decode("utf-8-sig")   # CRLF kept
            bm = (py_root / "bridge-map.txt").read_text(encoding="utf-8-sig")
            if name in ("fresh", "upgrade", "crlf"):
                assert idx.count("Ship.hotkey.ahk") == 1 and bm.count("Ship.hotkey.ahk") == 1, (name, idx, bm)
            if name == "upgrade":
                assert 'Mine.ahk"\n' + _SEED_INC_W in idx, idx
                assert (py_root / "hotkeys/Snippets.ahk").read_text(
                    encoding="utf-8-sig") == "; my snippets only\n", "Snippets.ahk is never merged"
            if name in ("commented", "rekeyed"):
                assert not r["merged"], (name, r)
            if name == "key_taken":
                assert "Ship.hotkey.ahk" not in idx, "a hotkey whose key is taken must be held back"
            if name == "crlf":
                assert _SEED_INC_S + "\r\n" + _SEED_INC_W + "\r\n" in idx, repr(idx)
            with _patched(REPO_ROOT=py_root):
                again = vk.seed_user_files()
            assert not again["created"] and not again["merged"], (name, again)


def test_live_files_seed_on_first_use_and_defaults_start():
    """A fresh clone has no live per-user files. The writer's path accessors
    make a missing one from its default on first use, the listing never
    shows a *.default file as a module, and — the build's promise — a tree
    holding only the shipped defaults starts: VoiceKit.ahk load-checks with
    the manifest seeded from hotkeys\\_index.default.ahk, which wires the
    RecordMySteps companion and nothing that doesn't ship (ToggleTimer is an
    example now, so key A is free on a new install)."""
    real_idx = (REAL_ROOT / "hotkeys" / "_index.default.ahk").read_text(encoding="utf-8-sig")
    real_map = (REAL_ROOT / "bridge-map.default.txt").read_text(encoding="utf-8-sig")
    assert "RecordMySteps.hotkey.ahk" in real_idx and "RecordMySteps.hotkey.ahk" in real_map
    assert "ToggleTimer" not in real_idx + real_map, "ToggleTimer moved to templates\\examples"
    for rel in vk.index_modules.__globals__["_INCLUDE_RE"].findall(real_idx):
        assert (REAL_ROOT / rel.replace("\\", "/")).exists(), f"default manifest includes missing {rel}"
    with _sandbox(libs=("_Common.ahk", "Theme.ahk")) as d:
        for rel in ("hotkeys/_index.default.ahk", "bridge-map.default.txt",
                    "hotkeys/Snippets.default.ahk", "hotkeys/RecordMySteps.hotkey.ahk", "VoiceKit.ahk"):
            shutil.copyfile(REAL_ROOT / rel, d / rel)
        assert not (d / "bridge-map.txt").exists()
        assert vk._bridge_map_file().exists(), "the accessor must seed a missing live file"
        assert (d / "bridge-map.txt").read_bytes() == (d / "bridge-map.default.txt").read_bytes()
        vk._bridge_map_file()                       # already there: no second seed, no second line
        created_log = (d / "logs" / "created.log").read_text(encoding="utf-8-sig")
        assert created_log.count("seeded | bridge-map.txt | created from its shipped default") == 1, \
            f"a first-use seed is logged like AHK SeedUserFilesLogged, once:\n{created_log}"
        assert vk.index_modules()["active"] == ["hotkeys\\Snippets.ahk",
                                                "hotkeys\\RecordMySteps.hotkey.ahk"], vk.index_modules()
        assert vk._snippets_file().exists()
        names = [e["base"] for e in vk.list_automations()["hotkey_modules"]]
        assert not any(n.lower().endswith(".default") for n in names), names
        ok, err = vk.validate_ahk(str(d / "VoiceKit.ahk"))
        assert ok, f"a tree of shipped defaults must start:\n{err}"
        assert vk.allocate_bridge_key() == "A", "key A is free on a new install"


def test_log_cap_matches_ahk():
    """logs\\errors.log and logs\\created.log are capped the way the engine
    caps workflow-runs.log: past the cap the oldest half goes, cut at a line
    boundary. created.log has two writers (AHK Log, Python log), so the two
    trims must agree byte for byte — run both on one oversized log (a small
    cap stands in for the real 1 MB) and compare."""
    lines = "".join(f"2026-09-30 10:{i:02d} | line {i} with some text — ünïcode\n"
                    for i in range(60))
    with _sandbox() as d:
        a, p = d / "ahk.log", d / "py.log"
        for f in (a, p):
            f.write_bytes(b"\xef\xbb\xbf" + lines.encode("utf-8"))
        _run_common_probe(d, "LogTrimIfOver(A_Args[1], 1000)", str(a))
        assert vk._log_trim_if_over(p, 1000)
        assert a.read_bytes() == p.read_bytes(), "the AHK and Python log trims drifted"
        kept = p.read_text(encoding="utf-8-sig")
        assert kept and lines.endswith(kept) and kept.startswith("2026-"), \
            "the newest half must survive, starting at a whole line"
        assert len(kept) < len(lines) * 0.6
        # Under the cap nothing is touched.
        before = p.read_bytes()
        assert not vk._log_trim_if_over(p, 10 ** 6) and p.read_bytes() == before
        # Crash-safe: the kept half is written beside the log and moved over
        # it, never delete-then-append. A log held open without delete-
        # sharing (Python's open() shares none) refuses the move — both
        # trims must then leave it byte-for-byte as it was, temp file gone.
        for f in (a, p):
            f.write_bytes(b"\xef\xbb\xbf" + lines.encode("utf-8"))
        full = p.read_bytes()
        with open(a, "rb"), open(p, "rb"):
            _run_common_probe(d, "LogTrimIfOver(A_Args[1], 1000)", str(a))
            assert not vk._log_trim_if_over(p, 1000)
        assert a.read_bytes() == full and p.read_bytes() == full, \
            "a refused trim must leave the log whole"
        assert not list(d.glob("*.tmp-*")), f"temp files left: {list(d.glob('*.tmp-*'))}"
    assert vk.LOG_CAP_BYTES == 1048576


def test_every_tool_is_in_the_readme():
    """mcp\\README.md is the tool list a human reads before wiring the server
    into a client — and its security section is the only place the
    immediate-effect tools are called out. It had drifted to 18 of 29 tools
    (the audit, 2026-09-29). Every @mcp.tool function name must appear in it,
    backticked; parsed from server.py's source, so this runs without fastmcp."""
    import ast
    src = (REAL_ROOT / "mcp" / "server.py").read_text(encoding="utf-8")
    tools = []
    for node in ast.walk(ast.parse(src)):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            for dec in node.decorator_list:
                target = dec.func if isinstance(dec, ast.Call) else dec
                if isinstance(target, ast.Attribute) and target.attr == "tool" \
                        and isinstance(target.value, ast.Name) and target.value.id == "mcp":
                    tools.append(node.name)
    assert len(tools) >= 29, f"found only {len(tools)} @mcp.tool functions — parser drift?"
    readme = (REAL_ROOT / "mcp" / "README.md").read_text(encoding="utf-8")
    missing = [t for t in tools if f"`{t}`" not in readme]
    assert not missing, f"mcp\\README.md doesn't document: {', '.join(missing)}"
    sec = readme[readme.index("## Security"):]
    for t in ("run_ahk_snippet", "stop_module", "update_macro", "edit_macro",
              "edit_hotkey_module", "update_hotkey_module", "run_automation", "press_hotkey"):
        assert f"`{t}`" in sec, f"the security section must name immediate-effect tool {t}"


def test_create_workflow_rollback_restores_the_previous_workflow():
    """A replace whose stub fails its load check (a lib file mid-edit is
    enough) used to DELETE the workflow it was replacing. Now both files go
    back to their exact previous bytes — also when the check itself raises —
    and a brand-new workflow leaves nothing behind."""
    def failing(*a, **k):
        return False, "boom (simulated load failure)"

    def raising(*a, **k):
        raise RuntimeError("simulated: AutoHotkey went away")

    with _sandbox() as d, _patched(make_shortcut=lambda *a, **k: None,
                                   validate_ahk=lambda *a, **k: (True, "")):
        first = vk.create_workflow("Roll Probe Zz", [("wait", "100", "", "")])
        assert first["replaced"] is False and "reloaded" not in first, first
        assert "nothing needs reloading" in first["note"], first
        steps, stub = d / "workflows" / "RollProbeZz.steps.txt", d / "macros" / "RollProbeZz.ahk"
        # Hand-touched bytes (CRLF, no BOM): the restore must be byte-exact.
        steps.write_bytes(b"; mine\r\nwait|250||\r\n")
        before = (steps.read_bytes(), stub.read_bytes())
        for bad in (failing, raising):
            vk.validate_ahk = bad
            try:
                vk.create_workflow("Roll Probe Zz", [("wait", "999", "", "")])
                assert False, "a failed load check must fail the create"
            except (vk.VoiceKitError, RuntimeError) as e:
                assert "boom" in str(e) or "went away" in str(e), e
            assert (steps.read_bytes(), stub.read_bytes()) == before, \
                f"the replaced workflow must come back byte for byte ({bad.__name__})"
        vk.validate_ahk = failing
        try:
            vk.create_workflow("Fresh Probe Zz", [("wait", "100", "", "")])
            assert False
        except vk.VoiceKitError:
            pass
        assert not (d / "workflows" / "FreshProbeZz.steps.txt").exists()
        assert not (d / "macros" / "FreshProbeZz.ahk").exists()

        # VoiceKit's own tool names are refused, whatever the case.
        (d / "macros" / "WorkflowStudio.ahk").write_text("; the Studio\n", encoding="utf-8")
        for nm in ("Workflow Studio", "workflow studio"):
            try:
                vk.create_workflow(nm, [("wait", "100", "", "")])
                assert False, "a tool name must be refused"
            except vk.VoiceKitError as e:
                assert "VoiceKit's own tools" in str(e), e


def test_create_workflow_owns_every_step_guard():
    """The MCP schema's to_abc used to be the only place some rules lived
    (click needs a name or a position, a drag needs four numbers, a move
    needs a position), so the directly imported writer let a malformed step
    through. Every guard is the writer's now, and names its step."""
    ok = ("wait", "100", "", "")
    cases = [
        (("run", "", "", ""), "needs a target"),
        (("focus", "", "x.exe", ""), "needs a window"),
        (("text", "", "", ""), "something to type"),
        (("keys", "", "", ""), "something to send"),
        (("click", "ahk_exe x.exe", "", ""), "element name"),
        (("click", "ahk_exe x.exe", "", "12;40"), "'x,y'"),
        (("hover", "", "Menu", ""), "needs a window"),
        (("drag", "ahk_exe x.exe", "", "1,2,3"), "x1,y1,x2,y2"),
        (("move", "ahk_exe x.exe", "middle", ""), "needs a position"),
        (("if", "ahk_exe x.exe", "", "clipboardchanged"), "only works on a waitfor"),
        (("capture", "", "echo hi", ""), "needs a name"),
        (("capture", "{{Inv}}", "echo hi", ""), "without braces"),
        (("capture", "Inv", "  ", ""), "needs a command"),
        (("capture", "Inv", "echo hi", "0"), "positive number"),
        (("capture", "Inv", "echo hi", "abc"), "positive number"),
        (("capture", "Inv", "echo hi", "-5"), "positive number"),
        (("capture", "Inv", "echo hi", "nan"), "positive number"),
        (("fill", "", "Amount", "1"), "needs a window"),
        (("fill", "w", "", "1"), "has no label"),
        (("fill", "w", "#3", "1"), "has no label"),
        (("fill", "w", "Amount#0", "1"), "'#0'"),
        (("fill", "w", "Amount", None), "needs a value"),
        (("fill", "w", "Amount"), "needs a value"),
        (("fill", "w", "Amount", "two\nlines"), "line break"),
        (("fill", "w", "Amount", "cr\rhere"), "line break"),
    ]
    for step, why in cases:
        try:
            vk.create_workflow("Guard Probe Zz", [ok, step])
            assert False, f"must refuse {step}"
        except vk.VoiceKitError as e:
            assert str(e).startswith("Step 2:") and why in str(e), (step, str(e))
    # ...and the well-formed versions pass the guards (then stop at the
    # sandbox's missing lib — any error but a step error is fine here).
    with _sandbox(), _patched(make_shortcut=lambda *a, **k: None,
                              validate_ahk=lambda *a, **k: (True, "")):
        got = vk.create_workflow("Guard Probe Zz", [
            ok, ("click", "ahk_exe x.exe", "", "12,40"), ("drag", "w", "", "1,2,3,4"),
            ("move", "w", "max", ""), ("text", " ", "", ""), ("hover", "w", "Menu", ""),
            ("capture", "Inv", "echo hi", ""), ("capture", "Inv2", "echo {{Inv}}", "1.5"),
            ("fill", "w", "Amount#2", "{{Inv2}}"), ("fill", "w", "Memo", "")])
        assert got["step_count"] == 10


def test_isolated_create_load_checks_the_body():
    """create_hotkey_module(isolate=True) used to check only the launcher, so
    a body that can never load (here: assigning to _Common's Log function)
    came back 'live now' and failed later as a dialog in a process nobody
    watched. The body is checked too, and a failed check changes nothing."""
    with _sandbox() as d:
        (d / "hotkeys" / "_index.ahk").write_text("#Requires AutoHotkey v2.0\n", encoding="utf-8")
        (d / "bridge-map.txt").write_text("; map\n", encoding="utf-8")
        idx0, map0 = (d / "hotkeys" / "_index.ahk").read_bytes(), (d / "bridge-map.txt").read_bytes()
        try:
            vk.create_hotkey_module("Body Probe Zz", 'LOG := "x"')
            assert False, "a body that can't load must be refused"
        except vk.VoiceKitError as e:
            assert "steps failed to load" in str(e), e
        assert not (d / "hotkeys" / "BodyProbeZz.ahk").exists()
        assert not (d / vk._body_rel("BodyProbeZz")).exists()
        assert (d / "hotkeys" / "_index.ahk").read_bytes() == idx0
        assert (d / "bridge-map.txt").read_bytes() == map0

        vk.create_hotkey_module("Body Probe Zz", 'TrayTip("fine")')
        files = (d / "hotkeys" / "BodyProbeZz.ahk", d / vk._body_rel("BodyProbeZz"))
        before = [f.read_bytes() for f in files]
        try:
            vk.create_hotkey_module("Body Probe Zz", 'LOG := "x"')
            assert False
        except vk.VoiceKitError:
            pass
        assert [f.read_bytes() for f in files] == before, "a failed replace changes nothing"
        assert not vk._module_backup_file("BodyProbeZz").exists(), \
            "a failed replace banks nothing"


def test_inprocess_clash_is_rolled_back_not_reported_live():
    """An in-process module loads fine ON ITS OWN and still breaks the master
    — a hotkey or function another loaded file also defines only clashes
    when VoiceKit.ahk compiles everything together. The reload's preflight
    parks it; the response used to say 'live now' anyway. Now the write is
    rolled back and the call fails with the load error. Runs the REAL
    reload (preflight included) against a sandbox master; only the final
    launch is stubbed."""
    launches: list = []
    with _sandbox(reload=None) as d, _patched(voicekit_running=lambda: True,
                                              _launch_master=launches.append):
        (d / "lib" / "fakecommon.ahk").write_text(
            "ClashNotifyZz(msg) {\n    return msg\n}\n", encoding="utf-8")
        (d / "VoiceKit.ahk").write_text(
            "#Requires AutoHotkey v2.0\n"
            '#Include "%A_ScriptDir%\\lib\\fakecommon.ahk"\n'
            '#Include "%A_ScriptDir%\\hotkeys\\_index.ahk"\n', encoding="utf-8")
        (d / "hotkeys" / "Other.ahk").write_text("^!+a:: {\n}\n", encoding="utf-8")
        (d / "hotkeys" / "_index.ahk").write_text(
            "#Requires AutoHotkey v2.0\n"
            '#Include "%A_ScriptDir%\\hotkeys\\Other.ahk"\n', encoding="utf-8")
        (d / "bridge-map.txt").write_text("; map\n", encoding="utf-8")

        # (1) Brand new, and its allocated key A duplicates Other.ahk's
        #     unregistered hotkey: nothing may be kept.
        try:
            vk.create_hotkey_module("Clash Probe Zz", 'MsgBox("x")', isolate=False)
            assert False, "a module the master can't load must not be reported live"
        except vk.VoiceKitError as e:
            assert "wouldn't load with it" in str(e) and "nothing was kept" in str(e), e
        assert not (d / "hotkeys" / "ClashProbeZz.ahk").exists()
        assert "ClashProbeZz" not in (d / "hotkeys" / "_index.ahk").read_text(encoding="utf-8-sig")
        assert "ClashProbeZz" not in (d / "bridge-map.txt").read_text(encoding="utf-8-sig")

        # (2) A working module updated into a function clash: the previous
        #     bytes come back and it is switched back on.
        made = vk.create_hotkey_module("Clash Upd Zz", 'MsgBox("ok")', isolate=False, key="Q")
        assert made["reloaded"] is True and "live now" in made["note"], made
        mod = d / "hotkeys" / "ClashUpdZz.ahk"
        before = mod.read_bytes()
        src = vk.read_hotkey_module("Clash Upd Zz")["body"]
        try:
            vk.update_hotkey_module("Clash Upd Zz",
                                    src + "\nClashNotifyZz(msg) {\n    return msg\n}\n")
            assert False, "an update that clashes must be rolled back"
        except vk.VoiceKitError as e:
            assert "wouldn't load with it" in str(e) and "put back" in str(e), e
        assert mod.read_bytes() == before, "the previous version must come back byte for byte"
        assert "hotkeys\\ClashUpdZz.ahk" in vk.index_modules()["active"], \
            "a module that was live before must be switched back on"
        assert not vk._module_backup_file("ClashUpdZz").exists(), "a rolled-back update banks nothing"
        assert vk.validate_ahk(str(d / "VoiceKit.ahk"))[0], "the sandbox master must load again"


def test_back_to_back_reloads_wait_for_the_launcher():
    """A rolled-back clash reloads twice within a fraction of a second. The
    second VoiceKitLauncher.ahk used to start while the first was still
    load-checking — and with no #SingleInstance directive (v2's default is
    Prompt) it sat on an 'already running, replace it?' modal on the user's
    desktop, holding up the reload the rollback depended on. Now a reload
    waits for the previous launcher, and the launcher says Off anyway. Runs
    REAL launchers (against a sandbox master that exits at once)."""
    assert re.search(r"(?im)^#SingleInstance\s+Off\b",
                     (REAL_ROOT / "VoiceKitLauncher.ahk").read_text(encoding="utf-8-sig")), \
        "the launcher must not default to #SingleInstance Prompt"

    # (1) The contract, with the process stubbed: the previous handle is
    #     waited on (bounded) BEFORE the next launch.
    events: list = []

    class FakeProc:
        def __init__(self, n):
            self.n = n

        def wait(self, timeout=None):
            events.append(("wait", self.n, timeout))

    def fake_launch(target):
        events.append(("launch", len([e for e in events if e[0] == "launch"]) + 1))
        return FakeProc(events[-1][1])

    with _sandbox(reload=None, libs=None) as d, \
            _patched(voicekit_running=lambda: True, _launch_master=fake_launch,
                     master_preflight=lambda reasons=None: (True, [], ""), _LAST_LAUNCH=None):
        (d / "VoiceKitLauncher.ahk").write_text("; stub\n", encoding="utf-8")
        assert vk.reload_voicekit() == "reloaded"
        assert vk.reload_voicekit() == "reloaded"
        assert events == [("launch", 1), ("wait", 1, vk.LAUNCHER_WAIT_S), ("launch", 2)], events

    # (2) For real: two launchers back to back; each must have exited
    #     before the next starts, and none may be left behind (a Prompt
    #     modal is exactly what would be left behind).
    real_launch = vk._launch_master
    procs: list = []

    def launch_and_check(target):
        assert all(p.poll() is not None for p in procs), \
            "the previous launcher must have exited before the next launch"
        p = real_launch(target)
        procs.append(p)
        return p

    with _sandbox(reload=None) as d, \
            _patched(voicekit_running=lambda: True, _launch_master=launch_and_check,
                     _LAST_LAUNCH=None):
        shutil.copyfile(REAL_ROOT / "VoiceKitLauncher.ahk", d / "VoiceKitLauncher.ahk")
        (d / "VoiceKit.ahk").write_text(
            "#Requires AutoHotkey v2.0\n#SingleInstance Force\n"
            '#Include "%A_ScriptDir%\\hotkeys\\_index.ahk"\nExitApp(0)\n', encoding="utf-8")
        (d / "hotkeys" / "_index.ahk").write_text("#Requires AutoHotkey v2.0\n",
                                                   encoding="utf-8")
        try:
            assert vk.reload_voicekit() == "reloaded"
            assert vk.reload_voicekit() == "reloaded"
            assert len(procs) == 2
            procs[-1].wait(timeout=vk.LAUNCHER_WAIT_S)
            assert [p.returncode for p in procs] == [0, 0], [p.returncode for p in procs]
            left = [p for p in vk._ahk_processes() if str(d).lower() in p["cmd"].lower()]
            assert not left, f"nothing from the sandbox may still be running: {left}"
        finally:
            for p in procs:
                if p.poll() is None:
                    p.kill()
            for p in vk._ahk_processes():         # the sandbox master, should it linger
                if str(d).lower() in p["cmd"].lower():
                    subprocess.run(["taskkill", "/f", "/pid", str(p["pid"])],
                                   capture_output=True)


def test_rollback_reload_that_parks_again_is_not_reported_live():
    """When the rollback's own reload parks the RESTORED version too, the
    clash is with some other file — the error must say the module is off,
    not 'VoiceKit reloaded with it'."""
    with _sandbox() as d:
        rel = "hotkeys\\Mod.ahk"
        mod = d / "hotkeys" / "Mod.ahk"
        mod.write_text("old\n", encoding="utf-8")
        (d / "hotkeys" / "_index.ahk").write_text(
            f'; #Include "%A_ScriptDir%\\{rel}"\n', encoding="utf-8")
        snap = vk._FileSnapshot(mod)
        mod.write_text("new\n", encoding="utf-8")
        vk._LAST_PARKED.clear()
        vk._LAST_PARKED[rel.lower()] = "clash"

        def reload_parks_again():
            vk._LAST_PARKED[rel.lower()] = "clash again"
            return "reloaded"
        vk.reload_voicekit = reload_parks_again     # _sandbox restores it
        try:
            vk._undo_if_parked(rel, snap, was_parked=False)
            assert False, "must raise"
        except vk.VoiceKitError as e:
            assert "turned off" in str(e) and "reloaded with it" not in str(e), e
        assert mod.read_text(encoding="utf-8") == "old\n"
        vk._LAST_PARKED.clear()


def test_stop_module_reports_a_process_it_cannot_reach():
    """A body pid that won't open for SYNCHRONIZE|TERMINATE (e.g. running
    elevated) but is still alive used to count as 'already gone': the flag
    was deleted before the body could see it and stopped=True came back.
    Simulated by denying OpenProcess for that pid only."""
    with _sandbox() as d:
        (d / "hotkeys" / "Probe.ahk").write_text("; stub\n", encoding="utf-8")
        probe = d / "sleeper.ahk"
        probe.write_text("#Requires AutoHotkey v2.0\n#SingleInstance Off\n"
                         'DllCall("Sleep", "UInt", 30000)\n', encoding="utf-8")
        p = subprocess.Popen([vk.AHK_EXE, str(probe)])
        real_k = vk._kernel32

        class Denying:
            def __init__(self):
                self._k = real_k()

            def __getattr__(self, name):
                return getattr(self._k, name)

            def OpenProcess(self, access, inherit, pid):
                if pid == p.pid and access & vk._PROCESS_TERMINATE:
                    import ctypes
                    ctypes.set_last_error(vk._ERROR_ACCESS_DENIED)   # what Windows says
                    return None
                return self._k.OpenProcess(access, inherit, pid)
        try:
            with _patched(_body_processes=lambda: [{"base": "Probe", "pid": p.pid,
                                                    "started": ""}],
                          _kernel32=Denying):
                t0 = time.perf_counter()
                r = vk.stop_module("Probe", 1)
                took = time.perf_counter() - t0
            assert r["stopped"] is False and r["still_running"] == [p.pid], r
            assert took >= 0.9, f"the grace period must still be waited out ({took:.2f} s)"
            assert not vk._body_stop_file("Probe").exists(), "the flag is cleaned up"
            assert p.poll() is None, "an unreachable pid is reported, not somehow ended"
            assert r.get("could_not_open") == [p.pid], r

            # Refused EVERY open (access denied, not "no such pid"): not
            # observable at all, so only the scan speaks for it — and it must
            # never read as gone, let alone as a graceful stop.
            import ctypes

            class Blind(Denying):
                def OpenProcess(self, access, inherit, pid):
                    if pid == p.pid:
                        ctypes.set_last_error(vk._ERROR_ACCESS_DENIED)
                        return None
                    return self._k.OpenProcess(access, inherit, pid)
            with _patched(_body_processes=lambda: [{"base": "Probe", "pid": p.pid,
                                                    "started": ""}],
                          _kernel32=Blind):
                r = vk.stop_module("Probe", 1)
            assert r["stopped"] is False and r["still_running"] == [p.pid], r
            assert "could not open process" in r["note"], r
            assert p.poll() is None
        finally:
            p.kill()
            p.wait(timeout=10)

        # An unreachable pid that goes away during the grace period stopped —
        # but this server couldn't watch it, so it is never called graceful.
        quick = d / "quick.ahk"
        quick.write_text("#Requires AutoHotkey v2.0\n#SingleInstance Off\n"
                         'DllCall("Sleep", "UInt", 1500)\n', encoding="utf-8")
        p = subprocess.Popen([vk.AHK_EXE, str(quick)])
        try:
            with _patched(_body_processes=lambda: [{"base": "Probe", "pid": p.pid,
                                                    "started": ""}],
                          _kernel32=Denying):
                r = vk.stop_module("Probe", 10)
            assert r["stopped"] is True and r["how"] == "exited", r
            assert "could not open process" in r["note"] and r["could_not_open"] == [p.pid], r
        finally:
            p.kill()
            p.wait(timeout=10)


def test_update_combo_check_is_a_real_hotkey_definition():
    """update_hotkey_module's 'must still define ^!+<key>' guard was a bare
    substring test: Send("^!+b"), a comment about the old binding and
    ^!+backspace:: all passed it, silently stranding the Voice Access pairing."""
    yes = ["^!+b:: {\n}", "!+^B::MsgBox(1)", "~$^!+b::", "  +^!b ::",
           'Hotkey("^!+b", Go)', "x := 1\nHotkey('~^+!b', Go, 'On')"]
    no = ['Send("^!+b")', "; ^!+b:: was the old binding", "^!+backspace::",
          '; Hotkey("^!+b", Go)', 'x := 1 ; Hotkey("^!+b", Go)',
          "/*\n^!+b::\n*/", "^!b::", "^!+c::", 's := "^!+b::"']
    for code in yes:
        assert vk._defines_combo(code, "B"), code
    for code in no:
        assert not vk._defines_combo(code, "B"), code
    assert vk._defines_combo("^!+[:: {\n}", "[") and vk._defines_combo("^!+;::", ";")


def test_rollbacks_and_backups_are_byte_exact():
    """A failed edit used to 'roll back' by re-writing a decoded snapshot, so
    a CRLF or BOM-less file came back re-encoded while the error said nothing
    changed; a load check that RAISED skipped the rollback entirely; and the
    banked backup was re-encoded too. All three are byte-exact now."""
    raw = b'#Requires AutoHotkey v2.0\r\nTrayTip("one")\r\nTrayTip("two")\r\n'
    with _sandbox() as d:
        mac = d / "macros" / "BytesProbeZz.ahk"
        mac.write_bytes(raw)
        try:
            vk.edit_macro("Bytes Probe Zz", 'TrayTip("one")', 'TrayTip("one"')
            assert False
        except vk.VoiceKitError as e:
            assert "rolled back" in str(e)
        assert mac.read_bytes() == raw, "a failed edit must leave the exact bytes"

        def raising(*a, **k):
            raise RuntimeError("simulated: load check blew up")
        with _patched(validate_ahk=raising):
            for call in (lambda: vk.update_macro("Bytes Probe Zz", 'TrayTip("x")'),
                         lambda: vk.edit_macro("Bytes Probe Zz", 'TrayTip("two")', 'TrayTip("2")')):
                try:
                    call()
                    assert False
                except RuntimeError:
                    pass
                assert mac.read_bytes() == raw, "a raising load check must still roll back"

        got = vk.update_macro("Bytes Probe Zz", '#Requires AutoHotkey v2.0\nTrayTip("new")\n')
        assert Path(got["previous_source_backup"]).read_bytes() == raw, \
            "the banked version must be the raw previous bytes"
        assert vk.read_macro_source("Bytes Probe Zz", previous=True)["source"] == \
            raw.decode("utf-8")


def test_batch_refuses_while_a_loop_runs_and_uses_its_own_file():
    """LoopRunner is #SingleInstance Force, so a second batch (an agent
    retrying after finished=False was enough) killed the running one, and
    both shared logs\\mcp-batch.csv. Now: refused while any loop runs,
    naming it; a per-run batch file; and an all-blank row is refused."""
    with _sandbox() as d:
        (d / "workflows" / "BatchProbeZz.steps.txt").write_text(
            "; probe\nask|Name||\ntext|{{Name}}||\n", encoding="utf-8")
        legacy = d / "logs" / "mcp-batch.csv"
        legacy.write_text("Name\r\nold\r\n", encoding="utf-8")
        loop_row = [{"pid": 4242, "started": "",
                     "cmd": f'"{vk.AHK_EXE}" "{d}\\lib\\LoopRunner.ahk" OtherFlowZz'}]
        with _patched(_ahk_processes=lambda: loop_row):
            for call in (lambda: vk.run_workflow_batch("Batch Probe Zz", [{"Name": "a"}]),
                         lambda: vk.run_automation("loop Batch Probe Zz")):
                try:
                    call()
                    assert False, "a loop launch must be refused while one runs"
                except vk.VoiceKitError as e:
                    assert "already running" in str(e) and "Other Flow Zz" in str(e), e
            assert legacy.exists(), "a refused batch must not touch anything"
        # Another tree's loop is not ours to protect.
        other = [{"pid": 4242, "started": "",
                  "cmd": f'"{vk.AHK_EXE}" "C:\\Elsewhere\\lib\\LoopRunner.ahk" X'}]
        launched: list = []

        def fake_launch(cmd, launched_as, file, wait, cwd=None):
            launched.append(cmd)
            return {"launched": launched_as, "file": file, "finished": False, "note": ""}
        with _patched(_ahk_processes=lambda: other, _launch_and_report=fake_launch):
            try:
                vk.run_workflow_batch("Batch Probe Zz", [{"Name": "a"}, {"name": "  "}])
                assert False, "an all-blank row must be refused"
            except vk.VoiceKitError as e:
                assert "Row 2 has no values" in str(e), e
            got = vk.run_workflow_batch("loop Batch Probe Zz", [{"Name": "a, b"}, {"NAME": "c"}])
        assert got["rows"] == 2, got
        batch = Path(launched[-1][3])
        assert re.fullmatch(r"mcp-batch-\d{14}-\d+\.csv", batch.name), batch.name
        assert batch.read_bytes() == b'\xef\xbb\xbfName\r\n"a, b"\r\nc\r\n', batch.read_bytes()
        assert not legacy.exists(), "an earlier batch file is cleared once no loop runs"


def test_run_automation_passes_args_and_cwd():
    """run_automation(args=[...]) (WP5): arguments reach the script as its
    A_Args verbatim (an argv list — spaces and & need no quoting), the launch
    runs in the macro's own folder like its voice shortcut, and a WORKFLOW's
    arguments stand in for the File Explorer selection — {{selected_file}} /
    {{selected_files}} resolve to them, so a selected-file workflow can be
    tested on a fixture with Explorer never consulted. Refusals happen before
    anything launches. Nothing here touches the desktop: the workflows only
    `run` (and `capture` the output of) throwaway AHK probes that record their
    own arguments — the Extract-Invoice shape, args -> capture -> a later step."""
    libs = ("_Common.ahk", "Workflow.ahk", "Acc.ahk", "UIA.ahk", "Clip.ahk", "ExplorerSel.ahk")
    with _sandbox(libs=libs) as d:
        # 1. a raw script: argv verbatim, cwd = macros\ (what its .lnk sets)
        vk.create_launch_macro("Args Probe Zz", ahk_body=(
            's := A_WorkingDir "`n"\nfor a in A_Args\n    s .= a "`n"\n'
            'FileAppend(s, A_ScriptDir "\\..\\logs\\args-probe.txt", "UTF-8")\n'))
        args = [r"C:\Some Folder\Invoice & Co.pdf", "plain", "two words"]
        r = vk.run_automation("Args Probe Zz", wait_seconds=20, args=args)
        assert r.get("finished") and r.get("exit_code") == 0, r
        lines = (d / "logs" / "args-probe.txt").read_text(encoding="utf-8-sig").splitlines()
        assert lines[0].lower() == str(d / "macros").lower(), f"cwd was {lines[0]}"
        assert lines[1:] == args, lines

        # 2. a workflow: its args ARE the selection
        fx = d / "fixtures"
        fx.mkdir()
        f1, f2 = fx / "Invoice & Co 2025.pdf", fx / "b.pdf"
        f1.write_text("x", encoding="utf-8")
        f2.write_text("y", encoding="utf-8")
        probe = d / "logs" / "sel-probe.ahk"
        seen = d / "logs" / "sel-probe.txt"
        probe.write_text(
            "#Requires AutoHotkey v2.0\n#SingleInstance Off\n"
            "s := A_Args.Length\nfor a in A_Args\n    s .= \"|\" a\n"
            "FileAppend(s \"`n\", A_ScriptDir \"\\sel-probe.txt\", \"UTF-8\")\n",
            encoding="utf-8")
        cmd = f'"{vk.AHK_EXE}" "{probe}" '
        vk.create_workflow("Sel One Zz", [("run", cmd + '"{{selected_file}}"', "", "")])
        vk.create_workflow("Sel Many Zz", [("run", cmd + "{{selected_files}}", "", "")])

        def probe_lines(n):
            deadline = time.time() + 15
            while time.time() < deadline:
                if seen.exists():
                    got = seen.read_text(encoding="utf-8-sig").splitlines()
                    if len(got) >= n:
                        return got
                time.sleep(0.2)
            return seen.read_text(encoding="utf-8-sig").splitlines() if seen.exists() else []

        # the %TEMP% form a response shows is accepted as the path
        tmp = tempfile.gettempdir().rstrip("\\")
        f1_arg = str(f1)
        if f1_arg.lower().startswith(tmp.lower() + "\\"):
            f1_arg = "%TEMP%" + f1_arg[len(tmp):]
        r = vk.run_automation("Sel One Zz", wait_seconds=30, args=[f1_arg])
        assert r.get("ok") is True and r["run"]["outcome"] == "ok", r
        assert [Path(a) for a in r["args"]] == [f1.resolve()], r["args"]
        got = probe_lines(1)
        assert got and got[0].lower() == f"1|{f1.resolve()}".lower(), got
        r = vk.run_automation("Sel Many Zz", wait_seconds=30, args=[str(f1), str(f2)])
        assert r.get("ok") is True and r["run"]["outcome"] == "ok", r
        got = probe_lines(2)
        assert len(got) >= 2 and got[1].lower() == \
            f"2|{f1.resolve()}|{f2.resolve()}".lower(), got

        # 2b. the Extract-Invoice shape (WP4): args -> {{selected_file}} -> a capture
        # command (bare placeholder: it arrives quoted, spaces and & intact) ->
        # its printed output -> a later step. set/capture/run only: nothing
        # touches the desktop.
        echo = d / "logs" / "echo-probe.ahk"
        echo.write_text(
            "#Requires AutoHotkey v2.0\n#SingleInstance Off\n"
            "FileAppend(\"got:\" A_Args.Length \":\" A_Args[1], \"*\", \"UTF-8\")\n",
            encoding="utf-8")
        vk.create_workflow("Sel Cap Zz", [
            ("capture", "Got", f'"{vk.AHK_EXE}" "{echo}" {{{{selected_file}}}}', "20"),
            ("run", cmd + '"{{Got}}"', "", "")])
        r = vk.run_automation("Sel Cap Zz", wait_seconds=30, args=[str(f1)])
        assert r.get("ok") is True and r["run"]["outcome"] == "ok", r
        got = probe_lines(3)
        assert len(got) >= 3 and got[2].lower() == f"1|got:1:{f1.resolve()}".lower(), got

        # (A FAILING capture is pinned by tests\engine-capture-selftest.ahk, in
        # quiet mode: through run_automation the stub would pop its failure
        # MsgBox on the user's desktop and sit there.)

        # 3. refusals — all before anything launches
        launched: list = []

        def fake_launch(*a, **k):
            launched.append(a)
            return {"finished": False}
        (d / "macros" / "WorkflowStudio.ahk").write_text("; dummy\n", encoding="utf-8")
        refusals = [
            (lambda: vk.run_automation("Sel One Zz", args=[str(fx / "missing.pdf")]),
             "isn't an existing file"),
            (lambda: vk.run_automation("Sel One Zz", args=["b.pdf"]), "full path"),
            # singular + 2 files: the engine would refuse with a desktop popup
            (lambda: vk.run_automation("Sel One Zz", args=[str(f1), str(f2)]),
             "exactly one file"),
            (lambda: vk.run_automation("Workflow Studio", args=[str(f1)]),
             "doesn't take arguments"),
            (lambda: vk.run_automation("loop Sel One Zz", args=[str(f1)]),
             "doesn't take arguments"),
            (lambda: vk.run_automation("Args Probe Zz", args=str(f1)), "list of strings"),
            (lambda: vk.run_automation("Args Probe Zz", args=["a\nb"]), "line break"),
        ]
        with _patched(_launch_and_report=fake_launch, _ahk_script_running=lambda *a, **k: False,
                      _ahk_processes=lambda: []):
            for call, why in refusals:
                try:
                    call()
                    assert False, f"must refuse: {why}"
                except vk.VoiceKitError as e:
                    assert why in str(e), (why, str(e))
        assert not launched, "a refused run must launch nothing"


def test_press_hotkey_resolves_a_companion_by_its_listed_name():
    """list_automations folds a companion hotkey onto its automation ('Record
    My Steps'), and press_hotkey is documented to press it — but resolved it
    only by the raw phrase or 'RecordMySteps.hotkey'."""
    entries = [{"combo": "Ctrl+Alt+Shift+W", "phrase": "record my steps (keyboard)",
                "file": "hotkeys\\RecordMySteps.hotkey.ahk", "created": ""},
               {"combo": "Ctrl+Alt+Shift+Q", "phrase": "toggle timer",
                "file": "hotkeys\\ToggleTimer.ahk", "created": ""}]
    with _patched(get_bridge_map=lambda: {"entries": entries, "free_keys": [],
                                          "reserved_keys": []}):
        assert vk.resolve_bridge("Record My Steps")["key"] == "W"
        assert vk.resolve_bridge("RecordMySteps")["key"] == "W"
        assert vk.resolve_bridge("Toggle Timer")["key"] == "Q"


def test_reload_notes_are_consistent():
    """'Did the reload happen' used to be worded four ways, two of them wrong
    ('not running' when the reload was refused), and tools that never reload
    returned a bare reloaded:false with no note."""
    subj = "the hotkey"
    assert "is live now" in vk._reload_note("reloaded", subj)
    assert "isn't running" in vk._reload_note("not_running", subj) and \
        "VoiceKitLauncher.ahk" in vk._reload_note("not_running", subj)
    assert "refused" in vk._reload_note("refused", subj)
    orig_err = vk._LAST_RELOAD_ERROR
    try:
        for stub, want in ((lambda: "refused", "refused"), (lambda: True, "reloaded"),
                           (lambda: "not_running", "not_running")):
            with _patched(reload_voicekit=stub):
                got = vk._reload_outcome(subj)
                assert got["reload_status"] == want and \
                    got["reloaded"] is (want == "reloaded"), got
        vk._LAST_RELOAD_ERROR = "files won't load"
        with _patched(reload_voicekit=lambda: False):       # an old-style bool stub
            assert vk._reload_outcome(subj)["reload_status"] == "refused"
        vk._LAST_RELOAD_ERROR = ""
        with _patched(reload_voicekit=lambda: False):
            assert vk._reload_outcome(subj)["reload_status"] == "not_running"
    finally:
        vk._LAST_RELOAD_ERROR = orig_err
    # Tools whose result never depends on a reload say so, instead of a bare
    # reloaded:false.
    with _sandbox(), _patched(make_shortcut=lambda *a, **k: None):
        got = vk.create_launch_macro("Note Probe Zz", ahk_body='TrayTip("x")')
        assert "reloaded" not in got and "nothing needs reloading" in got["note"], got


def test_server_schema_matches_the_writer():
    """server.py's WorkflowStep carries hand-written Literals for the step
    types, conditions and move positions — item 4 of the five-way lockstep,
    which no test compared. (Skipped when the MCP add-on isn't installed.)"""
    import typing
    try:
        import server
    except ImportError:
        print("      (skipped: fastmcp not importable — the MCP add-on isn't installed)")
        return

    def literal_values(ann):
        if typing.get_origin(ann) is typing.Literal:
            return set(typing.get_args(ann))
        for a in typing.get_args(ann):
            if typing.get_origin(a) is typing.Literal:
                return set(typing.get_args(a))
        raise AssertionError(f"no Literal in {ann}")

    f = server.WorkflowStep.model_fields
    assert literal_values(f["type"].annotation) == set(vk.STEP_TYPES)
    assert literal_values(f["condition"].annotation) == set(vk.WAIT_CONDS)
    assert literal_values(f["position"].annotation) == set(vk.MOVE_POSITIONS)
    # to_abc only maps; the writer owns the rules and names the step.
    abc = server.WorkflowStep(type="click", window="ahk_exe x.exe").to_abc()
    try:
        vk.create_workflow("Schema Probe Zz", [("wait", "100", "", ""), abc])
        assert False
    except vk.VoiceKitError as e:
        assert str(e).startswith("Step 2:") and "element name" in str(e), e
    assert server.WorkflowStep(type="drag", window="w", xy="1, 2, 3 ,4").to_abc() == \
        ("drag", "w", "", "1,2,3,4")
    assert server.WorkflowStep(type="waitfor", window="w", condition="textvisible",
                               element="Done", seconds=5).to_abc() == \
        ("waitfor", "w", "Done", "textvisible,5")
    assert server.WorkflowStep(type="capture", label="Invoice Data", command="python invoice.py",
                               seconds=90).to_abc() == ("capture", "Invoice Data", "python invoice.py", "90")
    assert server.WorkflowStep(type="capture", label="Inv", command="x").to_abc() == \
        ("capture", "Inv", "x", "")
    assert server.WorkflowStep(type="fill", window="w", element="Amount#2",
                               value="{{Box 1}}").to_abc() == ("fill", "w", "Amount#2", "{{Box 1}}")
    assert server.WorkflowStep(type="fill", window="w", element="Memo", value="").to_abc() == \
        ("fill", "w", "Memo", "")
    abc = server.WorkflowStep(type="fill", window="w", element="Amount").to_abc()   # no value
    try:
        vk.create_workflow("Schema Probe Zz", [abc])
        assert False, "a fill with no value must be refused, not turned into a clear"
    except vk.VoiceKitError as e:
        assert "needs a value" in str(e), e
    abc = server.WorkflowStep(type="capture", label="Inv").to_abc()      # no command
    try:
        vk.create_workflow("Schema Probe Zz", [abc])
        assert False
    except vk.VoiceKitError as e:
        assert "needs a command" in str(e), e


def test_master_liveness_reads_the_status_file():
    """health() rides on every tool response and used to spawn a PowerShell
    CIM scan each time (~0.5 s, measured). The master already publishes its
    pid and a 5-second heartbeat; _master_alive checks that pid through
    kernel32 instead, and only an ambiguous answer falls back to the scan."""
    with _sandbox() as d:
        ini = d / "logs" / "master-status.ini"

        def status(**kv):
            body = "[Master]\r\n" + "".join(f"{k}={v}\r\n" for k, v in kv.items())
            ini.write_bytes(b"\xff\xfe" + body.encode("utf-16-le"))

        assert vk._master_alive() is None, "no status file: can't tell"
        now = time.strftime("%Y%m%d%H%M%S")
        probe = d / "alive.ahk"
        probe.write_text("#Requires AutoHotkey v2.0\n#SingleInstance Off\n"
                         'DllCall("Sleep", "UInt", 30000)\n', encoding="utf-8")
        p = subprocess.Popen([vk.AHK_EXE, str(probe)])
        try:
            status(pid=p.pid, heartbeat=now, clean_exit=0)
            assert vk._master_alive() is True

            def no_scan():
                raise AssertionError("health() must not fall back to the process scan")
            with _patched(_ahk_processes=no_scan):
                t0 = time.perf_counter()
                h = vk.health()
                took = time.perf_counter() - t0
            # Loose bound: the real guarantee is "no scan"; this only catches
            # something heavy creeping back in (a scan costs ~0.5 s).
            assert h["master_running"] is True and took < 2.0, (h, took)
            status(pid=p.pid, heartbeat="20200101000000", clean_exit=0)
            assert vk._master_alive() is None, "alive pid, stale heartbeat: ambiguous"
            import os
            status(pid=os.getpid(), heartbeat=now, clean_exit=0)
            assert vk._master_alive() is False, "a pid running something else isn't the master"
        finally:
            p.kill()
            p.wait(timeout=10)
        status(pid=p.pid, heartbeat="20200101000000", clean_exit=0)
        assert vk._master_alive() is False, "a dead pid with a stale heartbeat is a dead master"
        status(pid=p.pid, heartbeat=now, clean_exit=1)
        assert vk._master_alive() is False, "a dead pid that exited cleanly is a dead master"
        # Dead pid, FRESH heartbeat, no clean exit: a reload mid-handover
        # (Force killed the old master, the new one hasn't written its pid).
        # Not "definitely down" — the scan decides, and finds no sandbox master.
        status(pid=p.pid, heartbeat=now, clean_exit=0)
        assert vk._master_alive() is None, "the reload handover must not read as 'not running'"
        assert vk.voicekit_running() is False
        # ...and it is NAMED: health says restarting, and a caller that needs
        # the hotkeys (press_hotkey) waits for the new master, then says
        # "restarting" rather than "not running" if it never reports in.
        assert vk._master_restarting() is True
        assert vk.health().get("master_restarting") is True
        t0 = time.perf_counter()
        assert vk._await_master_after_restart(0.5) is False
        assert time.perf_counter() - t0 >= 0.45, "the restart wait must actually wait"
        (d / "bridge-map.txt").write_text(
            "; map\nCtrl+Alt+Shift+Q|zz probe|hotkeys\\ZzProbe.ahk|x\n", encoding="utf-8")
        with _patched(_await_master_after_restart=lambda: False):
            try:
                vk.press_hotkey("Q")
                assert False, "press_hotkey mid-restart must not claim a press"
            except vk.VoiceKitError as e:
                assert "restarting" in str(e), e
        for kv in ({"clean_exit": 1}, {"heartbeat": "20200101000000"}):
            status(**{"pid": p.pid, "heartbeat": now, "clean_exit": 0, **kv})
            assert vk._master_restarting() is False, kv
            assert vk._await_master_after_restart(5) is None, "no restart: no wait"
            assert "master_restarting" not in vk.health()


def _run_bare_ahk(body: str, timeout: int = 30) -> str:
    """Run a script with NO VoiceKit includes (so only true built-ins exist)
    and return what it FileAppend-ed to its report file."""
    with tempfile.TemporaryDirectory() as d:
        rep = Path(d) / "out.txt"
        script = Path(d) / "bare.ahk"
        script.write_text("#Requires AutoHotkey v2.0\n#SingleInstance Off\n"
                          + body.replace("__OUT__", str(rep)) + "\nExitApp(0)\n",
                          encoding="utf-8-sig")
        subprocess.run([vk.AHK_EXE, "/ErrorStdOut=UTF-8", str(script)],
                       capture_output=True, timeout=timeout)
        return vk._read_text_any(rep) if rep.exists() else ""


def _builtins_resolved(names) -> dict:
    """{name_lower: Type} for each name, resolved in a bare AutoHotkey v2
    process — 'Func' / 'Class' for a real built-in, 'ERR' otherwise."""
    with tempfile.TemporaryDirectory() as d:
        lst = Path(d) / "names.txt"
        lst.write_text("\n".join(names), encoding="utf-8", newline="")
        out = _run_bare_ahk(
            'for n in StrSplit(FileRead("' + str(lst) + '", "UTF-8"), "`n", "`r ") {\n'
            '    if (n = "")\n'
            '        continue\n'
            '    try t := Type(%n%)\n'
            '    catch\n'
            '        t := "ERR"\n'
            '    FileAppend(n "|" t "`n", "__OUT__", "UTF-8")\n'
            '}')
    return {ln.split("|")[0].lower(): ln.split("|")[1]
            for ln in out.splitlines() if "|" in ln}


def test_load_error_explains_name_clash():
    """WP1 (the extraction feedback): a top-level `LOG := ...` failed with "This Func
    cannot be used as an output variable. Specifically: LOG" and nothing said
    Log is an AutoHotkey BUILT-IN — or that Notify comes from lib\\_Common.ahk,
    or that Out is run_ahk_snippet's own helper. Every user-facing load error
    now carries a Hint line naming the owner; the raw error stays first."""
    # The curated built-in list is real: every name resolves in a bare v2.
    got = _builtins_resolved(sorted(vk.AHK_BUILTINS))
    bad = [n for n in sorted(vk.AHK_BUILTINS) if got.get(n) not in ("Func", "Class")]
    assert not bad, f"not AutoHotkey built-ins on this interpreter: {bad}"
    for n in vk._AHK_BUILTIN_CLASSES:
        assert got.get(n) == "Class", (n, got.get(n))

    with _sandbox(), _patched(make_shortcut=lambda *a, **k: None):
        common = (vk.REPO_ROOT / "lib" / "_Common.ahk").read_text(encoding="utf-8-sig")
        notify_line = next(i for i, ln in enumerate(common.splitlines(), 1)
                           if ln.startswith("Notify("))
        # (1) Built-in: Log (which _Common also defines — the name is taken
        # either way, and the hint says both).
        try:
            vk.create_launch_macro("Clash Log Zz", ahk_body='LOG := "x"')
            assert False, "LOG := must fail to load"
        except vk.VoiceKitError as e:
            msg = str(e)
            assert "Specifically: LOG" in msg or "Specifically: Log" in msg, msg
            assert "built-in" in msg and "Hint:" in msg and "logText" in msg, msg
            assert msg.index("==>") < msg.index("Hint:"), "the raw error stays first"
        assert not (vk.REPO_ROOT / "macros" / "ClashLogZz.ahk").exists()
        # (2) Notify from lib\_Common.ahk, with its real line number.
        try:
            vk.create_launch_macro("Clash Notify Zz", ahk_body="Notify := 1")
            assert False
        except vk.VoiceKitError as e:
            assert f"lib\\_Common.ahk line {notify_line}" in str(e), str(e)
        # (3) A function-local variable of the same name still loads.
        got = vk.create_launch_macro("Local Log Zz",
                                     ahk_body="F() {\n    log := 1\n    return log\n}\nF()")
        assert got["type"] == "launch_macro"

    # A plain file with no includes: Log is purely the built-in.
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "bare.ahk"
        p.write_text("#Requires AutoHotkey v2.0\nLOG := 1\n", encoding="utf-8-sig")
        ok, err = vk.validate_ahk(str(p))
        assert not ok
        hint = vk._explain_ahk_error(err, script=p)
        assert "an AutoHotkey built-in function" in hint, hint
        # _error_module (the quarantine parser) still reads an explained error.
        fake = hint.replace(str(p), r"C:\VK\hotkeys\Zz.ahk")
        with _patched(REPO_ROOT=Path(r"C:\VK")):
            assert vk._error_module(fake) == r"hotkeys\Zz.ahk", fake
        assert vk._explain_ahk_error(hint) == hint, "explaining twice adds nothing"
        assert vk._explain_ahk_error("some other error") == "some other error"

    # (4) Out, the snippet prelude's helper.
    got = vk.run_ahk_snippet("out := 1")
    assert got["exit_code"] == 2, got
    assert "Out() helper run_ahk_snippet predefines" in got["errors"], got
    # A declaration clash names the OTHER definition, not the snippet's own.
    got = vk.run_ahk_snippet("Notify(a) {\n}")
    assert got["exit_code"] == 2 and "lib\\_Common.ahk line" in got["errors"], got
    # Function-local: fine.
    got = vk.run_ahk_snippet("F() {\n    log := 1\n    return log\n}\nOut(F())")
    assert got["exit_code"] == 0 and got["output"] == "1", got
    # (5) v1-style doubled quotes.
    for code in ('x := "a ""b"""', 'MsgBox "a""b"'):
        got = vk.run_ahk_snippet(code)
        assert got["exit_code"] == 2 and "backtick" in got["errors"], (code, got)
    # Runtime clashes (exit 3) are explained too.
    got = vk.run_ahk_snippet("x := Log + 1")
    assert got["exit_code"] == 3 and "Hint:" in got["errors"] and "'Log'" in got["errors"], got
    got = vk.run_ahk_snippet("F(max) {\n    return Max(1, max)\n}\nOut(F(2))")
    assert got["exit_code"] == 3 and "hides the function" in got["errors"], got
    # A load-time #Warn is a MODAL MsgBox: an unset variable or unreachable
    # code used to wedge the snippet for the whole timeout with no report.
    got = vk.run_ahk_snippet("Out(1)\nreturn", 10)
    assert got["exit_code"] == 0 and got["output"] == "1", got
    got = vk.run_ahk_snippet("Out(zzUnset)", 10)
    assert got["exit_code"] == 3 and "zzUnset" in got["errors"], got
    # A called string isn't a hidden function: no false Max-trap hint.
    got = vk.run_ahk_snippet('foo := "abc"\nfoo()')
    assert got["exit_code"] == 3 and "Hint:" not in got["errors"], got
    # The snippet's throwaway path (and with it the user name) is gone.
    for g in (vk.run_ahk_snippet("out := 1"), vk.run_ahk_snippet("x := Log + 1")):
        assert "snippet.ahk" in g["errors"] or "snippet line" in g["errors"], g
        assert tempfile.gettempdir().lower() not in g["errors"].lower(), g
        assert os.environ.get("USERPROFILE", "?").lower() not in g["errors"].lower(), g


def test_instructions_names_exist():
    """The FastMCP instructions list the names a caller's AHK code can't
    reuse. They are generated from the libraries at server start, and every
    name listed must really be taken: a VoiceKit definition in lib\\ (or the
    snippet's Out), or an AutoHotkey built-in on the installed interpreter."""
    text = vk.server_instructions()

    def listed(label: str) -> list[str]:
        line = next(ln for ln in text.splitlines() if ln.startswith(label))
        body = line[len(label):].split(" (and every")[0].rstrip(".")
        return [w.strip() for w in body.split(",") if w.strip()]
    builtins = listed("AutoHotkey built-ins:")
    ours = listed("VoiceKit functions:")
    fams = listed("VoiceKit prefix families (every name starting with):")
    assert "Log" in builtins and "Out" in ours and "Notify" in ours, text
    for f in ("Uia*", "Browser*", "Body*", "Hotkey*", "Bridge*", "Master*", "Wf*"):
        assert f in fams, (f, fams)
    got = _builtins_resolved(builtins)
    assert all(got.get(b.lower()) in ("Func", "Class") for b in builtins), got
    defs = set()
    for f in (REAL_ROOT / "lib").glob("*.ahk"):
        defs |= {m.group(1) for m in re.finditer(r"(?m)^([A-Za-z_]\w*)\(.*\)\s*(?:\{|=>)",
                                                 f.read_text(encoding="utf-8-sig"))}
    prelude = {m.group(1) for m in re.finditer(r"(?m)^([A-Za-z_]\w*)\(.*\)\s*\{",
                                               vk._SNIPPET_PRELUDE)}
    missing = [n for n in ours if n not in defs | prelude]
    assert not missing, f"instructions list names no library defines: {missing}"
    for f in fams:
        assert any(d.startswith(f[:-1]) for d in defs), f"no definition starts with {f}"
    for must in ("`\"", "single quotes", "edit_hotkey_module", "update_macro",
                 "server_info", "%USERPROFILE%", "_Common.ahk", "Browser.ahk"):
        assert must in text, must
    assert len(text) < 4000, "instructions load every session — keep them short"
    try:
        import server
    except ImportError:
        print("      (server half skipped: fastmcp not importable)")
        return
    assert server.mcp.instructions == text


def test_server_info_and_version():
    """WP10: which server is this? VERSION (build-stamped in an install, which
    has no .git), a stale-code flag when server.py / voicekit_writer.py
    changed after the server started, and server_info — other server
    processes and every client config that registers VoiceKit, the likely
    cause of "two sets of VoiceKit tool names"."""
    assert (REAL_ROOT / "VERSION").exists(), "the repo ships a VERSION file"
    assert vk.VERSION == vk.read_version(REAL_ROOT) and vk.VERSION != "unknown"
    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        assert vk.read_version(root) == "unknown"
        (root / "VERSION").write_text("9.9.9-test\n", encoding="utf-8")
        assert vk.read_version(root) == "9.9.9-test", "no .git: the file as written"
        (root / "VERSION").write_text("9.9.9+abc1234.20260930\n", encoding="utf-8")
        assert vk.read_version(root) == "9.9.9+abc1234.20260930", "a stamp is final"
    build = (REAL_ROOT / "build" / "Build-Installer.ps1").read_text(encoding="utf-8")
    assert "'VERSION'" in build and "rev-parse --short HEAD" in build, \
        "the installer build must stamp VERSION into the staged copy"

    with _sandbox():
        h = vk.health()
        assert h["version"] == vk.VERSION and h["install_root"] == str(vk.REPO_ROOT), h
        assert h["server_stale"] is False, h
        first = next(iter(vk._CODE_STAMPS))
        with _patched(_CODE_STAMPS={**vk._CODE_STAMPS, first: 1}):
            h = vk.health()
            assert h["server_stale"] is True and "reconnect" in h.get("note", ""), h
            assert vk.server_info()["server_stale"] is True

    # One server per registration: a venv launcher and its child fold into
    # one entry; this process is marked; a non-VoiceKit server.py is ignored.
    me, parent = os.getpid(), os.getppid()
    srv = str(REAL_ROOT / "mcp" / "server.py")
    rows = [
        {"pid": 101, "ppid": 1, "cmd": f"venv\\python.exe {srv}", "started": "a"},
        {"pid": 102, "ppid": 101, "cmd": f"python.exe {srv}", "started": "a"},
        {"pid": parent, "ppid": 1, "cmd": f'"venv\\python.exe" "{srv}"', "started": "b"},
        {"pid": me, "ppid": parent, "cmd": f'python.exe "{srv}"', "started": "b"},
        {"pid": 300, "ppid": 1, "cmd": r"python.exe C:\other\thing\server.py", "started": "c"},
    ]
    got = vk._mcp_server_processes(rows)
    assert [g["pids"] for g in got] == [[101, 102], sorted({parent, me})], got
    assert [g["this_server"] for g in got] == [False, True], got
    assert got[0]["script"] == srv, got

    # Client configs: every VoiceKit registration, with command + args only
    # (never env — it can hold secrets), and whether it points at this install.
    with tempfile.TemporaryDirectory() as d:
        code_cfg = Path(d) / ".claude.json"
        code_cfg.write_text(json.dumps({
            "mcpServers": {"voicekit": {"command": "py.exe", "args": [srv],
                                        "env": {"SECRET_TOKEN": "hunter2"}},
                           "other": {"command": "node", "args": ["x.js"]}},
            "projects": {"C:/p": {"mcpServers": {"VoiceKit-dev": {
                "command": "py.exe", "args": [r"D:\elsewhere\voicekit\mcp\server.py"]}}}},
        }), encoding="utf-8")
        desk = Path(d) / "claude_desktop_config.json"
        desk.write_text(json.dumps({"mcpServers": {"voicekit": {
            "command": "py.exe", "args": [srv]}}}), encoding="utf-8")
        with _patched(_client_config_files=lambda: [("Claude Code (user)", code_cfg),
                                                    ("Claude Desktop", desk),
                                                    ("missing", Path(d) / "nope.json")],
                      _python_processes=lambda: rows):
            info = vk.server_info()
    assert info["server_pid"] == me and info["version"] == vk.VERSION, info
    regs = info["registrations"]
    assert [r["name"] for r in regs] == ["voicekit", "VoiceKit-dev", "voicekit"], regs
    assert [r["this_install"] for r in regs] == [True, False, True], regs
    assert regs[1]["project"] == "C:/p", regs
    assert "hunter2" not in json.dumps(info) and "SECRET_TOKEN" not in json.dumps(info)
    assert "registered 3 times" in info["note"] and "two sets" in info["note"], info["note"]
    assert "%TEMP%" in regs[0]["config"], "server_info normalizes its own paths"
    try:
        import server
    except ImportError:
        print("      (server half skipped: fastmcp not importable)")
        return
    assert getattr(server.mcp, "version", vk.VERSION) == vk.VERSION
    assert callable(getattr(server, "server_info", None))


def test_paths_normalized_in_responses():
    """WP9 phase 1 (always on): no response carries C:\\Users\\<name>. The
    profile and temp prefixes become %USERPROFILE% / %TEMP% everywhere —
    except authored source a caller will send back — and tools that take a
    path accept that form again."""
    prof = os.environ.get("USERPROFILE") or str(Path.home())
    tmp = tempfile.gettempdir()
    sample = {
        "a": prof + r"\Documents\Invoice.pdf",
        "b": prof.replace("\\", "/").upper() + "/Desktop",
        "c": tmp + r"\vk\snippet.ahk (3)",
        "d": [prof + "2\\x", {"e": "see " + prof}],
        "esc": prof.replace("\\", "\\\\") + "\\\\x",
        prof + r"\key": 1,
        "n": 5,
    }
    got = vk.normalize_paths(sample)
    assert got["a"] == r"%USERPROFILE%\Documents\Invoice.pdf", got
    assert got["b"] == "%USERPROFILE%/Desktop", got
    assert got["c"] == r"%TEMP%\vk\snippet.ahk (3)", "temp before profile: most specific wins"
    assert got["d"] == [prof + "2\\x", {"e": "see %USERPROFILE%"}], \
        "a longer user name that merely starts with the profile's is not the profile"
    assert got["esc"].startswith("%USERPROFILE%"), got
    assert r"%USERPROFILE%\key" in got and got["n"] == 5, got
    kept = vk.normalize_paths({"source": prof, "file": prof}, exempt=("source",))
    assert kept == {"source": prof, "file": "%USERPROFILE%"}, kept
    assert vk.expand_user_paths(r"%userprofile%\Documents") == prof + r"\Documents"
    assert vk.expand_user_paths(r"%TEMP%\x") == tmp.rstrip("\\") + r"\x"
    assert vk.expand_user_paths(r"C:\real\path") == r"C:\real\path"

    with _sandbox(libs=tuple(p.name for p in (REAL_ROOT / "lib").glob("*.ahk"))) as d, \
            _patched(make_shortcut=lambda *a, **k: None, make_loop_shortcut=lambda *a, **k: ""):
        # Path-taking tools accept what a response showed them.
        vk.create_launch_macro("Open Docs Zz", opens=r"%USERPROFILE%\Documents")
        body = (d / "macros" / "OpenDocsZz.ahk").read_text(encoding="utf-8-sig")
        assert (prof + r"\Documents").lower() in body.lower() and "%USERPROFILE%" not in body
        vk.create_workflow("Temp Open Zz", [("run", r"%TEMP%\zz.txt", "", ""),
                                            ("focus", "ahk_exe notepad.exe",
                                             r"notepad %USERPROFILE%\a.txt", "")])
        steps = vk.read_workflow("Temp Open Zz")["steps"]
        assert steps[0]["a"] == tmp.rstrip("\\") + r"\zz.txt", steps
        assert steps[1]["b"] == f"notepad {prof}\\a.txt", steps
        # read_workflow_sheet shows a row's path normalized; handing that row
        # back to run_workflow_batch must type the REAL path.
        (d / "workflows" / "PathBatchZz.steps.txt").write_text(
            "; probe\nask|File||\ntext|{{File}}||\n", encoding="utf-8")
        launched: list = []

        def fake_launch(cmd, launched_as, file, wait, cwd=None):
            launched.append(cmd)
            return {"launched": launched_as, "file": file, "finished": False, "note": ""}
        with _patched(_ahk_processes=lambda: [], _launch_and_report=fake_launch):
            vk.run_workflow_batch("Path Batch Zz", [{"File": r"%USERPROFILE%\Invoice.pdf"}])
        sent = Path(launched[-1][3]).read_text(encoding="utf-8-sig")
        assert (prof + r"\Invoice.pdf").lower() in sent.lower() and "%USERPROFILE%" not in sent, sent
        # update_ai_prompt's previous_prompt is its undo text: exempt like read_ai_prompt.
        assert "previous_prompt" in vk.NORMALIZE_EXEMPT.get("update_ai_prompt", ()), \
            "update_ai_prompt's undo text must round-trip"
        # A source read keeps the real text; its path fields don't.
        try:
            import server
        except ImportError:
            print("      (server half skipped: fastmcp not importable)")
            return
        r = server._guard(vk.read_macro_source, "Open Docs Zz")
        assert r["source"] == vk.read_macro_source("Open Docs Zz")["source"], \
            "authored source must round-trip byte-for-byte"
        assert prof.lower() in r["source"].lower()
        assert r["file"].startswith("%TEMP%"), r["file"]
        assert "%TEMP%" in r["voicekit"]["install_root"], r["voicekit"]
        w = server._guard(vk.read_workflow, "Temp Open Zz")
        assert w["steps"][0]["a"] == steps[0]["a"], "read_workflow steps are exempt too"
        r = server._guard(vk.create_launch_macro, "Open Tmp Zz", None, tmp)
        assert prof.lower() not in json.dumps(r).lower(), r
        from fastmcp.exceptions import ToolError
        def boom():
            raise vk.VoiceKitError("couldn't open " + prof + r"\x.pdf")
        try:
            server._guard(boom)
            assert False
        except ToolError as e:
            assert "%USERPROFILE%" in str(e) and prof.lower() not in str(e).lower(), str(e)


# ---------------------------------------------------------------------------
# run_workflow_batch(source=...) — batch from a CSV or Excel range (WP7).
# Nothing here launches LoopRunner or touches the desktop: every run goes
# through a fake _launch_and_report, and the AHK half is the batchrows.ahk
# harness reading the file with the REAL WfLoopCsvRows.
# ---------------------------------------------------------------------------
def _batch_workflow(d: Path, base: str = "SendAmountZz", labels=("Amount", "Label"),
                    collect: bool = False) -> str:
    lines = ["; probe"] + [f"ask|{l}||" for l in labels] + ["text|{{Amount}}||"]
    if collect:
        lines.append("collect|Result||")
    (d / "workflows" / f"{base}.steps.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return base


@contextlib.contextmanager
def _fake_batch_launch(loops=()):
    """Every launch recorded, none performed; `loops` = LoopRunner cmd lines
    to pretend are running."""
    launched: list = []

    def fake(cmd, launched_as, file, wait, cwd=None):
        launched.append(cmd)
        return {"launched": launched_as, "file": file, "finished": False, "note": ""}
    procs = [{"pid": 4242, "started": "", "cmd": c} for c in loops]
    with _patched(_ahk_processes=lambda: procs, _launch_and_report=fake):
        yield launched


def _batch_records(path: Path) -> list:
    import csv as _csv
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(_csv.reader(f))


def _refused(fn, *needles) -> str:
    try:
        fn()
    except vk.VoiceKitError as e:
        for n in needles:
            assert n in str(e), (n, str(e))
        return str(e)
    raise AssertionError(f"expected a refusal mentioning {needles}")


def _openpyxl_or_skip():
    try:
        import openpyxl  # noqa: F401
        return openpyxl
    except ImportError:
        print("      (skipped: openpyxl not importable in this interpreter — run under "
              "mcp\\.venv, where requirements.txt installs it)")
        return None


def test_batch_source_csv_encodings_delimiters_and_leading_zeros():
    """A CSV is read as UTF-8 (with or without BOM) and falls back to
    Windows-1252 — Excel's plain 'CSV (Comma delimited)'; ';' (European
    Excel) and tab are detected from the first line; text passes exactly as
    written, so '00123' keeps its zeros."""
    with _sandbox() as d, _fake_batch_launch() as launched:
        base = _batch_workflow(d)
        src = d / "amounts.csv"
        src.write_bytes("Amount,Label\r\n00123,\"Jos\u00e9 \u2013 1,234\"\r\n".encode("utf-8-sig"))
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True)
        assert (r["encoding"], r["delimiter"]) == ("utf-8", "comma"), r
        assert r["preview"] == [{"source_row": 2, "Amount": "00123",
                                 "Label": "Jos\u00e9 \u2013 1,234"}], r["preview"]
        src.write_bytes("Amount,Label\r\n7,\"Jos\u00e9 \u2013 1,234\"\r\n".encode("cp1252"))
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True)
        assert r["encoding"] == "cp1252", r
        assert r["preview"][0]["Label"] == "Jos\u00e9 \u2013 1,234", r["preview"]
        src.write_text("Amount;Label\n1234,50;Box 1\n0042;Box 2\n", encoding="utf-8")
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True)
        assert r["delimiter"] == "semicolon", r
        assert [p["Amount"] for p in r["preview"]] == ["1234,50", "0042"], r["preview"]
        # A title line with no delimiter above a ';' table doesn't decide it.
        src.write_text("Invoice detail\n\nAmount;Label\n5;Box 1\n", encoding="utf-8")
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                  cell_range="A3:B")
        assert r["delimiter"] == "semicolon" and r["preview"][0]["Label"] == "Box 1", r
        # A one-column file stays one column.
        (d / "one.csv").write_text("Amount\n5\n", encoding="utf-8")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(d / "one.csv"),
                                               dry_run=True), "No column for the label 'Label'")
        # Bytes undefined in Windows-1252 are refused, never typed as '?'.
        src.write_bytes(b"Amount,Label\r\n7,\x81\x8d\xff\xfe\r\n")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True),
                 "isn't UTF-8 or Windows-1252")
        tsv = d / "amounts.tsv"
        tsv.write_text("Label\tAmount\nBox, one\t5\n", encoding="utf-8")
        r = vk.run_workflow_batch("Send Amount Zz", source=str(tsv), dry_run=True)
        assert r["delimiter"] == "tab" and r["preview"][0] == \
            {"source_row": 2, "Amount": "5", "Label": "Box, one"}, r
        assert not launched, "a dry run launches nothing"
        assert base


def test_batch_source_header_mapping_is_caseless_and_errors_list_headers():
    """columns maps a label to a header name case-insensitively; a label left
    out matches a header of its own name; an unresolvable label says which
    headers exist; a columns key that isn't a label is refused."""
    with _sandbox() as d, _fake_batch_launch():
        _batch_workflow(d)
        src = d / "amounts.csv"
        src.write_text("amt ,DONE,label\n5,,Box 1\n", encoding="utf-8")
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                  columns={"AMOUNT": "AMT"})
        assert r["columns"] == {"Amount": {"column": "A", "header": "amt"},
                                "Label": {"column": "C", "header": "label"}}, r["columns"]
        assert r["preview"][0] == {"source_row": 2, "Amount": "5", "Label": "Box 1"}
        e = _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src),
                                                   dry_run=True),
                     "No column for the label 'Amount'", "amt (A)", "DONE (B)", "label (C)")
        assert "5" not in e.split("Headers")[0], e
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                               columns={"Amount": "Total"}),
                 "No column for the label 'Amount'", "'Total'")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                               columns={"Amount": "amt", "Price": "col:A"}),
                 "'Price'", "isn't an ask label")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                               columns={"Amount": "amt"}, require=["Nope"]),
                 "require names 'Nope'")


def test_batch_source_col_letters_numbers_and_open_ended_range():
    """header=False maps by 'col:C' or column number (A = 1); an open-ended
    range under a title block runs to the last used row (trailing blank rows
    trimmed, not reported as skips); a column outside the range is refused."""
    with _sandbox() as d, _fake_batch_launch():
        _batch_workflow(d)
        src = d / "amounts.csv"
        src.write_text("Invoice amounts for client\n\nAmount,Done,Label\n10,,Box 1\n20,x,Box 2\n"
                       "30,,Box 3\n,,\n,,\n", encoding="utf-8")
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                  cell_range="A3:C")
        assert r["cell_range"] == "A3:C6" and r["header_row"] == 3, r
        assert r["source_rows"] == 3 and r["passes"] == 3 and r["pass_rows"] == "4-6", r
        assert r["skipped"] == {}, r["skipped"]
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                  cell_range="a4:c", header=False,
                                  columns={"Amount": "col:A", "Label": 3})
        assert r["columns"] == {"Amount": {"column": "A"}, "Label": {"column": "C"}}, r
        assert [(p["source_row"], p["Amount"], p["Label"]) for p in r["preview"]] == \
            [(4, "10", "Box 1"), (5, "20", "Box 2"), (6, "30", "Box 3")], r["preview"]
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                  cell_range="A:C", header=False,
                                  columns={"Amount": "col:A", "Label": "col:C"})
        assert r["cell_range"] == "A1:C6", r
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                               cell_range="A4:B", header=False,
                                               columns={"Amount": "col:A", "Label": "col:C"}),
                 "outside the range")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                               cell_range="A4:C", header=False,
                                               columns={"Amount": "col:A"}),
                 "'Label' has none")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                               cell_range="4:9"), "isn't a range")


def test_batch_source_skips_report_exact_source_rows():
    """skip_blank drops a row whose MAPPED cells are all blank even when an
    unmapped column is filled (the case WfLoopCsvRows alone would run);
    require and skip_if_filled drop theirs; every skip is reported by source
    row number, and pass N is the Nth kept row."""
    with _sandbox() as d, _fake_batch_launch() as launched:
        _batch_workflow(d)
        src = d / "amounts.csv"
        src.write_text("Amount,Done,Label,Note\n"
                       "10,,Box 1,\n"          # 2 runs
                       ",,,just a note\n"       # 3 blank (mapped cells)
                       "20,yes,Box 2,\n"       # 4 already done
                       ",,Box 3,\n"             # 5 required Amount blank
                       "40,,Box 4,\n"           # 6 runs
                       "50,x,,\n"               # 7 already done
                       ",,,\n"                  # 8 blank
                       "80,,,\n",               # 9 runs
                       encoding="utf-8")
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src),
                                  require=["amount"], skip_if_filled=["done"])
        assert r["pass_rows"] == "2, 6, 9" and r["passes"] == 3, r
        assert r["skipped"] == {"already_done": {"count": 2, "rows": "4, 7"},
                                "blank": {"count": 2, "rows": "3, 8"},
                                "required_blank": {"count": 1, "rows": "5"}}, r["skipped"]
        assert r["skip_if_filled"] == ["B"], r
        recs = _batch_records(Path(launched[-1][3]))
        assert recs[0] == ["Amount", "Label", vk.BATCH_META_COLUMN], recs
        assert recs[1:] == [["10", "Box 1", "2"], ["40", "Box 4", "6"], ["80", "", "9"]], recs
        # skip_blank=False keeps the all-blank rows (the metadata column keeps
        # WfLoopCsvRows from dropping them, so the pass numbering still holds).
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), skip_blank=False,
                                  skip_if_filled=["col:B"])
        assert r["pass_rows"] == "2-3, 5-6, 8-9", r
        recs = _batch_records(Path(launched[-1][3]))
        assert recs[2] == ["", "", "3"], recs


def test_batch_source_excel_errors_are_refused_by_row():
    """A cell holding an Excel error value (#N/A, #REF!, ...) the workflow
    would read refuses the whole batch, naming row and column — while one in
    a row that's skipped anyway doesn't."""
    with _sandbox() as d, _fake_batch_launch() as launched:
        _batch_workflow(d)
        src = d / "amounts.csv"
        src.write_text("Amount,Done,Label\n10,,Box 1\n#N/A,,Box 2\n#REF!,x,Box 3\n"
                       "5,,#DIV/0!\n", encoding="utf-8")
        e = _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src),
                                                   skip_if_filled=["Done"]),
                     "row 3 (column A: #N/A)", "row 5 (column C: #DIV/0!)", "Nothing was run")
        assert "row 4" not in e, "a row skipped as done isn't read, so its error is moot"
        assert "Box" not in e, "an error never carries cell values"
        assert not launched


def test_batch_source_xlsx_values_dates_and_errors():
    """xlsx via openpyxl (read_only, data_only): ints without '.0', floats at
    Excel's precision, dates ISO by default or per date_format, booleans,
    a '00000'-formatted number keeping its zeros, tabs by name, gap rows —
    and an error cell refused with its row. Skips politely without openpyxl."""
    openpyxl = _openpyxl_or_skip()
    if openpyxl is None:
        return
    import datetime as _dt
    with _sandbox() as d, _fake_batch_launch() as launched:
        _batch_workflow(d, labels=("Amount", "Label", "When"))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = "not this tab"
        inv = wb.create_sheet("Inv-2025")
        inv.append(["Invoice detail"])
        inv.append([])
        inv.append(["Amount", "Done", "Label", "When"])
        inv.append([1234.5, None, "Box 1", _dt.date(2026, 3, 15)])       # row 4
        inv.append([7.0, None, True, _dt.datetime(2026, 4, 1, 9, 30)])   # row 5
        inv.append([None, None, None, None])                             # row 6 gap
        inv.append([123, "x", "zip", None])                              # row 7 done
        inv["A7"].number_format = "00000"
        inv.append([0.1 + 0.2, None, "Box 3", None])                     # row 8
        inv.append([42, None, "Box 4", None])                            # row 9
        inv["A9"].number_format = "00000"
        inv.append(["#N/A", "x", "skipped anyway", None])               # row 10
        p = d / "Invoice workbook.xlsx"
        wb.save(p)
        r = vk.run_workflow_batch("Send Amount Zz", source=str(p), sheet="inv-2025",
                                  cell_range="A3:D", skip_if_filled=["col:B"], dry_run=True)
        assert r["source_kind"] == "xlsx" and r["sheet"] == "Inv-2025", r
        assert r["cell_range"] == "A3:D10" and r["source_rows"] == 7, r
        assert r["pass_rows"] == "4-5, 8-9", r
        assert r["skipped"] == {"already_done": {"count": 2, "rows": "7, 10"},
                                "blank": {"count": 1, "rows": "6"}}, r["skipped"]
        got = [(p_["Amount"], p_["Label"], p_["When"]) for p_ in r["preview"]]
        assert got == [("1234.5", "Box 1", "2026-03-15"), ("7", "TRUE", "2026-04-01 09:30:00"),
                       ("0.3", "Box 3", ""), ("00042", "Box 4", "")], got
        r = vk.run_workflow_batch("Send Amount Zz", source=str(p),
                                  cell_range="'Inv-2025'!A3:D5", date_format="%m/%d/%Y",
                                  dry_run=True)
        assert [p_["When"] for p_ in r["preview"]] == ["03/15/2026", "04/01/2026"], r["preview"]
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(p), sheet="Inv-2025",
                                               cell_range="A3:D", dry_run=True),
                 "row 10 (column A: #N/A)")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(p), sheet="Inv2025",
                                               dry_run=True), "no tab named 'Inv2025'", "Summary, Inv-2025")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(p), sheet="Inv-2025",
                                               date_format="mm/dd", dry_run=True),
                 "isn't a strftime pattern")
        # The default tab is the active one. Its lone cell can't be told from
        # data (no label named, nothing data-like under it), so the error
        # names letters only, never the cell (WP9: cells are client data).
        e = _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(p),
                                                   dry_run=True),
                     "No column for the label 'Amount'", "doesn't look like a header row",
                     "header=False", "col:A")
        assert "not this tab" not in e, e
        assert not launched
    assert vk._num_text(1234.4999999999998) == "1234.5"
    assert vk._num_text(1e-7) == "0.0000001" and vk._num_text(1.5e20) == "150000000000000000000"
    assert vk._num_text(-0.0) == "0" and vk._num_text(-2.50) == "-2.5"


def test_batch_source_xlsx_with_a_wrong_recorded_dimension_reads_every_row():
    """openpyxl's read-only mode trusts the sheet's recorded <dimension>, and
    some writers record just 'A1'. An open-ended range must still read to
    where the data really ends (reset_dimensions), not stop at row 1."""
    openpyxl = _openpyxl_or_skip()
    if openpyxl is None:
        return
    import zipfile
    with _sandbox() as d, _fake_batch_launch():
        _batch_workflow(d)
        wb = openpyxl.Workbook()
        for row in (["Amount", "Label"], [1, "a"], [2, "b"], [None, None], [4, "d"]):
            wb.active.append(row)
        good = d / "good.xlsx"
        wb.save(good)
        bad = d / "bad.xlsx"
        with zipfile.ZipFile(good) as zin, zipfile.ZipFile(bad, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("xl/worksheets/sheet"):
                    data, n = re.subn(rb'<dimension ref="[^"]*"', b'<dimension ref="A1"', data)
                    assert n == 1, "fixture: no <dimension> to break"
                zout.writestr(item, data)
        for p in (good, bad):
            r = vk.run_workflow_batch("Send Amount Zz", source=str(p), cell_range="A:B",
                                      dry_run=True)
            assert r["pass_rows"] == "2-3, 5" and r["cell_range"] == "A1:B5", (p.name, r)
            assert [x["Label"] for x in r["preview"]] == ["a", "b", "d"], r["preview"]
            r = vk.run_workflow_batch("Send Amount Zz", source=str(p), dry_run=True)
            assert r["passes"] == 3, (p.name, r)


def test_batch_source_missing_openpyxl_names_the_fix():
    """openpyxl is imported lazily: a CSV never needs it, and an .xlsx without
    it gets a clear 'pip install openpyxl into mcp\\.venv' — not a traceback."""
    import sys as _sys
    with _sandbox() as d, _fake_batch_launch():
        _batch_workflow(d)
        (d / "k.xlsx").write_bytes(b"PK\x03\x04 not really")
        csvf = d / "k.csv"
        csvf.write_text("Amount,Label\n1,a\n", encoding="utf-8")
        saved = _sys.modules.get("openpyxl", "absent")
        _sys.modules["openpyxl"] = None          # import openpyxl -> ImportError
        try:
            _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(d / "k.xlsx"),
                                                   dry_run=True),
                     "pip install openpyxl", "mcp\\.venv")
            assert vk.run_workflow_batch("Send Amount Zz", source=str(csvf),
                                         dry_run=True)["passes"] == 1
        finally:
            if saved == "absent":
                del _sys.modules["openpyxl"]
            else:
                _sys.modules["openpyxl"] = saved
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(d / "k.xls")),
                 "No file")
        (d / "k.xls").write_bytes(b"x")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(d / "k.xls")),
                 "save it as .xlsx")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source="k.csv"),
                 "full path")


def test_batch_source_dry_run_launches_nothing_and_hides_values_otherwise():
    """dry_run maps and previews — the ONLY response that carries cell values
    — and launches nothing, writes no batch file, and isn't refused by a
    running loop (it says a real run would be). A real run's response never
    contains a value. Through the server, the preview is path-normalized."""
    try:
        import server        # OUTSIDE the sandbox: its instructions are built at import
    except ImportError:
        server = None
    with _sandbox() as d:
        _batch_workflow(d, collect=True)
        prof = os.environ.get("USERPROFILE", "") or str(Path.home())
        src = d / "amounts.csv"
        src.write_text("Amount,Label\nSENTINEL-4471," + prof + "\\client.pdf\n11,Box\n"
                       + "".join(f"{i},Row {i}\n" for i in range(20)), encoding="utf-8")
        loop = [f'"{vk.AHK_EXE}" "{d}\\lib\\LoopRunner.ahk" OtherFlowZz']
        with _fake_batch_launch(loops=loop) as launched:
            r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                      preview_rows=3)
            assert r["dry_run"] is True and r["launched"] is False, r
            assert len(r["preview"]) == 3 and r["preview"][0]["Amount"] == "SENTINEL-4471"
            assert "would be refused" in r["note"], r["note"]
            assert not launched and not list((d / "logs").glob("mcp-batch*.csv"))
            _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src)),
                     "already running")
        with _fake_batch_launch() as launched:
            r = vk.run_workflow_batch("Send Amount Zz", source=str(src))
            assert launched and r["passes"] == 22, r
            assert "SENTINEL" not in json.dumps(r) and "client.pdf" not in json.dumps(r), r
            assert "never written" in r["results_land"] and "appended" in r["results_land"], r
            if server is None:
                print("      (server half skipped: fastmcp not importable)")
                return
            g = server._guard(vk.run_workflow_batch, "Send Amount Zz", None, 0,
                              source=str(src), dry_run=True)
            # Phase 1 normalizes the prefix; phase 2 (privacy masking, on by
            # default) tokenizes the name under it — and the token expands back.
            lab = g["preview"][0]["Label"]
            assert re.fullmatch(r"%USERPROFILE%\\<file#[0-9a-f]{4,12}>\.pdf", lab), lab
            assert vk.expand_user_paths(lab).lower() == (prof + "\\client.pdf").lower(), lab
            assert g["source"].startswith("%TEMP%"), g["source"]
            # The server passes every new parameter through by name.
            g = server.run_workflow_batch.fn("Send Amount Zz", source=str(src),
                                             cell_range="A1:B3", dry_run=True) \
                if hasattr(server.run_workflow_batch, "fn") else \
                server.run_workflow_batch("Send Amount Zz", source=str(src),
                                          cell_range="A1:B3", dry_run=True)
            assert g["pass_rows"] == "2-3" and g["cell_range"] == "A1:B3", g


def test_batch_source_next_cell_range_and_resume():
    """next_cell_range starts after the last row read (open-ended, or the
    original end while rows remain); when a waited run stops partway it
    starts at the first pass that didn't finish, and resume_with maps by
    column letter because a resumed range has no header row."""
    with _sandbox() as d:
        _batch_workflow(d)
        src = d / "amounts.csv"
        src.write_text("Amount,Done,Label\n" + "".join(f"{i},,Box {i}\n" for i in range(1, 9)),
                       encoding="utf-8")
        with _fake_batch_launch():
            r = vk.run_workflow_batch("Send Amount Zz", source=str(src), cell_range="A1:C",
                                      dry_run=True)
            assert r["next_cell_range"] == "A10:C", r
            r = vk.run_workflow_batch("Send Amount Zz", source=str(src), cell_range="A1:C20",
                                      dry_run=True)
            # A closed range past the file's end stops where the data does
            # (like xlsx): no phantom blank skips, and the unread remainder
            # of the asked-for range is what resumes.
            assert r["next_cell_range"] == "A10:C20" and r["skipped"] == {}, r
            assert r["resume_with"] == {"cell_range": "A10:C20", "header": False,
                                        "columns": {"Amount": "col:A", "Label": "col:C"}}, r
            r = vk.run_workflow_batch("Send Amount Zz", source=str(src), cell_range="A1:C5",
                                      dry_run=True, skip_blank=False, date_format="%d.%m.%Y")
            assert r["next_cell_range"] == "A6:C", r
            assert r["resume_with"]["date_format"] == "%d.%m.%Y" and                 r["resume_with"]["skip_blank"] is False, r["resume_with"]

        def partial(base, since=""):
            return {"outcome": "failed", "ok": False, "what_happened": "A step failed.",
                    "steps_total": 2, "passes_done": 3, "passes_total": 7}
        with _fake_batch_launch(), _patched(_run_report=partial):
            r = vk.run_workflow_batch("Send Amount Zz", source=str(src), cell_range="A1:C9",
                                      skip_if_filled=["Done"], require=["Amount"],
                                      wait_seconds=5)
        # Passes ran for rows 2..8 (row 9 = 8th data row... all 8 kept); 3 done.
        assert r["passes"] == 8 and r["next_cell_range"] == "A5:C9", r
        assert r["resume_with"]["skip_if_filled"] == ["col:B"], r["resume_with"]
        assert r["resume_with"]["require"] == ["Amount"], r["resume_with"]
        assert "source row 5" in r["resume_note"], r
        # Feeding resume_with back runs exactly the rest.
        with _fake_batch_launch():
            r2 = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                       **r["resume_with"])
        assert r2["pass_rows"] == "5-9", r2


def test_batch_source_own_sheet_runs_in_place():
    """The workflow's own inputs sheet goes straight to LoopRunner (no batch
    file), so the loop's fromSheet write-back fills collected values beside
    each row; filters are refused on it (they'd break that mapping), and an
    ANSI-saved sheet is refused because the loop reads it as UTF-8."""
    with _sandbox() as d, _fake_batch_launch() as launched:
        base = _batch_workflow(d, collect=True)
        sheet = d / "workflows" / f"{base}.inputs.csv"
        sheet.write_bytes("\ufeffAmount,Label,Result\r\n5,Box 1,\r\n,,\r\n6,Box 2,done\r\n"
                          .encode("utf-8"))
        r = vk.run_workflow_batch("Send Amount Zz", source=str(sheet).upper())
        assert launched and launched[-1][3] == os.path.abspath(sheet), launched
        assert not list((d / "logs").glob("mcp-batch*.csv")), "no copy of the sheet is made"
        # ...and the REAL loop, which spells the sheet "<root>\lib\..\workflows\",
        # recognizes what we passed as its own sheet (a plain string compare
        # didn't, so every "write back beside each row" run appended instead).
        def is_own(batch_path: str) -> str:
            out = d / "ownsheet.txt"
            p = subprocess.run([vk.AHK_EXE, "/ErrorStdOut", str(BATCH_HARNESS), "/ownsheet",
                                batch_path, str(out), str(d), base],
                               capture_output=True, text=True, timeout=30)
            assert p.returncode == 0, f"harness failed: {p.stdout}{p.stderr}"
            return out.read_text(encoding="utf-8-sig")
        assert is_own(launched[-1][3]) == "1"
        assert is_own(str(d / "amounts.csv")) == "0"
        assert r["source_kind"] == "workflow_sheet" and r["pass_rows"] == "2, 4", r
        assert r["skipped"] == {"blank": {"count": 1}}, r
        assert "beside each row" in r["results_land"], r
        assert "next_cell_range" not in r
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(sheet),
                                               require=["Amount"]),
                 "own inputs sheet", "require")
        sheet.write_bytes("Amount,Label\r\nJos\u00e9,Box\r\n".encode("cp1252"))
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(sheet)),
                 "isn't saved as UTF-8")


def test_batch_source_file_round_trips_through_real_wflooprows():
    """The Python-writes / AHK-reads seam: the batch file a source run writes,
    read by the REAL WfLoopCsvRows (mcp\\_conformance\\batchrows.ahk). Every
    row arrives — including one whose inputs are all blank (skip_blank=False)
    — in order, and commas, quotes, newlines, unicode, leading zeros and
    dates survive exactly."""
    openpyxl = _openpyxl_or_skip()
    import datetime as _dt
    with _sandbox() as d, _fake_batch_launch() as launched:
        _batch_workflow(d, labels=("Amount", "Label"))
        src = d / "amounts.csv"
        tricky = ['00123', 'a, "b"\nline two', 'Jos\u00e9 \u2603 \u2013', ' padded ', '#hash; semi']
        import csv as _csv
        with open(src, "w", encoding="utf-8-sig", newline="") as f:
            w = _csv.writer(f)
            w.writerow(["Amount", "Label"])
            for i, t in enumerate(tricky):
                if i == 2:
                    w.writerow(["", ""])       # an all-blank row mid-range (source row 4)
                w.writerow([t, t[::-1]])

        def check(rows, src_rows):
            # Read right away: the next batch run clears earlier batch files.
            batch = Path(launched[-1][3])
            out = d / "ahk-rows.txt"
            r = subprocess.run([vk.AHK_EXE, "/ErrorStdOut", str(BATCH_HARNESS), str(batch),
                                str(out), "Amount", "Label"],
                               capture_output=True, text=True, timeout=30)
            assert r.returncode == 0, f"harness failed: {r.stdout}{r.stderr}"
            lines = out.read_text(encoding="utf-8-sig").split("\n")[:-1]
            assert lines and not lines[0].startswith("ERR|"), lines
            got = [tuple(vk.wf_decode(v) for v in ln.split("|")[1:]) for ln in lines]
            assert got == rows, (got, rows)
            # srcRec is the batch record; the metadata column names the source row.
            assert [ln.split("|")[0] for ln in lines] == [str(i) for i in range(2, len(rows) + 2)]
            assert [r_[-1] for r_ in _batch_records(batch)[1:]] == src_rows

        vk.run_workflow_batch("Send Amount Zz", source=str(src), skip_blank=False)
        want = [(t, t[::-1]) for t in tricky]
        check(want[:2] + [("", "")] + want[2:], [str(n) for n in range(2, 8)])
        if openpyxl is not None:
            wb = openpyxl.Workbook()
            wb.active.append(["Amount", "Label"])
            wb.active.append([_dt.date(2026, 3, 15), "0042"])
            wb.active.append([1234.5, "x"])
            p = d / "k.xlsx"
            wb.save(p)
            vk.run_workflow_batch("Send Amount Zz", source=str(p), date_format="%m/%d/%Y")
            check([("03/15/2026", "0042"), ("1234.5", "x")], ["2", "3"])


def test_batch_rows_and_source_are_exclusive_and_rows_mode_is_unchanged():
    """Exactly one of rows / source; the source-only options are refused with
    rows (rows are filtered by the caller); rows mode writes the same file as
    before (no metadata column)."""
    with _sandbox() as d, _fake_batch_launch() as launched:
        _batch_workflow(d)
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz"), "exactly one of rows")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", [{"Amount": 1, "Label": 2}],
                                               source=str(d / "x.csv")), "exactly one of rows")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", [{"Amount": 1, "Label": 2}],
                                               require=["Amount"], dry_run=True),
                 "require, dry_run only apply with source=")
        vk.run_workflow_batch("Send Amount Zz", [{"amount": "1,5", "LABEL": "x"}])
        assert Path(launched[-1][3]).read_bytes() == b'\xef\xbb\xbfAmount,Label\r\n"1,5",x\r\n'


# ---------------------------------------------------------------------------
# WP9 phase 2 — privacy masking (mcp/privacy.py). Every test runs against a
# FAKE profile (privacy.PROFILE_OVERRIDE) with VoiceKit installed inside it —
# a typical installed copy's layout — so nothing here depends on, or writes tokens for,
# the real profile.
# ---------------------------------------------------------------------------
import privacy as pv  # noqa: E402

_T = r"[0-9a-f]{4,12}"


@contextlib.contextmanager
def _privacy_sandbox(settings: str | None = None, libs=("_Common.ahk",)):
    """A fake user profile: VoiceKit INSTALLED under it
    (AppData\\Local\\Programs\\VoiceKit, like %LOCALAPPDATA%\\Programs\\VoiceKit
    on an installed copy), a client's invoice under Documents, an org OneDrive.
    REPO_ROOT and privacy's profile/temp point there for the duration."""
    from types import SimpleNamespace
    tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
    base = Path(tmp.name)
    prof = base / "Users" / "TestUser"
    temp = prof / "AppData" / "Local" / "Temp"
    root = prof / "AppData" / "Local" / "Programs" / "VoiceKit"
    orig = (vk.REPO_ROOT, vk.VOICE_MACROS, vk.reload_voicekit,
            pv.PROFILE_OVERRIDE, pv.TEMP_OVERRIDE)
    try:
        for sub in ("macros", "hotkeys", "workflows", "prompts", "logs", "startmenu", "lib"):
            (root / sub).mkdir(parents=True)
        temp.mkdir(parents=True, exist_ok=True)
        for lib in libs:
            shutil.copyfile(REAL_ROOT / "lib" / lib, root / "lib" / lib)
        client = prof / "Documents" / "Clients" / "Smith John"
        client.mkdir(parents=True)
        inv = client / "2025 Invoice Smith.pdf"
        inv.write_bytes(b"%PDF-1.4 fixture")
        (prof / "Downloads").mkdir()
        (prof / "OneDrive - Example Corp" / "Documents" / "Doe LLC").mkdir(parents=True)
        if settings is not None:
            (root / "logs" / "settings.ini").write_text(settings, encoding="utf-8")
        vk.REPO_ROOT, vk.VOICE_MACROS = root, root / "startmenu"
        vk.reload_voicekit = lambda: "not_running"
        pv.PROFILE_OVERRIDE, pv.TEMP_OVERRIDE = str(prof), str(temp)
        yield SimpleNamespace(base=base, prof=prof, temp=temp, root=root, client=client, inv=inv)
    finally:
        (vk.REPO_ROOT, vk.VOICE_MACROS, vk.reload_voicekit,
         pv.PROFILE_OVERRIDE, pv.TEMP_OVERRIDE) = orig
        pv._MAPS.clear()
        pv._KEY_CACHE.clear()
        pv._SETTINGS_CACHE.clear()
        tmp.cleanup()


def _pv(obj, **kw):
    """What server._guard does to a response (without needing fastmcp)."""
    return pv.protect(obj, root=vk.REPO_ROOT, normalize=vk.normalize_paths_text, **kw)


def _raises(fn, *needles, exc=Exception) -> str:
    try:
        fn()
    except exc as e:          # noqa: BLE001
        for n in needles:
            assert n in str(e), (n, str(e))
        return str(e)
    raise AssertionError(f"expected an error mentioning {needles}")


def _server_or_none():
    try:
        import server
        return server
    except ImportError:
        print("      (server half skipped: fastmcp not importable)")
        return None


def test_privacy_paths_mode_masks_names_under_the_profile():
    """The default mode (no [Privacy] section at all): every file and folder
    name under the profile becomes a token that keeps the extension; the
    well-known folders, the VoiceKit install root (itself under the profile
    on an installed copy), the Voice Macros Start Menu folder and %TEMP%'s own
    name stay readable; prose after a path is left alone."""
    with _privacy_sandbox() as s:
        inv = str(s.inv)
        vm = s.prof / "AppData" / "Roaming" / "Microsoft" / "Windows" / "Start Menu" / \
            "Programs" / "Voice Macros" / "open send invoice.lnk"
        r = _pv({"file": inv, "dl": str(s.prof / "Downloads"),
                 "od": str(s.prof / "OneDrive - Example Corp" / "Documents" / "Doe LLC" / "t.xlsx"),
                 "vk": str(s.root / "macros" / "Send Invoice.ahk"),
                 "tmp": str(s.temp / "vk-snip" / "out.txt"),
                 "lnk": str(vm), "n": 5})
        assert re.fullmatch(rf"%USERPROFILE%\\Documents\\<dir#{_T}>\\<dir#{_T}>\\<file#{_T}>\.pdf",
                            r["file"]), r["file"]
        blob = json.dumps(r)
        assert "Smith" not in blob and "Clients" not in blob and "Doe" not in blob, blob
        assert r["dl"] == "%USERPROFILE%\\Downloads", r["dl"]
        assert re.fullmatch(rf"%USERPROFILE%\\OneDrive - Example Corp\\Documents\\<dir#{_T}>"
                            rf"\\<file#{_T}>\.xlsx", r["od"]), r["od"]
        assert r["vk"] == "%USERPROFILE%\\AppData\\Local\\Programs\\VoiceKit\\macros\\Send Invoice.ahk", \
            "the install root and automation names stay readable"
        assert re.fullmatch(rf"%TEMP%\\<dir#{_T}>\\<file#{_T}>\.txt", r["tmp"]), r["tmp"]
        assert r["lnk"].endswith("\\Voice Macros\\open send invoice.lnk"), r["lnk"]
        assert r["n"] == 5
        # Prose around a path: an existing file ends where its name ends; a
        # quoted folder ends at the closing quote; a missing file at its extension.
        note = _pv(f"Saved {inv} to disk; '{s.client}' is the folder. "
                   f"{s.client / '2026 Draft Smith.docx'} was not found.")
        assert re.fullmatch(rf"Saved %USERPROFILE%\\Documents\\<dir#{_T}>\\<dir#{_T}>\\<file#{_T}>\.pdf"
                            rf" to disk; '%USERPROFILE%\\Documents\\<dir#{_T}>\\<dir#{_T}>' is the "
                            rf"folder\. %USERPROFILE%\\Documents\\<dir#{_T}>\\<dir#{_T}>\\<file#{_T}>"
                            rf"\.docx was not found\.", note), note
        # Forward slashes, and a JSON-escaped (doubled backslash) spelling.
        fwd = _pv(inv.replace("\\", "/"))
        assert re.fullmatch(rf"%USERPROFILE%/Documents/<dir#{_T}>/<dir#{_T}>/<file#{_T}>\.pdf", fwd), fwd
        esc = _pv(json.dumps({"p": inv}))
        assert "Smith" not in esc and "<file#" in esc, esc
        # Stable: the same name, the same token, every call; masking a masked
        # response again changes nothing.
        assert _pv({"x": inv})["x"] == r["file"]
        assert _pv(r["file"]) == r["file"]
        assert _pv({"k": inv.upper()})["k"].split("\\")[-1].lower() == \
            r["file"].split("\\")[-1].lower(), \
            "case-folded: the same name in another case is the same token"


def test_privacy_tokens_are_stable_across_processes_and_the_key_is_made_once():
    """A token handed out yesterday must still mean the same file today: the
    per-install key is created once (and only when something is masked), the
    map persists, and a fresh process derives the very same token."""
    with _privacy_sandbox() as s:
        key = s.root / "logs" / pv.KEY_FILE
        tmap = s.root / "logs" / pv.MAP_FILE
        _pv({"a": "no paths here", "b": str(s.prof / "Documents"), "c": str(s.root / "macros")})
        assert not key.exists() and not tmap.exists(), \
            "nothing to mask -> no key and no map written"
        first = _pv(str(s.inv))
        assert key.exists() and tmap.exists()
        key_bytes = key.read_bytes()
        assert re.fullmatch(rb"[0-9a-f]{64}\r?\n", key_bytes), key_bytes
        pv._MAPS.clear()
        pv._KEY_CACHE.clear()               # as a restarted server would start
        assert _pv(str(s.inv)) == first
        assert key.read_bytes() == key_bytes, "the key is reused, never rewritten"
        import sys
        code = ("import sys; sys.path.insert(0, sys.argv[1]); import privacy as pv; "
                "pv.PROFILE_OVERRIDE, pv.TEMP_OVERRIDE = sys.argv[2], sys.argv[3]; "
                "print(pv.protect(sys.argv[4], root=sys.argv[5]))")
        out = subprocess.run([sys.executable, "-c", code, str(REAL_ROOT / "mcp"), str(s.prof),
                              str(s.temp), str(s.inv), str(s.root)],
                             capture_output=True, text=True, timeout=60)
        assert out.stdout.strip() == first, (out.stdout, out.stderr, first)
        assert key.read_bytes() == key_bytes
        data = json.loads(tmap.read_text(encoding="utf-8"))
        names = {v["t"] for v in data["tokens"].values()}
        assert {"Clients", "Smith John", "2025 Invoice Smith"} <= names, names
        # A different install (another key) gets different tokens.
        with _privacy_sandbox() as s2:
            assert _pv(str(s2.inv)) != first


def test_privacy_bare_names_are_masked_in_the_same_response():
    """A snippet that prints the path and then just the file name must not
    leak the name the path already hid — in any field of the same response.
    A longer word that merely starts with the name is left alone."""
    with _privacy_sandbox() as s:
        r = _pv({"output": f"{s.inv}\n2025 Invoice Smith.pdf\nfolder: Smith John\nSmith Johnson "
                           f"stays\n2025 Invoice SMITH",
                 "note": "Opened Smith John's file", "Smith John": "keys are left alone"})
        lines = r["output"].split("\n")
        toks = re.findall(rf"<(?:dir|file)#{_T}>", lines[0])
        assert len(toks) == 3, lines[0]
        assert lines[1] == toks[2] + ".pdf", lines
        assert lines[2] == "folder: " + toks[1], lines
        assert lines[3] == "Smith Johnson stays", "whole names only"
        assert lines[4] == toks[2], "caseless"
        assert r["note"] == f"Opened {toks[1]}'s file", r["note"]
        assert "Smith John" in r, "dict keys never get the bare-name pass"
        # A name that only ever appears bare is NOT caught (documented limit).
        assert _pv({"t": "Client: Smith John"})["t"] == "Client: Smith John"


def test_privacy_mask_roots_and_settings_reload():
    """MaskRoots adds roots outside the profile (a drive folder, a UNC share):
    the root stays, every name under it is masked; a sibling that merely
    starts with the root's name is not. The mode is re-read when the file
    changes; an unknown Mode falls back to paths and says so."""
    ini = "[Privacy]\nMode=paths\nMaskRoots=Q:\\Clients; \\\\srv\\share\\clients ;relative\\x\n"
    with _privacy_sandbox(settings=ini) as s:
        cfg = pv.settings(s.root)
        assert cfg["mask_roots"] == ["Q:\\Clients", "\\\\srv\\share\\clients"], cfg
        assert "isn't a full path" in cfg["note"], cfg
        r = _pv({"a": "Q:\\Clients\\Doe LLC\\2025 statement.pdf",
                 "b": "\\\\srv\\share\\clients\\Lee Family\\statement.pdf",
                 "c": "Q:\\Clientsx\\y.pdf", "d": "Q:\\Other\\z.pdf",
                 "e": "q:/clients/Doe LLC/notes.txt"})
        assert re.fullmatch(rf"Q:\\Clients\\<dir#{_T}>\\<file#{_T}>\.pdf", r["a"]), r["a"]
        assert re.fullmatch(rf"\\\\srv\\share\\clients\\<dir#{_T}>\\<file#{_T}>\.pdf", r["b"]), r["b"]
        assert r["c"] == "Q:\\Clientsx\\y.pdf" and r["d"] == "Q:\\Other\\z.pdf", r
        assert re.fullmatch(rf"q:/clients/<dir#{_T}>/<file#{_T}>\.txt", r["e"]), r["e"]
        assert r["a"].split("\\")[2] == r["e"].split("/")[2], "same folder, same token"
        ini_p = s.root / "logs" / "settings.ini"
        ini_p.write_text("[privacy]\nmode = OFF\n", encoding="utf-16")   # AHK IniWrite's encoding
        os.utime(ini_p, ns=(time.time_ns(), time.time_ns() + 10_000_000))
        assert pv.mode(s.root) == "off"
        assert _pv(str(s.inv)) == "%USERPROFILE%\\Documents\\Clients\\Smith John\\2025 Invoice Smith.pdf", \
            "off: phase-1 normalization only"
        ini_p.write_text("[Privacy]\nMode=none\n", encoding="utf-8")
        os.utime(ini_p, ns=(time.time_ns(), time.time_ns() + 20_000_000))
        assert pv.mode(s.root) == "paths" and "using paths" in pv.settings(s.root)["note"]


def test_privacy_off_mode_writes_nothing():
    with _privacy_sandbox(settings="[Privacy]\nMode=off\n") as s:
        r = _pv({"f": str(s.inv), "t": "SSN 123-45-6789"})
        assert r == {"f": "%USERPROFILE%\\Documents\\Clients\\Smith John\\2025 Invoice Smith.pdf",
                     "t": "SSN 123-45-6789"}, r
        assert not (s.root / "logs" / pv.KEY_FILE).exists()


def test_privacy_strict_masks_ssn_and_ein_one_way():
    """strict = paths + SSN/EIN. Conservative on purpose: bare nine digits
    only right after an ID label; dates, phone numbers, amounts and longer
    digit runs are left alone. ID tokens are never stored, revealed or
    expanded."""
    with _privacy_sandbox(settings="[Privacy]\nMode=strict\n") as s:
        txt = ("SSN 123-45-6789, EIN 12-3456789, SSN: 987654321, TIN 123 45 6780, "
               "Social Security Number 222-33-4444. Not these: 2025-10-15, 555-123-4567, "
               "12-3456, $1,234.56, invoice 123456789, 123-45-67890, 12-34567890, "
               "2025-12-31T10:00, ZIP 12345-6789")
        r = _pv({"t": txt, "f": str(s.inv)})
        t = r["t"]
        for gone in ("123-45-6789,", "EIN 12-3456789", "987654321", "123 45 6780", "222-33-4444"):
            assert gone not in t, (gone, t)
        assert t.count("<ssn#") == 4 and t.count("<ein#") == 1, t
        for kept in ("2025-10-15", "555-123-4567", "12-3456,", "$1,234.56", "invoice 123456789",
                     "123-45-67890", "12-34567890", "2025-12-31T10:00", "12345-6789"):
            assert kept in t, (kept, t)
        assert re.fullmatch(rf"%USERPROFILE%\\Documents\\<dir#{_T}>\\<dir#{_T}>\\<file#{_T}>\.pdf",
                            r["f"]), "strict includes paths mode"
        assert _pv("123-45-6789") == _pv("id 123-45-6789").split(" ")[1], "stable per number"
        ssn = re.search(r"<ssn#[0-9a-f]+>", t).group(0)
        _raises(lambda: vk.reveal(ssn), "one-way")
        _raises(lambda: vk.expand_user_paths(f"x {ssn}"), "one-way")
        assert "6789" not in (s.root / "logs" / pv.MAP_FILE).read_text(encoding="utf-8")
        # A labelled ID is masked even with the label in another case.
        assert "<ssn#" in _pv("ssn#123456789")


def test_privacy_tokens_round_trip_through_tool_inputs():
    """What a masked response showed, handed back as input, acts on the REAL
    file: run_automation args (raw script and workflow), run_workflow_batch
    rows (a read_workflow_sheet row as shown), create_launch_macro opens and
    a workflow run step. An unknown token is an error, never a guess; a
    token in saved code or a non-path step field is refused."""
    libs = tuple(p.name for p in (REAL_ROOT / "lib").glob("*.ahk"))
    with _privacy_sandbox(libs=libs) as s, \
            _patched(make_shortcut=lambda *a, **k: None, make_loop_shortcut=lambda *a, **k: ""):
        shown = _pv(str(s.inv))
        assert "<file#" in shown
        launched: list = []

        def fake(cmd, launched_as, file, wait, cwd=None):
            launched.append(cmd)
            return {"launched": launched_as, "file": file, "finished": False, "note": ""}
        (s.root / "macros" / "ArgsZz.ahk").write_text("; raw\n", encoding="utf-8")
        (s.root / "workflows" / "SelZz.steps.txt").write_text(
            "; probe\ntext|{{selected_file}}||\n", encoding="utf-8")
        (s.root / "macros" / "SelZz.ahk").write_text("; stub\n", encoding="utf-8")
        with _patched(_launch_and_report=fake):
            vk.run_automation("Args Zz", args=[shown, "plain", "%TEMP%\\x"])
            assert launched[-1][2:] == [str(s.inv), "plain", str(s.temp) + "\\x"], launched[-1]
            vk.run_automation("Sel Zz", args=[shown])
            assert Path(launched[-1][2]).resolve() == s.inv.resolve(), launched[-1]
            server = _server_or_none()
            if server:
                g = server._guard(vk.run_automation, "Args Zz", 0, [shown])
                assert g["args"] == [shown], "the reply masks the args it echoes"
                assert g["voicekit"]["privacy"] == "paths", g["voicekit"]
        # Batch rows — exactly what read_workflow_sheet shows for a row.
        _batch_workflow(s.root)
        (s.root / "workflows" / "SendAmountZz.inputs.csv").write_text(
            f"Amount,Label\n5,{s.inv}\n", encoding="utf-8-sig")
        sheet_rows = _pv(vk.read_workflow_sheet("Send Amount Zz"))["sheet"]["rows"]
        assert sheet_rows == [{"Amount": "5", "Label": shown}], sheet_rows
        with _fake_batch_launch() as bl:
            vk.run_workflow_batch("Send Amount Zz", sheet_rows)
        sent = Path(bl[-1][3]).read_text(encoding="utf-8-sig")
        assert str(s.inv) in sent and "<file#" not in sent, sent
        # opens + a workflow's run step (both land in saved files, expanded).
        vk.create_launch_macro("Open Invoice Zz", opens=shown)
        body = (s.root / "macros" / "OpenInvoiceZz.ahk").read_text(encoding="utf-8-sig")
        assert str(s.inv) in body and "<file#" not in body
        vk.create_workflow("Open Invoice Flow Zz", [("run", shown, "", ""),
                                               ("capture", "Out", f'type "{shown}"', "")])
        steps = vk.read_workflow("Open Invoice Flow Zz")["steps"]
        assert steps[0]["a"] == str(s.inv) and steps[1]["b"] == f'type "{s.inv}"', steps
        # Unknown / forged tokens: an error, never a guess.
        _raises(lambda: vk.run_automation("Args Zz", args=["%USERPROFILE%\\<file#ffff>.pdf"]),
                "Unknown privacy token <file#ffff>")
        tok = re.search(rf"<file#{_T}>", shown).group(0)
        tmap = s.root / "logs" / pv.MAP_FILE
        data = json.loads(tmap.read_text(encoding="utf-8"))
        data["tokens"][tok[6:-1]]["t"] = "Someone Else"          # a tampered map entry
        tmap.write_text(json.dumps(data), encoding="utf-8")
        _raises(lambda: vk.expand_user_paths(shown), "Unknown privacy token")
        pv._MAPS.clear()
        # Tokens never get baked into saved code, text or other step fields.
        _raises(lambda: vk.create_launch_macro("Bad Zz", ahk_body=f'Run("{shown}")'),
                "privacy token", "ahk_body")
        _raises(lambda: vk.create_workflow("Bad Flow Zz", [("text", shown, "", "")]),
                "privacy token")
        _raises(lambda: vk.create_snippet("/bz", shown), "privacy token")
        assert not (s.root / "macros" / "BadZz.ahk").exists()


def test_privacy_snippet_code_tokens_unmask_and_reveal_are_audited():
    """run_ahk_snippet: a masked path inside a quoted string is expanded
    (escaped — names may hold ' ; or a backtick) and the code really opens the
    file; a token outside a string is refused. The output comes back masked,
    or unmasked with unmask=True — audited, like every reveal, by token and
    tool, never by the name."""
    libs = tuple(p.name for p in (REAL_ROOT / "lib").glob("*.ahk"))
    with _privacy_sandbox(libs=libs) as s:
        odd = s.prof / "Documents" / "O'Brien; Pat`s"
        odd.mkdir()
        (odd / "notes.txt").write_text("x", encoding="utf-8")
        shown = _pv(str(s.inv))
        shown_odd = _pv(str(odd / "notes.txt"))
        assert "Brien" not in shown_odd
        code = (f'Out(FileExist("{shown}") ? "found" : "missing")\n'
                f"Out(FileExist('{shown_odd}') ? \"found2\" : \"missing2\")  ; {shown}\n"
                f'Out("{shown}")\n'
                f'SplitPath("{shown}", &nm)\nOut(nm)')
        got = vk.run_ahk_snippet(code, timeout_s=30)
        lines = got["output"].splitlines()
        assert lines[:2] == ["found", "found2"], got
        assert lines[2] == str(s.inv) and lines[3] == s.inv.name, got
        _raises(lambda: vk.run_ahk_snippet(f"x := 1\nOut({shown})", timeout_s=5),
                "outside a string literal")
        server = _server_or_none()
        audit = s.root / "logs" / pv.AUDIT_FILE
        if server:
            g = server._guard(vk.run_ahk_snippet, code, 30)
            assert "Smith" not in json.dumps(g), g["output"]
            gl = g["output"].splitlines()
            assert gl[2] == shown and gl[3] == shown.split("\\")[-1], gl
            assert not audit.exists(), "a masked call audits nothing"
            u = server._guard(vk.run_ahk_snippet, code, 30, _unmask=True)
            assert "%USERPROFILE%\\Documents\\Clients\\Smith John\\2025 Invoice Smith.pdf" in u["output"]
            a = audit.read_text(encoding="utf-8")
            assert "run_ahk_snippet | unmask" in a and "Smith" not in a, a
        # reveal: one token, or a whole masked path.
        ftok = re.search(rf"<file#{_T}>", shown).group(0)
        assert vk.reveal(ftok)["text"] == "2025 Invoice Smith"
        r = vk.reveal(shown)
        assert r["text"] == "%USERPROFILE%\\Documents\\Clients\\Smith John\\2025 Invoice Smith.pdf", r
        assert len(r["tokens"]) == 3
        a = audit.read_text(encoding="utf-8")
        assert f"| reveal | {ftok}" in a and "Smith" not in a and "2025 Invoice" not in a, a
        _raises(lambda: vk.reveal("%USERPROFILE%\\<dir#abcdef12>"), "Unknown privacy token")
        _raises(lambda: vk.reveal("plain text"), "No privacy token")
        if server:
            g = server._guard(vk.reveal, shown)
            assert g["text"] == r["text"], "reveal's answer is not masked again"


def test_privacy_exempt_fields_and_logs_through_the_server():
    """Authored source round-trips byte-for-byte even when it holds a client
    path; everything around it is masked — including read_log('errors'),
    whose AHK error lines carry full file paths."""
    server = _server_or_none()
    if server is None:
        return
    with _privacy_sandbox() as s:
        src = f'; hand-written\nRun("{s.inv}")\n'
        (s.root / "macros" / "RawZz.ahk").write_text(src, encoding="utf-8")
        (s.root / "workflows" / "RawFlowZz.steps.txt").write_text(
            f"; p\ntext|{s.inv}||\n", encoding="utf-8")
        r = server._guard(vk.read_macro_source, "Raw Zz")
        assert r["source"] == vk.read_macro_source("Raw Zz")["source"] and str(s.inv) in r["source"]
        w = server._guard(vk.read_workflow, "Raw Flow Zz")
        assert w["steps"][0]["a"] == str(s.inv), "read_workflow steps are exempt"
        (s.root / "logs" / "errors.log").write_text(
            f"2026-09-30 10:00:00 | SplitPages.ahk | Error: Couldn't open {s.inv} "
            f"[{s.client}\\split.ahk line 12]\n", encoding="utf-8")
        e = server._guard(vk.read_log, "errors", 10, None)
        blob = json.dumps(e)
        assert "Smith" not in blob and "<file#" in blob, blob
        # A ToolError's text goes through the same pipeline.
        _batch_workflow(s.root)
        missing = s.client / "missing.csv"
        ex = _raises(lambda: server._guard(vk.run_workflow_batch, "Send Amount Zz", None, 0,
                                           source=str(missing)), "No file at")
        assert "Smith" not in ex and "<dir#" in ex, ex
        assert "voicekit" in e and e["voicekit"]["privacy"] == "paths", e.get("voicekit")


def test_privacy_adversarial_leaks_stay_closed():
    """The review's leak hunt, pinned: a missing name beside a sibling that
    merely starts the same ('Smith' vs 'Smith Jane.pdf'), a quoted path, a
    file:/// URL, %APPDATA% and '~' spellings, a Path object, redirected
    shell folders, a readable tree that would contain the profile, a name
    that came IN as a token and goes out bare, an exception outside the
    writer's own error types, a non-dict result, trimmed tokens never
    reissued, a damaged map never overwritten, and the bare-name pass staying
    fast with thousands of names."""
    import sys
    libs = tuple(p.name for p in (REAL_ROOT / "lib").glob("*.ahk"))
    with _privacy_sandbox(libs=libs) as s:
        (s.client / "Smith").write_text("x", encoding="utf-8")
        leaky = ("Jane", "Smith", "Clients", "Jr", "Invoice", "Brown")

        def clean(v):
            blob = json.dumps(v, default=str)
            for w in leaky:
                assert w not in blob, (w, blob)
            return v
        clean(_pv({"f": f"{s.client / 'Smith Jane.pdf'} was not found",
                   "d": f"{s.client / 'Smith Jane'} Moved",
                   "q": f'Folder "{s.client.parent / "Smith John Jr"}" is missing',
                   "e": str(s.client.parent / "Smith John Jr"),
                   "u": "file:///" + str(s.inv).replace("\\", "/").replace(" ", "%20"),
                   "a": "%APPDATA%\\Clients\\Smith John\\x.pdf",
                   "t": "~\\Documents\\Clients\\Smith John",
                   "p": s.inv}))
        # Existing names followed by prose still end where the name ends.
        r = _pv(f"{s.client / 'Smith'} is here")
        assert r.endswith(" is here") and "Smith" not in r, r
        # A Documents folder redirected to a share / another drive is masked
        # like a MaskRoot (the root itself stays readable).
        pv.KNOWN_FOLDERS_OVERRIDE = ["\\\\filer\\home$\\testuser\\Documents", "R:\\Docs"]
        try:
            r = clean(_pv({"a": "\\\\filer\\home$\\testuser\\Documents\\Clients\\Smith John\\x.pdf",
                           "b": "R:\\Docs\\Brown LLC\\statement.pdf"}))
            assert r["a"].startswith("\\\\filer\\home$\\testuser\\Documents\\<dir#"), r
            assert r["b"].startswith("R:\\Docs\\<dir#"), r
        finally:
            pv.KNOWN_FOLDERS_OVERRIDE = None
        # A "readable" tree that contains the profile must not unmask it.
        old_ahk = os.environ.get("VOICEKIT_AHK")
        os.environ["VOICEKIT_AHK"] = str(s.prof / "AutoHotkey64.exe")
        try:
            clean(_pv(str(s.inv)))
        finally:
            if old_ahk is None:
                os.environ.pop("VOICEKIT_AHK", None)
            else:
                os.environ["VOICEKIT_AHK"] = old_ahk
        shown = _pv(str(s.inv))
        ftok = re.search(rf"<file#{_T}>", shown).group(0)
        assert len(ftok) - len("<file#>") >= 6, "6+ hex: collisions stay rare at the map cap"
        # A name that came IN as a token is masked where it goes out bare.
        pv.start_call()
        vk.expand_user_paths(f"{ftok}.pdf")
        assert _pv("cannot open 2025 Invoice Smith.pdf", known=pv.take_expanded()) == \
            f"cannot open {ftok}.pdf"
        server = _server_or_none()
        if server:
            from fastmcp.exceptions import ToolError

            def boom():
                raise KeyError(str(s.inv))
            clean(_raises(lambda: server._guard(boom), "Internal error", exc=ToolError))
            clean(server._guard(lambda: [str(s.inv)]))
            clean(server._guard(lambda: str(s.inv)))
            got = server._guard(vk.run_ahk_snippet, f'Out("{ftok}.pdf")', 30)
            assert got["output"].strip() == f"{ftok}.pdf", got["output"]
        # Trimmed ids are retired, never handed to a new name.
        tmap = s.root / "logs" / pv.MAP_FILE
        old_cap = pv.MAP_CAP
        pv.MAP_CAP = 5
        try:
            for i in range(8):
                _pv(str(s.prof / "Documents" / f"Brown {i:02d}"))
            data = json.loads(tmap.read_text(encoding="utf-8"))
            assert data["retired"] and not set(data["retired"]) & set(data["tokens"]), data
            gone = data["retired"][0]
            _raises(lambda: vk.expand_user_paths(f"<dir#{gone}>"), "Unknown privacy token")
        finally:
            pv.MAP_CAP = old_cap
        # A damaged map fails closed and is never overwritten with a fresh one.
        tmap.write_text("{damaged", encoding="utf-8")
        pv._MAPS.clear()
        _raises(lambda: _pv(str(s.prof / "Documents" / "Brown New")), "can't be read",
                exc=pv.PrivacyError)
        assert tmap.read_text(encoding="utf-8") == "{damaged"
        tmap.unlink()
        pv._MAPS.clear()
        # Thousands of distinct masked names in one big log stay cheap.
        big = "\n".join(f"2026-09-30 | x | {s.client}\\Client {i:04d}.pdf missing (Client {i:04d})"
                        for i in range(5000))
        t0 = time.perf_counter()
        out = _pv({"lines": big})["lines"]
        took = time.perf_counter() - t0
        assert "Client 0" not in out, out[:300]
        assert took < 10, f"masking 5000 paths took {took:.1f} s"
        print(f"      (5000-path log masked in {took * 1000:.0f} ms)")
    sys.stdout.flush()


def test_batch_header_errors_never_quote_a_data_row():
    """WP9 7b: when the row read as the header is really the first client's
    row, the mapping error names column letters only — never its cells."""
    with _sandbox() as d, _fake_batch_launch():
        _batch_workflow(d)
        src = d / "amounts.csv"
        src.write_text("Smith John,1234.50\nDoe LLC,6\n", encoding="utf-8")
        e = _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src),
                                                   dry_run=True),
                     "No column for the label 'Amount'", "doesn't look like a header row",
                     "header=False", "A-B")
        assert "Smith" not in e and "1234" not in e and "Doe" not in e, e
        # A real header (titles over numbers) is still listed, to help map it.
        src.write_text("Client,Amt\nSmith John,5\n", encoding="utf-8")
        _refused(lambda: vk.run_workflow_batch("Send Amount Zz", source=str(src),
                                               dry_run=True),
                 "Headers in this range: Client (A), Amt (B)")
        # ...and a data-row "header" mapped by letter isn't echoed as a header.
        src.write_text("Smith John,1234.50\nDoe LLC,6\n", encoding="utf-8")
        r = vk.run_workflow_batch("Send Amount Zz", source=str(src), dry_run=True,
                                  columns={"Amount": "col:B", "Label": "col:A"})
        assert r["columns"] == {"Amount": {"column": "B"}, "Label": {"column": "A"}}, r["columns"]
    assert vk._value_like("1,234.50") and vk._value_like("2025-10-15") and \
        vk._value_like("123-45-6789") and not vk._value_like("Amount") and \
        not vk._value_like("Box 1")


def _run_all():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for t in tests:
        try:
            t()
            print(f"PASS  {t.__name__}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            print(f"FAIL  {t.__name__}\n      {e}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    return failed


if __name__ == "__main__":
    import sys
    sys.exit(1 if _run_all() else 0)
